#!/usr/bin/env python3
"""Nonintrusive native observations and the explicitly labelled Linux font fixture.

The only application override is the authored workbook renderer's font-path list.
There are no recognizer, geometry, ownership, persistence or receipt doubles.
"""
import argparse
import atexit
from dataclasses import asdict
from datetime import date, datetime
from hashlib import file_digest, sha256
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import stat
import sys
import traceback

MODELS = {
    'text-detector-dbnet-v2': ('217e61488f38bbe0e5ae9b281bad7f80a5f7b4f4',
        '7a9bf8063ac3947473a3d31e1d2b40fad6aa6d5e669ab4b32650578fb40a2b75'),
    'text-recognizer-parseq-middle-v2': ('1bcaac6ff8eb37edd80b20536e497391f234c038',
        '0fbc62e8654c1ed3d565cce0471b998e42a3bcb40b3dcc8d936a59b480a8b139'),
    'layout-parser-rtdtrv2-v2': ('6772f1b6ea5b694da2e8fe19cf6d38f453822616',
        'd6118e5d0ec70d5eafc5cf1b830977a199a12b3223277ba9419f834ddc05f044'),
    'table-structure-recognizer-rtdtrv2-open-beta': ('91103cbcefc7f9d0ba8234fd3caba28b6d564a46',
        '264fa489afadfc81279f6d880df4643c96133a95289673f2068b276569587eb0'),
}
OUT = None
AUDIT = {'violations': [], 'connections': [], 'foreign_keys_connections': [],
         'case_bindings': [], 'observer_errors': []}
OBSERVED = []
INPUTS = []
MODEL_LOADS = []
FONT = None
GLYPHS = None
ISOLATION = None
BIND_SOCKETS = []


def digest(path):
    with Path(path).open('rb') as handle:
        return file_digest(handle, 'sha256').hexdigest()


def save(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')


def download_models(cache, output):
    from huggingface_hub import hf_hub_download
    from huggingface_hub.errors import EntryNotFoundError

    cache.mkdir(parents=True, exist_ok=False)
    recorded = {}
    try:
        for name, (revision, expected) in MODELS.items():
            repo = 'KotaroKinoshita/yomitoku-' + name
            model = Path(hf_hub_download(repo, 'model.safetensors', revision=revision,
                         cache_dir=cache / 'hub', token=False))
            if digest(model) != expected:
                raise ValueError(f'model byte mismatch: {repo}@{revision}')
            entry = {'repo': repo, 'revision': revision, 'model_sha256': expected,
                     'model_path': str(model.relative_to(cache))}
            recorded[name] = entry
            try:
                config = Path(hf_hub_download(repo, 'config.json', revision=revision,
                              cache_dir=cache / 'hub', token=False))
            except EntryNotFoundError as exc:
                response = exc.response
                if response is None or response.status_code != 404:
                    raise
                entry['config'] = {'status': 404, 'url': response.url,
                                   'error_type': type(exc).__name__}
                marker = model.parents[2] / '.no_exist' / revision / 'config.json'
                if not marker.is_file():
                    raise ValueError('actual pinned config 404 cache marker missing')
            else:
                entry['config'] = {'status': 200, 'sha256': digest(config)}
                raise ValueError('config now exists; frozen reference recorded absence, review required')
            # Yomitoku asks for the default ref. Bind ONLY this owned offline cache
            # to the already downloaded, hash-verified historical revision.
            reference = model.parents[2] / 'refs/main'
            reference.parent.mkdir(exist_ok=True)
            reference.write_text(revision)
            model.resolve().chmod(0o444)
            entry['offline_main_binding'] = revision
    finally:
        save(output, recorded)


def descriptor_path(fd, *, directory=False):
    """Resolve this live descriptor, never a cached/reusable descriptor number."""
    before = os.fstat(fd)
    if directory and not stat.S_ISDIR(before.st_mode):
        raise ValueError('dir_fd is not a directory')
    if sys.platform == 'linux':
        raw = os.readlink(f'/proc/self/fd/{fd}')
    elif sys.platform == 'darwin':
        import fcntl

        # Darwin SDK sys/fcntl.h: F_GETPATH=50, buffer size MAXPATHLEN=1024.
        raw = os.fsdecode(fcntl.fcntl(fd, 50, bytes(1024)).split(b'\0', 1)[0])
    else:
        raise ValueError('descriptor paths unsupported on this platform')
    if not Path(raw).is_absolute():
        raise ValueError('descriptor has no physical filesystem path')
    path = Path(raw).resolve(strict=True)
    physical, after = path.stat(), os.fstat(fd)
    identity = lambda item: (item.st_dev, item.st_ino, stat.S_IFMT(item.st_mode))
    if identity(before) != identity(physical) or identity(before) != identity(after):
        raise ValueError('descriptor physical identity changed')
    return path


def audit(event, args):
    def refuse(reason):
        violation = {'event': event, 'reason': reason}
        AUDIT['violations'].append(violation)
        # Omit source lines/locals and address text; retain the actual caller chain.
        violation['stack'] = [{'file': frame.filename, 'line': frame.lineno, 'function': frame.name}
                              for frame in traceback.extract_stack()[:-1]]
        if event.startswith('socket.'):
            family = args[2] if event == 'socket.getaddrinfo' else getattr(args[0] if args else None, 'family', None)
            address = args[:2] if event == 'socket.getaddrinfo' else (args[-1] if args else None)
            violation['socket_family'] = getattr(family, 'name', str(family))
            violation['address'] = {'kind': type(address).__name__, 'value': '<redacted>'}
            if isinstance(address, tuple) and len(address) > 1 and type(address[1]) is int:
                violation['address']['port'] = address[1]
        raise RuntimeError(reason)

    def owned(path, dir_fd=-1, *, follow_final=True, write_sink=False):
        original = str(path) if isinstance(path, int) else os.fsdecode(path)
        try:
            if isinstance(path, int):
                resolved = descriptor_path(path)
            else:
                resolved = Path(os.fsdecode(path))
                if not resolved.is_absolute() and dir_fd not in (None, -1):
                    base = descriptor_path(dir_fd, directory=True)
                    if not base.is_relative_to(OUT):
                        raise ValueError('dir_fd outside owned native evidence root')
                    resolved = base / resolved
                # Entry mutations do not follow the final symlink. Trailing slash
                # and dot components require directory resolution by the OS.
                if follow_final or original.endswith(('/', '/.', '/..')) or original in ('.', '..'):
                    resolved = resolved.resolve()
                else:
                    resolved = resolved.parent.resolve() / resolved.name
        except (OSError, ValueError, RuntimeError) as exc:
            refuse('unresolved native write ownership: ' + original + ': ' + str(exc))
        # Only opening the existing null device for writing is a permitted sink;
        # it is not an exemption for chmod, removal, linking, or other mutations.
        if (write_sink and resolved == Path(os.devnull)
                and (isinstance(path, int) or original == os.devnull)):
            return resolved
        if not resolved.is_relative_to(OUT):
            refuse('write outside owned native evidence root: ' + original)
        return resolved

    if event.startswith('socket.'):
        if event == 'socket.__new__':
            return
        if event == 'socket.bind' and ISOLATION is not None and isolation.capability_bind(*args):
            record = {'operation': 'IPv6 loopback ephemeral capability reservation',
                      'requested': ['::1', 0], 'os_result': 'pending', 'closed': False}
            AUDIT['allowed_capability_binds'].append(record)
            BIND_SOCKETS.append((args[0], record))
            return
        refuse('native network/external process forbidden')
    elif event in ('subprocess.Popen', 'os.system', 'os.exec', 'os.posix_spawn', 'os.fork', 'os.forkpty'):
        refuse('native network/external process forbidden')
    elif event == 'sqlite3.connect':
        path = Path(args[0]).resolve()
        if not path.is_relative_to(OUT):
            refuse('SQLite outside owned native evidence root')
        AUDIT['connections'].append(str(path))
    elif event == 'open':
        _path, mode, flags = args
        if ((mode and any(x in mode for x in 'wax+'))
                or flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC | os.O_APPEND)):
            # CPython's open audit event omits dir_fd: a relative os.open path
            # cannot safely be attributed to cwd rather than an unknown dir_fd.
            if not isinstance(_path, int) and not Path(os.fsdecode(_path)).is_absolute():
                refuse('relative open write has no auditable dir_fd')
            owned(_path, write_sink=True)
    elif event in ('os.remove', 'os.rmdir'):
        owned(args[0], args[1], follow_final=False)
    elif event in ('os.mkdir', 'os.chmod'):
        owned(args[0], args[2], follow_final=event == 'os.chmod')
    elif event == 'os.truncate':
        owned(args[0])
    elif event in ('os.rename', 'os.link'):
        source = owned(args[0], args[2], follow_final=False)
        owned(args[1], args[3], follow_final=False)
        # link's audit event omits follow_symlinks. Validate both interpretations
        # of a symlink source, not just the directory entry being linked.
        if event == 'os.link' and source.is_symlink():
            owned(source)
    elif event == 'os.symlink':
        destination = owned(args[1], args[2], follow_final=False)
        owned(destination.parent / os.fsdecode(args[0]))


def finish():
    sys.setprofile(None)
    for sock, record in BIND_SOCKETS:
        record['closed'] = sock.fileno() == -1
    save(OUT / 'audit.json', AUDIT)
    save(OUT / 'native-observed.json', OBSERVED)
    save(OUT / 'inputs.json', INPUTS)
    save(OUT / 'model-loads.json', MODEL_LOADS)


def observe_bind(_frame, event, function):
    if event not in ('c_return', 'c_exception') or getattr(function, '__name__', None) != 'bind':
        return
    for sock, record in BIND_SOCKETS:
        if getattr(function, '__self__', None) is sock and record['os_result'] == 'pending':
            record['os_result'] = 'bound' if event == 'c_return' else 'os_error'
            if event == 'c_return':
                record['actual_address'] = list(sock.getsockname())


def install_audit(isolation_args=None):
    global OUT, ISOLATION, isolation
    OUT = Path(os.environ['OWNED_EVIDENCE_ROOT']).resolve(strict=True)
    if OUT.is_symlink() or OUT == Path('/'):
        raise ValueError('physical owned output root required')
    # Only real /proc and kernel syscall verification can enable the exception.
    # The environment alone, or a supplied precomputed state dictionary, cannot.
    if isolation_args is not None:
        import native_network_isolation as isolation

        ISOLATION = isolation.verify_arguments(isolation_args)
    AUDIT['isolation'] = ISOLATION
    AUDIT['allowed_capability_binds'] = []
    sys.addaudithook(audit)
    atexit.register(finish)
    if ISOLATION is not None:
        sys.setprofile(observe_bind)


def value(item):
    import numpy as np
    from PIL import Image

    if isinstance(item, np.ndarray):
        return {'type': 'ndarray', 'dtype': str(item.dtype), 'shape': list(item.shape),
                'pixels_sha256': sha256(item.tobytes()).hexdigest()}
    if isinstance(item, Image.Image):
        return {'type': 'PIL.Image', 'mode': item.mode, 'size': list(item.size),
                'pixels_sha256': sha256(item.tobytes()).hexdigest()}
    if isinstance(item, dict):
        return {key: value(val) for key, val in item.items()}
    if isinstance(item, (list, tuple)):
        return {'type': type(item).__name__, 'items': [value(val) for val in item]}
    if isinstance(item, np.generic):
        return {'type': str(type(item)), 'value': item.item()}
    if isinstance(item, (date, datetime)):
        return {'type': type(item).__name__, 'value': item.isoformat()}
    if item is None or type(item) in (bool, int, float, str):
        return item
    raise TypeError(f'unobserved native value type: {type(item)}')


def font_fixture():
    global FONT, GLYPHS
    from fontTools.ttLib import TTFont
    from PIL import ImageFont
    from src.services import workbook_pdf_renderer as renderer
    from src.services import master_order_form_template_service as master
    from openpyxl import load_workbook

    path = Path(os.environ['NATIVE_CI_FONT'])
    expected = os.environ['NATIVE_CI_FONT_SHA256']
    if path.resolve(strict=True) != path or digest(path) != expected:
        raise ValueError('explicit Noto font identity changed')
    font = ImageFont.truetype(str(path), 18)
    if not isinstance(font, ImageFont.FreeTypeFont) or str(font.path) != str(path):
        raise ValueError('explicit FreeType font was not loaded')
    with TTFont(path, fontNumber=0) as table:
        GLYPHS = set(table.getBestCmap())
    # Validate the immutable workbook and the existing fixture's generated text
    # before either renderer executes; capture_input also checks every final cell.
    workbook = load_workbook(master.MASTER_TEMPLATE_PATH)
    try:
        texts = [renderer._cell_text(cell.value, cell.number_format)
                 for sheet in workbook for row in sheet for cell in row]
        texts.append('Synthetic native unit facility 0123456789/-.\u732e\u7acb\u4e3b\u671d\u663c\u5915\u65e5\u6708\u706b\u6c34\u6728\u91d1\u571f\u5e38\u98df\u533a\u5206\u66dc\u4ed8')
        missing = {ord(char) for text in texts for char in text if not char.isspace()} - GLYPHS
        if missing:
            raise ValueError(f'explicit font lacks authored source glyphs: {sorted(missing)}')
    finally:
        workbook.close()
    FONT = {'path': str(path), 'sha256': expected,
            'package_version': os.environ['NATIVE_CI_FONT_PACKAGE'],
            'label': 'Linux-authored-input font override', 'font_index': 0}
    original = renderer._DEFAULT_FONT_PATHS
    renderer._DEFAULT_FONT_PATHS = [str(path)]
    return lambda: setattr(renderer, '_DEFAULT_FONT_PATHS', original)


def capture_input(fixture):
    from openpyxl import load_workbook
    import fitz
    from src.services import master_order_form_template_service as master
    from src.services import workbook_pdf_renderer as renderer
    from test_modernization_owned_ocr_workflow import FACILITY, WEEK

    f, identity = fixture
    folder = OUT / f'input-{len(INPUTS)}'
    folder.mkdir()
    workbook_path = f.root / 'synthetic-numeric-fax.xlsx'
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        schema = {row[0].value: row[1].value for row in workbook['generated_template_schema']
                  if len(row) > 1 and row[0].value}
        if (schema['facility_id'] != FACILITY or schema['week_value'] != WEEK
                or identity['week_id'] != WEEK or identity['physical_menu_row_count'] != 56
                or len(f.expected) != 56 or len(f.plan) != 56):
            raise ValueError('authored facility/week/physical identity mismatch')
        if {cell: str(sheet[cell].value) for cell in f.expected} != f.expected:
            raise ValueError('workbook numeric oracle differs from existing fixture')
        cells = []
        all_text = ''
        for row in sheet:
            for cell in row:
                text = renderer._cell_text(cell.value, cell.number_format)
                all_text += text
                cells.append({'cell': cell.coordinate, 'value': value(cell.value),
                              'style': list(cell._style or ()), 'number_format': cell.number_format})
        if f.facility['facility_name'] not in all_text:
            raise ValueError('authored facility name not in rendered worksheet')
        absent = sorted({ord(char) for char in all_text if not char.isspace()} - GLYPHS)
        if absent:
            raise ValueError(f'explicit font lacks worksheet glyphs: {absent}')
        scenario = {'facility': FACILITY, 'week': WEEK, 'expected': f.expected,
                    'physical_identity': value(identity), 'target_plan': value(f.plan),
                    'worksheet_cells': cells, 'print_area': str(sheet.print_area),
                    'physical_row_plan': json.loads(schema['week_menu_physical_row_plan']),
                    'glyph_coverage_verified': True}
    finally:
        workbook.close()
    target = folder / 'authored.pdf'
    shutil.copyfile(f.pdf, target)
    shutil.copyfile(workbook_path, folder / 'authored.xlsx')
    if digest(target) != f.receipt['sha256']:
        raise ValueError('actual selected PDF receipt differs from authored PDF')
    pages = []
    with fitz.open(target) as document:
        pdf_metadata = document.metadata
        if len(document) != 1:
            raise ValueError('existing numeric fixture must render one page')
        for index, page in enumerate(document):
            pixels = page.get_pixmap(dpi=144, colorspace=fitz.csRGB, alpha=False)
            pixels.save(folder / f'page-{index + 1}.png')
            pages.append({'width': pixels.width, 'height': pixels.height, 'channels': pixels.n,
                          'dpi': 144, 'pixels_sha256': sha256(pixels.samples).hexdigest()})
    INPUTS.append({'scenario': scenario, 'font': FONT, 'page_pixels': pages,
                   'pdf_file': str(target.relative_to(OUT)), 'pdf_sha256': digest(target),
                   'pdf_metadata': pdf_metadata,
                   'workbook_sha256': digest(folder / 'authored.xlsx'),
                   'generation': 'unchanged numeric_fax -> render_workbook_path_to_pdf; font-path override only'})
    save(OUT / 'inputs.json', INPUTS)


def observe_model(frame, event, result):
    if (event != 'return' or result is None or frame.f_code.co_name != '_load_as_safetensor'
            or not frame.f_code.co_filename.endswith('huggingface_hub/hub_mixin.py')):
        return
    path = Path(frame.f_locals['model_file'])
    prefix = 'models--KotaroKinoshita--yomitoku-'
    repo = path.parents[2].name
    name = repo.removeprefix(prefix)
    if not repo.startswith(prefix) or name not in MODELS:
        raise ValueError('unapproved native model file')
    revision, expected = MODELS[name]
    if (path.name != 'model.safetensors' or path.parent.name != revision
            or digest(path) != expected
            or not path.resolve().is_relative_to(Path(os.environ['HF_HOME']).resolve())):
        raise ValueError('loaded native model differs from fixed source')
    MODEL_LOADS.append({'name': name, 'revision': revision, 'sha256': expected})


def observe(frame, event, result):
    observe_bind(frame, event, result)
    if event != 'return' or result is None:
        return
    name, filename = frame.f_code.co_name, frame.f_code.co_filename
    try:
        observe_model(frame, event, result)
        if name == '_load_font' and filename.endswith('workbook_pdf_renderer.py'):
            from PIL import ImageFont
            if not isinstance(result, ImageFont.FreeTypeFont) or str(result.path) != FONT['path']:
                raise ValueError('renderer fell through explicit font to fallback')
        elif name == '_build_preprocess_for_ocr' and filename.endswith('hakodate_cell_ocr_batch_service.py'):
            import cv2
            original = frame.f_locals['original']
            index = len(OBSERVED)
            if not cv2.imwrite(str(OUT / f'actual-pre-ocr-input-{index}.png'), original):
                raise ValueError('could not preserve actual pre-OCR pixels')
            OBSERVED.append({'function': name, 'input_page_pixels': value(original),
                             'result': value({key: result[key] for key in
                                ('raw_rectified', 'target_regions', 'expected_sheet_cells')})})
        elif name == '_run_text_recognizer_records' and filename.endswith('render_best_method_overlay_all_facilities.py'):
            OBSERVED.append({'function': name, 'input': value({key: frame.f_locals[key] for key in
                             ('raw_rectified_bgr', 'regions', 'truth')}), 'result': value(result)})
        elif name == '_apply_hakodate_sheet_output_to_sheet_payload' and filename.endswith('order_service.py'):
            OBSERVED.append({'function': name, 'result': value(result)})
    except Exception:
        # Observer failure invalidates proof; it must not replace the product result.
        AUDIT['observer_errors'].append(traceback.format_exc())


def run_backend(modules):
    import pytest
    from sqlalchemy import event
    from sqlalchemy.engine import Engine

    def enable_fk(connection, _record):
        connection.execute('PRAGMA foreign_keys=ON')
        fk = connection.execute('PRAGMA foreign_keys').fetchone()[0]
        databases = connection.execute('PRAGMA database_list').fetchall()
        if fk != 1 or any(not Path(row[2]).resolve().is_relative_to(OUT) for row in databases):
            raise ValueError('every fixture connection requires owned SQLite and FK=1')
        AUDIT['foreign_keys_connections'].append({'databases': databases, 'foreign_keys': fk})

    # Install before imports/fixture creation, not only after fixture setup.
    event.listen(Engine, 'connect', enable_fk)
    restore_font = None

    @pytest.hookimpl(hookwrapper=True)
    def pytest_fixture_setup(fixturedef, request):
        outcome = yield
        if fixturedef.argname == 'numeric_fax' and outcome.excinfo is None:
            capture_input(outcome.get_result())

    @pytest.hookimpl(trylast=True)
    def pytest_runtest_setup(item):
        from src import db
        with db.SessionLocal() as session:
            engine = session.get_bind()
            for bound in (db.engine, engine):
                if (bound.url.drivername != 'sqlite'
                        or not Path(bound.url.database).resolve().is_relative_to(OUT)):
                    raise ValueError('native test used a non-owned database')
            connection = session.connection()
            connection.exec_driver_sql('PRAGMA foreign_keys=ON')
            fk = connection.exec_driver_sql('PRAGMA foreign_keys').scalar_one()
            if fk != 1:
                raise ValueError('SQLite foreign keys are not enforced')
            AUDIT['case_bindings'].append({'id': item.nodeid, 'database': str(engine.url), 'foreign_keys': fk})

    class Plugin:
        pass

    plugin = Plugin()
    plugin.pytest_fixture_setup = pytest_fixture_setup
    plugin.pytest_runtest_setup = pytest_runtest_setup
    try:
        restore_font = font_fixture()
        sys.setprofile(observe)
        return pytest.main(['-p', 'no:cacheprovider', '-q', '--tb=short',
                            '--basetemp', str(OUT / 'tests'), '--junitxml', str(OUT / 'results.xml'),
                            *modules], plugins=[plugin])
    finally:
        sys.setprofile(None)
        if restore_font is not None:
            restore_font()
        event.remove(Engine, 'connect', enable_fk)


def run_ocr():
    from app import pdf_render, preprocess, yomitoku_runner
    from PIL import Image, ImageFilter

    pdf_path = Path(os.environ['OWNED_PDF'])
    if digest(pdf_path) != os.environ['OWNED_PDF_SHA256']:
        raise ValueError('verified backend input artifact hash mismatch')
    pdf = pdf_path.read_bytes()
    pages = pdf_render.render_pdf_to_page_images(pdf, 200)
    if not pages:
        raise ValueError('native renderer returned no input pages')
    result = {'rendered': value(pages), 'preprocessed': [], 'pillow': []}
    for index, bgr in pages:
        result['preprocessed'].append(value(preprocess.build_images_for_match_and_ocr_from_bgr(bgr)))
        img = Image.fromarray(bgr[:, :, ::-1])
        img.save(OUT / f'pipeline-input-{index}.png')
        buf = BytesIO()
        img.save(buf, format='PNG')
        reopened = Image.open(BytesIO(buf.getvalue()))
        sample = reopened.crop((0, 0, min(512, img.width), min(512, img.height)))
        result['pillow'].append(value({'png_roundtrip': reopened, 'crop': sample,
            'lanczos': img.resize((img.width // 2, img.height // 2), Image.Resampling.LANCZOS),
            'gray': img.convert('L'), 'rank': sample.filter(ImageFilter.RankFilter(3, 4))}))
        produced = yomitoku_runner._save_pdf([bgr])
        result['pillow'][-1]['pdf_rendered'] = value(pdf_render.render_pdf_to_page_images(produced, 72))
    save(OUT / 'pixels.json', result)
    analysis = []

    def observe_analysis(frame, event, returned):
        try:
            observe_bind(frame, event, returned)
            observe_model(frame, event, returned)
            if (event == 'return' and frame.f_code.co_name == '_serialize_results'
                    and frame.f_code.co_filename.endswith('yomitoku_runner.py')):
                analysis.append(value(returned))
        except Exception:
            AUDIT['observer_errors'].append(traceback.format_exc())

    sys.setprofile(observe_analysis)
    try:
        output = yomitoku_runner.run_yomitoku(pdf_bytes=None, dpi=200, device='cpu', visualize=False,
            ignore_line_break=False, no_figure=True, figure_width=200, figure_dir='figures', page_images=pages)
        if not output[0] or not analysis:
            raise ValueError('missing actual four-model native output')
        save(OUT / 'native-result.json', {'pages': value([asdict(page) for page in output[0]]),
                                         'analysis': analysis, 'pdf_outputs': value(output[1:])})
    finally:
        sys.setprofile(None)
    return 0


def main():
    import native_network_isolation as isolation

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('mode', choices=('models', 'backend', 'ocr', 'network-check'))
    parser.add_argument('modules', nargs='*')
    parser.add_argument('--models', type=Path)
    parser.add_argument('--output', type=Path)
    isolation.add_arguments(parser)
    args = parser.parse_args()
    if args.mode == 'models':
        if args.models is None or args.output is None:
            parser.error('model cache and manifest paths required')
        download_models(args.models, args.output)
        return 0
    install_audit(args)
    try:
        if ISOLATION is None:
            raise RuntimeError('verified kernel isolation required; no native fallback')
        if args.mode == 'network-check':
            # Fresh process: this is the real library import under both guards,
            # not an import before auditing or a substituted capability result.
            import urllib3
            from urllib3.util.connection import HAS_IPV6

            finish()
            records = AUDIT['allowed_capability_binds']
            if (AUDIT['violations'] or AUDIT['observer_errors'] or len(records) != 1
                    or not records[0]['closed'] or records[0]['os_result'] == 'pending'
                    or HAS_IPV6 != (records[0]['os_result'] == 'bound')):
                raise RuntimeError('actual urllib3 capability reservation proof incomplete')
            save(OUT / 'capability.json', {'urllib3': urllib3.__version__, 'HAS_IPV6': HAS_IPV6,
                                          'actual_bind': records[0]})
            return 0
        if args.mode == 'backend':
            if not args.modules:
                raise ValueError('explicit complete native test modules required')
            return int(run_backend(args.modules))
        return run_ocr()
    except Exception as exc:
        save(OUT / 'native-error.json', {'type': type(exc).__name__, 'message': str(exc),
                                         'traceback': traceback.format_exc()})
        raise
    finally:
        finish()


if __name__ == '__main__':
    raise SystemExit(main())
