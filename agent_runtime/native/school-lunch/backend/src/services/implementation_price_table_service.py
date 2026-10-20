from __future__ import annotations

import os

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_DOWN, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
import math
import re
import unicodedata
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from src.db import session_scope
from src.models.implementation_price_draft import ImplementationPriceDraft
from src.services.implementation_price_scope_service import require_facility, require_same_facility, validate_price_scope, menu_source_ids, PriceScopeError
from src.services.implementation_price_invoice_normalization import normalize_invoice_cost, unit as invoice_unit


OUTPUT_DIR = Path(os.environ.get("OUTPUT_DIR", "/tmp/hospital_order_outputs"))
_COUNT_UNITS = {"個", "本", "袋", "piece", "pieces", "count", "unit"}
_GRAM_UNITS = {"g", "ｇ", "グラム"}
_KILOGRAM_UNITS = {"kg", "ｋｇ", "キログラム"}
_CONFIRMED_COST_STATUS = "confirmed"


@dataclass(frozen=True)
class PriceTableItem:
    name: str
    amount_per_person: Decimal
    unit: str
    unit_price: Decimal
    price_unit: str
    tax_included: bool
    origin_or_maker: str
    supplier: str
    source: str = ""
    purchased_quantity: Decimal | None = None
    purchased_unit: str | None = None
    yield_rate: Decimal | None = None
    total_quantity_precision: int | None = None
    delivered_quantity: Decimal | None = None
    delivered_unit: str | None = None
    invoice_amount: Decimal | None = None
    amount_per_person_provided: bool = True
    quantity_basis: str = "per_person"
    actual_quantity: Decimal | None = None
    amount_per_person_precision: int | None = None
    amount_per_person_rounding: str | None = None
    invoice_tax_included: bool | None = None
    invoice_amount_rounding: str | None = None
    invoice_original_price: Decimal | None = None
    invoice_original_unit: str | None = None
    invoice_original_tax_included: bool | None = None


@dataclass(frozen=True)
class PriceTableSchool:
    name: str
    recipient: str
    headcount: int
    menu_name: str
    items: tuple[PriceTableItem, ...]
    allocation_headcount: int | None = None
    reserve_meals: int | None = None


@dataclass(frozen=True)
class PriceTableRequest:
    implementation_date: date
    corporation_name: str
    representative_name: str
    schools: tuple[PriceTableSchool, ...]
    tax_rate: Decimal | None = None


@dataclass(frozen=True)
class CalculatedPriceTableRow:
    name: str
    amount_per_person: Decimal | None
    unit: str
    total_quantity: Decimal
    total_unit: str
    unit_price_tax_included: Decimal
    price_unit: str
    total_price: int
    per_person_price: Decimal
    origin_or_maker: str
    supplier: str
    source: str
    quantity_basis: str = "per_person"


def _decimal(value: object, field_name: str) -> Decimal:
    if value is None or value == "":
        raise ValueError(f"{field_name} is required")
    try:
        parsed = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field_name}は数値で指定してください。") from exc
    if not parsed.is_finite():
        raise ValueError(f"{field_name}は有限の数値で指定してください。")
    return parsed


def _integer(value: object, field_name: str, *, minimum: int | None = None, maximum: int | None = None) -> int:
    parsed = _decimal(value, field_name)
    if parsed != parsed.to_integral_value():
        raise ValueError(f"{field_name}は整数で指定してください。")
    result = int(parsed)
    if minimum is not None and result < minimum:
        raise ValueError(f"{field_name}は{minimum}以上で指定してください。")
    if maximum is not None and result > maximum:
        raise ValueError(f"{field_name}は{maximum}以下で指定してください。")
    return result


def _required_text(payload: dict, field_name: str) -> str:
    value = str(payload.get(field_name) or "").strip()
    if not value:
        raise ValueError(f"{field_name} is required")
    return value


def _optional_text(payload: dict, field_name: str) -> str:
    return str(payload.get(field_name) or "").strip()


def parse_price_table_request(payload: dict) -> PriceTableRequest:
    if not isinstance(payload, dict):
        raise ValueError("payload must be an object")
    raw_date = _required_text(payload, "implementation_date")
    try:
        implementation_date = date.fromisoformat(raw_date)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("implementation_date must be YYYY-MM-DD") from exc

    tax_rate = payload.get("tax_rate")
    parsed_tax_rate = _decimal(tax_rate, "tax_rate") if tax_rate not in (None, "") else None
    raw_schools = payload.get("schools")
    if not isinstance(raw_schools, list) or not raw_schools:
        raise ValueError("schools must be a non-empty list")

    schools: list[PriceTableSchool] = []
    for school_idx, raw_school in enumerate(raw_schools, start=1):
        if not isinstance(raw_school, dict):
            raise ValueError(f"schools[{school_idx}] must be an object")
        headcount = _integer(raw_school.get("headcount"), f"schools[{school_idx}].headcount", minimum=1)
        allocation_headcount = ( _integer(raw_school.get("allocation_headcount"), f"schools[{school_idx}].allocation_headcount", minimum=1)
            if raw_school.get("allocation_headcount") not in (None, "") else None )
        reserve_meals = ( _integer(raw_school.get("reserve_meals"), f"schools[{school_idx}].reserve_meals", minimum=0)
            if raw_school.get("reserve_meals") not in (None, "") else None )
        raw_items = raw_school.get("items")
        if not isinstance(raw_items, list) or not raw_items:
            raise ValueError(f"schools[{school_idx}].items must be a non-empty list")
        items: list[PriceTableItem] = []
        for item_idx, raw_item in enumerate(raw_items, start=1):
            if not isinstance(raw_item, dict):
                raise ValueError(f"schools[{school_idx}].items[{item_idx}] must be an object")
            price_tax_included = raw_item.get("unit_price_tax_included")
            price_tax_excluded = raw_item.get("unit_price_tax_excluded")
            if price_tax_included not in (None, "") and price_tax_excluded not in (None, ""):
                raise ValueError(
                    f"schools[{school_idx}].items[{item_idx}] must not define both tax-included and tax-excluded prices"
                )
            if price_tax_included in (None, "") and price_tax_excluded in (None, ""):
                raise ValueError(f"schools[{school_idx}].items[{item_idx}].unit_price is required")
            if price_tax_excluded not in (None, "") and parsed_tax_rate is None:
                raise ValueError("tax_rate is required when any unit_price_tax_excluded is provided")
            tax_included = price_tax_included not in (None, "")
            unit_price = _decimal(
                price_tax_included if tax_included else price_tax_excluded,
                f"schools[{school_idx}].items[{item_idx}].unit_price",
            )
            if unit_price < 0:
                raise ValueError(f"schools[{school_idx}].items[{item_idx}].unit_priceは0以上で指定してください。")
            quantity_basis = str(raw_item.get("quantity_basis") or "per_person").strip()
            if quantity_basis not in {"per_person", "actual_count", "delivery_allocation"}:
                raise ValueError(f"schools[{school_idx}].items[{item_idx}].quantity_basisが不正です。")
            amount_provided = raw_item.get("amount_per_person") not in (None, "")
            purchased_provided = raw_item.get("purchased_quantity") not in (None, "")
            actual_provided = raw_item.get("actual_quantity") not in (None, "")
            invoice_tax_included = raw_item.get("invoice_tax_included")
            invoice_amount_rounding = raw_item.get("invoice_amount_rounding") or None
            if invoice_amount_rounding not in {None, "down", "half_up"}:
                raise ValueError(f"{raw_item.get('name', '')}: unsupported invoice_amount_rounding")
            if invoice_tax_included is not None and not isinstance(invoice_tax_included, bool):
                raise ValueError("納品金額の税区分を税抜または税込で指定してください。")
            for field in ("amount_per_person", "purchased_quantity", "delivered_quantity", "invoice_amount"):
                if raw_item.get(field) not in (None, "") and _decimal(raw_item[field], field) < 0:
                    raise ValueError(f"{raw_item.get('name', '')}: {field}は0以上で指定してください。")
            if quantity_basis == "actual_count":
                if not actual_provided:
                    raise ValueError(f"{raw_item.get('name', '')}: 実個数・本数を入力してください。")
                if _normalize_quantity_unit(_required_text(raw_item, "unit")) == "kg":
                    raise ValueError(f"schools[{school_idx}].items[{item_idx}]: actual_countは個数単位のみ指定できます。")
                if purchased_provided or raw_item.get("yield_rate") not in (None, ""):
                    raise ValueError(f"schools[{school_idx}].items[{item_idx}]: actual_countではpurchased_quantityまたはyield_rateを指定できません。")
                _integer(raw_item.get("actual_quantity"), f"schools[{school_idx}].items[{item_idx}].actual_quantity", minimum=0)
            elif quantity_basis == "delivery_allocation":
                if allocation_headcount is None or reserve_meals is None:
                    raise ValueError(f"schools[{school_idx}].items[{item_idx}]: delivery_allocationにはallocation_headcountとreserve_mealsが必要です。")
                if raw_item.get("delivered_quantity") in (None, "") or not _optional_text(raw_item, "delivered_unit"):
                    raise ValueError(f"schools[{school_idx}].items[{item_idx}]: delivery_allocationにはdelivered_quantityとdelivered_unitが必要です。")
                if raw_item.get("amount_per_person_precision") in (None, "") or not _optional_text(raw_item, "amount_per_person_rounding") or raw_item.get("total_quantity_precision") in (None, ""):
                    raise ValueError(f"schools[{school_idx}].items[{item_idx}]: delivery_allocationには各精度と丸め指定が必要です。")
                if purchased_provided or raw_item.get("yield_rate") not in (None, ""):
                    raise ValueError(f"schools[{school_idx}].items[{item_idx}]: delivery_allocationではpurchased_quantityまたはyield_rateを指定できません。")
            elif not amount_provided and not purchased_provided:
                raise ValueError(
                    f"schools[{school_idx}].items[{item_idx}] requires amount_per_person or purchased_quantity"
                )
            items.append(
                PriceTableItem(
                    name=_required_text(raw_item, "name"),
                    amount_per_person=(
                        _decimal(raw_item.get("amount_per_person"), f"schools[{school_idx}].items[{item_idx}].amount_per_person")
                        if amount_provided else Decimal("0")
                    ),
                    unit=_required_text(raw_item, "unit"),
                    unit_price=unit_price,
                    price_unit=_required_text(raw_item, "price_unit"),
                    tax_included=tax_included,
                    origin_or_maker=_required_text(raw_item, "origin_or_maker"),
                    supplier=_required_text(raw_item, "supplier"),
                    source=str(raw_item.get("source") or "").strip(),
                    purchased_quantity=(
                        _decimal(
                            raw_item.get("purchased_quantity"),
                            f"schools[{school_idx}].items[{item_idx}].purchased_quantity",
                        )
                        if raw_item.get("purchased_quantity") not in (None, "")
                        else None
                    ),
                    purchased_unit=(
                        str(raw_item.get("purchased_unit") or "").strip()
                        if raw_item.get("purchased_quantity") not in (None, "")
                        else None
                    ),
                    yield_rate=(
                        _decimal(raw_item.get("yield_rate"), f"schools[{school_idx}].items[{item_idx}].yield_rate")
                        if raw_item.get("yield_rate") not in (None, "")
                        else None
                    ),
                    total_quantity_precision=(
                        _integer(
                            raw_item.get("total_quantity_precision"),
                            f"schools[{school_idx}].items[{item_idx}].total_quantity_precision",
                            minimum=0, maximum=6,
                        )
                        if raw_item.get("total_quantity_precision") not in (None, "")
                        else None
                    ),
                    delivered_quantity=(
                        _decimal(
                            raw_item.get("delivered_quantity"),
                            f"schools[{school_idx}].items[{item_idx}].delivered_quantity",
                        )
                        if raw_item.get("delivered_quantity") not in (None, "")
                        else None
                    ),
                    delivered_unit=(
                        str(raw_item.get("delivered_unit") or "").strip()
                        if raw_item.get("delivered_quantity") not in (None, "")
                        else None
                    ),
                    invoice_amount=(
                        _decimal(
                            raw_item.get("invoice_amount"),
                            f"schools[{school_idx}].items[{item_idx}].invoice_amount",
                        )
                        if raw_item.get("invoice_amount") not in (None, "")
                        else None
                    ),
                    amount_per_person_provided=amount_provided,
                    quantity_basis=quantity_basis,
                    invoice_tax_included=invoice_tax_included,
                    invoice_amount_rounding=invoice_amount_rounding,
                    invoice_original_price=(_decimal(raw_item.get("invoice_original_price"), "invoice_original_price") if raw_item.get("invoice_original_price") not in (None, "") else None),
                    invoice_original_unit=raw_item.get("invoice_original_unit"),
                    invoice_original_tax_included=raw_item.get("invoice_original_tax_included"),
                    actual_quantity=(
                        _decimal(raw_item.get("actual_quantity"), f"schools[{school_idx}].items[{item_idx}].actual_quantity")
                        if actual_provided else None
                    ),
                    amount_per_person_precision=(
                        _integer(raw_item.get("amount_per_person_precision"), f"schools[{school_idx}].items[{item_idx}].amount_per_person_precision", minimum=0, maximum=6)
                        if raw_item.get("amount_per_person_precision") not in (None, "") else None
                    ),
                    amount_per_person_rounding=(
                        _optional_text(raw_item, "amount_per_person_rounding") or None
                    ),
                )
            )
        schools.append(
            PriceTableSchool(
                name=_required_text(raw_school, "name"),
                recipient=_required_text(raw_school, "recipient"),
                headcount=headcount,
                menu_name=_required_text(raw_school, "menu_name"),
                items=tuple(items),
                allocation_headcount=allocation_headcount,
                reserve_meals=reserve_meals,
            )
        )

    return PriceTableRequest(
        implementation_date=implementation_date,
        corporation_name=_required_text(payload, "corporation_name"),
        representative_name=_required_text(payload, "representative_name"),
        tax_rate=parsed_tax_rate,
        schools=tuple(schools),
    )


def build_price_table_payload_from_manual_sources(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("payload must be an object")
    menu_source_ids(payload)
    source_menu_execution_id = payload.get("source_menu_execution_id")
    if source_menu_execution_id is not None and (
        type(source_menu_execution_id) is not int or source_menu_execution_id <= 0
    ):
        raise ValueError("source_menu_execution_id must identify a selected menu execution")
    raw_costs = payload.get("cost_items")
    if not isinstance(raw_costs, list) or not raw_costs:
        raise ValueError("cost_items must be a non-empty list")

    cost_by_key: dict[str, dict] = {}
    for cost_idx, raw_cost in enumerate(raw_costs, start=1):
        if not isinstance(raw_cost, dict):
            raise ValueError(f"cost_items[{cost_idx}] must be an object")
        key = _optional_text(raw_cost, "key") or _required_text(raw_cost, "name")
        if key in cost_by_key:
            raise ValueError(f"duplicate cost key: {key}")
        cost_by_key[key] = raw_cost

    raw_schools = payload.get("schools")
    if not isinstance(raw_schools, list) or not raw_schools:
        raise ValueError("schools must be a non-empty list")

    schools: list[dict] = []
    for school_idx, raw_school in enumerate(raw_schools, start=1):
        if not isinstance(raw_school, dict):
            raise ValueError(f"schools[{school_idx}] must be an object")
        if raw_school.get("headcount_confirmed") is False:
            raise ValueError("手配表の予定食数です。実際の生徒食数を確認してから計算してください。")
        raw_quantity_items = raw_school.get("quantity_items")
        if not isinstance(raw_quantity_items, list) or not raw_quantity_items:
            raise ValueError(f"schools[{school_idx}].quantity_items must be a non-empty list")

        items: list[dict] = []
        for item_idx, raw_item in enumerate(raw_quantity_items, start=1):
            if not isinstance(raw_item, dict):
                raise ValueError(f"schools[{school_idx}].quantity_items[{item_idx}] must be an object")
            item_name = _required_text(raw_item, "name")
            cost_key = _optional_text(raw_item, "cost_key") or item_name
            cost = cost_by_key.get(cost_key)
            if cost is None:
                raise ValueError(f"missing cost item for key: {cost_key}")
            if _optional_text(cost, "status") != _CONFIRMED_COST_STATUS:
                raise ValueError(f"cost item must be confirmed before use: {cost_key}")
            if cost.get("price_basis", "direct") == "invoice":
                cost = normalize_invoice_cost(cost, implementation_date=payload.get("implementation_date"), tax_rate=payload.get("tax_rate"))
            elif cost.get("price_basis", "direct") != "direct":
                raise ValueError(f"{item_name}: unsupported price_basis")
            price_tax_basis = cost.get("price_tax_basis")
            if price_tax_basis is not None:
                if price_tax_basis not in {"included", "excluded"}:
                    raise ValueError(f"{item_name}: unsupported price_tax_basis")
                selected_price = cost.get(f"unit_price_tax_{price_tax_basis}")
                opposite_price = cost.get(f"unit_price_tax_{'excluded' if price_tax_basis == 'included' else 'included'}")
                if selected_price in (None, "") or opposite_price not in (None, ""):
                    raise ValueError(f"{item_name}: unit price must match the selected tax basis")
            merged = {
                "name": item_name,
                "amount_per_person": raw_item.get("amount_per_person"),
                "unit": raw_item.get("unit"),
                "unit_price_tax_included": cost.get("unit_price_tax_included"),
                "unit_price_tax_excluded": cost.get("unit_price_tax_excluded"),
                "price_unit": cost.get("price_unit"),
                "origin_or_maker": cost.get("origin_or_maker"),
                "supplier": cost.get("supplier"),
                "source": " / ".join(
                    part
                    for part in [
                        _optional_text(raw_item, "source"),
                        _optional_text(cost, "source"),
                    ]
                    if part
                ),
                "purchased_quantity": raw_item.get("purchased_quantity"),
                "purchased_unit": raw_item.get("purchased_unit"),
                "yield_rate": raw_item.get("yield_rate"),
                "total_quantity_precision": raw_item.get("total_quantity_precision"),
                "delivered_quantity": cost.get("delivered_quantity"),
                "delivered_unit": cost.get("delivered_unit"),
                "invoice_amount": cost.get("invoice_amount"),
                "invoice_tax_included": cost.get("invoice_tax_included"),
                "invoice_amount_rounding": cost.get("invoice_amount_rounding"),
                "invoice_original_price": cost.get("invoice_original_price"),
                "invoice_original_unit": cost.get("invoice_original_unit"),
                "invoice_original_tax_included": cost.get("invoice_original_tax_included"),
                "quantity_basis": raw_item.get("quantity_basis"),
                "actual_quantity": raw_item.get("actual_quantity"),
                "amount_per_person_precision": raw_item.get("amount_per_person_precision"),
                "amount_per_person_rounding": raw_item.get("amount_per_person_rounding"),
            }
            items.append({key: value for key, value in merged.items() if value not in (None, "")})
        schools.append(
            {
                "name": _required_text(raw_school, "name"),
                "recipient": _required_text(raw_school, "recipient"),
                "headcount": raw_school.get("headcount"),
                "allocation_headcount": raw_school.get("allocation_headcount"),
                "reserve_meals": raw_school.get("reserve_meals"),
                "menu_name": _required_text(raw_school, "menu_name"),
                "items": items,
            }
        )

    return {
        **{key: payload[key] for key in ("facility_id", "source_ocr_job_ids", "source_menu_execution_ids") if key in payload},
        **(
            {"source_menu_execution_id": source_menu_execution_id}
            if source_menu_execution_id is not None
            else {}
        ),
        "implementation_date": payload.get("implementation_date"),
        "corporation_name": payload.get("corporation_name"),
        "representative_name": payload.get("representative_name"),
        "tax_rate": payload.get("tax_rate"),
        "schools": schools,
    }


def parse_price_table_request_from_manual_sources(payload: dict) -> PriceTableRequest:
    return parse_price_table_request(build_price_table_payload_from_manual_sources(payload))


def save_manual_source_draft(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("payload must be an object")
    draft_id = uuid4().hex
    school_names = "・".join(str(school.get("name") or "").strip() for school in payload.get("schools", []) if school.get("name"))
    label = " / ".join(part for part in [str(payload.get("implementation_date") or ""), school_names, str(payload.get("schools", [{}])[0].get("menu_name") or "").strip()] if part)
    now = datetime.utcnow()
    with session_scope() as session:
        implementation_date = validate_price_scope(session, payload, for_adoption=True)
        session.add(ImplementationPriceDraft(id=draft_id, facility_id=payload["facility_id"], implementation_date=implementation_date, label=label or draft_id, payload_json=payload, created_at=now, updated_at=now))
    return {"draft_id": draft_id}


def list_manual_source_drafts(*, facility_id: int, month: str, implementation_date: date | None = None) -> list[dict]:
    if not isinstance(month, str) or not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise PriceScopeError("対象月をYYYY-MMで指定してください。")
    try:
        first_day = date.fromisoformat(f"{month}-01")
        next_month = date(first_day.year + (first_day.month == 12), first_day.month % 12 + 1, 1)
    except ValueError as exc:
        raise PriceScopeError("対象月が不正です。") from exc
    if implementation_date is not None and not first_day <= implementation_date < next_month:
        raise PriceScopeError("実施日filterが対象月と一致しません。")
    with session_scope() as session:
        require_facility(session, facility_id)
        scoped = ((ImplementationPriceDraft.facility_id == facility_id)
                  & (ImplementationPriceDraft.implementation_date >= first_day)
                  & (ImplementationPriceDraft.implementation_date < next_month))
        if implementation_date is not None:
            scoped = scoped & (ImplementationPriceDraft.implementation_date == implementation_date)
        drafts = session.execute(
            select(ImplementationPriceDraft).where(scoped | ImplementationPriceDraft.facility_id.is_(None)).order_by(ImplementationPriceDraft.updated_at.desc())
        ).scalars().all()
        return [
            {
                "draft_id": draft.id,
                "facility_id": draft.facility_id,
                "scope_status": "needs_review" if draft.facility_id is None else "assigned",
                "label": draft.label,
                "implementation_date": draft.payload_json.get("implementation_date", ""),
                "updated_at": draft.updated_at,
            }
            for draft in drafts
        ]


def load_manual_source_draft(draft_id: str, *, facility_id: int) -> dict:
    normalized = str(draft_id or "").strip()
    if len(normalized) != 32 or not all(character in "0123456789abcdef" for character in normalized.lower()):
        raise ValueError("invalid draft_id")
    with session_scope() as session:
        require_facility(session, facility_id)
        draft = session.get(ImplementationPriceDraft, normalized)
        if draft is None:
            raise FileNotFoundError("draft not found")
        require_same_facility(draft.facility_id, facility_id)
        require_same_facility(draft.payload_json.get("facility_id"), facility_id)
        implementation_date = validate_saved_manual_source_draft_scope(session, draft.payload_json)
        if implementation_date != draft.implementation_date:
            raise PriceScopeError("保存済み実施日とpayloadが一致しません。要確認で停止しました。")
        return dict(draft.payload_json)


def validate_saved_manual_source_draft_scope(session, payload: dict) -> date:
    """Validate saved draft identity without re-executing its captured OCR sources."""
    facility = require_facility(session, payload.get("facility_id"))
    schools = payload.get("schools")
    if not isinstance(schools, list) or len(schools) != 1:
        raise PriceScopeError("実施価格表は選択学校1校が必須です。複数校原本は分割せず要確認で停止します。")
    if not isinstance(schools[0], dict) or schools[0].get("name") != facility.name:
        raise PriceScopeError("実施価格表の学校名が選択学校のFacilityと一致しません。")
    try:
        return date.fromisoformat(payload["implementation_date"])
    except (KeyError, TypeError, ValueError) as exc:
        raise PriceScopeError("実施日をYYYY-MM-DDで指定してください。") from exc


def validate_manual_source_scope(payload: dict) -> None:
    with session_scope() as session:
        validate_price_scope(session, payload, for_adoption=False)


def _tax_included_price(item: PriceTableItem, tax_rate: Decimal | None) -> Decimal:
    if item.tax_included:
        return item.unit_price
    if tax_rate is None:
        raise ValueError("tax_rate is required when tax-excluded prices are used")
    return (item.unit_price * (Decimal("1") + tax_rate)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _quantity_quantum(precision: int | None, default: str = "0.01") -> Decimal:
    if precision is None:
        return Decimal(default)
    if precision < 0 or precision > 6:
        raise ValueError("total_quantity_precision must be between 0 and 6")
    return Decimal("1") if precision == 0 else Decimal("0." + ("0" * (precision - 1)) + "1")


def _normalize_quantity_unit(unit: str) -> str:
    normalized = str(unit or "").strip().lower()
    if normalized in _GRAM_UNITS:
        return "kg"
    if normalized in _KILOGRAM_UNITS:
        return "kg"
    if normalized in _COUNT_UNITS or unit in _COUNT_UNITS:
        return str(unit).strip()
    raise ValueError(f"unsupported unit: {unit}")


def _total_quantity(
    amount_per_person: Decimal,
    unit: str,
    headcount: int,
    *,
    precision: int | None = None,
) -> tuple[Decimal, str]:
    normalized = str(unit or "").strip().lower()
    raw_total = amount_per_person * Decimal(headcount)
    if normalized in _GRAM_UNITS:
        return (raw_total / Decimal("1000")).quantize(_quantity_quantum(precision), rounding=ROUND_HALF_UP), "kg"
    if normalized in _KILOGRAM_UNITS:
        return raw_total.quantize(_quantity_quantum(precision), rounding=ROUND_HALF_UP), "kg"
    if normalized in _COUNT_UNITS or unit in _COUNT_UNITS:
        return raw_total.quantize(_quantity_quantum(precision, "0.01"), rounding=ROUND_HALF_UP), unit
    raise ValueError(f"unsupported unit: {unit}")


def _normalize_price_unit(price_unit: str) -> str:
    normalized = str(price_unit or "").strip().lower()
    if not normalized:
        raise ValueError("price_unit is required")
    if "kg" in normalized or "ｋｇ" in normalized:
        return "kg"
    for count_unit in ("個", "本", "袋"):
        if count_unit in str(price_unit):
            return count_unit
    if normalized in _COUNT_UNITS:
        return str(price_unit).strip()
    raise ValueError(f"unsupported price_unit: {price_unit}")


def _effective_total_quantity(item: PriceTableItem, headcount: int) -> tuple[Decimal, str]:
    if item.purchased_quantity is not None:
        if not item.purchased_unit:
            raise ValueError(f"{item.name}: purchased_unit is required when purchased_quantity is provided")
        total_unit = _normalize_quantity_unit(item.purchased_unit)
        purchased_quantity = item.purchased_quantity
        if str(item.purchased_unit).strip().lower() in _GRAM_UNITS:
            purchased_quantity /= Decimal("1000")
        purchased = purchased_quantity.quantize(
            _quantity_quantum(item.total_quantity_precision), rounding=ROUND_HALF_UP
        )
        if item.amount_per_person_provided:
            expected, expected_unit = _total_quantity(item.amount_per_person, item.unit, headcount, precision=item.total_quantity_precision)
            if item.yield_rate is not None:
                if expected_unit != "kg" or item.yield_rate <= 0 or item.yield_rate > 1:
                    raise ValueError(f"{item.name}: yield_rate must be valid for weight items")
                expected = (expected / item.yield_rate).quantize(_quantity_quantum(item.total_quantity_precision), rounding=ROUND_HALF_UP)
            if _normalize_quantity_unit(expected_unit) != total_unit or expected != purchased:
                raise ValueError(f"{item.name}: purchased_quantity does not match amount_per_person and headcount")
        return purchased, total_unit

    total_quantity, total_unit = _total_quantity(
        item.amount_per_person,
        item.unit,
        headcount,
        precision=item.total_quantity_precision,
    )
    if item.yield_rate is None:
        return total_quantity, total_unit
    if total_unit != "kg":
        raise ValueError(f"{item.name}: yield_rate can only be applied to weight items")
    if item.yield_rate <= 0 or item.yield_rate > 1:
        raise ValueError(f"{item.name}: yield_rate must be greater than 0 and less than or equal to 1")
    return (total_quantity / item.yield_rate).quantize(
        _quantity_quantum(item.total_quantity_precision), rounding=ROUND_HALF_UP
    ), total_unit


def _amount_quantum(precision: int) -> Decimal:
    return Decimal("1") if precision == 0 else Decimal("0." + ("0" * (precision - 1)) + "1")


def _allocation_amount_per_person(item: PriceTableItem, school: PriceTableSchool) -> Decimal:
    if school.allocation_headcount is None or school.reserve_meals is None:
        raise ValueError(f"{item.name}: delivery_allocationにはallocation_headcountとreserve_mealsが必要です。")
    if item.delivered_quantity is None or not item.delivered_unit:
        raise ValueError(f"{item.name}: delivery_allocationにはdelivered_quantityとdelivered_unitが必要です。")
    if item.amount_per_person_precision is None or item.total_quantity_precision is None or item.amount_per_person_rounding is None:
        raise ValueError(f"{item.name}: delivery_allocationには各精度と丸め指定が必要です。")
    if item.amount_per_person_rounding not in {"half_up", "down"}:
        raise ValueError(f"{item.name}: 配分した1人量の丸め方は四捨五入または切り捨てを選択してください。")
    if _normalize_quantity_unit(item.unit) != "kg" or _normalize_quantity_unit(item.delivered_unit) != "kg":
        raise ValueError(f"{item.name}: delivery_allocationは重量単位のみ指定できます。")
    if item.delivered_quantity < 0:
        raise ValueError(f"{item.name}: delivered_quantityは0以上で指定してください。")
    # Convert delivery into the item's declared g/kg basis before deriving a per-person value.
    delivered_kg = item.delivered_quantity / (Decimal("1000") if str(item.delivered_unit).strip().lower() in _GRAM_UNITS else Decimal("1"))
    delivered = delivered_kg * (Decimal("1000") if str(item.unit).strip().lower() in _GRAM_UNITS else Decimal("1"))
    rounding = ROUND_HALF_UP if item.amount_per_person_rounding == "half_up" else ROUND_DOWN
    return (delivered / Decimal(school.allocation_headcount + school.reserve_meals)).quantize(
        _amount_quantum(item.amount_per_person_precision), rounding=rounding
    )


def calculate_price_table_rows(
    school: PriceTableSchool,
    tax_rate: Decimal | None,
) -> tuple[CalculatedPriceTableRow, ...]:
    rows: list[CalculatedPriceTableRow] = []
    for item in school.items:
        calculated_amount = item.amount_per_person
        if item.quantity_basis == "actual_count":
            if item.actual_quantity is None:
                raise ValueError(f"{item.name}: actual_quantityが必要です。")
            if item.actual_quantity < 0:
                raise ValueError(f"{item.name}: actual_quantityは0以上で指定してください。")
            if _normalize_quantity_unit(item.unit) == "kg":
                raise ValueError(f"{item.name}: actual_countは個数単位のみ指定できます。")
            total_quantity = item.actual_quantity.quantize(_quantity_quantum(item.total_quantity_precision), rounding=ROUND_HALF_UP)
            total_unit = item.unit
        elif item.quantity_basis == "delivery_allocation":
            calculated_amount = _allocation_amount_per_person(item, school)
            total_quantity, total_unit = _total_quantity(calculated_amount, item.unit, school.headcount, precision=item.total_quantity_precision)
        else:
            total_quantity, total_unit = _effective_total_quantity(item, school.headcount)
        price_unit = _normalize_price_unit(item.price_unit)
        normalized_total_unit = _normalize_quantity_unit(total_unit)
        if price_unit != normalized_total_unit and price_unit != total_unit:
            raise ValueError(f"{item.name}: price_unit {item.price_unit} does not match total unit {total_unit}")
        unit_price = _tax_included_price(item, tax_rate)
        total_price = _money(unit_price * total_quantity)
        rows.append(
            CalculatedPriceTableRow(
                name=item.name,
                amount_per_person=calculated_amount if (item.amount_per_person_provided or item.quantity_basis == "delivery_allocation") else None,
                unit=item.unit,
                total_quantity=total_quantity,
                total_unit=total_unit,
                unit_price_tax_included=unit_price,
                price_unit=item.price_unit,
                total_price=total_price,
                per_person_price=_per_person_price(total_price, school.headcount),
                origin_or_maker=item.origin_or_maker,
                supplier=item.supplier,
                source=item.source,
                quantity_basis=item.quantity_basis,
            )
        )
    return tuple(rows)


def build_price_table_preview_from_manual_sources(payload: dict) -> dict:
    request = parse_price_table_request_from_manual_sources(payload)
    schools: list[dict] = []
    for school in request.schools:
        rows = calculate_price_table_rows(school, request.tax_rate)
        total_price = sum(row.total_price for row in rows)
        schools.append(
            {
                "name": school.name,
                "menu_name": school.menu_name,
                "headcount": school.headcount,
                "allocation_headcount": school.allocation_headcount,
                "reserve_meals": school.reserve_meals,
                "rows": [
                    {
                        "name": row.name,
                        "amount_per_person": row.amount_per_person,
                        "unit": row.unit,
                        "total_quantity": row.total_quantity,
                        "total_unit": row.total_unit,
                        "unit_price_tax_included": row.unit_price_tax_included,
                        "price_unit": row.price_unit,
                        "total_price": row.total_price,
                        "per_person_price": row.per_person_price,
                        "origin_or_maker": row.origin_or_maker,
                        "supplier": row.supplier,
                        "source": row.source,
                        "quantity_basis": row.quantity_basis,
                        "invoice_reconciliation": _invoice_reconciliation(
                            item,
                            row.unit_price_tax_included,
                        ),
                    }
                    for item, row in zip(school.items, rows, strict=True)
                ],
                "total_price": total_price,
                "per_person_price": _per_person_price(total_price, school.headcount),
            }
        )
    return {"implementation_date": request.implementation_date, "schools": schools}


def _invoice_reconciliation(item: PriceTableItem, unit_price_tax_included: Decimal) -> dict:
    if item.delivered_quantity is None or item.invoice_amount is None:
        return {"status": "not_available"}
    if item.delivered_quantity < 0 or item.invoice_amount < 0:
        raise ValueError(f"{item.name}: delivered quantity and invoice amount must not be negative")
    if not item.delivered_unit:
        raise ValueError(f"{item.name}: delivered_unit is required when delivered_quantity is provided")
    delivered_unit = invoice_unit(item.delivered_unit) if item.invoice_original_price is not None else _normalize_quantity_unit(item.delivered_unit)
    price_unit = item.invoice_original_unit if item.invoice_original_price is not None else _normalize_price_unit(item.price_unit)
    if delivered_unit != price_unit:
        return {
            "status": "unit_mismatch",
            "delivered_quantity": item.delivered_quantity,
            "delivered_unit": item.delivered_unit,
            "invoice_amount": item.invoice_amount,
        }
    if item.invoice_tax_included is None:
        return {"status": "tax_basis_required"}
    original_tax_included = item.invoice_original_tax_included if item.invoice_original_price is not None else item.tax_included
    if item.invoice_tax_included != original_tax_included and item.invoice_original_price is not None:
        return {"status": "tax_basis_mismatch"}
    if not item.invoice_tax_included and original_tax_included:
        # A rounded tax-inclusive price cannot uniquely recover the source price.
        return {"status": "tax_basis_mismatch"}
    reconciliation_price = item.invoice_original_price if item.invoice_original_price is not None else (unit_price_tax_included if item.invoice_tax_included else item.unit_price)
    delivered_quantity = item.delivered_quantity
    if item.invoice_original_price is None and str(item.delivered_unit).strip().lower() in _GRAM_UNITS:
        delivered_quantity /= Decimal("1000")
    unrounded_amount = reconciliation_price * delivered_quantity
    if unrounded_amount != unrounded_amount.to_integral_value() and item.invoice_amount_rounding is None:
        return {"status": "rounding_required", "unrounded_amount": unrounded_amount, "invoice_amount": item.invoice_amount}
    if item.invoice_amount_rounding not in {None, "down", "half_up"}:
        raise ValueError(f"{item.name}: unsupported invoice_amount_rounding")
    calculated_amount = int(unrounded_amount.quantize(Decimal("1"), rounding=ROUND_HALF_UP if item.invoice_amount_rounding == "half_up" else ROUND_DOWN))
    invoice_amount = item.invoice_amount
    difference = invoice_amount - calculated_amount
    return {
        "status": "matched" if difference == 0 else "mismatch",
        "delivered_quantity": item.delivered_quantity,
        "delivered_unit": item.delivered_unit,
        "invoice_amount": invoice_amount,
        "calculated_amount": calculated_amount,
        "difference": difference,
    }


def _money(value: Decimal) -> int:
    return int(value.quantize(Decimal("1"), rounding=ROUND_DOWN))


def _per_person_price(total_price: int, headcount: int) -> Decimal:
    return (Decimal(total_price) / Decimal(headcount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _safe_sheet_title(value: str, fallback: str) -> str:
    invalid_chars = set('[]:*?/\\')
    title = "".join("_" if ch in invalid_chars else ch for ch in str(value or "").strip())
    return (title or fallback)[:31]


def _price_text_height(text: str, column_width: float, font_size: float = 11) -> float:
    # Excel widths are digit-based, while CJK glyphs take roughly two digits.
    # Reserve padding and account for word wrapping as well as explicit newlines.
    capacity = max(1, math.floor((column_width * 7 * 0.75 - 8) / (font_size * 0.55)))
    lines = 0
    for paragraph in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        lines += 1
        used = 0
        for token in re.findall(r"[A-Za-z0-9]+|[^\n]", paragraph):
            widths = [0 if unicodedata.combining(ch) else 2 if unicodedata.east_asian_width(ch) in "WF" or ch in "MW" else 1 for ch in token]
            if used and sum(widths) <= capacity and used + sum(widths) > capacity:
                lines += 1
                used = 0
            for width in widths:
                if used and used + width > capacity:
                    lines += 1
                    used = 0
                used += width
    height = max(24.95, lines * font_size * 1.3 + 4)
    if height > 409:
        raise ValueError("実施価格表の食品名・生産地/製造元・購入先が長すぎて、Excelの1行に全文を表示できません。出力する文字数を確認してください。")
    return height


def build_price_table_workbook(request: PriceTableRequest) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="666666")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = [27.375, 11, 10, 17.75, 11.875, 12.375, 14, 10]

    for school_idx, school in enumerate(request.schools, start=1):
        worksheet = workbook.create_sheet(_safe_sheet_title(school.name, f"学校{school_idx}"))
        worksheet.merge_cells("A1:H1")
        worksheet["A1"] = "実施価格一覧表"
        worksheet["A1"].font = Font(bold=True, size=16)
        worksheet["A1"].alignment = Alignment(horizontal="center")
        worksheet.merge_cells("A3:H3")
        recipient = school.recipient.rstrip()
        worksheet["A3"] = recipient if recipient.endswith(("殿", "様", "御中")) else f"{recipient}　殿"
        worksheet.merge_cells("A6:D6")
        worksheet["A6"] = f"{request.implementation_date.year} 年 {request.implementation_date.month} 月 {request.implementation_date.day} 日"
        worksheet["E5"] = "法人名"
        worksheet.merge_cells("F5:H5")
        worksheet["F5"] = request.corporation_name
        worksheet["E6"] = "代表者名"
        worksheet.merge_cells("F6:G6")
        worksheet["F6"] = request.representative_name
        worksheet["H6"] = "印"
        worksheet["A7"] = "実施人数"
        worksheet["B7"] = school.headcount
        worksheet["C7"] = "名"
        worksheet.merge_cells("A8:H8")
        worksheet["A8"] = f"献立名: {school.menu_name}"
        worksheet["A8"].alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.row_dimensions[8].height = max(44.25, 16 * ((len(school.menu_name) + 49) // 50))

        headers = [
            "食品名",
            "1人当たり\n購入量",
            "全体量",
            "税込単価",
            "全体の価格\n(円)",
            "1人当たり価格\n(円)",
            "生産地/製造元",
            "購入先",
        ]
        for col_idx, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=9, column=col_idx)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        calculated_rows = calculate_price_table_rows(school, request.tax_rate)
        total_price_sum = 0
        for offset, calculated in enumerate(calculated_rows):
            row_idx = 10 + offset
            total_price_sum += calculated.total_price
            values = [
                calculated.name,
                float(calculated.amount_per_person) if calculated.amount_per_person is not None else None,
                float(calculated.total_quantity),
                float(calculated.unit_price_tax_included),
                calculated.total_price,
                None,
                calculated.origin_or_maker,
                calculated.supplier,
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.value = value
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        total_row = max(41, 10 + len(school.items))
        for row_idx in range(10, total_row):
            worksheet.row_dimensions[row_idx].height = max(
                _price_text_height(str(worksheet.cell(row_idx, col).value or ""), widths[col - 1])
                for col in (1, 7, 8)
            )
            for col_idx in range(1, 9):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.border = border
                cell.alignment = Alignment(horizontal="right" if col_idx in (2, 3, 4, 5, 6) else "left", vertical="center", wrap_text=True)
                if col_idx in (2, 4):
                    cell.number_format = '#,##0.###'
                elif col_idx == 3:
                    cell.number_format = '#,##0.00#'
                elif col_idx == 5:
                    cell.number_format = '#,##0'
                elif col_idx == 6:
                    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin, diagonal=thin, diagonalUp=True)
        for row_idx, calculated in enumerate(calculated_rows, start=10):
            for col_idx, value, unit in (
                (2, calculated.amount_per_person, calculated.unit),
                (3, calculated.total_quantity, calculated.total_unit),
                (4, calculated.unit_price_tax_included, f"円/{calculated.total_unit}"),
            ):
                decimals = max(0, -value.normalize().as_tuple().exponent) if value is not None else 0
                number_format = '#,##0' + ('.' + '0' * decimals if decimals else '')
                worksheet.cell(row=row_idx, column=col_idx).number_format = f'{number_format}" {unit}"'
        worksheet.cell(row=total_row, column=1).value = "合計"
        worksheet.cell(row=total_row, column=5).value = total_price_sum
        worksheet.cell(row=total_row, column=5).number_format = '#,##0'
        worksheet.cell(row=total_row, column=6).value = float(_per_person_price(total_price_sum, school.headcount))
        worksheet.cell(row=total_row, column=6).number_format = '#,##0.00'
        for col_idx in range(1, len(headers) + 1):
            cell = worksheet.cell(row=total_row, column=col_idx)
            cell.font = Font(bold=True)
            cell.border = border
        worksheet.row_dimensions[total_row].height = 33.75
        worksheet.cell(row=total_row + 2, column=1).value = "※単価は，税込みで記入する。"
        worksheet.merge_cells(start_row=total_row + 2, start_column=1, end_row=total_row + 2, end_column=8)
        for col_idx, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(col_idx)].width = width
        worksheet.row_dimensions[1].height = 27.75
        worksheet.row_dimensions[9].height = max(37.5, *(
            _price_text_height(header, width) for header, width in zip(headers, widths)
        ))
        worksheet.print_area = f"A1:H{total_row + 2}"
        worksheet.sheet_view.showGridLines = False
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.orientation = "portrait"
        worksheet.page_setup.fitToWidth = 1
        expanded_rows = any(worksheet.row_dimensions[row].height > 24.95 for row in range(10, total_row))
        worksheet.page_setup.fitToHeight = 1 if total_row <= 41 and not expanded_rows else 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.blackAndWhite = True
        worksheet.page_margins = PageMargins(left=0.7874, right=0.7874, top=0.7874, bottom=0.3937, header=0.5118, footer=0.5118)
        worksheet.print_title_rows = "1:9"
        worksheet.freeze_panes = "A10"

    for worksheet in workbook:
        for row in worksheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"

    return workbook


def build_price_table_workbook_bytes(payload: dict) -> bytes:
    request = parse_price_table_request(payload)
    workbook = build_price_table_workbook(request)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def build_price_table_workbook_file(payload: dict) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    request = parse_price_table_request(payload)
    workbook = build_price_table_workbook(request)
    safe_date = request.implementation_date.isoformat()
    output_path = OUTPUT_DIR / f"implementation_price_table_{safe_date}_{uuid4().hex}.xlsx"
    try:
        workbook.save(output_path)
    except Exception:
        output_path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return output_path


def build_scoped_price_table_workbook_file(validation_payload: dict, payload: dict, *, from_sources=False) -> Path:
    from copy import deepcopy
    original = deepcopy(validation_payload)
    materialization = deepcopy(payload)
    output_path = None
    try:
        with session_scope() as session:
            validate_price_scope(session, original, for_adoption=True)
            if from_sources:
                materialization = build_price_table_payload_from_manual_sources(materialization)
            output_path = build_price_table_workbook_file(materialization)
        return output_path
    except Exception:
        if output_path is not None:
            output_path.unlink(missing_ok=True)
        raise


def build_scoped_price_table_preview(validation_payload: dict, payload: dict) -> dict:
    from copy import deepcopy
    original, materialization = deepcopy(validation_payload), deepcopy(payload)
    with session_scope() as session:
        validate_price_scope(session, original, for_adoption=False)
        return build_price_table_preview_from_manual_sources(materialization)
