from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any


_DOCUMENT_TYPES = {"invoice", "quantity_sheet", "cooking_arrangement"}
_FULLWIDTH_TRANSLATION = str.maketrans(
    {
        "０": "0",
        "１": "1",
        "２": "2",
        "３": "3",
        "４": "4",
        "５": "5",
        "６": "6",
        "７": "7",
        "８": "8",
        "９": "9",
        "．": ".",
        "，": ",",
        "－": "-",
        "　": " ",
    }
)


@dataclass(frozen=True)
class YomitokuPage:
    page_index: int
    markdown_text: str
    tables: list[dict[str, Any]]
    orientation: dict[str, Any] | None = None
    words: list[dict[str, Any]] = field(default_factory=list)


def _workspace_root() -> Path:
    return Path(__file__).resolve().parents[3]


def _load_yomitoku_runner():
    ocr_pipeline_root = _workspace_root() / "ocr_pipeline"
    if str(ocr_pipeline_root) not in sys.path:
        sys.path.insert(0, str(ocr_pipeline_root))
    try:
        from app.yomitoku_runner import run_yomitoku  # noqa: PLC0415
    except ModuleNotFoundError as exc:
        missing_name = getattr(exc, "name", "") or str(exc)
        raise RuntimeError(
            "Yomitoku runtime dependency missing: "
            f"{missing_name}. Install the canonical OCR/runtime dependencies for implementation-price OCR."
        ) from exc
    except ImportError as exc:
        raise RuntimeError(
            "Yomitoku runtime failed to import. Install the canonical OCR/runtime dependencies for implementation-price OCR."
        ) from exc

    return run_yomitoku


def normalize_document_type(document_type: str) -> str:
    normalized_type = str(document_type or "").strip()
    if normalized_type not in _DOCUMENT_TYPES:
        raise ValueError("document_type must be invoice, quantity_sheet, or cooking_arrangement")
    return normalized_type


def _run_yomitoku_pdf(
    *,
    pdf_bytes: bytes,
    dpi: int,
    device: str,
    orient_invoice: bool = False,
) -> list[YomitokuPage]:
    run_yomitoku = _load_yomitoku_runner()
    page_images = None
    orientations = {}
    if orient_invoice:
        from app.pdf_render import render_pdf_to_page_images
        from app.yomitoku_runner import _get_analyzer
        import cv2
        import numpy as np

        page_images = []
        analyzer = _get_analyzer(device, False)
        # Use the public components of the pinned worker runtime (0.9.5).
        detector, recognizer = analyzer.text_detector, analyzer.text_recognizer
        for index, image in render_pdf_to_page_images(pdf_bytes, dpi):
            scale = min(1.0, 1100 / max(image.shape[:2]))
            thumb = cv2.resize(image, None, fx=scale, fy=scale) if scale < 1 else image
            scores = {}
            # Orientation is chosen from actual recognition, never a filename,
            # invoice amount reverse calculation or page-specific rotation rule.
            for angle in (0, 90, 180, 270):
                sample = np.ascontiguousarray(np.rot90(thumb, -angle // 90))
                detections, _ = detector(sample)
                # Only orientation uses representative text boxes. After the
                # rotation is selected, normal OCR reads ALL cells at full DPI.
                points = sorted(list(getattr(detections, "points", [])), key=lambda box: cv2.contourArea(np.asarray(box, dtype=np.float32)), reverse=True)[:32]
                if points:
                    prediction, _ = recognizer(sample, points)
                    contents = list(getattr(prediction, "contents", []))
                    confidences = list(getattr(prediction, "scores", []))
                else:
                    contents, confidences = [], []
                text = " ".join(str(content) for content in contents)
                compact = re.sub(r"\s+", "", text)
                headers = sum(compact.count(header) for header in ("納品", "品名", "商品名", "数量", "単価", "金額", "伝票", "税抜"))
                reliable = sum(min(len(str(content)), 20) * float(confidence) for content, confidence in zip(contents, confidences) if float(confidence) >= .7)
                # The recognizer can read individually rotated crops even when
                # the entire page is sideways. Header recognition alone must
                # not select a sideways table. Weight actual horizontal text
                # geometry; no numeric value or supplier-specific rule is used.
                horizontal = sum(float(np.ptp(np.asarray(box)[:, 0])) > float(np.ptp(np.asarray(box)[:, 1])) * 1.2 for box in points)
                horizontal_share = horizontal / len(points) if points else 0
                scores[angle] = round((headers * 100 + reliable) * horizontal_share, 3)
            rotation = max(scores, key=scores.get)
            orientations[index] = {"rotation_clockwise": rotation, "recognition_scores": scores, "review_required": True}
            page_images.append((index, np.ascontiguousarray(np.rot90(image, -rotation // 90))))
    page_results, _ocr_pdf, _layout_pdf = run_yomitoku(
        pdf_bytes=pdf_bytes,
        dpi=dpi,
        device=device,
        visualize=False,
        ignore_line_break=False,
        no_figure=True,
        figure_width=1200,
        figure_dir="figures",
        **({"page_images": page_images} if page_images is not None else {}),
    )
    pages: list[YomitokuPage] = []
    for page in page_results:
        pages.append(
            YomitokuPage(
                page_index=int(getattr(page, "page_index", len(pages))),
                markdown_text=str(getattr(page, "markdown_text", "") or ""),
                tables=list(getattr(page, "tables", []) or []),
                orientation=orientations.get(int(getattr(page, "page_index", len(pages)))),
                words=list(getattr(page, "words", []) or []),
            )
        )
    return pages


def _split_markdown_cells(line: str) -> list[str]:
    content = line.strip()
    if content.startswith("|"):
        content = content[1:]
    if content.endswith("|"):
        content = content[:-1]
    return [cell.strip().replace("<br>", " ").replace("<br/>", " ") for cell in content.split("|")]


def _looks_like_separator(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", cell.replace(" ", "")) for cell in cells)


def _extract_markdown_tables(markdown: str) -> list[list[list[str]]]:
    tables: list[list[list[str]]] = []
    current: list[list[str]] = []
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if "|" not in line:
            if current:
                tables.append(current)
                current = []
            continue
        cells = _split_markdown_cells(line)
        if not any(cells) or _looks_like_separator(cells):
            continue
        current.append(cells)
    if current:
        tables.append(current)
    return tables


def _explicit_cooking_facility_name(rows: list[list[str]], tokens=()) -> str:
    """Accept only one complete school-name cell/token, never a selected school."""
    match = _find_header_row(rows)
    preamble = rows[:match[0]] if match is not None else []
    candidates = {
        value.strip()
        for value in [*(cell for row in preamble for cell in row), *tokens]
        if isinstance(value, str) and re.fullmatch(r"[^\n]+(?:学校|学園)", value.strip())
    }
    return next(iter(candidates)) if len(candidates) == 1 else ""


def _normalize_cooking_header_spans(table: dict[str, Any]) -> dict[str, Any]:
    """Normalize explicit header spans once, preserving raw cells and row indexes."""
    if table.get("header_spans_normalized"):
        return table
    rows = table["rows"]
    match = _find_header_row(rows)
    if match is None or match[1]["menu"] != match[1]["ingredient"]:
        return table
    header_index, columns = match
    starts = [i for i, value in enumerate(rows[header_index]) if value.strip()]
    spans = [(start, starts[i + 1] if i + 1 < len(starts) else len(rows[header_index]))
             for i, start in enumerate(starts)]
    collapsed, unit_rows, sources = [], [], dict(table.get("cell_sources", {}))
    for row_index, row in enumerate(rows):
        if row_index <= header_index:
            collapsed.append(row if row_index < header_index else [row[start] for start, _ in spans])
            continue
        values = [" ".join(value for value in row[start:end] if value) for start, end in spans]
        # Require an explicit unit label immediately following the table header.
        if (row_index == header_index + 1 and "単位" in values
                and all(not value or value in {"単位", "g", "ｇ"} for value in values)):
            unit_rows.append(row_index + 1)
            values = [""] * len(values)
        collapsed.append(values)
        source_data = dict(sources.get(row_index + 1, {}))

        def address(column):
            if table.get("source") == "structured_excel":
                from openpyxl.utils import get_column_letter
                return f"{table['table_id']}!{get_column_letter(column + 1)}{row_index + 1}"
            return f"{table.get('source')}:page{table.get('page_index')}:table{table.get('table_id')}:row{row_index + 1}:column{column + 1}"

        for (start, _), value in zip(spans, values):
            label = re.sub(r"\s+", "", rows[header_index][start])
            key = next((key for key, label_part in (("amount_per_person", "一人分量"),
                ("order_quantity", "発注量"), ("supplier", "業者"), ("cutting", "切り方"))
                if label_part in label), None)
            if key:
                source_data[key] = value
                source_data[f"{key}_source"] = address(start)
        source_data["source_cell"] = address(columns["ingredient"])
        if columns["process"] is not None:
            source_data["cooking_method_source"] = address(columns["process"])
        sources[row_index + 1] = source_data
    return {**table, "raw_rows": table.get("raw_rows", rows), "rows": collapsed, "cell_sources": sources,
            "explicit_unit_header_rows": unit_rows, "header_spans_normalized": True}


def _tables_from_pages(pages: list[YomitokuPage]) -> list[dict[str, Any]]:
    tables: list[dict[str, Any]] = []
    for page in pages:
        school_tokens = [word["text"] for word in page.words
                         if isinstance(word, dict) and isinstance(word.get("text"), str)]
        for index, raw_table in enumerate(page.tables, start=1):
            if isinstance(raw_table, dict) and isinstance(raw_table.get("rows"), list):
                tables.append(
                    {
                        "page_index": page.page_index,
                        "source": raw_table.get("source") or "yomitoku_table_structure",
                        "table_id": raw_table.get("table_id") or f"p{page.page_index}_structured{index}",
                        "rows": raw_table.get("rows") or [],
                        "page_text": page.markdown_text,
                        "facility_name": _explicit_cooking_facility_name(raw_table["rows"], school_tokens),
                    }
                )
        for index, rows in enumerate(_extract_markdown_tables(page.markdown_text), start=1):
            tables.append(
                {
                    "page_index": page.page_index,
                    "source": "yomitoku_markdown",
                    "table_id": f"p{page.page_index}_md{index}",
                    "rows": rows,
                    "page_text": page.markdown_text,
                    "facility_name": _explicit_cooking_facility_name(rows, school_tokens),
                }
            )
    return tables


def _normalize_text(value: object) -> str:
    return str(value or "").translate(_FULLWIDTH_TRANSLATION).strip()


def _compact_key(value: object) -> str:
    return re.sub(r"\s+", "", _normalize_text(value)).lower()


def _first_numeric(value: object) -> str:
    tokens = _numeric_tokens(value)
    return tokens[0] if tokens else ""


def _numeric_tokens(value: object) -> list[str]:
    text = _normalize_text(value).replace("\\.", ".")
    text = re.sub(r"(?<=\d)\s*\.\s*(?=\d)", ".", text)
    text = re.sub(r"(?<=\d),\s*(?=\d{3}(?:\D|$))", "", text)
    return re.findall(r"-?\d+(?:\.\d+)?", text)




def _decimal_or_none(value: object) -> Decimal | None:
    raw = _first_numeric(value)
    if not raw:
        return None
    try:
        return Decimal(raw)
    except Exception:
        return None


def _explicit_decimal_or_none(value: object) -> Decimal | None:
    raw = _normalize_text(value).replace(",", "")
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", raw):
        return None
    try:
        return Decimal(raw)
    except Exception:
        return None




def _decimal_string(value: Decimal | None) -> str:
    if value is None:
        return ""
    return str(value)


def _infer_price_unit(row: list[str], unit_index: int | None) -> str:
    source = _normalize_text(row[unit_index]) if unit_index is not None and unit_index < len(row) else ""
    lowered = source.lower()
    if "kg" in lowered or "ｋｇ" in lowered:
        return "円/kg"
    if "g" in lowered or "ｇ" in lowered:
        return "円/kg"
    if "本" in source:
        return "円/本"
    if "個" in source:
        return "円/個"
    if "袋" in source:
        return "円/袋"
    return ""


def _find_column(header: list[str], keywords: tuple[str, ...]) -> int | None:
    for index, cell in enumerate(header):
        key = _compact_key(cell).replace("·", "").replace("・", "")
        if any(_compact_key(keyword).replace("·", "").replace("・", "") in key for keyword in keywords):
            return index
    return None






def _cell(row: list[str], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return row[index]


def _row_has_product_name(value: str) -> bool:
    text = _normalize_text(value)
    if not text:
        return False
    compact = _compact_key(text)
    if compact in {
        "備考",
        "合計",
        "小計",
        "商品代",
        "内消費税",
        "税対象",
        "10%対象",
        "8%対象",
        "品名",
        "商品名",
        "名称",
    }:
        return False
    if "消費税" in compact or "対象" in compact or "合計" in compact or "小計" in compact:
        return False
    if compact.endswith(("学校", "中学校", "小学校", "支援学校")):
        return False
    if re.fullmatch(r"\d+", compact):
        return False
    if re.fullmatch(r"[a-z]{1,3}", compact):
        return False
    return bool(re.search(r"[A-Za-zぁ-んァ-ヶ一-龥]", text))


def _clean_item_name(value: str) -> str:
    text = _normalize_text(value)
    text = re.sub(r"^\d{4,}\s+", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _candidate_invoice_line_from_row(
    *,
    table: dict[str, Any],
    row: list[str],
    row_index: int,
    name_idx: int,
    quantity_idx: int | None,
    unit_idx: int | None,
    unit_price_idx: int | None,
    amount_idx: int | None,
    origin_or_maker_idx: int | None,
    supplier_idx: int | None,
) -> dict[str, Any] | None:
    name = _clean_item_name(_cell(row, name_idx))
    if not _row_has_product_name(name):
        return None
    values = [_cell(row, column) for column in (quantity_idx, unit_price_idx, amount_idx)]
    parsed = [_numeric_tokens(value) for value in values]
    quantity, unit_price, amount = [Decimal(tokens[0]) if len(tokens) == 1 else None for tokens in parsed]
    if all(value is None for value in (quantity, unit_price, amount)) and not any(values):
        return None
    unit = _cell(row, unit_idx)
    if not unit:
        match = re.search(r"(?i)(?<![A-Za-z])(kg|g|個|本|袋|缶|枚|尾)(?![A-Za-z])", values[0])
        unit = match.group(1).lower() if match else ""
    warnings = [field + "_ambiguous" for field, tokens in zip(("quantity", "unit_price", "amount"), parsed) if len(tokens) > 1]
    if not unit:
        warnings.append("invoice_unit_missing")
    return {
        "key": "",
        "name": name,
        "status": "needs_review",
        "quantity": _decimal_string(quantity),
        "unit": unit,
        "unit_price": _decimal_string(unit_price),
        "unit_price_tax_included": "",
        "amount": _decimal_string(amount),
        "raw_quantity": _cell(row, quantity_idx),
        "raw_unit_price": _cell(row, unit_price_idx),
        "raw_amount": _cell(row, amount_idx),
        "raw_cells": row,
        "page_index": table.get("page_index"),
        "table_id": table.get("table_id") or "",
        "table_source": table.get("source") or "",
        "price_unit": _infer_price_unit([unit], 0),
        "origin_or_maker": _cell(row, origin_or_maker_idx),
        "supplier": _cell(row, supplier_idx),
        "source": f"yomitoku:{table.get('source') or 'table'}:page{table.get('page_index')}:row{row_index}",
        "review_reason": "yomitoku_candidate_requires_human_confirmation",
        "review_issues": warnings,
        "source_row_index": row_index,
    }


def _invoice_line_items_from_table(table: dict[str, Any]) -> list[dict[str, Any]]:
    rows = table.get("rows")
    if not isinstance(rows, list) or len(rows) < 2:
        return []
    normalized_rows = [[_normalize_text(cell) for cell in row] for row in rows if isinstance(row, list)]
    if len(normalized_rows) < 2:
        return []
    def name_column(row):
        column = _find_column(row, ("品名", "商品名", "名称", "摘要"))
        if column is None:
            column = next((index for index, cell in enumerate(row) if _compact_key(cell) in {"品", "商品"}), None)
        return column

    header_index = next((index for index, row in enumerate(normalized_rows[:30]) if name_column(row) is not None and _find_column(row, ("数量", "単価", "金額", "原単価")) is not None), None)
    if header_index is None:
        return []
    header = normalized_rows[header_index]
    body_rows = normalized_rows[header_index + 1:]
    name_idx = name_column(header)
    if name_idx is None:
        name_idx = _find_column(header, ("商 品 名",))
    if name_idx is None:
        return []
    quantity_idx = _find_column(header, ("数量", "数 量", "訂正後数量", "訂正後改量"))
    unit_price_idx = _find_column(header, ("原単価", "原 単 価", "単価", "単 価", "売単価", "野東"))
    amount_idx = _find_column(header, ("原価金額", "原 価 金 額", "金額", "額金", "金 額", "元価金額", "売価金額"))
    unit_idx = _find_column(header, ("単位",))
    origin_or_maker_idx = _find_column(header, ("産地", "メーカー", "等級"))
    supplier_idx = _find_column(header, ("業者", "仕入先", "納入", "販売"))
    if quantity_idx is None and unit_price_idx is None and amount_idx is None:
        return []
    candidates: list[dict[str, Any]] = []
    for row_index, row in enumerate(body_rows, start=header_index + 1):
        item = _candidate_invoice_line_from_row(
            table=table,
            row=row,
            row_index=row_index,
            name_idx=name_idx,
            quantity_idx=quantity_idx,
            unit_idx=unit_idx,
            unit_price_idx=unit_price_idx,
            amount_idx=amount_idx,
            origin_or_maker_idx=origin_or_maker_idx,
            supplier_idx=supplier_idx,
        )
        if item is not None:
            candidates.append(item)
    return candidates


def _invoice_item_signature(item: dict[str, Any]) -> tuple[str, str, str, str, str]:
    name = _compact_key(item.get("name")).replace("\\", "")
    return (
        name,
        str(item.get("quantity") or ""),
        _compact_key(item.get("unit")),
        str(item.get("unit_price") or item.get("unit_price_tax_included") or ""),
        str(item.get("amount") or ""),
    )


def _dedupe_cross_source_invoice_tables(invoice_tables: list[dict[str, Any]]) -> list[dict[str, Any]]:
    kept: list[dict[str, Any]] = []
    structured_signatures: set[tuple[int, tuple[tuple[str, str, str, str, str], ...]]] = set()
    for table in invoice_tables:
        signature = (
            int(table.get("page_index") or 0),
            tuple(_invoice_item_signature(item) for item in table.get("line_items") or []),
        )
        is_markdown = table.get("source") == "yomitoku_markdown"
        if is_markdown and signature in structured_signatures:
            continue
        kept.append(table)
        if not is_markdown:
            structured_signatures.add(signature)
    return kept




def _invoice_tables_with_line_items(tables: list[dict[str, Any]]) -> list[dict[str, Any]]:
    invoice_tables: list[dict[str, Any]] = []
    for table in tables:
        line_items = _invoice_line_items_from_table(table)
        if not line_items:
            continue
        invoice_tables.append(
            {
                "page_index": table.get("page_index"),
                "table_id": table.get("table_id") or "",
                "source": table.get("source") or "",
                "line_items": line_items,
            }
        )
    return _dedupe_cross_source_invoice_tables(invoice_tables)


def _quantity_candidates_from_tables(tables: list[dict[str, Any]]) -> list[dict[str, Any]]:
    parsed = [(table, _quantity_candidates_from_table(table)) for table in tables]
    structured_counts: dict[tuple, int] = {}

    def signature(table: dict[str, Any], items: list[dict[str, Any]]) -> tuple:
        return (table.get("page_index"), tuple(
            tuple(item[field] for field in ("name", "amount_per_person", "unit", "cost_key"))
            for item in items
        ))

    # Pair representations one-to-one, including when markdown appears first.
    for table, items in parsed:
        if table.get("source") == "yomitoku_table_structure" and items:
            key = signature(table, items)
            structured_counts[key] = structured_counts.get(key, 0) + 1
    candidates: list[dict[str, Any]] = []
    for table, items in parsed:
        key = signature(table, items)
        if table.get("source") == "yomitoku_markdown" and structured_counts.get(key, 0):
            structured_counts[key] -= 1
            continue
        candidates.extend(items)
    return candidates


def _quantity_candidates_from_table(table: dict[str, Any]) -> list[dict[str, Any]]:
    rows = table.get("rows")
    if not isinstance(rows, list) or len(rows) < 2:
        return []
    normalized_rows = [[_normalize_text(cell) for cell in row] for row in rows if isinstance(row, list)]
    if len(normalized_rows) < 2:
        return []
    header = normalized_rows[0]
    name_idx = _find_column(header, ("品名", "食品", "材料", "商品", "名称"))
    amount_idx = _find_column(header, ("一人", "1人", "使用量", "分量", "数量"))
    unit_idx = _find_column(header, ("単位",))
    if name_idx is None:
        name_idx = 0
    if amount_idx is None:
        amount_idx = next(
            (
                index
                for index, _cell in enumerate(header)
                if index != name_idx and any(_decimal_or_none(row[index]) is not None for row in normalized_rows[1:] if index < len(row))
            ),
            None,
        )
    if amount_idx is None:
        return []
    candidates: list[dict[str, Any]] = []
    for row_index, row in enumerate(normalized_rows[1:], start=1):
        if name_idx >= len(row) or amount_idx >= len(row):
            continue
        name = row[name_idx].strip()
        amount = _decimal_or_none(row[amount_idx])
        if not name or amount is None:
            continue
        unit = row[unit_idx] if unit_idx is not None and unit_idx < len(row) else ""
        candidates.append(
            {
                "name": name,
                "amount_per_person": str(amount),
                "unit": unit,
                "cost_key": "",
                "source": f"yomitoku:{table.get('source') or 'table'}:page{table.get('page_index')}:row{row_index}",
                "status": "needs_review",
                "review_reason": "yomitoku_candidate_requires_human_confirmation",
            }
        )
    return candidates


def _normalize_explicit_date(value: object) -> tuple[str, str | None]:
    raw = _normalize_text(value)
    if not raw:
        return "", "service_date_missing"
    match = re.fullmatch(r"(\d{4})\s*(?:年|[-/.])\s*(\d{1,2})\s*(?:月|[-/.])\s*(\d{1,2})\s*日?", raw)
    if not match:
        if re.fullmatch(r"\d{1,2}\s*(?:月|[-/.])\s*\d{1,2}\s*日?", raw):
            return "", "service_date_year_required"
        return "", "service_date_invalid"
    year, month, day = (int(item) for item in match.groups())
    try:
        from datetime import date

        return date(year, month, day).isoformat(), None
    except ValueError:
        return "", "service_date_invalid"


def _find_header_row(rows: list[list[str]]) -> tuple[int, dict[str, int | None]] | None:
    keyword_map: dict[str, tuple[str, ...]] = {
        "date": ("提供日", "実施日", "調理日", "日付", "年月日"),
        "menu": ("献立名", "料理名", "メニュー名", "献立"),
        "ingredient": ("食材名", "食品名", "材料名", "食材", "食品", "材料"),
        "quantity": ("使用量", "必要量", "数量"),
        "unit": ("単位",),
        "servings": ("提供食数", "予定食数", "食数", "人数"),
        "process": ("調理作業内容", "調理作業", "調理方法", "作業内容", "工程"),
        "location": ("作業場所", "調理場所", "場所"),
        "time": ("開始時刻", "作業時刻", "時刻", "時間"),
        "work_zone": ("作業区分", "汚染区分"),
        "line_style": ("線種", "導線区分"),
    }
    best: tuple[int, int, dict[str, int | None]] | None = None
    for row_index, row in enumerate(rows[:30]):
        columns = {key: _find_column(row, keywords) for key, keywords in keyword_map.items()}
        if columns["menu"] is None or columns["ingredient"] is None:
            continue
        score = sum(index is not None for index in columns.values())
        if best is None or score > best[0]:
            best = (score, row_index, columns)
    return (best[1], best[2]) if best is not None else None


def _explicit_work_zone(value: object) -> str | None:
    key = _compact_key(value)
    if key in {"汚染", "汚染作業", "contaminated"}:
        return "contaminated"
    if key in {"非汚染", "非汚染作業", "noncontaminated", "non_contaminated"}:
        return "non_contaminated"
    return None


def _explicit_line_style(value: object) -> str | None:
    key = _compact_key(value)
    if key in {"実線", "solid"}:
        return "solid"
    if key in {"点線", "破線", "dashed"}:
        return "dashed"
    return None


def _explicit_time(value: object) -> str | None:
    raw = _normalize_text(value)
    if not raw:
        return None
    match = re.fullmatch(r"(\d{1,2})[:：時](\d{1,2})分?", raw)
    if not match:
        return None
    hour, minute = (int(item) for item in match.groups())
    if hour > 23 or minute > 59:
        return None
    return f"{hour:02d}:{minute:02d}"


def _explicit_date_from_page_text(value: object) -> tuple[str, str, str | None]:
    candidates: list[tuple[str, str]] = []
    for line in str(value or "").splitlines():
        normalized_line = _normalize_text(line)
        if "印刷" in normalized_line:
            continue
        for match in re.finditer(r"\d{4}\s*(?:年|[-/.])\s*\d{1,2}\s*(?:月|[-/.])\s*\d{1,2}\s*日?", normalized_line):
            raw_date = match.group(0)
            service_date, issue = _normalize_explicit_date(raw_date)
            if not issue:
                candidates.append((service_date, raw_date))
    unique = list(dict.fromkeys(candidates))
    if len(unique) == 1:
        return unique[0][0], unique[0][1], None
    if not unique:
        return "", "", "service_date_missing"
    return "", "", "service_date_ambiguous"


def _explicit_servings_from_preamble(rows: list[list[str]], header_index: int) -> int | None:
    candidates: list[int] = []
    for row_index, row in enumerate(rows[:header_index]):
        for column_index, value in enumerate(row):
            if _compact_key(value) not in {"換算人", "提供食数", "予定食数", "食数"}:
                continue
            for following in rows[row_index + 1 : header_index]:
                raw_value = _normalize_text(_cell(following, column_index)).replace(",", "")
                raw_value = re.sub(r"(?<=\d)\s*\.\s*(?=\d)", ".", raw_value)
                decimal_value = _explicit_decimal_or_none(raw_value)
                if decimal_value is None or decimal_value <= 0 or decimal_value != int(decimal_value):
                    continue
                candidates.append(int(decimal_value))
                break
    unique = list(dict.fromkeys(candidates))
    return unique[0] if len(unique) == 1 else None


def _explicit_quantity_and_unit(value: object) -> tuple[Decimal | None, str]:
    raw = _normalize_text(value).replace("\n", " ")
    raw = re.sub(r"(?<=\d)\s*\.\s*(?=\d)", ".", raw)
    raw = raw.replace(",", "").strip()
    if any(marker in raw for marker in ("約", "程度", "目安", "～", "〜", "~")):
        return None, ""
    match = re.fullmatch(r"(-?\d+(?:\.\d+)?)\s*([^\d\s][^\d]*)", raw)
    if not match:
        return None, ""
    try:
        quantity = Decimal(match.group(1))
    except Exception:
        return None, ""
    unit = re.sub(r"\s+", "", match.group(2)).strip()
    if quantity < 0 or not unit:
        return None, ""
    return quantity, unit


def _combined_menu_name(value: object) -> str | None:
    raw = _normalize_text(value)
    match = re.fullmatch(r"[【\[]\s*([\s\S]+?)\s*[】\]]", raw)
    return re.sub(r"\s+", " ", match.group(1)).strip() if match else None


def _extract_combined_cooking_table(
    *,
    table: dict[str, Any],
    rows: list[list[str]],
    header_index: int,
    columns: dict[str, int | None],
    source_prefix: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    issues: list[dict[str, Any]] = []
    date_source = table.get("page_text") or "\n".join(
        _normalize_text(cell)
        for row in rows[:header_index]
        for cell in row
    )
    service_date, raw_service_date, date_issue = _explicit_date_from_page_text(date_source)
    if date_issue:
        issues.append({"code": date_issue, "message": "提供日を年月日で確認してください。", "source": source_prefix})
    servings = _explicit_servings_from_preamble(rows, header_index)
    if servings is None:
        issues.append({"code": "servings_unresolved", "message": "提供食数を確認してください。", "source": source_prefix})
    quantity_column = columns["quantity"]
    process_column = columns["process"]
    menu_column = columns["menu"]
    if quantity_column is None:
        issues.append({"code": "quantity_header_missing", "message": "使用量列を特定できませんでした。", "source": source_prefix})
    if process_column is None:
        issues.append({"code": "process_header_missing", "message": "調理作業内容の列を特定できませんでした。", "source": source_prefix})

    menus: list[dict[str, Any]] = []
    current_menu: dict[str, Any] | None = None
    pending_quantity = ""
    after_separator = False
    for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if row_index in table.get("explicit_unit_header_rows", []):
            continue
        raw_name = _cell(row, menu_column).strip()
        process_fragment = _cell(row, process_column).strip()
        if raw_name.startswith("※"):
            issues.append({"code": "source_note_review", "message": "原本の注記を確認してください。", "source": f"{source_prefix}:row{row_index}", "raw_value": " ".join(cell for cell in row if cell)})
            continue
        row_has_content = any(_normalize_text(cell) for cell in row)
        if not row_has_content:
            # Direct cell extraction has no OCR heading loss. Blank rows do not
            # replace the last explicit heading; OCR gaps still require review.
            after_separator = table.get("source") != "structured_excel"
            continue
        menu_name = _combined_menu_name(raw_name)
        if menu_name:
            current_menu = {
                "facility_name": table.get("facility_name", ""),
                "service_date": service_date,
                "raw_service_date": raw_service_date,
                "name": menu_name,
                "servings": servings,
                "status": "needs_review",
                "ingredients": [],
            }
            menus.append(current_menu)
            current_menu["cooking_instructions"] = process_fragment
            current_menu["cooking_instruction_sources"] = ([{"text": process_fragment, "source": f"{source_prefix}:row{row_index}"}] if process_fragment else [])
            pending_quantity = _cell(row, quantity_column).strip()
            after_separator = False
            continue
        if current_menu is None:
            issues.append({
                "code": "menu_heading_missing",
                "message": "先頭区間の献立名を特定できませんでした。空欄の未解決献立として保持します。",
                "source": f"{source_prefix}:row{row_index}",
                "raw_value": raw_name,
            })
            current_menu = {
                "facility_name": table.get("facility_name", ""),
                "service_date": service_date,
                "raw_service_date": raw_service_date,
                "name": "",
                "servings": servings,
                "status": "needs_review",
                "ingredients": [],
            }
            menus.append(current_menu)
            pending_quantity = ""
            after_separator = False
        if after_separator and any(_normalize_text(cell) for index, cell in enumerate(row) if index != process_column):
            issues.append({
                "code": "menu_heading_missing",
                "message": "区切り後の献立名を特定できませんでした。空欄の未解決献立として保持します。",
                "source": f"{source_prefix}:row{row_index}",
                "raw_value": raw_name,
            })
            current_menu = {
                "facility_name": table.get("facility_name", ""),
                "service_date": service_date,
                "raw_service_date": raw_service_date,
                "name": "",
                "servings": servings,
                "status": "needs_review",
                "ingredients": [],
            }
            menus.append(current_menu)
            pending_quantity = ""
            after_separator = False
        if process_fragment:
            current_menu["cooking_instructions"] = "\n".join(filter(None, [current_menu.get("cooking_instructions", ""), process_fragment]))
            current_menu.setdefault("cooking_instruction_sources", []).append({"text": process_fragment, "source": f"{source_prefix}:row{row_index}"})
        if not raw_name:
            raw_quantity = _cell(row, quantity_column).strip()
            has_nonprocess_content = any(
                _normalize_text(cell)
                for column_index, cell in enumerate(row)
                if column_index != process_column
            )
            if has_nonprocess_content:
                row_source = f"{source_prefix}:row{row_index}"
                issues.append({
                    "code": "ingredient_name_missing",
                    "message": "食材名が空の行です。空欄の未解決食材として保持します。",
                    "source": row_source,
                    "raw_cells": row,
                })
                quantity, unit = _explicit_quantity_and_unit(raw_quantity)
                if quantity is None:
                    issues.append({
                        "code": "quantity_unresolved",
                        "message": "使用量と単位を確認してください。",
                        "source": row_source,
                        "raw_value": raw_quantity,
                    })
                pending_quantity = ""
                ingredient = {
                    "name": "",
                    "required_quantity": _decimal_string(quantity),
                    "raw_quantity": raw_quantity,
                    "unit": unit,
                    "process_text": "",
                    "process_steps": [],
                    "source": row_source,
                    "status": "needs_review",
                }
                current_menu["ingredients"].append(ingredient)
                continue
            continue

        row_source = f"{source_prefix}:row{row_index}"
        raw_quantity = _cell(row, quantity_column).strip()
        if not raw_quantity and pending_quantity:
            raw_quantity = pending_quantity
        pending_quantity = ""
        quantity, unit = _explicit_quantity_and_unit(raw_quantity)
        if quantity is None:
            issues.append({
                "code": "quantity_unresolved",
                "message": "使用量と単位を確認してください。",
                "source": row_source,
                "raw_value": raw_quantity,
            })
        ingredient = {
            "name": raw_name,
            "required_quantity": _decimal_string(quantity),
            "raw_quantity": raw_quantity,
            "unit": unit,
            "process_text": "",
            "process_steps": [],
            "source": row_source,
            "status": "needs_review",
            **table.get("cell_sources", {}).get(row_index, {}),
        }
        current_menu["ingredients"].append(ingredient)

    for menu in menus:
        menu.setdefault("cooking_instructions", "")
        menu.setdefault("cooking_instruction_sources", [])
        if not menu["ingredients"]:
            issues.append({
                "code": "menu_ingredients_missing",
                "message": "献立に食材行がありません。",
                "source": source_prefix,
                "raw_value": menu["name"],
            })
        names = [item["name"].strip() for item in menu["ingredients"]]
        duplicate_names = sorted({name for name in names if name and names.count(name) > 1})
        for duplicate_name in duplicate_names:
            issues.append({
                "code": "duplicate_ingredient_name",
                "message": "同じ献立内に同名の食材が複数あります。数量と工程を確認してください。",
                "source": source_prefix,
                "raw_value": duplicate_name,
            })
    return menus, issues


def _cooking_arrangement_from_tables(
    tables: list[dict[str, Any]],
    *,
    filename: str,
    provider: str,
) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    grouped: dict[tuple[str, ...], dict[str, Any]] = {}
    seen_rows: set[tuple[str, ...]] = set()

    structured_pages = {
        table.get("page_index")
        for table in tables
        if table.get("source") == "yomitoku_table_structure"
    }
    selected_tables = [
        table
        for table in tables
        if not (
            table.get("source") == "yomitoku_markdown"
            and table.get("page_index") in structured_pages
        )
    ]

    for table in selected_tables:
        raw_rows = table.get("rows")
        if not isinstance(raw_rows, list):
            continue
        rows = [[_normalize_text(cell) for cell in row] for row in raw_rows if isinstance(row, list)]
        table = _normalize_cooking_header_spans({
            **table, "raw_rows": table.get("raw_rows", raw_rows), "rows": rows,
        })
        rows = table["rows"]
        header_match = _find_header_row(rows)
        source_prefix = (
            f"{provider}:{table.get('source') or 'table'}:"
            f"page{table.get('page_index')}:table{table.get('table_id') or ''}"
        )
        if header_match is None:
            issues.append({
                "code": "required_headers_unresolved",
                "message": "献立名と食材名の列見出しを特定できませんでした。",
                "source": source_prefix,
            })
            continue
        header_index, columns = header_match
        if columns["menu"] == columns["ingredient"]:
            combined_menus, combined_issues = _extract_combined_cooking_table(
                table=table,
                rows=rows,
                header_index=header_index,
                columns=columns,
                source_prefix=source_prefix,
            )
            issues.extend(combined_issues)
            for combined_index, menu in enumerate(combined_menus):
                for instruction_source in menu["cooking_instruction_sources"]:
                    instruction_source["source"] = f"{filename}:{instruction_source['source']}"
                grouping_name = menu["name"] or f"__unresolved__:{source_prefix}:{combined_index}"
                group_key = (menu.get("facility_name", ""), menu["service_date"], menu["raw_service_date"], grouping_name)
                existing = grouped.get(group_key)
                if existing is None:
                    grouped[group_key] = menu
                    continue
                existing["cooking_instructions"] = "\n".join(filter(None, [existing.get("cooking_instructions", ""), menu["cooking_instructions"]]))
                existing.setdefault("cooking_instruction_sources", []).extend(menu["cooking_instruction_sources"])
                existing_signatures = {
                    (item["name"], item["required_quantity"], item["unit"], item["process_text"])
                    for item in existing["ingredients"]
                }
                existing["ingredients"].extend(
                    item for item in menu["ingredients"]
                    if (item["name"], item["required_quantity"], item["unit"], item["process_text"]) not in existing_signatures
                )
            continue
        if columns["date"] is None:
            issues.append({"code": "date_header_missing", "message": "日付列を特定できませんでした。", "source": source_prefix})
        if columns["quantity"] is None:
            issues.append({"code": "quantity_header_missing", "message": "数量列を特定できませんでした。", "source": source_prefix})
        if columns["process"] is None:
            issues.append({"code": "process_header_missing", "message": "調理作業内容の列を特定できませんでした。", "source": source_prefix})

        for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            menu_name = _cell(row, columns["menu"]).strip()
            ingredient_name = _cell(row, columns["ingredient"]).strip()
            if not menu_name and not ingredient_name:
                continue
            row_source = f"{source_prefix}:row{row_index}"
            if not menu_name or not ingredient_name:
                issues.append({
                    "code": "menu_or_ingredient_missing",
                    "message": "献立名または食材名が空の行です。自動では補完しません。",
                    "source": row_source,
                    "raw_cells": row,
                })
                continue

            raw_date = _cell(row, columns["date"]).strip()
            service_date, date_issue = _normalize_explicit_date(raw_date)
            if date_issue:
                issues.append({
                    "code": date_issue,
                    "message": "提供日を年月日で確認してください。",
                    "source": row_source,
                    "raw_value": raw_date,
                })
            raw_quantity = _cell(row, columns["quantity"]).strip()
            quantity = _explicit_decimal_or_none(raw_quantity)
            if quantity is None or quantity < 0:
                issues.append({
                    "code": "quantity_unresolved",
                    "message": "使用量を確認してください。",
                    "source": row_source,
                    "raw_value": raw_quantity,
                })
                quantity = None
            unit = _cell(row, columns["unit"]).strip()
            if not unit:
                issues.append({"code": "unit_missing", "message": "単位を入力してください。", "source": row_source})

            process_text = _cell(row, columns["process"]).strip()
            location = _cell(row, columns["location"]).strip()
            scheduled_time_raw = _cell(row, columns["time"]).strip()
            scheduled_time = _explicit_time(scheduled_time_raw)
            if scheduled_time_raw and scheduled_time is None:
                issues.append({
                    "code": "scheduled_time_invalid",
                    "message": "開始時刻を確認してください。",
                    "source": row_source,
                    "raw_value": scheduled_time_raw,
                })
            work_zone_raw = _cell(row, columns["work_zone"]).strip()
            line_style_raw = _cell(row, columns["line_style"]).strip()
            process_steps: list[dict[str, Any]] = []
            if process_text and location:
                process_steps.append({
                    "sequence": 1,
                    "location": location,
                    "operation": process_text,
                    "scheduled_time": scheduled_time or None,
                    "work_zone": _explicit_work_zone(work_zone_raw),
                    "route_line_style": _explicit_line_style(line_style_raw),
                })
            elif process_text and not location:
                issues.append({
                    "code": "process_location_required",
                    "message": "調理作業の場所を確認してください。作業内容だけを保持しました。",
                    "source": row_source,
                })

            servings_raw = _cell(row, columns["servings"]).strip()
            servings_decimal = _explicit_decimal_or_none(servings_raw)
            servings = int(servings_decimal) if servings_decimal is not None and servings_decimal == int(servings_decimal) else None
            if servings is None or servings <= 0:
                issues.append({
                    "code": "servings_unresolved",
                    "message": "提供食数を確認してください。",
                    "source": row_source,
                    "raw_value": servings_raw,
                })
                servings = None

            signature = (
                table.get("facility_name", ""),
                service_date,
                raw_date,
                menu_name,
                ingredient_name,
                _decimal_string(quantity),
                unit,
                process_text,
                location,
            )
            if signature in seen_rows:
                continue
            seen_rows.add(signature)
            group_key = (table.get("facility_name", ""), service_date, raw_date, menu_name)
            menu = grouped.setdefault(
                group_key,
                {
                    "facility_name": table.get("facility_name", ""),
                    "service_date": service_date,
                    "raw_service_date": raw_date,
                    "name": menu_name,
                    "servings": servings,
                    "status": "needs_review",
                    "ingredients": [],
                },
            )
            if menu["servings"] is None and servings is not None:
                menu["servings"] = servings
            menu["ingredients"].append({
                "name": ingredient_name,
                "required_quantity": _decimal_string(quantity),
                "raw_quantity": raw_quantity,
                "unit": unit,
                "process_text": process_text,
                "process_steps": process_steps,
                "source": row_source,
                "status": "needs_review",
            })

    menus = list(grouped.values())
    if not menus:
        issues.append({
            "code": "no_cooking_rows",
            "message": "確認画面へ取り込める献立・食材行がありませんでした。",
            "source": filename,
        })
    return {
        "document_type": "cooking_arrangement",
        "filename": str(filename or "").strip(),
        "status": "needs_review",
        "provider": provider,
        "review_required": True,
        "menus": menus,
        "issues": issues,
    }


def extract_cooking_arrangement_candidates_from_tables(
    tables: list[dict[str, Any]],
    *,
    filename: str = "",
    provider: str = "structured_excel",
) -> dict[str, Any]:
    return _cooking_arrangement_from_tables(tables, filename=filename, provider=provider)


def extract_yomitoku_candidates(
    *,
    pdf_bytes: bytes,
    document_type: str,
    filename: str = "",
    dpi: int = 200,
    device: str = "cpu",
) -> dict[str, Any]:
    normalized_type = normalize_document_type(document_type)
    if not pdf_bytes:
        raise ValueError("pdf_bytes is required")
    pages = _run_yomitoku_pdf(pdf_bytes=pdf_bytes, dpi=dpi, device=device, orient_invoice=normalized_type == "invoice")
    tables = _tables_from_pages(pages)
    if normalized_type == "invoice":
        invoice_tables = _invoice_tables_with_line_items(tables)
        invoice_line_items = [item for table in invoice_tables for item in table["line_items"]]
        cost_items = invoice_line_items
        quantity_items: list[dict[str, Any]] = []
    elif normalized_type == "quantity_sheet":
        invoice_tables = []
        invoice_line_items = []
        cost_items = []
        quantity_items = _quantity_candidates_from_tables(tables)
    else:
        cooking_result = _cooking_arrangement_from_tables(
            tables,
            filename=filename,
            provider="yomitoku",
        )
        return {
            **cooking_result,
            "pages": [
                {
                    "page_index": page.page_index,
                    "markdown_text": page.markdown_text,
                    "tables": [table for table in tables if table.get("page_index") == page.page_index],
                }
                for page in pages
            ],
        }
    return {
        "document_type": normalized_type,
        "filename": str(filename or "").strip(),
        "status": "needs_review",
        "provider": "yomitoku",
        "review_required": True,
        "cost_items": cost_items,
        "invoice_tables": invoice_tables,
        "invoice_line_items": invoice_line_items,
        "quantity_items": quantity_items,
        "pages": [
            {
                "page_index": page.page_index,
                "markdown_text": page.markdown_text,
                "tables": [table for table in tables if table.get("page_index") == page.page_index],
                "orientation": page.orientation,
            }
            for page in pages
        ],
    }
