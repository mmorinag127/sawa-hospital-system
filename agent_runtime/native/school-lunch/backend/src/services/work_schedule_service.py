from __future__ import annotations

from copy import deepcopy

from datetime import date, time
from calendar import monthrange
import re
import unicodedata
from io import BytesIO
from math import ceil, floor
from typing import Any

from sqlalchemy import select, update
from sqlalchemy.orm import Session, selectinload, object_session
from sqlalchemy.orm.attributes import set_committed_value
from src.services.school_lunch_hygiene import inspect_hygiene, missing_task_times, require_task_times, task_assignees, ownership_rows

from src.models.school_lunch import (
    DailyWorkSchedule,
    Facility,
    MenuExecution,
    MenuIngredient,
    WorkScheduleAssignee,
    WorkScheduleRow,
    WorkScheduleTask,
)


class WorkScheduleNotFound(LookupError):
    pass


class WorkScheduleConflict(ValueError):
    pass


def _load_schedule(session: Session, schedule_id: int) -> DailyWorkSchedule:
    statement = (
        select(DailyWorkSchedule)
        .where(DailyWorkSchedule.id == schedule_id)
        .options(
            selectinload(DailyWorkSchedule.facility),
            selectinload(DailyWorkSchedule.rows).selectinload(WorkScheduleRow.assignees),
            selectinload(DailyWorkSchedule.rows).selectinload(WorkScheduleRow.tasks),
        )
    )
    schedule = session.execute(statement).scalar_one_or_none()
    if schedule is None:
        raise WorkScheduleNotFound(f"work schedule {schedule_id} was not found")
    return schedule


def _serialize_schedule(schedule: DailyWorkSchedule) -> dict[str, Any]:
    session = object_session(schedule)
    current = _menu_versions(session, schedule.facility_id, schedule.service_date)
    from src.services.process_work_schedule import saved_patterns, revision_map, missing_execution
    patterns = saved_patterns(session, schedule.facility_id, schedule.service_date)
    process_revisions = revision_map(patterns)
    result = {
        "id": schedule.id,
        "facility": {"id": schedule.facility.id, "name": schedule.facility.name},
        "service_date": schedule.service_date,
        "plating_start": schedule.plating_start,
        "plating_end": schedule.plating_end,
        "timing_profile_id": schedule.timing_profile_id,
        "timing_profile_version": schedule.timing_profile_version,
        "timeline_ticks": schedule.timeline_ticks,
        "delivery_assignees": schedule.delivery_assignees or [],
        "source_notes": schedule.source_notes or [],
        "status": schedule.status,
        "version": schedule.version,
        "time_profile": schedule.time_profile,
        "source_menu_versions": schedule.source_menu_versions,
        "source_process_revisions": schedule.source_process_revisions,
        "current_process_revisions": process_revisions,
        "pattern_missing": missing_execution(patterns,session),
        "current_menu_versions": current,
        "stale": schedule.source_menu_versions != current or (schedule.source_process_revisions or {}) != process_revisions,
        "rows": [
            {
                "id": row.id,
                "position": row.position,
                "menu_execution_id": row.menu_execution_id,
                "menu_version": row.menu_version,
                "menu_name": row.menu_name,
                "assignees": [assignee.name for assignee in row.assignees],
                "tasks": [
                    {
                        "id": task.id,
                        "sequence": task.sequence,
                        "assignees": task.assignees,
                        "start_time": task.start_time,
                        "end_time": task.end_time,
                        "duration_minutes": task.duration_minutes,
                        "offset_before_plating_minutes": task.offset_before_plating_minutes,
                        "operation": task.operation,
                        "work_zone": task.work_zone,
                        "ingredient_name": task.ingredient_name,
                        "phase": task.phase,
                        "kettle": task.kettle,
                        "kettle_use": task.kettle_use,
                        "uses_raw_meat": task.uses_raw_meat,
                    }
                    for task in row.tasks
                ],
            }
            for row in schedule.rows
        ],
    }
    try:
        result["hygiene"] = inspect_hygiene(result)
    except ValueError as exc:
        result["hygiene"] = {"missing": [str(exc)], "events": []}
    return result


def _menu_versions(session: Session, facility_id: int, service_date: date) -> dict[str, int]:
    return {str(row.id): row.version for row in session.execute(select(MenuExecution).where(
        MenuExecution.facility_id == facility_id, MenuExecution.service_date == service_date,
    )).scalars()}


def list_work_schedules(session: Session, facility_id=None, start=None, end=None) -> list[dict[str, Any]]:
    statement = (
        select(DailyWorkSchedule)
        .options(
            selectinload(DailyWorkSchedule.facility),
            selectinload(DailyWorkSchedule.rows).selectinload(WorkScheduleRow.assignees),
            selectinload(DailyWorkSchedule.rows).selectinload(WorkScheduleRow.tasks),
        )
        .order_by(DailyWorkSchedule.service_date.desc(), DailyWorkSchedule.id.desc())
    )
    if facility_id is not None:
        statement = statement.where(DailyWorkSchedule.facility_id == facility_id)
    if start is not None:
        statement = statement.where(DailyWorkSchedule.service_date >= start, DailyWorkSchedule.service_date < end)
    return [_serialize_schedule(schedule) for schedule in session.execute(statement).scalars().all()]


def get_work_schedule(session: Session, schedule_id: int) -> dict[str, Any]:
    return _serialize_schedule(_load_schedule(session, schedule_id))


def build_work_schedule_draft(session: Session, facility_name: str, service_date: date, refresh_sources: bool = False, *, rebuild: bool = False, source_import: bool = False) -> dict[str, Any]:
    facility = session.execute(select(Facility).where(Facility.name == facility_name)).scalar_one_or_none()
    if facility is None:
        raise WorkScheduleNotFound("the selected facility was not found")

    existing = session.execute(
        select(DailyWorkSchedule.id).where(
            DailyWorkSchedule.facility_id == facility.id,
            DailyWorkSchedule.service_date == service_date,
        )
    ).scalar_one_or_none()
    if existing is not None and not (refresh_sources or rebuild or source_import):
        return get_work_schedule(session, existing)

    menus = session.execute(
        select(MenuExecution)
        .where(MenuExecution.facility_id == facility.id, MenuExecution.service_date == service_date)
        .options(
            selectinload(MenuExecution.assignees),
            selectinload(MenuExecution.menu_ingredients).selectinload(MenuIngredient.ingredient),
            selectinload(MenuExecution.menu_ingredients).selectinload(MenuIngredient.process_steps),
        )
        .order_by(MenuExecution.id)
    ).scalars().all()
    if not menus and not source_import:
        raise WorkScheduleNotFound("the selected date has no saved menus")

    rows: list[dict[str, Any]] = []
    for row_position, menu in enumerate(menus, start=1):
        tasks: list[dict[str, Any]] = []
        for ingredient_row in menu.menu_ingredients:
            if ingredient_row.process_steps:
                for step in ingredient_row.process_steps:
                    tasks.append({
                        "id": None,
                        "sequence": len(tasks) + 1,
                        "start_time": step.scheduled_time,
                        "end_time": None,
                        "operation": step.operation,
                        "work_zone": step.work_zone,
                        "ingredient_name": ingredient_row.ingredient.name,
                    })
            elif ingredient_row.process_text.strip():
                tasks.append({
                    "id": None,
                    "sequence": len(tasks) + 1,
                    "start_time": None,
                    "end_time": None,
                    "operation": ingredient_row.process_text.strip(),
                    "work_zone": None,
                    "ingredient_name": ingredient_row.ingredient.name,
                })
        if not tasks:
            tasks.append({
                "id": None,
                "sequence": 1,
                "start_time": None,
                "end_time": None,
                "operation": "",
                "work_zone": None,
                "ingredient_name": "",
            })
        rows.append({
            "id": None,
            "position": row_position,
            "menu_execution_id": menu.id,
            "menu_version": menu.version,
            "menu_name": menu.name,
            "assignees": [assignee.name for assignee in menu.assignees],
            "tasks": tasks,
        })

    from src.services.process_work_schedule import saved_patterns, project_tasks, revision_map, missing_execution
    patterns = saved_patterns(session, facility.id, service_date)
    from src.services.shared_menu_review import school_direct_document
    rows=[row for row in rows if not (str(row['menu_execution_id']) in patterns
        and school_direct_document(session,patterns[str(row['menu_execution_id'])][1].document))]
    previous = get_work_schedule(session, existing) if existing is not None else {}
    for row in rows:
        for task in row["tasks"]:
            task["assignees"] = []
        pattern = patterns.get(str(row['menu_execution_id']))
        if pattern and not source_import:
            row['tasks'] = project_tasks(pattern[1].document)
    return {
        "id": previous.get("id"),
        "facility": {"id": facility.id, "name": facility.name},
        "service_date": service_date,
        "plating_start": previous.get("plating_start"),
        "plating_end": previous.get("plating_end"),
        "timeline_ticks": previous.get("timeline_ticks"),
        "timing_profile_id": previous.get("timing_profile_id"),
        "timing_profile_version": previous.get("timing_profile_version"),
        "delivery_assignees": previous.get("delivery_assignees", []),
        "source_notes": [] if source_import else previous.get("source_notes", []),
        "status": "draft",
        "version": _load_schedule(session, existing).version if existing else 1,
        "time_profile": previous.get("time_profile"),
        "source_menu_versions": _menu_versions(session, facility.id, service_date),
        "source_process_revisions": revision_map(patterns) if not source_import else None,
        "pattern_missing": missing_execution(patterns,session),
        "current_menu_versions": _menu_versions(session, facility.id, service_date),
        "stale": False,
        "rows": rows,
    }


def _validate_payload(session: Session, payload: dict[str, Any], *, facility: Facility, existing: DailyWorkSchedule | None = None) -> None:
    from src.services.process_work_schedule import saved_patterns, assert_pattern_rows, missing_execution
    patterns = saved_patterns(session, facility.id, payload['service_date'])
    try:
        assert_pattern_rows(payload, patterns,session)
    except ValueError as exc:
        raise WorkScheduleConflict(str(exc)) from exc
    if payload.get('status') != 'draft' and missing_execution(patterns,session):
        raise WorkScheduleConflict('工程グラフの実績が未入力です。下書きとして保存し、実績を入力してください。')
    if payload.get("status") != "draft" and missing_task_times(payload):
        raise WorkScheduleConflict("時刻未入力の工程表は下書き（draft）として保存してください。")
    delivery = payload.get("delivery_assignees", [])
    if len(delivery) != len(set(delivery)):
        raise WorkScheduleConflict("配送担当者が重複しています。")
    profile_id, profile_version = payload.get("timing_profile_id"), payload.get("timing_profile_version")
    ticks = payload.get("timeline_ticks")
    if (profile_id is None) != (profile_version is None):
        raise WorkScheduleConflict("時間設定のIDと版を両方指定してください。")
    if ticks is not None and (len(ticks) < 2 or any(not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", tick) for tick in ticks)
                              or any(a >= b for a, b in zip(ticks, ticks[1:]))):
        raise WorkScheduleConflict("時刻目盛は同日のHH:MMを2個以上、重複なく昇順で入力してください。")
    if profile_id is not None:
        from src.models.school_lunch_timing import TimingProfile
        # Saved documents keep the version/ticks used at creation even after settings change.
        preserved = existing is not None and (profile_id, profile_version, ticks) == (
            existing.timing_profile_id, existing.timing_profile_version, existing.timeline_ticks)
        if not preserved:
            profile = session.get(TimingProfile, profile_id)
            if profile is None or profile.facility_id != facility.id:
                raise WorkScheduleConflict("この学校の時間設定が見つかりません。")
            if profile.version != profile_version or profile.timeline_ticks != ticks:
                raise WorkScheduleConflict("時間設定が更新されています。再読み込みして適用してください。")
    if payload.get("time_profile") not in {None, "normal", "early"}:
        raise WorkScheduleConflict("時間区分は通常・早いから選択してください")
    if payload.get("source_menu_versions") is not None and payload["source_menu_versions"] != _menu_versions(session, facility.id, payload["service_date"]):
        raise WorkScheduleConflict("献立が変更されています。最新の献立を読み込み、工程を確認してください")
    if payload["plating_end"] <= payload["plating_start"]:
        raise WorkScheduleConflict("plating end time must be later than plating start time")
    rows = payload["rows"]
    positions = [row["position"] for row in rows]
    if len(positions) != len(set(positions)):
        raise WorkScheduleConflict("work schedule row positions must be unique")
    menu_names = [row["menu_name"].strip() for row in rows]
    if len(menu_names) != len(set(menu_names)):
        raise WorkScheduleConflict("menu names must be unique within a work schedule")

    assignments: dict[str, list[tuple[time, time, str]]] = {}
    for row in rows:
        assignees = row.get("assignees", [])
        if len(assignees) != len(set(assignees)):
            raise WorkScheduleConflict(f"assignees must be unique: {row['menu_name']}")
        menu_execution_id = row.get("menu_execution_id")
        if menu_execution_id is not None:
            menu = session.get(MenuExecution, menu_execution_id)
            if menu is None:
                raise WorkScheduleConflict(f"linked menu was not found: {row['menu_name']}")
            if menu.facility_id != facility.id or menu.service_date != payload["service_date"]:
                raise WorkScheduleConflict(f"linked menu does not match the facility and date: {row['menu_name']}")
            if row.get("menu_version") is not None and row["menu_version"] != menu.version:
                raise WorkScheduleConflict(f"献立が変更されています。献立から下書きを作り直してください: {row['menu_name']}")
        tasks = row["tasks"]
        if not tasks:
            raise WorkScheduleConflict(f"at least one work task is required: {row['menu_name']}")
        sequences = [task["sequence"] for task in tasks]
        if len(sequences) != len(set(sequences)):
            raise WorkScheduleConflict(f"work task sequences must be unique: {row['menu_name']}")
        for task in tasks:
            duration, offset = task.get("duration_minutes"), task.get("offset_before_plating_minutes")
            if (duration is None) != (offset is None):
                raise WorkScheduleConflict("逆算する作業は所要分数と盛付開始何分前に終えるかを両方入力してください。")
            complete = task.get("start_time") is not None and task.get("end_time") is not None
            if complete and task["end_time"] <= task["start_time"]:
                raise WorkScheduleConflict(f"work task end time must be later than start time: {row['menu_name']}")
            owners = task_assignees(row, task)
            if len(owners) != len(set(owners)):
                raise WorkScheduleConflict(f"assignees must be unique: {row['menu_name']}・{task['operation']}")
            for assignee in owners:
                normalized_assignee = assignee.strip()
                if not normalized_assignee:
                    raise WorkScheduleConflict(f"assignee name must not be blank: {row['menu_name']}")
                if complete:
                    assignments.setdefault(normalized_assignee, []).append(
                        (task["start_time"], task["end_time"], row["menu_name"])
                    )

    for assignee, intervals in assignments.items():
        ordered = sorted(intervals, key=lambda item: (item[0], item[1], item[2]))
        for previous, current in zip(ordered, ordered[1:]):
            if current[0] < previous[1]:
                raise WorkScheduleConflict(
                    f"assignee has overlapping work tasks: {assignee} "
                    f"({previous[2]} and {current[2]})"
                )
    try:
        inspect_hygiene(payload, require_complete=payload.get("status") != "draft")
    except ValueError as exc:
        raise WorkScheduleConflict(str(exc)) from exc


def _replace_rows(session: Session, schedule: DailyWorkSchedule, rows: list[dict[str, Any]]) -> None:
    schedule.rows.clear()
    session.flush()
    for row_data in sorted(rows, key=lambda item: item["position"]):
        row = WorkScheduleRow(
            schedule=schedule,
            menu_execution_id=row_data.get("menu_execution_id"),
            position=row_data["position"],
            menu_name=row_data["menu_name"].strip(),
            menu_version=row_data.get("menu_version"),
        )
        session.add(row)
        session.flush()
        for position, name in enumerate(row_data.get("assignees", []), start=1):
            session.add(WorkScheduleAssignee(row=row, position=position, name=name.strip()))
        for task_data in sorted(row_data["tasks"], key=lambda item: item["sequence"]):
            session.add(WorkScheduleTask(
                row=row,
                sequence=task_data["sequence"],
                assignees=task_data.get("assignees"),
                start_time=task_data["start_time"],
                end_time=task_data["end_time"],
                duration_minutes=task_data.get("duration_minutes"),
                offset_before_plating_minutes=task_data.get("offset_before_plating_minutes"),
                operation=task_data["operation"].strip(),
                work_zone=task_data["work_zone"],
                ingredient_name=task_data.get("ingredient_name", "").strip(),
                phase=task_data.get("phase"), kettle=task_data.get("kettle"),
                kettle_use=task_data.get("kettle_use"), uses_raw_meat=task_data.get("uses_raw_meat"),
            ))


def create_work_schedule(session: Session, payload: dict[str, Any]) -> dict[str, Any]:
    facility = session.execute(select(Facility).where(Facility.name == payload["facility_name"])).scalar_one_or_none()
    if facility is None:
        raise WorkScheduleConflict("the selected facility was not found")
    duplicate = session.execute(
        select(DailyWorkSchedule.id).where(
            DailyWorkSchedule.facility_id == facility.id,
            DailyWorkSchedule.service_date == payload["service_date"],
        )
    ).scalar_one_or_none()
    if duplicate is not None:
        raise WorkScheduleConflict("the work schedule already exists")
    _validate_payload(session, payload, facility=facility)
    schedule = DailyWorkSchedule(
        facility=facility,
        service_date=payload["service_date"],
        plating_start=payload["plating_start"],
        plating_end=payload["plating_end"],
        delivery_assignees=payload.get("delivery_assignees", []),
        source_notes=payload.get("source_notes", []),
        timing_profile_id=payload.get("timing_profile_id"),
        timing_profile_version=payload.get("timing_profile_version"),
        timeline_ticks=payload.get("timeline_ticks"),
        status=payload["status"],
        version=1,
        time_profile=payload.get("time_profile"),
        source_menu_versions=_menu_versions(session, facility.id, payload["service_date"]),
        source_process_revisions=payload.get('source_process_revisions'),
    )
    session.add(schedule)
    session.flush()
    _replace_rows(session, schedule, payload["rows"])
    session.flush()
    return get_work_schedule(session, schedule.id)


def update_work_schedule(session: Session, schedule_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    schedule = _load_schedule(session, schedule_id)
    if payload["version"] != schedule.version:
        raise WorkScheduleConflict("the work schedule was changed by another operation")
    result = session.execute(update(DailyWorkSchedule).where(
        DailyWorkSchedule.id == schedule_id, DailyWorkSchedule.version == payload["version"]
    ).values(version=payload["version"] + 1).execution_options(synchronize_session=False))
    if result.rowcount != 1:
        raise WorkScheduleConflict("the work schedule was changed by another operation")
    set_committed_value(schedule, "version", payload["version"] + 1)
    if any(task.assignees is not None for row in schedule.rows for task in row.tasks) and any(
        "assignees" not in task for row in payload["rows"] for task in row["tasks"]
    ):
        raise WorkScheduleConflict("作業別担当の指定が送信されていません。画面を再読み込みし、担当を確認してください。保存済みの担当は変更していません。")
    if payload["facility_name"] != schedule.facility.name or payload["service_date"] != schedule.service_date:
        raise WorkScheduleConflict("facility and service date cannot be changed after saving")
    current_versions = _menu_versions(session, schedule.facility_id, schedule.service_date)
    if schedule.source_menu_versions != current_versions and payload.get("source_menu_versions") != current_versions:
        raise WorkScheduleConflict("献立変更後の工程を確認してから保存してください")
    _validate_payload(session, payload, facility=schedule.facility, existing=schedule)
    schedule.plating_start = payload["plating_start"]
    schedule.plating_end = payload["plating_end"]
    schedule.delivery_assignees = payload.get("delivery_assignees", [])
    if "source_notes" in payload:
        schedule.source_notes = payload["source_notes"]
    schedule.timing_profile_id = payload.get("timing_profile_id")
    schedule.timing_profile_version = payload.get("timing_profile_version")
    schedule.timeline_ticks = payload.get("timeline_ticks")
    schedule.status = payload["status"]
    schedule.time_profile = payload.get("time_profile")
    schedule.source_menu_versions = current_versions
    schedule.source_process_revisions = payload.get('source_process_revisions')
    _replace_rows(session, schedule, payload["rows"])
    session.flush()
    return get_work_schedule(session, schedule.id)


def _minutes(value: time) -> int:
    return value.hour * 60 + value.minute


def reuse_work_schedule_draft(session: Session, payload: dict[str, Any]) -> dict[str, Any]:
    from copy import deepcopy
    from src.services import school_lunch_common_service, school_lunch_timing_service

    source = get_work_schedule(session, payload["source_schedule_id"])
    if source.get('source_process_revisions'):
        raise WorkScheduleConflict('パターン工程は対象日の工程グラフで再利用し、当日の設備・担当・時刻を入力してください。実績のコピーは行いません。')
    if source["version"] != payload["source_version"]:
        raise WorkScheduleConflict("元の工程表が変更されています。再読み込みしてください。")
    if source["stale"]:
        raise WorkScheduleConflict("元献立が変更されています。元工程表を確認してください。")
    require_task_times_for_output(source)
    target = build_work_schedule_draft(session, payload["facility_name"], payload["service_date"])
    if target["id"] is not None:
        raise WorkScheduleConflict("移行先には工程表が存在します。上書きは行いません。")

    def content(value):
        if isinstance(value, dict):
            return {key: content(item) for key, item in value.items() if key != "id"}
        if isinstance(value, list):
            return [content(item) for item in value]
        return value

    def signature(menu_id):
        menu = school_lunch_common_service.get_menu_execution(session, menu_id)
        return menu, content({key: menu[key] for key in ("name", "servings", "ingredients")})

    targets = [(row, signature(row["menu_execution_id"])[1]) for row in target["rows"]]
    rows, used, timed = [], set(), []
    for original in source["rows"]:
        row = deepcopy(original)
        row["id"] = None
        if original["menu_execution_id"] is None:
            if not payload.get("include_manual_rows", False):
                raise WorkScheduleConflict("献立に紐付かない手入力行があります。手入力行の再利用を明示的に選択してください。")
            row["menu_version"] = None
        else:
            if original.get("menu_version") is None:
                raise WorkScheduleConflict("元献立と版が未確認の行は再利用できません。")
            menu, expected = signature(original["menu_execution_id"])
            if original["menu_version"] != menu["version"]:
                raise WorkScheduleConflict("元献立が変更されています。元工程表を確認してください。")
            matches = [candidate for candidate, actual in targets if actual == expected]
            if len(matches) != 1 or matches[0]["menu_execution_id"] in used:
                raise WorkScheduleConflict(f"移行先の献立・食材が完全一致で一意に対応しません: {original['menu_name']}")
            matched = matches[0]
            used.add(matched["menu_execution_id"])
            row.update(menu_execution_id=matched["menu_execution_id"], menu_version=matched["menu_version"])
        for task in row["tasks"]:
            if not task["start_time"] or not task["end_time"] or not source["plating_start"]:
                raise WorkScheduleConflict("元工程表の開始・終了・盛付開始時刻を入力してください。")
            if any(value.second or value.microsecond for value in (task["start_time"], task["end_time"], source["plating_start"])):
                raise WorkScheduleConflict("分未満の時刻がある工程表は再利用できません。時刻を確認してください。")
            duration = _minutes(task["end_time"]) - _minutes(task["start_time"])
            offset = _minutes(source["plating_start"]) - _minutes(task["end_time"])
            if duration <= 0:
                raise WorkScheduleConflict("元工程表の開始・終了時刻を確認してください。日跨ぎは再利用できません。")
            task.update(duration_minutes=duration, offset_before_plating_minutes=offset)
            task["id"] = None
            timed.append({"key": str(len(timed)), "duration_minutes": duration, "offset_before_plating_minutes": offset})
        rows.append(row)
    if len(used) != len(targets):
        raise WorkScheduleConflict("移行先に元工程表と対応しない献立があります。")
    try:
        calculated = school_lunch_timing_service.calculate(session, {**payload, "tasks": timed})
    except (ValueError, LookupError) as exc:
        raise WorkScheduleConflict(str(exc)) from exc
    for task, result in zip((task for row in rows for task in row["tasks"]), calculated["tasks"]):
        task.update(start_time=time.fromisoformat(result["start_time"]), end_time=time.fromisoformat(result["end_time"]))
    profile = calculated["profile"]
    target.update(rows=rows, plating_start=profile["plating_start"], plating_end=profile["plating_end"],
                  timing_profile_id=profile["id"], timing_profile_version=profile["version"],
                  timeline_ticks=profile["timeline_ticks"], delivery_assignees=deepcopy(source["delivery_assignees"]))
    _validate_payload(session, target, facility=session.get(Facility, target["facility"]["id"]))
    target["hygiene"] = inspect_hygiene(target)
    return target


def require_task_times_for_output(schedule):
    if schedule.get('pattern_missing'):
        raise WorkScheduleConflict('工程グラフの実績入力を確認してください: ' + '、'.join(schedule['pattern_missing']))
    try:
        require_task_times(schedule)
    except ValueError as exc:
        raise WorkScheduleConflict(str(exc)) from exc
    if any(row.get("menu_execution_id") is not None and row.get("menu_version") is None for row in schedule["rows"]):
        raise WorkScheduleConflict("元献立の版が未確認です。献立から下書きを作り直してください。")


def validate_work_schedule_export(session: Session, schedule: dict[str, Any]) -> None:
    require_task_times_for_output(schedule)
    if schedule.get("plating_start") is None or schedule.get("plating_end") is None:
        raise WorkScheduleConflict("盛付開始・終了時刻を入力してから出力してください。")
    facility = session.get(Facility, schedule["facility"]["id"])
    if facility is None:
        raise WorkScheduleConflict("工程表の学校が見つかりません。")
    _validate_payload(
        session, {**schedule, "facility_name": facility.name}, facility=facility,
        existing=_load_schedule(session, schedule["id"]),
    )


def generate_work_schedule_workbook(schedule: dict[str, Any], *, original: bool = False) -> bytes:
    from openpyxl import Workbook

    require_task_times_for_output(schedule)
    if schedule.get("stale"):
        raise WorkScheduleConflict("献立変更後の工程表を確認・保存してから出力してください")
    if original:
        from src.services.work_schedule_template_export import generate_original_work_schedule
        try:
            return generate_original_work_schedule(schedule)
        except ValueError as exc:
            raise WorkScheduleConflict(str(exc)) from exc
    workbook = Workbook()
    _populate_work_schedule_sheet(workbook.active, schedule)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def begin_schedule_export(session):
    from src.services.route_export_snapshot import begin_export_read
    if session.get_transaction() is None and session.identity_map:
        raise WorkScheduleConflict("出力は空のidentity mapを持つclean Sessionで開始してください")
    begin_export_read(session)


def check_schedule_export_sources(schedules, expected_sources):
    if expected_sources is None:
        return
    if len(expected_sources) != len(schedules) or len({source["id"] for source in expected_sources}) != len(expected_sources):
        raise WorkScheduleConflict("工程表の対象集合が変更されています")
    by_id = {source["id"]: source for source in expected_sources}
    for schedule in schedules:
        actual = {"id": schedule["id"], "version": schedule["version"], "facility_id": schedule["facility"]["id"], "service_date": schedule["service_date"]}
        if by_id.get(schedule["id"]) != actual:
            raise WorkScheduleConflict("工程表の学校・提供日・版が変更されています")


def generate_monthly_work_schedule_workbook(session: Session, facility_name: str, month: str, *, original: bool = False, expected_sources=None) -> bytes:
    from openpyxl import Workbook

    begin_schedule_export(session)
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise WorkScheduleConflict("対象月をYYYY-MMで指定してください")
    start = date.fromisoformat(month + "-01")
    end = start.replace(day=monthrange(start.year, start.month)[1])
    facility = session.execute(select(Facility).where(Facility.name == facility_name)).scalar_one_or_none()
    if facility is None:
        raise WorkScheduleNotFound("対象の学校がありません")
    dates = set(session.execute(select(MenuExecution.service_date).where(
        MenuExecution.facility_id == facility.id, MenuExecution.service_date.between(start, end),
    )).scalars())
    schedules = [item for item in list_work_schedules(session) if item["facility"]["id"] == facility.id and start <= item["service_date"] <= end]
    schedules = deepcopy(schedules)
    check_schedule_export_sources(schedules, expected_sources)
    missing = dates - {item["service_date"] for item in schedules}
    if missing:
        raise WorkScheduleConflict("工程表が未作成です: " + ", ".join(str(day) for day in sorted(missing)))
    if not schedules:
        raise WorkScheduleNotFound("対象月の工程表がありません")
    for schedule in schedules:
        validate_work_schedule_export(session, schedule)
    if original:
        from src.services.work_schedule_template_export import generate_monthly_original_work_schedule
        for schedule in schedules:
            if schedule["stale"]:
                raise WorkScheduleConflict(f"{schedule['service_date']}の献立変更を工程表に反映してください")
        try:
            return generate_monthly_original_work_schedule(sorted(schedules, key=lambda item: item["service_date"]))
        except ValueError as exc:
            raise WorkScheduleConflict(str(exc)) from exc
    workbook = Workbook()
    workbook.remove(workbook.active)
    for schedule in sorted(schedules, key=lambda item: item["service_date"]):
        if schedule["stale"]:
            raise WorkScheduleConflict(f"{schedule['service_date']}の献立変更を工程表に反映してください")
        _populate_work_schedule_sheet(workbook.create_sheet(), schedule)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _wrap_schedule_text(text: str, column_width: float, font_size: float) -> list[str]:
    """Conservative printed-width estimate; CJK characters occupy two units."""
    capacity = max(1, floor((column_width * 7 * 0.75 - 8) / (font_size * 0.55)))
    lines = []
    for paragraph in text.split("\n"):
        line = ""
        width = 0
        for character in paragraph:
            units = 0 if unicodedata.combining(character) else (2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1)
            if line and width + units > capacity:
                lines.append(line)
                line, width = "", 0
            line += character
            width += units
        lines.append(line)
    return lines


def _work_schedule_lanes(tasks, position, width_points, extra_labels=None):
    """Pack each ingredient's sequential work horizontally, never merge work.

    Text occupancy and time occupancy are separate: widening a short label
    must not widen the duration band. True overlaps get distinct lanes.
    """
    result = []
    # Prefer an existing ingredient lane, but a changed material set is not
    # itself a reason to allocate a physical row in the same menu workflow.
    for index, task in sorted(enumerate(tasks), key=lambda item: (_minutes(item[1]["start_time"]), item[0])):
        start, end = _minutes(task["start_time"]), _minutes(task["end_time"])
        if start >= end:
            raise WorkScheduleConflict("工程の開始・終了時刻を確認してください")
        left, right = position(start), position(end)
        label = (f"{task.get('ingredient_name') + '：' if task.get('ingredient_name') else ''}"
                 f"{task['operation']}\n{start//60:02d}:{start%60:02d}～{end//60:02d}:{end%60:02d}")
        if task.get("kettle"):
            label += f" 釜{task['kettle']}-{task['kettle_use']}"
        if extra_labels and index in extra_labels:
            label += "\n" + extra_labels[index]
        # Long labels use more horizontal space before more physical rows.
        units = max(sum(2 if unicodedata.east_asian_width(c) in {"W", "F"} else 1 for c in line)
                    for line in label.splitlines())
        span = min(1, max(right-left, 82 / width_points, units * 5 / 4 / width_points))
        text_left = max(0, min(left, 1-span))
        text_right = text_left + span
        lines = _wrap_schedule_text(label, span * width_points / 5.25, 10)
        for offset in range(0, len(lines), 20):
            chunk = lines[offset:offset+20]
            entry = {"task": task, "index": index, "start": left, "end": right,
                     "text_start": text_left, "text_end": text_right,
                     "text": "\n".join(chunk), "duration": offset == 0}
            compatible = [lane for lane in result if all(
                text_right <= other["text_start"] + 1e-8 or text_left >= other["text_end"] - 1e-8
                for other in lane["entries"])]
            lane = next((lane for lane in compatible if any(
                other["task"].get("ingredient_name", "") == task.get("ingredient_name", "")
                for other in lane["entries"])), compatible[0] if compatible else None)
            if lane is None:
                lane = {"entries": [], "height": 42}
                result.append(lane)
            lane["entries"].append(entry)
            lane["height"] = max(lane["height"], len(chunk) * 14 + 16)
    return result


def _populate_work_schedule_sheet(sheet, schedule: dict[str, Any]) -> None:
    from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    # Every standard export (daily/monthly, Excel/PDF) must use the same
    # complete hygiene check as the original-format renderer.
    try:
        inspect_hygiene(schedule, require_complete=True)
    except ValueError as exc:
        raise WorkScheduleConflict(str(exc)) from exc
    sheet.title = str(schedule["service_date"].day)
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.2, right=0.2, top=0.3, bottom=0.3, header=0.1, footer=0.1)

    task_start_minutes = [_minutes(task["start_time"]) for row in schedule["rows"] for task in row["tasks"]]
    task_end_minutes = [_minutes(task["end_time"]) for row in schedule["rows"] for task in row["tasks"]]
    explicit_ticks = schedule.get("timeline_ticks")
    grid_start = _minutes(time.fromisoformat(explicit_ticks[0])) if explicit_ticks else floor(min(task_start_minutes) / 30) * 30
    grid_end = _minutes(time.fromisoformat(explicit_ticks[-1])) if explicit_ticks else ceil(max(task_end_minutes) / 30) * 30
    if any(value < grid_start for value in task_start_minutes) or any(value > grid_end for value in task_end_minutes):
        raise WorkScheduleConflict("作業時刻が時間目盛の範囲外です")
    timeline = [_minutes(time.fromisoformat(tick)) for tick in explicit_ticks] if explicit_ticks else list(range(grid_start, grid_end + 1, 30))
    minute_columns = grid_end - grid_start
    last_col = 2 + minute_columns
    axis_width = 140
    minute_width = axis_width / minute_columns
    middle_col = 3 + minute_columns // 2
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = sheet.cell(1, 1, "調理作業工程表")
    title.font = Font(name="Yu Gothic", size=18, bold=True)
    title.alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 24
    sheet.row_dimensions[3].height = 20
    sheet.row_dimensions[4].height = 8
    sheet.row_dimensions[5].height = 24
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=middle_col-1)
    sheet.cell(2, 1, f"{schedule['service_date']:%Y年%m月%d日}  {schedule['facility']['name']}")
    sheet.merge_cells(start_row=2, start_column=middle_col, end_row=2, end_column=last_col)
    profile = {"normal": "通常", "early": "早い"}.get(schedule.get("time_profile"), "未選択")
    sheet.cell(2, middle_col, f"{profile} / 盛り付け時間 {schedule['plating_start']:%H:%M}〜{schedule['plating_end']:%H:%M}")
    # Match the supplied July workbook: Office accent6 (F79646), light 40%.
    zone_fills = {
        "contaminated_meat_fish": PatternFill("solid", fgColor="FFFFFF99"),
        "contaminated": PatternFill("solid", fgColor=Color(theme=9, tint=0.3999450666829432)),
        "non_contaminated": PatternFill("solid", fgColor="FFCCFFFF"),
    }
    legend = [("contaminated_meat_fish", "■ 汚染作業（肉・魚下処理）"), ("contaminated", "■ 汚染作業（野菜等）"), ("non_contaminated", "■ 非汚染作業")]
    for index, (zone, label) in enumerate(legend):
        start = 1 + index * last_col // 3
        end = (index + 1) * last_col // 3
        sheet.merge_cells(start_row=3, start_column=start, end_row=3, end_column=end)
        sheet.cell(3, start, label).fill = zone_fills[zone]

    headers = ["献立名", "担当者"]
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(5, column, label)
        cell.font = Font(name="Yu Gothic", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill("solid", fgColor="E2E8F0")
    for start, end in zip(timeline, timeline[1:]):
        column = 3 + start - grid_start
        sheet.merge_cells(start_row=5, start_column=column, end_row=5, end_column=2 + end - grid_start)
        cell = sheet.cell(5, column, f"{start//60:02d}:{start%60:02d}～{end//60:02d}:{end%60:02d}")
        cell.font = Font(name="Yu Gothic", size=10)
        cell.alignment = Alignment(horizontal="center")

    thin = Side(style="thin", color="64748B")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    current_row = 6
    label_groups = []
    label_references = []
    for schedule_row in ownership_rows(schedule["rows"]):
        tasks = schedule_row["tasks"]
        row_start = current_row
        plan = _work_schedule_lanes(tasks, lambda value: (value-grid_start)/minute_columns, axis_width * 5.25)
        for lane in plan:
            sheet.row_dimensions[current_row].height = lane["height"] - 8
            sheet.row_dimensions[current_row + 1].height = 8
            for entry in lane["entries"]:
                start_column = 3 + round(entry["text_start"] * minute_columns)
                end_column = min(last_col, 2 + round(entry["text_end"] * minute_columns))
                if end_column > start_column:
                    sheet.merge_cells(start_row=current_row, start_column=start_column, end_row=current_row, end_column=end_column)
                cell = sheet.cell(current_row, start_column, entry["text"])
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = Font(name="Yu Gothic", size=9)
                if entry["duration"]:
                    for column in range(3 + round(entry["start"] * minute_columns), 3 + round(entry["end"] * minute_columns)):
                        sheet.cell(current_row + 1, column).fill = zone_fills[entry["task"]["work_zone"]]
            current_row += 2
        if current_row == row_start:
            sheet.row_dimensions[current_row].height = 42
            current_row += 1
        menu_lines = _wrap_schedule_text(schedule_row["menu_name"], 23, 10)
        assignee_lines = _wrap_schedule_text("\n".join(schedule_row["assignees"]), 18, 10)
        if len(menu_lines) > 22:
            reference = f"M{len(label_references) + 1}"
            label_references.append((reference, "献立名全文", schedule_row["menu_name"]))
            menu_lines = _wrap_schedule_text(f"献立名 [{reference}]\n末尾の共通一覧参照", 23, 10)
        if len(assignee_lines) > 22:
            reference = f"A{len(label_references) + 1}"
            label_references.append((reference, "担当者一覧", "\n".join(schedule_row["assignees"])))
            assignee_lines = _wrap_schedule_text(f"担当 [{reference}]\n全員・末尾一覧参照", 18, 10)
        label_height = max(len(menu_lines), len(assignee_lines)) * 15 + 8
        total_height = sum(sheet.row_dimensions[row].height for row in range(row_start, current_row))
        remaining_height = max(0, label_height - total_height)
        for row in range(row_start, current_row):
            extra = min(remaining_height, 344 - sheet.row_dimensions[row].height)
            sheet.row_dimensions[row].height += extra
            remaining_height -= extra
        while remaining_height > 0:
            height = min(344, max(42, remaining_height))
            sheet.row_dimensions[current_row].height = height
            remaining_height -= height
            current_row += 1
        row_end = current_row - 1
        label_groups.append((row_start, row_end, menu_lines, assignee_lines))

    # Merged vertical labels spanning pages are printed only around their
    # midpoint. Keep the compact original grouping on short sheets, but repeat
    # ownership on every physical row of a multi-page sheet.
    repeat_labels = bool(label_references) or sum(sheet.row_dimensions[row].height for row in range(6, current_row)) > 600
    for row_start, row_end, menu_lines, assignee_lines in label_groups:
        if not repeat_labels and row_end > row_start:
            sheet.merge_cells(start_row=row_start, start_column=1, end_row=row_end, end_column=1)
            sheet.merge_cells(start_row=row_start, start_column=2, end_row=row_end, end_column=2)
        label_rows = range(row_start, row_end + 1) if repeat_labels else (row_start,)
        for row in label_rows:
            displayed_lines = []
            for column, lines in ((1, menu_lines), (2, assignee_lines)):
                displayed_lines.append(len(lines))
                cell = sheet.cell(row, column, "\n".join(lines))
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = Font(name="Yu Gothic", size=10)
            if repeat_labels:
                sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height, max(displayed_lines) * 15 + 8)

    # A long ownership label is one complete set, never a different subset on
    # each task. Every task uses the same reference; the common list retains
    # the full text and marks any continuation as part of that same reference.
    for reference, title, text in label_references:
        lines = _wrap_schedule_text(text, axis_width * 0.8, 10)
        chunks = [lines[index:index + 20] for index in range(0, len(lines), 20)]
        for index, chunk in enumerate(chunks, start=1):
            caption = f"[{reference}] {title}\nこの番号の全工程に共通"
            if title == "担当者一覧":
                caption += "（全員）"
            caption += f"\n一覧 {index}/{len(chunks)}"
            caption_lines = _wrap_schedule_text(caption, 41, 10)
            sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
            sheet.merge_cells(start_row=current_row, start_column=3, end_row=current_row, end_column=last_col)
            for column, value in ((1, "\n".join(caption_lines)), (3, "\n".join(chunk))):
                cell = sheet.cell(current_row, column, value)
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell.font = Font(name="Yu Gothic", size=10)
                cell.fill = PatternFill("solid", fgColor="F1F5F9")
            sheet.row_dimensions[current_row].height = max(42, max(len(caption_lines), len(chunk)) * 15 + 8)
            current_row += 1

    for note in schedule.get("source_notes", []):
        text = f"原本から採用した注記｜{note['source']}\n{note['text']}"
        lines = _wrap_schedule_text(text, axis_width, 10)
        for offset in range(0, len(lines), 20):
            chunk = lines[offset:offset + 20]
            sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=last_col)
            cell = sheet.cell(current_row, 1, "\n".join(chunk))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name="Yu Gothic", size=10)
            sheet.row_dimensions[current_row].height = len(chunk) * 15 + 12
            current_row += 1
    for row in sheet.iter_rows(min_row=5, max_row=max(current_row - 1, 5), min_col=1, max_col=last_col):
        for cell in row:
            # A minute is a coordinate, not a visible table column.
            cell.border = Border(bottom=thin) if cell.column > 2 else border
    sheet.column_dimensions["A"].width = 23
    sheet.column_dimensions["B"].width = 18
    for column in range(3, last_col + 1):
        sheet.column_dimensions[get_column_letter(column)].width = minute_width
    sheet.freeze_panes = "C6"
    sheet.print_area = f"A1:{get_column_letter(last_col)}{max(current_row - 1, 6)}"
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = "s"
