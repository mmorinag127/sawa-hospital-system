from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from typing import Any
import re
import posixpath
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageSequence

from src.services import implementation_price_table_ocr_service
from src.services.school_lunch_cooking_parser import extract_tables


MAX_UPLOAD_BYTES = 20 * 1024 * 1024
PDF_EXTENSIONS = {".pdf"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}
EXCEL_EXTENSIONS = {".xls", ".xlsx", ".xlsm"}
MAX_EXCEL_CELLS = 500_000
MAX_EXCEL_SHEETS = 100
MAX_EXCEL_EXPANDED_BYTES = 80 * 1024 * 1024
ALLOWED_EXTENSIONS = PDF_EXTENSIONS | IMAGE_EXTENSIONS | EXCEL_EXTENSIONS


def validate_upload(*, filename: str, content: bytes) -> str:
    suffix = Path(str(filename or "").strip()).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("file must be PDF, PNG, JPEG, TIFF, XLS, XLSX, or XLSM")
    if not content:
        raise ValueError("file is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("file exceeds 20MB limit")
    return suffix


def image_to_pdf(*, filename: str, image_bytes: bytes) -> bytes:
    validate_upload(filename=filename, content=image_bytes)
    try:
        source = Image.open(BytesIO(image_bytes))
        pages = [frame.convert("RGB") for frame in ImageSequence.Iterator(source)]
    except Exception as exc:  # noqa: BLE001
        raise ValueError("image could not be read") from exc
    if not pages:
        raise ValueError("image has no pages")
    output = BytesIO()
    pages[0].save(output, format="PDF", save_all=True, append_images=pages[1:], resolution=200)
    return output.getvalue()


def normalize_ocr_upload(*, filename: str, content: bytes) -> tuple[str, bytes]:
    suffix = validate_upload(filename=filename, content=content)
    if suffix in PDF_EXTENSIONS:
        if not content.startswith(b"%PDF-"):
            raise ValueError("file is not a valid PDF")
        return filename, content
    if suffix in IMAGE_EXTENSIONS:
        return f"{Path(filename).stem}.pdf", image_to_pdf(filename=filename, image_bytes=content)
    raise ValueError("Excel files must use structured extraction")


def _excel_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="minutes")
    return str(value).strip()


def _structured_table(rows: list[list[str]], title: str, index: int) -> dict[str, Any]:
    """Collapse explicit header spans, retaining original Excel cell coordinates."""
    table: dict[str, Any] = {"page_index": index, "source": "structured_excel", "table_id": title, "rows": rows}
    match = implementation_price_table_ocr_service._find_header_row(rows)
    if match is None:
        return table
    table["facility_name"] = implementation_price_table_ocr_service._explicit_cooking_facility_name(rows)
    return implementation_price_table_ocr_service._normalize_cooking_header_spans(table)


def _drawing_notes(workbook_bytes: bytes) -> list[dict[str, str]]:
    notes: list[dict[str, str]] = []
    with ZipFile(BytesIO(workbook_bytes)) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_EXCEL_EXPANDED_BYTES:
            raise ValueError("Excel expanded content exceeds limit")
        drawing_sheets: dict[str, list[str]] = {}
        relationship_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        def relationships(part: str) -> dict[str, str]:
            relation_path = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
            if relation_path not in archive.namelist():
                return {}
            return {node.attrib["Id"]: posixpath.normpath(posixpath.join(posixpath.dirname(part), node.attrib["Target"])).lstrip("/") for node in ElementTree.fromstring(archive.read(relation_path)) if node.attrib.get("TargetMode") != "External"}
        workbook_parts = relationships("xl/workbook.xml")
        for sheet in ElementTree.fromstring(archive.read("xl/workbook.xml")).iter():
            if sheet.tag.rsplit("}", 1)[-1] != "sheet":
                continue
            sheet_path = workbook_parts.get(sheet.attrib.get(relationship_ns + "id", ""))
            if sheet_path:
                for drawing_path in relationships(sheet_path).values():
                    if drawing_path.startswith("xl/drawings/"):
                        drawing_sheets.setdefault(drawing_path, []).append(sheet.attrib["name"])
        for name in archive.namelist():
            if not (name.startswith("xl/drawings/") and name.endswith((".xml", ".vml"))):
                continue
            root = ElementTree.fromstring(archive.read(name))
            texts = [node.text.strip() for node in root.iter() if node.tag.rsplit("}", 1)[-1] in {"t", "div"} and node.text and node.text.strip()]
            if texts:
                sheet_names = "・".join(drawing_sheets.get(name, []))
                notes.append({"code": "drawing_text_review", "message": "図形内の注記です。アレルギー対応等を確認してください。", "source": f"{sheet_names}: {name}" if sheet_names else name, "raw_value": "\n".join(texts)})
    return notes


def extract_excel_candidates(*, filename: str, workbook_bytes: bytes) -> dict[str, Any]:
    suffix = validate_upload(filename=filename, content=workbook_bytes)
    if suffix not in EXCEL_EXTENSIONS:
        raise ValueError("structured extraction requires XLS, XLSX, or XLSM")
    tables: list[dict[str, Any]] = []
    notes: list[dict[str, str]] = []
    total_cells = 0
    try:
        if suffix == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=workbook_bytes, on_demand=True)
            try:
                if workbook.nsheets > MAX_EXCEL_SHEETS:
                    raise ValueError("Excel sheet count exceeds limit")
                for index in range(workbook.nsheets):
                    sheet = workbook.sheet_by_index(index)
                    total_cells += sheet.nrows * sheet.ncols
                    if total_cells > MAX_EXCEL_CELLS:
                        raise ValueError("Excel cell count exceeds limit")
                    rows = []
                    for row in range(sheet.nrows):
                        values = []
                        for cell in sheet.row(row):
                            value = cell.value
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate_as_datetime(value, workbook.datemode)
                                if 0 <= cell.value < 1:
                                    value = value.time()
                            values.append(_excel_value(value))
                        rows.append(values)
                    tables.append(_structured_table(rows, sheet.name, index + 1))
                    workbook.unload_sheet(index)
            finally:
                workbook.release_resources()
        else:
            notes = _drawing_notes(workbook_bytes)
            workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
            try:
                if len(workbook.worksheets) > MAX_EXCEL_SHEETS:
                    raise ValueError("Excel sheet count exceeds limit")
                for index, sheet in enumerate(workbook.worksheets, start=1):
                    if total_cells + (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_EXCEL_CELLS:
                        raise ValueError("Excel cell count exceeds limit")
                    rows = []
                    for row in sheet.iter_rows(values_only=True):
                        total_cells += max(len(row), 1)
                        if total_cells > MAX_EXCEL_CELLS:
                            raise ValueError("Excel cell count exceeds limit")
                        rows.append([_excel_value(value) for value in row])
                    tables.append(_structured_table(rows, sheet.title, index))
            finally:
                workbook.close()
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ValueError("Excel file could not be read") from exc
    result = extract_tables(
        tables,
        filename=filename,
        provider="structured_excel",
    )
    result["issues"].extend(notes)
    return result
