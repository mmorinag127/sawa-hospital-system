"""Read workbook schedule evidence, not inferred executable operations.

Draft sheets contain explicit header values, ticks, menu groups and their raw
grid cells. Every cell has its original coordinate, merge and fill. Drawings
are resolved through OOXML relationships and keep zero-based anchor markers.
No group is a recipe and no grid width is a duration.
"""
from __future__ import annotations

from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
import posixpath
import re
import unicodedata
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import range_boundaries

from src.services import cooking_arrangement_service as upload


def _raw(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _text(value):
    return "" if value is None else str(value)


def _color(color):
    return {key: getattr(color, key) for key in ("type", "tint", "auto", "indexed", "rgb", "theme")
            if key in color.__dict__}


def _cell(cell, merges):
    return {"cell": cell.coordinate, "value": _raw(cell.value),
            "merged_range": next((str(span) for span in merges if cell.coordinate in span), None),
            "fill": {"pattern_type": cell.fill.patternType,
                     "foreground": _color(cell.fill.fgColor), "background": _color(cell.fill.bgColor)}}


def _date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    text = re.sub(r"\s+", "", unicodedata.normalize("NFKC", _text(value)))
    match = re.search(r"(?:令和(元|\d+)年|(\d{4})年)(\d+)月(\d+)日", text)
    if not match:
        return ""
    year = (2018 + (1 if match[1] == "元" else int(match[1]))) if match[1] else int(match[2])
    try:
        return date(year, int(match[3]), int(match[4])).isoformat()
    except ValueError:
        return ""


def _clock(value):
    if isinstance(value, time):
        return value.strftime("%H:%M")
    match = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?::00)?\s*", unicodedata.normalize("NFKC", _text(value)))
    if match:
        try:
            return time(int(match[1]), int(match[2])).strftime("%H:%M")
        except ValueError:
            pass
    return ""


def _drawings(content):
    by_sheet = {}
    rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    with ZipFile(BytesIO(content)) as archive:
        def relationships(part):
            path = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
            if path not in archive.namelist():
                return {}
            return {node.attrib["Id"]: posixpath.normpath(posixpath.join(posixpath.dirname(part), node.attrib["Target"])).lstrip("/")
                    for node in ET.fromstring(archive.read(path)) if node.attrib.get("TargetMode") != "External"}

        workbook_rels = relationships("xl/workbook.xml")
        for sheet in ET.fromstring(archive.read("xl/workbook.xml")).iter():
            if sheet.tag.rsplit("}", 1)[-1] != "sheet":
                continue
            part = workbook_rels.get(sheet.attrib.get(rel_id))
            if not part:
                continue
            sheet_rels = relationships(part)
            notes = []
            for node in ET.fromstring(archive.read(part)).iter():
                if node.tag.rsplit("}", 1)[-1] != "drawing":
                    continue
                drawing = sheet_rels.get(node.attrib.get(rel_id))
                if not drawing:
                    continue
                for anchor in ET.fromstring(archive.read(drawing)):
                    texts = [node.text for node in anchor.iter() if node.tag.rsplit("}", 1)[-1] == "t" and node.text]
                    markers = {}
                    for marker in anchor:
                        tag = marker.tag.rsplit("}", 1)[-1]
                        if tag in {"from", "to"}:
                            markers[tag] = {child.tag.rsplit("}", 1)[-1]: int(child.text) for child in marker if child.text}
                        elif tag in {"pos", "ext"}:
                            markers[tag] = dict(marker.attrib)
                    # Preserve the complete XML evidence, including unrecognised shapes.
                    # Coordinates and connector ends are not executable task times.
                    objects = [child for child in anchor
                               if child.tag.rsplit("}", 1)[-1] not in {"from", "to", "pos", "ext", "clientData"}]
                    identities = [dict(child.attrib) for child in anchor.iter()
                                  if child.tag.rsplit("}", 1)[-1] == "cNvPr"]
                    notes.append({"source": drawing, "anchor_type": anchor.tag.rsplit("}", 1)[-1],
                                  "anchor": markers, "text": "\n".join(texts),
                                  "object_types": [child.tag.rsplit("}", 1)[-1] for child in objects],
                                  "identities": identities,
                                  "source_xml": ET.tostring(anchor, encoding="unicode")})
            by_sheet[sheet.attrib["name"]] = notes
    return by_sheet


def extract_work_schedule_draft(*, filename: str, workbook_bytes: bytes) -> dict:
    suffix = upload.validate_upload(filename=filename, content=workbook_bytes)
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("work schedule import requires XLSX or XLSM")
    try:
        with ZipFile(BytesIO(workbook_bytes)) as archive:
            if sum(item.file_size for item in archive.infolist()) > upload.MAX_EXCEL_EXPANDED_BYTES:
                raise ValueError("Excel expanded content exceeds limit")
            declared_cells = 0
            for name in archive.namelist():
                if not re.fullmatch(r"xl/worksheets/[^/]+\.xml", name):
                    continue
                for _, element in ET.iterparse(BytesIO(archive.read(name)), events=("start",)):
                    if element.tag.rsplit("}", 1)[-1] == "dimension":
                        left, top, right, bottom = range_boundaries(element.attrib["ref"])
                        declared_cells += (right - left + 1) * (bottom - top + 1)
                        if declared_cells > upload.MAX_EXCEL_CELLS:
                            raise ValueError("Excel cell count exceeds limit")
                        break
        drawings = _drawings(workbook_bytes)
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Excel file could not be read") from exc
    try:
        if len(workbook.worksheets) > upload.MAX_EXCEL_SHEETS:
            raise ValueError("Excel sheet count exceeds limit")
        if sum(sheet.max_row * sheet.max_column for sheet in workbook) > upload.MAX_EXCEL_CELLS:
            raise ValueError("Excel cell count exceeds limit")
        result = {"document_type": "work_schedule", "filename": filename,
                  "source_sha256": sha256(workbook_bytes).hexdigest(), "status": "needs_review",
                  "review_required": True, "sheets": [], "auxiliary_sheets": [], "issues": []}
        for sheet in workbook:
            merges = list(sheet.merged_cells.ranges)
            raw_cells = [_cell(cell, merges) for row in sheet for cell in row
                         if not isinstance(cell, MergedCell)
                         and (cell.value is not None or cell.fill.patternType is not None)]
            drawing_objects = drawings.get(sheet.title, [])
            parsed = {"name": sheet.title, "raw_cells": raw_cells,
                      "merged_ranges": [str(span) for span in merges],
                      "drawing_objects": drawing_objects,
                      "drawing_notes": [item for item in drawing_objects if item["text"]]}
            header = next((row[0].row for row in sheet.iter_rows(min_col=1, max_col=2)
                           if _text(row[0].value).strip() == "献立名" and _text(row[1].value).strip() == "担当者"), None)
            if header is None or sheet.title.strip().lower() in {"原本", "テンプレート", "template"}:
                parsed["kind"] = "template" if header else "auxiliary"
                result["auxiliary_sheets"].append(parsed)
                continue
            parsed.update({"raw_date": _raw(sheet["A1"].value), "date_source": "A1", "service_date": _date(sheet["A1"].value),
                           "facility_name": _text(sheet["C1"].value).strip(), "facility_source": "C1",
                           "groups": [], "time_ticks": [], "delivery_cells": [], "issues": []})
            def issue(code, source, message):
                entry = {"code": code, "source": f"{sheet.title}!{source}", "message": message}
                parsed["issues"].append(entry)
                result["issues"].append(entry)
            if not parsed["service_date"]:
                issue("service_date_unresolved", "A1", "原本の提供日を確認してください。シート名から補完しません。")
            if not parsed["facility_name"]:
                issue("facility_unresolved", "C1", "学校名を確認してください。")
            candidates = []
            for row_number in range(6, header):
                cell = sheet.cell(row_number, 2)
                raw = _text(cell.value)
                normalized = unicodedata.normalize("NFKC", raw)
                if row_number != 6 and not re.search("[～〜~]", normalized):
                    continue
                pieces = re.split("[～〜~]", normalized)
                clocks = [_clock(piece) for piece in pieces]
                valid = len(clocks) == 2 and all(clocks) and clocks[0] < clocks[1]
                if raw:
                    candidates.append({"raw": raw, "source": cell.coordinate, "start": clocks[0] if valid else "", "end": clocks[1] if valid else ""})
            parsed["plating_candidates"] = candidates
            distinct = {(item["start"], item["end"]) for item in candidates}
            valid_plating = len(distinct) == 1 and all(next(iter(distinct)))
            parsed["plating"] = (dict(candidates[0]) if valid_plating else {
                "raw": " / ".join(f"{item['source']}: {item['raw']}" for item in candidates),
                "source": ",".join(item["source"] for item in candidates), "start": "", "end": ""})
            if len(distinct) > 1:
                issue("plating_time_ambiguous", parsed["plating"]["source"], "異なる盛付時刻が複数あります。原本の候補を確認して開始・終了を選んでください。")
            if not valid_plating:
                issue("plating_time_unresolved", "B6", "盛り付け開始・終了時刻を確認してください。")
            for column in range(3, sheet.max_column + 1):
                cell = sheet.cell(header, column)
                if cell.value is not None:
                    parsed["time_ticks"].append({"cell": cell.coordinate, "raw": _raw(cell.value), "time": _clock(cell.value)})
            # Only a cell explicitly saying 配送 is classified as delivery.
            tick_columns = [sheet[tick["cell"]].column for tick in parsed["time_ticks"]]
            group_rows = [row for row in range(header + 1, sheet.max_row + 1)
                          if not isinstance(sheet.cell(row, 1), MergedCell) and sheet.cell(row, 1).value is not None]
            for index, start in enumerate(group_rows):
                end = group_rows[index + 1] - 1 if index + 1 < len(group_rows) else sheet.max_row
                group_cell = sheet.cell(start, 1)
                group = {"name": _text(group_cell.value), "source_cell": group_cell.coordinate,
                         "merged_range": next((str(span) for span in merges if group_cell.coordinate in span), None),
                         "row_start": start, "row_end": end, "assignees_raw": _text(sheet.cell(start, 2).value),
                         "assignees_source": sheet.cell(start, 2).coordinate, "cells": []}
                if group["assignees_raw"]:
                    issue("assignees_review_required", group["assignees_source"], "担当者の原文です。氏名の区切りと在籍情報を確認してください。")
                for row in range(start, end + 1):
                    for column in tick_columns:
                        cell = sheet.cell(row, column)
                        if isinstance(cell, MergedCell) or cell.value is None:
                            continue
                        evidence = _cell(cell, merges)
                        if "配送" in _text(cell.value):
                            parsed["delivery_cells"].append(evidence)
                        else:
                            group["cells"].append(evidence)
                parsed["groups"].append(group)
            if parsed["drawing_notes"]:
                issue("drawing_notes_review_required", "drawing", "図形の注記を原位置と照合してください。工程へ自動適用していません。")
            if any(not item["text"] for item in drawing_objects):
                issue("drawing_objects_review_required", "drawing", "矢印・色見本など文字のない図形を原位置と照合してください。時刻や工程には変換していません。")
            result["sheets"].append(parsed)
        return result
    finally:
        workbook.close()
