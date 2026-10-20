"""A monthly projection from canonical menus, never a stock transaction ledger."""
from collections import defaultdict
from datetime import date
from decimal import Decimal, ROUND_CEILING, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
import json
import math
import unicodedata

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import selectinload

from src.models.monthly_procurement import MonthlyProcurement
from src.models.school_lunch import Facility, MenuExecution, MenuIngredient
from src.services import school_lunch_ingredient_settings as ingredient_settings


def period(month):
    try:
        start = date.fromisoformat(month + "-01")
        if start.strftime("%Y-%m") != month:
            raise ValueError()
        end = date(start.year + (start.month == 12), start.month % 12 + 1, 1)
        return start, end
    except ValueError:
        raise HTTPException(422, "月はYYYY-MMで指定してください。")


def source(session, facility_id, month):
    start, end = period(month)
    facility = session.get(Facility, facility_id)
    if not facility:
        raise HTTPException(404, "学校が見つかりません。")
    menus = session.scalars(select(MenuExecution).where(
        MenuExecution.facility_id == facility_id,
        MenuExecution.service_date >= start, MenuExecution.service_date < end,
    ).options(selectinload(MenuExecution.menu_ingredients).selectinload(MenuIngredient.ingredient))
        .order_by(MenuExecution.service_date, MenuExecution.id)).all()
    versions = {}
    for menu in menus:
        # Include shared ingredient units: editing another menu can change that master.
        values = []
        for row in menu.menu_ingredients:
            ingredient = ingredient_settings.for_row(row)
            rate = ingredient.conversion_rate
            values.append((row.ingredient_id, ingredient.name, str(Decimal(str(row.required_quantity)).normalize()), row.unit,
                           ingredient.stock_unit, ingredient.order_unit,
                           str(Decimal(str(rate)).normalize()) if rate is not None else None))
        digest = sha256(json.dumps([menu.name, menu.servings, values], ensure_ascii=False).encode()).hexdigest()
        versions[str(menu.id)] = f"{menu.version}:{digest}"
    return facility, menus, versions


def inventory_setting_issue(ingredient):
    missing = [label for name, label in (("stock_unit", "在庫単位"), ("order_unit", "発注単位"),
                                        ("conversion_rate", "換算率")) if getattr(ingredient, name) is None]
    if missing:
        return f"{ingredient.name}: {'・'.join(missing)}が未設定です。献立入力で設定するまで在庫・発注は未計算です。"
    rate = Decimal(str(ingredient.conversion_rate))
    if not ingredient.stock_unit.strip() or not ingredient.order_unit.strip() or not rate.is_finite() or rate <= 0:
        return f"{ingredient.name}: 在庫単位・発注単位・換算率が不正です。在庫・発注は未計算です。"
    return None


def quantity_in_stock(quantity, unit, ingredient):
    issue = inventory_setting_issue(ingredient)
    if issue:
        raise ValueError(issue)
    if unit == ingredient.stock_unit:
        return Decimal(str(quantity))
    if unit == ingredient.order_unit and ingredient.order_unit != ingredient.stock_unit:
        return Decimal(str(quantity)) * Decimal(str(ingredient.conversion_rate))
    raise ValueError(f"{ingredient.name}: {unit}から{ingredient.stock_unit}への換算が未登録です。")


def money(value):
    return str(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def compute(session, facility_id, month):
    facility, menus, versions = source(session, facility_id, month)
    saved = session.scalar(select(MonthlyProcurement).where(
        MonthlyProcurement.facility_id == facility_id, MonthlyProcurement.month == month))
    inputs = saved.inputs if saved else {}
    stale = bool(saved and saved.menu_versions != versions)
    ingredients = ingredient_settings.for_menus(menus)
    inventory_issues = [issue for ingredient in ingredients.values() if (issue := inventory_setting_issue(ingredient))]
    issues = []
    if not menus:
        issues.append("対象月の献立が登録されていません。")
        inventory_issues.append("対象月の献立が登録されていません。")
    if stale:
        issues.append("献立または食材単位が変更されています。再計算結果を確認し、入力を保存してください。")
    demand = defaultdict(Decimal)
    cost_rows = []
    material_rows = []
    orders = []
    stock_rows = []
    month_cost = Decimal(0)
    cost_complete = True
    daily_cost = defaultdict(Decimal)
    daily_complete = defaultdict(lambda: True)
    bad_demand = set()
    for menu in menus:
        for row in menu.menu_ingredients:
            ingredient = ingredients[row.ingredient_id]
            config = inputs.get(str(ingredient.id), {})
            qty = None
            amount = None
            try:
                qty = quantity_in_stock(row.required_quantity, row.unit, ingredient)
                demand[(menu.service_date, ingredient.id)] += qty
            except ValueError as exc:
                issues.append(str(exc))
                inventory_issues.append(str(exc))
                bad_demand.add(ingredient.id)
            price = config.get("unit_price")
            price_unit = config.get("price_unit")
            if price is not None and price_unit and qty is not None:
                try:
                    units = quantity_in_stock(1, price_unit, ingredient)
                    amount = (qty / units * Decimal(price)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                    month_cost += amount
                    daily_cost[menu.service_date] += amount
                except ValueError as exc:
                    issues.append(str(exc))
            if amount is None:
                cost_complete = False
                daily_complete[menu.service_date] = False
            cost_rows.append({"date": menu.service_date.isoformat(), "menu_id": menu.id, "menu": menu.name,
                "servings": menu.servings, "ingredient_id": ingredient.id, "ingredient": ingredient.name,
                "quantity": str(row.required_quantity), "unit": row.unit,
                "unit_price": price, "price_unit": price_unit, "cost": money(amount) if amount is not None else None,
                "per_person_cost": money(amount / menu.servings) if amount is not None else None})
    for ingredient_id, ingredient in sorted(ingredients.items(), key=lambda pair: pair[1].name):
        config = inputs.get(str(ingredient_id), {})
        opening = config.get("opening_stock")
        price = config.get("unit_price")
        price_unit = config.get("price_unit")
        supplier = config.get("supplier", "")
        increment = config.get("order_increment")
        material_rows.append({"ingredient_id": ingredient_id, "name": ingredient.name,
            "stock_unit": ingredient.stock_unit, "order_unit": ingredient.order_unit,
            "conversion_rate": str(ingredient.conversion_rate) if ingredient.conversion_rate is not None else None,
            "opening_stock": opening, "unit_price": price, "price_unit": price_unit,
            "supplier": supplier, "order_increment": increment, "incoming": config.get("incoming", [])})
        if price is None or not price_unit:
            issues.append(f"{ingredient.name}: 単価・単価単位を入力してください。")
        if opening is None:
            issues.append(f"{ingredient.name}: 月初在庫を入力してください（在庫なしは0）。")
            inventory_issues.append(f"{ingredient.name}: 月初在庫を入力してください（在庫なしは0）。")
            continue
        if ingredient_id in bad_demand:
            continue
        balance = Decimal(opening)
        incoming = defaultdict(Decimal)
        for item in config.get("incoming", []):
            incoming[date.fromisoformat(item["date"])] += Decimal(item["quantity"])
        dates = sorted(set(incoming) | {day for day, key in demand if key == ingredient_id})
        for day in dates:
            before = balance
            received = incoming[day]
            required = demand[(day, ingredient_id)]
            available = balance + received
            shortage = max(Decimal(0), required - available)
            # One order unit contains conversion_rate stock units, explicitly registered in the master.
            rate = Decimal(str(ingredient.conversion_rate))
            if ingredient.order_unit == ingredient.stock_unit:
                rate = Decimal(1)
            if shortage and increment is None:
                issues.append(f"{ingredient.name}: 発注刻みを入力してください。{day.isoformat()}以降の発注・残在庫は未計算です。")
                inventory_issues.append(f"{ingredient.name}: 発注刻みが未設定のため在庫・発注は未計算です。")
                break
            ordered = Decimal(0)
            if shortage:
                step = Decimal(increment)
                ordered = (shortage / rate / step).to_integral_value(rounding=ROUND_CEILING) * step
            balance = available + ordered * rate - required
            if ordered:
                if not supplier:
                    issues.append(f"{ingredient.name}: 発注先を入力してください。")
                    inventory_issues.append(f"{ingredient.name}: 発注先を入力してください。")
                orders.append({"date": day.isoformat(), "ingredient_id": ingredient_id,
                    "ingredient": ingredient.name, "supplier": supplier, "quantity": str(ordered),
                    "unit": ingredient.order_unit, "stock_quantity": str(ordered * rate)})
            stock_rows.append({"date": day.isoformat(), "ingredient_id": ingredient_id,
                "ingredient": ingredient.name, "unit": ingredient.stock_unit, "opening": str(before),
                "incoming": str(received), "required": str(required), "shortage": str(shortage),
                "planned_order_stock": str(ordered * rate), "projected_closing": str(balance)})
    # Never show a partially calculated balance/order list as the month's plan.
    if inventory_issues:
        stock_rows = []
        orders = []
    issues.extend(inventory_issues)
    return {"facility": {"id": facility.id, "name": facility.name}, "month": month,
        "version": saved.version if saved else 0, "menu_versions": versions, "stale": stale,
        "menu_count": len(menus), "ingredients": material_rows, "cost_rows": cost_rows,
        "daily_costs": [{"date": day.isoformat(), "cost": money(daily_cost[day]) if daily_complete[day] else None}
                        for day in sorted({menu.service_date for menu in menus})],
        "month_cost": money(month_cost) if cost_complete and menus else None,
        "stock_rows": stock_rows, "orders": orders, "issues": sorted(set(issues)),
        "inventory_ready": not inventory_issues, "inventory_issues": sorted(set(inventory_issues))}


def save(session, facility_id, month, payload):
    _, menus, versions = source(session, facility_id, month)
    if payload["menu_versions"] != versions:
        raise HTTPException(409, "献立が更新されました。再読込して確認してください。")
    allowed = {row.ingredient_id for menu in menus for row in menu.menu_ingredients}
    items = payload["ingredients"]
    if len({item["ingredient_id"] for item in items}) != len(items):
        raise HTTPException(422, "食材が重複しています。")
    if {item["ingredient_id"] for item in items} != allowed:
        raise HTTPException(422, "対象月の食材一式を指定してください。")
    start, end = period(month)
    canonical = ingredient_settings.for_menus(menus)
    for item in items:
        if item.get("price_unit"):
            try:
                quantity_in_stock(1, item["price_unit"], canonical[item["ingredient_id"]])
            except ValueError as exc:
                raise HTTPException(422, str(exc))
        for incoming in item["incoming"]:
            if not start <= date.fromisoformat(incoming["date"]) < end:
                raise HTTPException(422, "入荷日は対象月内を指定してください。")
    inputs = {str(item["ingredient_id"]): item for item in items}
    expected = payload["version"]
    try:
        if expected == 0:
            session.add(MonthlyProcurement(facility_id=facility_id, month=month, version=1,
                                          inputs=inputs, menu_versions=versions))
            session.flush()
        else:
            result = session.execute(update(MonthlyProcurement).where(
                MonthlyProcurement.facility_id == facility_id, MonthlyProcurement.month == month,
                MonthlyProcurement.version == expected).values(inputs=inputs, menu_versions=versions, version=expected + 1))
            if result.rowcount != 1:
                raise HTTPException(409, "別の操作で更新されました。再読込してください。")
            session.flush()
        session.expire_all()
    except IntegrityError:
        session.rollback()
        raise HTTPException(409, "別の操作で登録されました。再読込してください。")
    return compute(session, facility_id, month)


def excel(result):
    if result["issues"]:
        raise HTTPException(409, {"message": "未確認項目を解消してから出力してください。", "issues": result["issues"]})
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = [
        ("実施価格", ["日付", "献立", "食数", "食材", "必要量", "単位", "税込単価", "単価単位", "金額", "1人当たり金額"],
         [[r[k] for k in ["date", "menu", "servings", "ingredient", "quantity", "unit", "unit_price", "price_unit", "cost", "per_person_cost"]] for r in result["cost_rows"]]),
        ("在庫計画", ["日付", "食材", "単位", "繰越予定", "登録入荷", "必要量", "不足量", "追加発注予定量", "残在庫予定"],
         [[r[k] for k in ["date", "ingredient", "unit", "opening", "incoming", "required", "shortage", "planned_order_stock", "projected_closing"]] for r in result["stock_rows"]]),
        ("発注案", ["必要日（納品期限）", "食材", "発注先", "発注数", "発注単位", "在庫単位換算量"],
         [[r[k] for k in ["date", "ingredient", "supplier", "quantity", "unit", "stock_quantity"]] for r in result["orders"]]),
    ]
    for name, headers, rows in sheets:
        sheet = workbook.create_sheet(name)
        sheet.append([f"{result['facility']['name']} {result['month']} {name}"])
        sheet.append(["発注案を必要日までに入荷する想定の計画です。発注・入荷の確定処理は行いません。"])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        if name == "実施価格":
            sheet.append(["月合計", None, None, None, None, None, None, None, result["month_cost"]])
        numeric_columns = {"実施価格": (3, 5, 7, 9, 10), "在庫計画": (4, 5, 6, 7, 8, 9), "発注案": (4, 6)}[name]
        for row_index in range(4, sheet.max_row + 1):
            for column in numeric_columns:
                cell = sheet.cell(row_index, column)
                if cell.value is not None:
                    cell.value = Decimal(str(cell.value))
                    cell.number_format = "0.00" if name == "実施価格" and column in (7, 9, 10) else "0.######"
        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:{sheet.cell(3, len(headers)).column_letter}{max(3, 3 + len(rows))}"
        for cell in sheet[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 20
        for title_row in (1, 2):
            sheet.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=len(headers))
            sheet.row_dimensions[title_row].height = 30
        for cells in sheet.iter_rows(min_row=3):
            line_count = 1
            for cell in cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                # Conservative CJK width; preserve the original cell value in full.
                text = str(cell.value) if cell.value is not None else ""
                lines = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in part) / 18)) for part in text.split("\n"))
                line_count = max(line_count, lines)
            sheet.row_dimensions[cells[0].row].height = max(24, line_count * 15 + 6)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:3"
        # Treat operator strings as text, not spreadsheet formulas.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
    summary = workbook.create_sheet("日別合計・参照情報")
    summary.append(["学校", result["facility"]["name"]])
    summary.append(["対象月", result["month"]])
    summary.append(["入力バージョン", result["version"]])
    summary.append(["日付", "実施価格合計（税込）"])
    for row in result["daily_costs"]:
        summary.append([row["date"], Decimal(row["cost"])])
    summary.append(["月合計", Decimal(result["month_cost"])])
    summary.append(["共通献立ID", "参照バージョン・内容識別子"])
    for key, value in result["menu_versions"].items():
        summary.append([key, value])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 76
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.orientation = "landscape"
    summary.page_setup.paperSize = summary.PAPERSIZE_A4
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.print_title_rows = "1:3"
    for row in summary:
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
