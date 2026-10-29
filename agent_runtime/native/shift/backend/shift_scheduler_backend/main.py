import csv
import hashlib
import json
from copy import copy
from datetime import date, timedelta
from io import BytesIO, StringIO
import logging
import os
from pathlib import Path
from types import SimpleNamespace
from typing import Literal

from fastapi import Depends, FastAPI, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field, StrictInt, StrictStr
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy import text, tuple_
from sqlalchemy.dialects.postgresql import insert as postgres_insert
from sqlalchemy.dialects.sqlite import insert as sqlite_insert
from sqlalchemy.orm import Session, joinedload, load_only
from starlette.middleware.cors import CORSMiddleware
from starlette.responses import JSONResponse

from .db import Base, get_engine, get_session
from .settings_api_schemas import (
    OptimizationSettingsResponse, RulePreflightResponse,
    RuleAssistantCandidateResponse, RuleAssistantApplyResponse,
)
from .models import (
    Assignment,
    AssignmentExplanation,
    Availability,
    CandidateStaffStatus,
    CandidateStaffStatusExplanation,
    PositionAssignmentRule,
    CandidateSchedule,
    Demand,
    EditorRevision,
    Position,
    PositionPriority,
    ReferenceAssignment,
    RuleConfig,
    RotationPolicy,
    RotationQueueState,
    ScheduleConfirmation,
    ScheduleException,
    ShiftTemplate,
    SolveRun,
    Staff,
    StaffPairRule,
    StaffPositionPreference,
    StaffPositionRule,
    StaffSkill,
    Violation,
    WorkdayRule,
)
from .editor_revision_service import (
    begin as begin_editor_write, begin_multi_resource_write_snapshot, begin_read_snapshot, claim_collection, claim_item,
    dev_seed_write_snapshot, command_payload, conflict, end_read_snapshot, finalize_command, read_collection,
    require_revision, technical_revision,
)
from .periods import (
    build_active_settings,
    resolve_period_request,
    scale_decimal_units,
)
from .rule_config_utils import PAID_LEAVE_HALF_DAY_WORK_KIND
from .rotation_queue import validate_rotation_period
from .rule_preflight import analyze_rule_feasibility
from .rule_assistant import build_rule_candidates
from .planning_api import router as planning_router
from .sample_data import seed_sample_data
from .rule_profiles import effective_rule_config, initialize_initial_profile
from .pdf_parity import compare_reference_cells, reference_status_count, group_candidate_assignments
from .pdf_seed_data import (
    PDF_PAGE5_END_EXCLUSIVE,
    PDF_PAGE5_START,
    PDF_REFERENCE_SOURCE_PAGE5_FULL,
    PDF_STAFF_NAMES,
    PRACTICAL_TRAINEES,
    SCHOOL_AD_STAFF,
    seed_pdf_page5_full_manual_data,
)
from .storage import StorageError, store_bytes
from .schema_contract import assert_runtime_schema
from .constraint_api_schemas import (
    ConstraintCollectionResponse, ConstraintCommandResponse,
    ConstraintDeleteResponse, ConstraintBulkResponse,
)
from .staff_api_schemas import (
    StaffCollectionResponse, StaffCommandResponse, StaffDeleteResponse,
    StaffDisplayOrderResponse,
)


SPREADSHEET_FORMULA_PREFIXES = ("=", "+", "-", "@")
RevisionToken = StrictInt | None


def _neutralize_spreadsheet_cell(value):
    """Return text that spreadsheet applications cannot interpret as a formula."""
    if isinstance(value, str) and value.startswith(SPREADSHEET_FORMULA_PREFIXES):
        return f"'{value}"
    return value
from .jobs import SolveJobRequest, serialize_job_status
from .validation_snapshot import begin_solver_input_snapshot
from .candidate_command import CandidatePrecondition, candidate_snapshot, candidate_precondition, execute_candidate_command
from .candidate_validation import validate_candidate as validate_candidate_state
from .candidate_validation import _load_prior_statuses
from .schedule_selection import confirmed_candidate, direct_predecessor_run, prior_history_candidate, prior_history_coverage_errors, unresolved_candidate_detail
from .security import require_service_token, service_auth_enabled
from .solve_contract import (
    AssignmentLockResponse, CandidateValidation, ScheduleBulkResponse,
    ScheduleCellResponse, SolveResponse, SolveRunDetail, SolveRunSummary,
)

app = FastAPI(title="Shift Scheduler Backend")
from .staging_data import DatasetMiddleware
app.add_middleware(DatasetMiddleware)
logger = logging.getLogger(__name__)
app.middleware("http")(require_service_token)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://localhost:3001", "http://127.0.0.1:3000", "http://127.0.0.1:3001"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.include_router(planning_router)
from .staging_copy_migration import router as staging_copy_migration_router
app.include_router(staging_copy_migration_router)


@app.get("/staging-dataset-ready")
def staging_dataset_ready():
    from . import staging_data as datasets
    if not datasets.staging_runtime():
        raise HTTPException(404, "Not found")
    from alembic.config import Config
    from alembic.script import ScriptDirectory
    config = Config(str(Path(__file__).resolve().parents[1] / "alembic.ini"))
    with get_engine().connect() as connection:
        assert_runtime_schema(connection, ScriptDirectory.from_config(config))
    return {"dataset": datasets.current_dataset(), "ready": True}


class PriorStatusInput(BaseModel):
    staff_id: int
    work_date: date
    status: str = Field(pattern=r"^(working|off)$")


class PriorTimeSlotInput(BaseModel):
    start_minute: int = Field(ge=0, lt=1440)
    end_minute: int = Field(gt=0, le=2880)
    work_minutes: int = Field(ge=0, le=1440)


class PriorTimeInput(BaseModel):
    staff_id: int
    work_date: date
    status: str = Field(pattern=r"^(working|off)$")
    slots: list[PriorTimeSlotInput]


class SolveRequest(BaseModel):
    target_month: str | None = Field(default=None, pattern=r"^\d{4}-(0[1-9]|1[0-2])$")
    start_date: date | None = None
    end_date_exclusive: date | None = None
    strategy: str = "balanced"
    candidate_count: int = Field(default=3, ge=1, le=5)
    locked_assignment_ids: list[int] = Field(default_factory=list)
    prior_statuses: list[PriorStatusInput] = Field(default_factory=list)
    prior_time_history: list[PriorTimeInput] = Field(default_factory=list)


class RulePreflightRequest(BaseModel):
    target_month: str | None = Field(default=None, pattern=r"^\d{4}-(0[1-9]|1[0-2])$")
    start_date: date | None = None
    end_date_exclusive: date | None = None


class RuleAssistantInput(BaseModel):
    text: str = Field(min_length=1, max_length=20000)


class StaffInput(BaseModel):
    name: str
    staff_group: str = "general"
    max_daily_work_minutes: int | None = Field(default=None, ge=1, le=1440)
    max_weekly_work_minutes: int | None = Field(default=None, ge=1, le=10080)
    max_monthly_work_minutes: int | None = Field(default=None, ge=1, le=44640)
    min_monthly_work_minutes: int | None = Field(default=None, ge=0, le=44640)
    target_monthly_work_minutes: int | None = Field(default=None, ge=0, le=44640)
    paid_leave_eligible: bool = False
    half_day_shift_template_id: int | None = None
    active_from: date | None = None
    active_to: date | None = None
    notes: str | None = None
    revision: RevisionToken = None
    collection_revision: RevisionToken = None


class StaffDisplayOrderInput(BaseModel):
    staff_ids: list[int]
    collection_revision: RevisionToken = None


class ConstraintInput(BaseModel):
    staff_id: int
    work_date: date
    status: Literal["request_off", "unavailable", "half_day", "request_work", "preferred_work"]
    strength: Literal["hard", "soft"] = "soft"
    weight: int = Field(default=10, ge=0)
    shift_template_id: int | None = None
    reason: str | None = None
    revision: RevisionToken = None
    collection_revision: RevisionToken = None


class ConstraintBulkExistingInput(BaseModel):
    id: StrictInt
    revision: RevisionToken = None


class ConstraintBulkCellInput(BaseModel):
    staff_id: int
    work_date: date
    existing: list[ConstraintBulkExistingInput] | None = None


class ConstraintBulkInput(BaseModel):
    cells: list[ConstraintBulkCellInput] = Field(min_length=1)
    operation: str
    strength: str = "hard"
    weight: int = Field(default=20, ge=1)
    shift_template_id: int | None = None
    reason: str | None = None
    collection_revision: RevisionToken = None


class HolidayCountedWorkKindsInput(BaseModel):
    off: float = Field(default=1.0, ge=0, le=366)
    half_day: float = Field(default=0.5, ge=0, le=366)
    regular: float = Field(default=0.0, ge=0, le=366)
    working: float = Field(default=0.0, ge=0, le=366)
    position: float = Field(default=0.0, ge=0, le=366)


class StaffWeekendRuleInput(BaseModel):
    staff_id: int
    minimum_weekend_off_days: int = Field(default=0, ge=0, le=2)
    fixed_weekday_off: int | None = Field(default=None, ge=0, le=4)
    allowed_weekday_offs: list[int] | None = None
    fixed_weekday_off_priority: Literal["hard", "prefer"] = "hard"
    consecutive_weekend_off_count: int = Field(default=0, ge=0, le=31)
    consecutive_weekend_off_priority: Literal["hard", "prefer"] = "hard"
    preference_weight: int = Field(default=100, ge=1, le=1000)


class DayCategoryStaffingRangeInput(BaseModel):
    day_category: Literal["weekday", "saturday", "sunday", "holiday"]
    min_working: int | None = Field(default=None, ge=0, le=100)
    max_working: int | None = Field(default=None, ge=0, le=100)


class GroupDayCategoryStaffingRangeInput(DayCategoryStaffingRangeInput):
    staff_group: str = Field(min_length=1, max_length=80)


class StaffDayCategoryStaffingRangeInput(DayCategoryStaffingRangeInput):
    staff_ids: list[int] = Field(min_length=1)


class PositionDayCategoryProhibitionInput(BaseModel):
    position_id: int
    day_categories: list[Literal["weekday", "saturday", "sunday", "holiday"]]


class PositionStaffingGroupInput(BaseModel):
    position_id: int
    staff_ids: list[int] = Field(min_length=1)
    min_count: int = Field(default=1, ge=1, le=100)


class OptimizationSettingsInput(BaseModel):
    settings_revision: str | None = None
    strategy: str = "balanced"
    standard_period_mode: str = "calendar_month"
    standard_period_start_day: int | None = Field(default=None, ge=1, le=28)
    max_period_days: int = Field(default=62, ge=1, le=366)
    holiday_count_unit: int = Field(default=2, ge=1, le=100)
    holiday_counted_work_kinds: HolidayCountedWorkKindsInput = Field(
        default_factory=HolidayCountedWorkKindsInput
    )
    monthly_holiday_exact_days: float | None = Field(default=None, ge=0, le=366)
    monthly_holiday_min_days: float | None = Field(default=None, ge=0, le=366)
    monthly_holiday_max_days: float | None = Field(default=None, ge=0, le=366)
    auto_apply_japanese_holidays: bool = True
    forbid_half_day_on_weekdays: bool = False
    forbid_half_day_on_saturdays: bool = False
    auto_half_day_allowed_weekdays: list[int] | None = None
    max_daily_work_minutes: int | None = Field(default=None, ge=1, le=1440)
    max_weekly_work_minutes: int | None = Field(default=None, ge=1, le=10080)
    max_monthly_work_minutes: int | None = Field(default=None, ge=1, le=44640)
    min_rest_minutes: int = Field(default=0, ge=0, le=10080)
    min_monthly_work_minutes: int | None = Field(default=None, ge=0, le=44640)
    half_day_work_minutes: int = Field(default=240, ge=1, le=1440)
    full_day_work_minutes: int = Field(default=480, ge=1, le=1440)
    weekly_min_off_days: int | None = Field(default=None, ge=0, le=7)
    require_prior_period_continuity: bool = False
    monthly_required_off_days: int | None = Field(default=None, ge=0, le=31)
    monthly_min_off_days: int | None = Field(default=None, ge=0, le=31)
    monthly_max_off_days: int | None = Field(default=None, ge=0, le=31)
    max_consecutive_work_days: int | None = Field(default=None, ge=1, le=31)
    max_consecutive_off_days: int | None = Field(default=None, ge=1, le=31)
    require_monthly_consecutive_off: bool = False
    monthly_consecutive_off_strength: Literal["hard", "soft"] = "hard"
    monthly_consecutive_off_weight: int = Field(default=1, ge=1, le=1000)
    coverage_weight: int | None = Field(default=None, ge=0, le=1000)
    balanced_shift_types: list[str] = Field(default_factory=list)
    weekday_off_spread_weight: int = Field(default=0, ge=0, le=100)
    weekday_off_target_min: int = Field(default=0, ge=0, le=31)
    weekday_off_target_max: int = Field(default=3, ge=0, le=31)
    avoid_unplanned_off_weight: int = Field(default=0, ge=0, le=1000)
    consecutive_work_weight: int = Field(default=0, ge=0, le=1000)
    target_work_minutes_weight: int = Field(default=0, ge=0, le=1000)
    weekend_fairness_weight: int = Field(default=0, ge=0, le=1000)
    shift_fairness_weight: int = Field(default=0, ge=0, le=1000)
    regular_work_display_blank: bool = False
    weekday_off_hard_max: int | None = Field(default=None, ge=0, le=31)
    weekday_off_hard_min: int | None = Field(default=None, ge=0, le=31)
    weekday_off_exempt_staff_ids: list[int] = Field(default_factory=list)
    weekend_min_staff_count: int | None = Field(default=None, ge=0, le=100)
    weekend_max_staff_count: int | None = Field(default=None, ge=0, le=100)
    weekend_preferred_staff_ids: list[int] = Field(default_factory=list)
    weekend_comparison_staff_ids: list[int] = Field(default_factory=list)
    weekend_preferred_min_advantage: int = Field(default=0, ge=0, le=100)
    staff_weekend_rules: list[StaffWeekendRuleInput] = Field(default_factory=list)
    weekend_coverage_staff_ids: list[int] = Field(default_factory=list)
    weekend_coverage_min_working: int | None = Field(default=None, ge=0, le=100)
    weekend_coverage_max_working: int | None = Field(default=None, ge=0, le=100)
    weekend_coverage_target_working: int | None = Field(default=None, ge=0, le=100)
    weekend_coverage_priority: Literal["hard", "prefer"] = "hard"
    weekend_coverage_weight: int = Field(default=100, ge=1, le=1000)
    weekend_coverage_target_weight: int = Field(default=100, ge=1, le=1000)
    weekend_adjustment_policy: Literal["ask", "allow_minimum_weekday_offs"] = "ask"
    additional_weekday_off_weight: int = Field(default=100, ge=0, le=1000)
    day_category_staffing_ranges: list[DayCategoryStaffingRangeInput] = Field(default_factory=list)
    group_day_category_staffing_ranges: list[GroupDayCategoryStaffingRangeInput] = Field(default_factory=list)
    staff_day_category_staffing_ranges: list[StaffDayCategoryStaffingRangeInput] = Field(default_factory=list)
    position_day_category_prohibitions: list[PositionDayCategoryProhibitionInput] = Field(default_factory=list)
    position_staffing_groups: list[PositionStaffingGroupInput] = Field(default_factory=list)


class RuleAssistantPairRuleInput(BaseModel):
    first_staff_id: int
    second_staff_id: int
    kind: Literal["avoid_same_day_off"]
    strength: Literal["hard", "soft"] = "hard"
    weight: int = Field(default=100, ge=1, le=1000)
    reason: str | None = None


class RuleAssistantStaffUpdateInput(BaseModel):
    staff_id: int
    paid_leave_eligible: bool


class RuleAssistantItemVersionInput(BaseModel):
    id: StrictInt
    revision: RevisionToken = None


class RuleAssistantCollectionVersionsInput(BaseModel):
    collection_revision: RevisionToken = None
    items: list[RuleAssistantItemVersionInput] | None = None


class RuleAssistantInputVersionsInput(BaseModel):
    staff: RuleAssistantCollectionVersionsInput | None = None
    pair_rules: RuleAssistantCollectionVersionsInput | None = None
    settings_revision: StrictStr | None = None


class RuleAssistantApplyInput(BaseModel):
    settings: OptimizationSettingsInput
    pair_rule_candidates: list[RuleAssistantPairRuleInput] = Field(default_factory=list)
    staff_updates: list[RuleAssistantStaffUpdateInput] = Field(default_factory=list)
    input_versions: RuleAssistantInputVersionsInput | None = None


class AssignmentLockInput(BaseModel):
    locked: bool
    command_precondition: CandidatePrecondition | None = None


class ScheduleCellInput(BaseModel):
    command_precondition: CandidatePrecondition | None = None
    work_kind: str
    shift_template_id: int | None = None
    position_id: int | None = None
    reason: str | None = Field(default=None, min_length=1, max_length=500)


class ScheduleBulkCellInput(ScheduleCellInput):
    staff_id: int
    work_date: date


class ScheduleBulkInput(BaseModel):
    command_precondition: CandidatePrecondition | None = None
    cells: list[ScheduleBulkCellInput] = Field(min_length=1)


class PdfPage5SeedInput(BaseModel):
    confirmation: str
    reset_data: bool = True


def _require_dev_seed_schema(session: Session) -> None:
    from alembic.config import Config
    from alembic.script import ScriptDirectory

    config = Config(str(Path(__file__).resolve().parents[1] / "alembic.ini"))
    assert_runtime_schema(session.connection(), ScriptDirectory.from_config(config))


@app.on_event("startup")
def startup() -> None:
    from alembic.config import Config
    from alembic.script import ScriptDirectory

    config = Config(str(Path(__file__).resolve().parents[1] / "alembic.ini"))
    with get_engine().connect() as connection:
        assert_runtime_schema(connection, ScriptDirectory.from_config(config))


@app.exception_handler(ValueError)
def value_error_handler(_request: Request, error: ValueError) -> JSONResponse:
    return JSONResponse(status_code=422, content={"detail": str(error)})


@app.exception_handler(StorageError)
def storage_error_handler(_request: Request, error: StorageError) -> JSONResponse:
    return JSONResponse(status_code=503, content={"detail": str(error)})


@app.exception_handler(SQLAlchemyError)
def database_error_handler(_request: Request, _error: SQLAlchemyError) -> JSONResponse:
    return JSONResponse(
        status_code=503,
        content={"detail": "データベースに接続できません。設定を確認して再試行してください"},
    )


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/ready")
def ready(session: Session = Depends(get_session)) -> dict[str, str | bool]:
    session.execute(text("SELECT 1"))
    return {
        "status": "ready",
        "database": "ok",
        "service_auth": service_auth_enabled(),
    }


def _staff_half_day_template_ids(session: Session) -> dict[int, int]:
    raw = _active_settings(session).get("staff_half_day_shift_template_ids") or {}
    return {int(staff_id): int(template_id) for staff_id, template_id in raw.items()}


def _staff_display_order_ids(session: Session) -> list[int]:
    existing_ids = {row[0] for row in session.query(Staff.id).all()}
    raw = _active_settings(session).get("staff_display_order_ids") or []
    configured: list[int] = []
    for value in raw:
        staff_id = int(value)
        if staff_id in existing_ids and staff_id not in configured:
            configured.append(staff_id)
    return configured + sorted(existing_ids - set(configured))


def _save_staff_display_order(session: Session, staff_ids: list[int]) -> None:
    existing_ids = {row[0] for row in session.query(Staff.id).all()}
    if len(staff_ids) != len(set(staff_ids)) or set(staff_ids) != existing_ids:
        raise HTTPException(
            status_code=422,
            detail="表示順には登録済みの人員を重複なく全員指定してください",
        )
    row = _locked_active_config(session)
    if row is None:
        row = RuleConfig(name="active", config_json={})
        session.add(row)
    config = dict(row.config_json or {})
    config["staff_display_order_ids"] = staff_ids
    row.config_json = config


def _remove_staff_from_display_order(session: Session, staff_id: int) -> None:
    row = _locked_active_config(session)
    if row is None:
        return
    config = dict(row.config_json or {})
    config["staff_display_order_ids"] = [
        int(value)
        for value in (config.get("staff_display_order_ids") or [])
        if int(value) != staff_id
    ]
    row.config_json = config


def _save_staff_half_day_template(
    session: Session, staff_id: int, template_id: int | None
) -> None:
    row = _locked_active_config(session)
    if row is None:
        row = RuleConfig(name="active", config_json={})
        session.add(row)
    config = dict(row.config_json or {})
    mapping = dict(config.get("staff_half_day_shift_template_ids") or {})
    if template_id is None:
        mapping.pop(str(staff_id), None)
    else:
        mapping[str(staff_id)] = template_id
    config["staff_half_day_shift_template_ids"] = mapping
    row.config_json = config


def _staff_payload(
    staff: Staff,
    half_day_template_id: int | None = None,
    template_by_id: dict[int, ShiftTemplate] | None = None,
) -> dict:
    half_day_template = (template_by_id or {}).get(half_day_template_id)
    return {
        "id": staff.id,
        "name": staff.name,
        "staff_group": staff.staff_group,
        "max_daily_work_minutes": staff.max_daily_work_minutes,
        "max_weekly_work_minutes": staff.max_weekly_work_minutes,
        "max_monthly_work_minutes": staff.max_monthly_work_minutes,
        "min_monthly_work_minutes": staff.min_monthly_work_minutes,
        "target_monthly_work_minutes": staff.target_monthly_work_minutes,
        "paid_leave_eligible": staff.paid_leave_eligible,
        "half_day_shift_template_id": half_day_template_id,
        "half_day_shift_template_name": _template_display_name(half_day_template),
        "active_from": staff.active_from.isoformat() if staff.active_from else None,
        "active_to": staff.active_to.isoformat() if staff.active_to else None,
        "notes": staff.notes,
        "revision": staff.revision,
    }


def _staff_items_payload(session: Session, staff_ids: list[int]) -> list[dict]:
    requested = set(staff_ids)
    mapping = {key: value for key, value in _staff_half_day_template_ids(session).items() if key in requested}
    templates = {
        row.id: row
        for row in session.query(ShiftTemplate)
        .filter(ShiftTemplate.id.in_(set(mapping.values())))
        .all()
    } if mapping else {}
    staff_by_id = {row.id: row for row in session.query(Staff).filter(Staff.id.in_(staff_ids)).all()}
    return [
        _staff_payload(staff_by_id[staff_id], mapping.get(staff_id), templates)
        for staff_id in staff_ids
    ]


def _constraint_payload(constraint: Availability) -> dict:
    return {
        "id": constraint.id,
        "staff_id": constraint.staff_id,
        "staff_name": constraint.staff.name,
        "work_date": constraint.work_date.isoformat(),
        "status": constraint.status,
        "strength": constraint.strength,
        "weight": constraint.weight,
        "shift_template_id": constraint.shift_template_id,
        "reason": constraint.reason,
        "revision": constraint.revision,
    }


def _ensure_staff_ids_exist(session: Session, staff_ids: list[int]) -> None:
    existing = {row[0] for row in session.query(Staff.id).filter(Staff.id.in_(staff_ids)).all()} if staff_ids else set()
    missing = sorted(set(staff_ids) - existing)
    if missing:
        raise HTTPException(status_code=422, detail=f"スタッフが見つかりません: {missing}")


def _active_settings(session: Session) -> dict:
    row = session.query(RuleConfig).filter(RuleConfig.name == "active").one_or_none()
    config = row.config_json if row else {}
    return {**build_active_settings(config), "settings_revision": _settings_revision(config)}


def _settings_revision(config: dict) -> str:
    return hashlib.sha256(json.dumps(config, sort_keys=True, separators=(",", ":"), ensure_ascii=True).encode()).hexdigest()


def _locked_active_config(session: Session) -> RuleConfig:
    session.flush()
    insert = sqlite_insert if session.get_bind().dialect.name == "sqlite" else postgres_insert
    session.execute(insert(RuleConfig).values(name="active", config_json={}).on_conflict_do_nothing(index_elements=[RuleConfig.name]))
    return session.query(RuleConfig).filter(RuleConfig.name == "active").with_for_update().populate_existing().one()


def _validated_optimization_settings_payload(payload: OptimizationSettingsInput) -> dict:
    data = payload.model_dump(exclude={"settings_revision"})
    if data["standard_period_mode"] not in {"calendar_month", "fixed_start_day"}:
        raise HTTPException(
            status_code=422,
            detail="standard_period_mode は calendar_month または fixed_start_day を指定してください",
        )
    if data["standard_period_mode"] == "fixed_start_day" and data["standard_period_start_day"] is None:
        raise HTTPException(
            status_code=422,
            detail="fixed_start_day では standard_period_start_day を指定してください",
        )
    if data["standard_period_mode"] == "calendar_month":
        data["standard_period_start_day"] = None

    if (
        data["monthly_min_off_days"] is not None
        and data["monthly_max_off_days"] is not None
        and data["monthly_min_off_days"] > data["monthly_max_off_days"]
    ):
        raise HTTPException(status_code=422, detail="月の休み最小日数は最大日数以下にしてください")
    if data["weekday_off_target_min"] > data["weekday_off_target_max"]:
        raise HTTPException(
            status_code=422,
            detail="平日の休み目安の最小人数は最大人数以下にしてください",
        )
    if (
        data["weekday_off_hard_min"] is not None
        and data["weekday_off_hard_max"] is not None
        and data["weekday_off_hard_min"] > data["weekday_off_hard_max"]
    ):
        raise HTTPException(status_code=422, detail="平日の休み人数の最小は最大以下にしてください")
    for key in ("day_category_staffing_ranges", "group_day_category_staffing_ranges", "staff_day_category_staffing_ranges"):
        seen = set()
        for row in data[key]:
            identity = (
                tuple(sorted(row["staff_ids"])) if "staff_ids" in row else row.get("staff_group"),
                row["day_category"],
            )
            if identity in seen:
                raise HTTPException(status_code=422, detail="同じ対象・曜日区分の人数設定を重複登録できません")
            seen.add(identity)
            if "staff_ids" in row and len(set(row["staff_ids"])) != len(row["staff_ids"]):
                raise HTTPException(status_code=422, detail="出勤人数の対象者を重複指定できません")
            if row["min_working"] is None and row["max_working"] is None:
                raise HTTPException(status_code=422, detail="勤務人数の最小または最大を指定してください")
            if row["min_working"] is not None and row["max_working"] is not None and row["min_working"] > row["max_working"]:
                raise HTTPException(status_code=422, detail="勤務人数の最小は最大以下にしてください")
    seen_positions = set()
    for group in data["position_staffing_groups"]:
        if group["position_id"] in seen_positions:
            raise HTTPException(status_code=422, detail="同じ持ち場の主担当条件を重複登録できません")
        seen_positions.add(group["position_id"])
        if len(set(group["staff_ids"])) != len(group["staff_ids"]) or group["min_count"] > len(group["staff_ids"]):
            raise HTTPException(status_code=422, detail="主担当の対象者・最低人数を確認してください")
    seen_positions = set()
    for row in data["position_day_category_prohibitions"]:
        if row["position_id"] in seen_positions:
            raise HTTPException(status_code=422, detail="同じ担当の禁止曜日を重複登録できません")
        seen_positions.add(row["position_id"])
        if not row["day_categories"]:
            raise HTTPException(status_code=422, detail="担当を禁止する曜日区分を指定してください")
        if len(set(row["day_categories"])) != len(row["day_categories"]):
            raise HTTPException(status_code=422, detail="担当を禁止する曜日区分を重複指定できません")
    if (
        data["weekend_min_staff_count"] is not None
        and data["weekend_max_staff_count"] is not None
        and data["weekend_min_staff_count"] > data["weekend_max_staff_count"]
    ):
        raise HTTPException(status_code=422, detail="土日の勤務人数の最小は最大以下にしてください")
    if set(data["weekend_preferred_staff_ids"]) & set(data["weekend_comparison_staff_ids"]):
        raise HTTPException(status_code=422, detail="土日の人数比較で同じ人を両方のグループに登録できません")
    if data["weekend_preferred_min_advantage"] and (
        not data["weekend_preferred_staff_ids"] or not data["weekend_comparison_staff_ids"]
    ):
        raise HTTPException(status_code=422, detail="土日の人数差を設定する場合は両方のグループに人員を登録してください")
    weekend_rule_staff_ids = [row["staff_id"] for row in data["staff_weekend_rules"]]
    if len(set(weekend_rule_staff_ids)) != len(weekend_rule_staff_ids):
        raise HTTPException(status_code=422, detail="個人別の曜日・週末ルールに同じ人を重複登録できません")
    for row in data["staff_weekend_rules"]:
        raw_allowed_weekdays = row.get("allowed_weekday_offs")
        if raw_allowed_weekdays is None:
            row["allowed_weekday_offs"] = None
        else:
            allowed_weekdays = [int(value) for value in raw_allowed_weekdays]
            if any(weekday < 0 or weekday > 4 for weekday in allowed_weekdays):
                raise HTTPException(status_code=422, detail="平日休みを許可する曜日は月曜から金曜だけ指定できます")
            if len(set(allowed_weekdays)) != len(allowed_weekdays):
                raise HTTPException(status_code=422, detail="平日休みを許可する曜日を重複指定できません")
            row["allowed_weekday_offs"] = sorted(set(allowed_weekdays))
    raw_auto_half_day_allowed_weekdays = data.get("auto_half_day_allowed_weekdays")
    data["forbid_half_day_on_weekdays"] = bool(data.get("forbid_half_day_on_weekdays"))
    data["forbid_half_day_on_saturdays"] = bool(data.get("forbid_half_day_on_saturdays"))
    if raw_auto_half_day_allowed_weekdays is None:
        allowed_auto_half_day_weekdays = [
            weekday
            for weekday in range(7)
            if not (
                (weekday < 5 and data["forbid_half_day_on_weekdays"])
                or (weekday == 5 and data["forbid_half_day_on_saturdays"])
            )
        ]
        data["auto_half_day_allowed_weekdays"] = (
            None if allowed_auto_half_day_weekdays == list(range(7)) else allowed_auto_half_day_weekdays
        )
    else:
        auto_half_day_allowed_weekdays = [int(value) for value in raw_auto_half_day_allowed_weekdays]
        if any(weekday < 0 or weekday > 6 for weekday in auto_half_day_allowed_weekdays):
            raise HTTPException(status_code=422, detail="自動半日勤務を許可する曜日は日曜まで指定できます")
        if len(set(auto_half_day_allowed_weekdays)) != len(auto_half_day_allowed_weekdays):
            raise HTTPException(status_code=422, detail="自動半日勤務を許可する曜日を重複指定できません")
        data["auto_half_day_allowed_weekdays"] = sorted(set(auto_half_day_allowed_weekdays))
        data["forbid_half_day_on_weekdays"] = not any(
            weekday in data["auto_half_day_allowed_weekdays"] for weekday in range(5)
        )
        data["forbid_half_day_on_saturdays"] = 5 not in data["auto_half_day_allowed_weekdays"]
    if data["weekend_coverage_min_working"] is not None:
        if not data["weekend_coverage_staff_ids"]:
            raise HTTPException(status_code=422, detail="週末の最低出勤人数を設定する場合は対象者を選択してください")
        if data["weekend_coverage_min_working"] > len(set(data["weekend_coverage_staff_ids"])):
            raise HTTPException(status_code=422, detail="週末の最低出勤人数は対象者数以下にしてください")
    if (
        data["weekend_coverage_max_working"] is not None
        and data["weekend_coverage_staff_ids"]
        and data["weekend_coverage_max_working"] > len(set(data["weekend_coverage_staff_ids"]))
    ):
        raise HTTPException(status_code=422, detail="休日グループの最大出勤人数は対象者数以下にしてください")
    if (
        data["weekend_coverage_min_working"] is not None
        and data["weekend_coverage_max_working"] is not None
        and data["weekend_coverage_min_working"] > data["weekend_coverage_max_working"]
    ):
        raise HTTPException(status_code=422, detail="休日グループの最小出勤人数は最大以下にしてください")
    if (
        data["weekend_coverage_target_working"] is not None
        and data["weekend_coverage_staff_ids"]
        and data["weekend_coverage_target_working"] > len(set(data["weekend_coverage_staff_ids"]))
    ):
        raise HTTPException(status_code=422, detail="休日グループの目標出勤人数は対象者数以下にしてください")
    if (
        data["weekend_coverage_target_working"] is not None
        and data["weekend_coverage_min_working"] is not None
        and data["weekend_coverage_target_working"] < data["weekend_coverage_min_working"]
    ):
        raise HTTPException(status_code=422, detail="休日グループの目標出勤人数は最小以上にしてください")
    if (
        data["weekend_coverage_target_working"] is not None
        and data["weekend_coverage_max_working"] is not None
        and data["weekend_coverage_target_working"] > data["weekend_coverage_max_working"]
    ):
        raise HTTPException(status_code=422, detail="休日グループの目標出勤人数は最大以下にしてください")
    if (
        data["monthly_holiday_min_days"] is not None
        and data["monthly_holiday_max_days"] is not None
        and data["monthly_holiday_min_days"] > data["monthly_holiday_max_days"]
    ):
        raise HTTPException(
            status_code=422,
            detail="休日換算の最小日数は最大日数以下にしてください",
        )

    try:
        scale_decimal_units(
            data["monthly_holiday_exact_days"],
            data["holiday_count_unit"],
            "monthly_holiday_exact_days",
        )
        scale_decimal_units(
            data["monthly_holiday_min_days"],
            data["holiday_count_unit"],
            "monthly_holiday_min_days",
        )
        scale_decimal_units(
            data["monthly_holiday_max_days"],
            data["holiday_count_unit"],
            "monthly_holiday_max_days",
        )
        for name, value in data["holiday_counted_work_kinds"].items():
            scale_decimal_units(value, data["holiday_count_unit"], f"holiday_counted_work_kinds.{name}")
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    return data


def _solve_run_period_payload(solve_run: SolveRun) -> dict:
    from .period_metadata_v1 import display_period
    return display_period(solve_run.target_month, solve_run.config_json)


def _automatic_prior_statuses(
    session: Session,
    *,
    start_date: date,
    settings: dict,
    time_history: list[dict] | None = None,
) -> tuple[tuple[int, date, str], ...]:
    from .operational_constraints import history_days
    from .validation_snapshot import historical_time_rows
    weekly_window = 7 if int(settings.get("weekly_min_off_days", 0) or 0) else 0
    consecutive_window = int(settings.get("max_consecutive_work_days", 0) or 0) + 1
    consecutive_off_window = int(settings.get("max_consecutive_off_days", 0) or 0) + 1
    required_days = max(max(weekly_window, consecutive_window, consecutive_off_window, 7) - 1,
                        history_days(settings, session.query(Staff).all()))
    history_start = start_date - timedelta(days=required_days)
    from .schedule_selection import predecessor_runs, unresolved_predecessor_detail
    runs = predecessor_runs(session, start_date)
    if len(runs) > 1:
        confirmed_runs = [row for row in runs if session.query(ScheduleConfirmation).filter_by(solve_run_id=row.id).first()]
        selected = [row for row in runs if (row.config_json or {}).get("prior_candidate_for_start_date") == start_date.isoformat()
                    and (row.config_json or {}).get("prior_candidate_id")]
        if len(selected) > 1 or (not selected and len(confirmed_runs) != 1):
            raise HTTPException(status_code=409, detail=unresolved_predecessor_detail(session, start_date))
    run = direct_predecessor_run(session, start_date)
    if run is not None:
        period = _solve_run_period_payload(run)
        candidates = (
            session.query(CandidateSchedule)
            .filter(CandidateSchedule.solve_run_id == run.id)
            .order_by(CandidateSchedule.candidate_no)
            .all()
        )
        try:
            candidate = prior_history_candidate(session, run, start_date)
        except HTTPException as error:
            if isinstance(error.detail, str) and error.detail.startswith(("前期間の参照先", "選択済みの前期間候補", "前期間の案が選択後")):
                raise HTTPException(status_code=409, detail=unresolved_predecessor_detail(session, start_date)) from error
            raise
        if candidate is None and len(candidates) == 1:
            candidate = candidates[0]
        elif candidate is None and len(candidates) > 1:
            raise HTTPException(status_code=409, detail=unresolved_candidate_detail(run, candidates))
        if candidate is None:
            return ()
        historical_settings = (candidate.score_breakdown_json or {}).get("rule_config")
        if not isinstance(historical_settings, dict):
            return ()
        from .validation_snapshot import historical_candidate_validation
        validation = historical_candidate_validation(session, candidate) if (
            confirmed_candidate(session, run) is not None
            or (run.config_json or {}).get("prior_candidate_id") == candidate.id
        ) else validate_candidate_state(candidate.id, session)
        if validation["blocking"] or prior_history_coverage_errors(session, candidate, start_date):
            return ()
        statuses = (
            session.query(CandidateStaffStatus)
            .filter(CandidateStaffStatus.candidate_id == candidate.id)
            .filter(CandidateStaffStatus.work_date >= history_start)
            .filter(CandidateStaffStatus.work_date < start_date)
            .order_by(CandidateStaffStatus.work_date, CandidateStaffStatus.staff_id)
            .all()
        )
        history = {
            key: status for key, status in _load_prior_statuses(run.config_json or {}).items()
            if history_start <= key[1] < date.fromisoformat(period["start_date"])
        }
        history.update({
            (row.staff_id, row.work_date): "off" if row.status == "off" else "working"
            for row in statuses
        })
        if history:
            if session.info.get("solver_input_snapshot_transaction") is session.get_transaction():
                session.info["solver_prior_source"] = {
                    "solve_run_id": run.id, "candidate_id": candidate.id,
                    "history_start": history_start.isoformat(),
                    "history_end_exclusive": start_date.isoformat(),
                }
            if time_history is not None:
                period_start = date.fromisoformat(period["start_date"])
                time_history.extend(
                    row for row in (run.config_json or {}).get("prior_time_history", [])
                    if history_start <= date.fromisoformat(row["work_date"]) < period_start
                )
                time_history.extend(
                    row for row in historical_time_rows(session, candidate)
                    if max(history_start, period_start) <= date.fromisoformat(row["work_date"]) < start_date
                )
            return tuple((staff_id, work_date, status) for (staff_id, work_date), status in sorted(history.items()))
        return ()
    return ()


@app.get("/staff", responses={200: {"model": StaffCollectionResponse}})
def list_staff(session: Session = Depends(get_session)) -> dict:
    return read_collection(
        session,
        "staff",
        lambda: _staff_items_payload(session, _staff_display_order_ids(session)),
    )


@app.put("/staff/display-order", responses={200: {"model": StaffDisplayOrderResponse}})
def save_staff_display_order(
    payload: StaffDisplayOrderInput, session: Session = Depends(get_session)
) -> dict:
    begin_editor_write(session)
    claim_collection(session, "staff", payload.collection_revision)
    _save_staff_display_order(session, payload.staff_ids)
    return finalize_command(
        session,
        "staff",
        lambda: {"items": _staff_items_payload(session, payload.staff_ids)},
    )


@app.post("/staff", status_code=201, responses={201: {"model": StaffCommandResponse}})
def create_staff(payload: StaffInput, session: Session = Depends(get_session)) -> dict:
    begin_editor_write(session)
    claim_collection(session, "staff", payload.collection_revision)
    if session.query(Staff).filter(Staff.name == payload.name).first():
        raise HTTPException(status_code=409, detail="staff name already exists")
    if payload.half_day_shift_template_id is not None:
        template = session.get(ShiftTemplate, payload.half_day_shift_template_id)
        if template is None or template.shift_type != "half_day":
            raise HTTPException(status_code=422, detail="半日勤務用の勤務時間を選択してください")
    staff = Staff(**payload.model_dump(exclude={"half_day_shift_template_id", "revision", "collection_revision"}))
    session.add(staff)
    session.flush()
    _save_staff_half_day_template(session, staff.id, payload.half_day_shift_template_id)
    template = (
        session.get(ShiftTemplate, payload.half_day_shift_template_id)
        if payload.half_day_shift_template_id is not None
        else None
    )
    return finalize_command(session, "staff", lambda: _staff_payload(
        staff,
        payload.half_day_shift_template_id,
        {template.id: template} if template is not None else {},
    ))


WEEKDAY_LABELS = ["月", "火", "水", "木", "金", "土", "日"]


def _template_display_name(template: ShiftTemplate | None) -> str:
    if template is None:
        return ""
    return template.display_name or template.name or ""


def _position_display_name(position: Position | None) -> str:
    if position is None:
        return ""
    return position.display_name or position.name or ""


def _format_matrix_value(
    status: CandidateStaffStatus | None,
    assignment: Assignment | None,
    *,
    regular_work_display_blank: bool = False,
) -> str:
    if status is None:
        if assignment is None:
            return ""
        if assignment.shift_template.shift_type == "half_day":
            return ""
        if assignment.shift_template.shift_type == "regular":
            return "" if regular_work_display_blank else _template_display_name(assignment.shift_template)
        return _position_display_name(assignment.position)
    if status.status == "off":
        return "○"
    if status.work_kind == "position":
        return _position_display_name(assignment.position if assignment is not None else None)
    if status.work_kind == PAID_LEAVE_HALF_DAY_WORK_KIND:
        return "有0.5"
    if status.work_kind == "half_day":
        return ""
    if status.work_kind == "regular":
        return "" if regular_work_display_blank else _template_display_name(status.shift_template)
    if status.shift_template is not None and status.shift_template.shift_type == "half_day":
        return ""
    if status.shift_template is not None and status.shift_template.shift_type == "regular":
        return "" if regular_work_display_blank else _template_display_name(status.shift_template)
    return ""


def _is_half_day_entry(status: CandidateStaffStatus | None, assignment: Assignment | None) -> bool:
    return bool(
        (status is not None and status.work_kind == "half_day")
        or (status is None and assignment is not None and assignment.shift_template.shift_type == "half_day")
    )


def _period_label(dates: list[date]) -> str:
    first = dates[0]
    last = dates[-1]
    return f"{first.month}/{first.day}-{last.month}/{last.day}"


def _build_shift_matrix_sheet(
    sheet,
    solve_run,
    assignments: list[Assignment],
    statuses: list[CandidateStaffStatus],
    session,
    assignment_cells=None,
) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    MATRIX_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
    MATRIX_BORDER = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    DIAGONAL_BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
        diagonal=Side(style="thin", color="FF000000"),
        diagonalDown=True,
    )
    WEEKEND_FILL = PatternFill(fill_type="solid", fgColor="FFF9E79F")
    TRAINEE_ROW_FILL = PatternFill(fill_type="solid", fgColor="FFFDEAAE")
    SCHOOL_AD_ROW_FILL = PatternFill(fill_type="solid", fgColor="FFB9D9EB")
    SUNDAY_FONT = Font(color="FFFF0000")
    SATURDAY_FONT = Font(color="FF0000FF")

    regular_work_display_blank = bool(_active_settings(session).get("regular_work_display_blank"))
    run_period = _solve_run_period_payload(solve_run)
    assignment_by_staff_date = assignment_cells if assignment_cells is not None else group_candidate_assignments(assignments)
    status_by_staff_date = {(row.staff_id, row.work_date): row for row in statuses}
    result_staff_ids = {row.staff_id for row in assignments} | {row.staff_id for row in statuses}
    staff_ids = [
        staff_id for staff_id in _staff_display_order_ids(session) if staff_id in result_staff_ids
    ]
    dates = sorted({row.work_date for row in assignments} | {row.work_date for row in statuses})
    staff_rows = [
        (staff_id, slot)
        for staff_id in staff_ids
        for slot in range(max(1, max((len(assignment_by_staff_date.get((staff_id, day), [])) for day in dates), default=0)))
    ]
    multiple_dates = {day for day in dates if any(len(assignment_by_staff_date.get((staff_id, day), [])) > 1 for staff_id in staff_ids)}

    staffs = (
        {
            staff.id: staff
            for staff in (
                session.query(Staff)
                .filter(Staff.id.in_(staff_ids))
                .order_by(Staff.id)
                .all()
                if staff_ids
                else []
            )
        }
        if staff_ids
        else {}
    )

    if dates:
        sheet["A1"] = _period_label(dates)
    else:
        if (
            (solve_run.config_json or {}).get("start_date")
            and run_period["start_date"]
            and run_period["end_date_exclusive"]
        ):
            start = date.fromisoformat(run_period["start_date"])
            end = date.fromisoformat(run_period["end_date_exclusive"]) - date.resolution
            sheet["A1"] = f"{start.month}/{start.day}-{end.month}/{end.day}"
        else:
            sheet["A1"] = f"{solve_run.target_month} シフト"
    sheet.merge_cells("A1:A2")
    sheet["A1"].alignment = MATRIX_ALIGNMENT

    for column_index, work_date in enumerate(dates, start=2):
        column_letter = get_column_letter(column_index)
        sheet[f"{column_letter}1"] = f"{work_date.month}/{work_date.day}"
        sheet[f"{column_letter}2"] = WEEKDAY_LABELS[work_date.weekday()]
        if work_date.weekday() == 5:
            for row in (1, 2):
                sheet[f"{column_letter}{row}"].fill = WEEKEND_FILL
                sheet[f"{column_letter}{row}"].font = SATURDAY_FONT
        elif work_date.weekday() == 6:
            for row in (1, 2):
                sheet[f"{column_letter}{row}"].fill = WEEKEND_FILL
                sheet[f"{column_letter}{row}"].font = SUNDAY_FONT
    wrapped_rows = set()
    for row_offset, (staff_id, slot) in enumerate(staff_rows):
        staff = staffs[staff_id]
        row_number = 3 + row_offset
        staff_row_fill = (
            TRAINEE_ROW_FILL
            if staff.name in PRACTICAL_TRAINEES
            else SCHOOL_AD_ROW_FILL if staff.name in SCHOOL_AD_STAFF else None
        )
        sheet.cell(row=row_number, column=1, value=_neutralize_spreadsheet_cell(staff.name))
        if staff_row_fill is not None:
            sheet.cell(row=row_number, column=1).fill = staff_row_fill

        for column_index, work_date in enumerate(dates, start=2):
            status = status_by_staff_date.get((staff_id, work_date))
            cell_assignments = assignment_by_staff_date.get((staff_id, work_date), [])
            assignment = cell_assignments[slot] if slot < len(cell_assignments) else None
            if len(cell_assignments) > 1 and assignment is not None:
                value = f"{_position_display_name(assignment.position)} / {_template_display_name(assignment.shift_template)}"
                wrapped_rows.add(row_number)
            elif slot == 0:
                value = _format_matrix_value(status, assignment, regular_work_display_blank=regular_work_display_blank)
            else:
                value = ""
                status = None
            if len(cell_assignments) > 1 and slot == 0 and status is not None and status.work_kind != "position":
                status_label = _format_matrix_value(status, None, regular_work_display_blank=regular_work_display_blank)
                if status_label:
                    value = f"{status_label}\n{value}"
            matrix_cell = sheet.cell(
                row=row_number,
                column=column_index,
                value=_neutralize_spreadsheet_cell(value),
            )
            if _is_half_day_entry(status, assignment):
                if len(cell_assignments) <= 1:
                    matrix_cell.value = ""
                matrix_cell.border = DIAGONAL_BORDER
            if work_date.weekday() == 5:
                matrix_cell.fill = WEEKEND_FILL
            elif work_date.weekday() == 6:
                matrix_cell.fill = WEEKEND_FILL
            elif staff_row_fill is not None:
                matrix_cell.fill = staff_row_fill

    max_row = 2 + len(staff_rows)
    max_column = 1 + len(dates)
    for row in range(1, max_row + 1):
        for column in range(1, max_column + 1):
            cell = sheet.cell(row=row, column=column)
            cell.alignment = MATRIX_ALIGNMENT
            compact_font = copy(cell.font)
            compact_font.name = "Arial Unicode MS"
            compact_font.sz = 8
            cell.font = compact_font
            if cell.border.diagonal.style is None:
                cell.border = MATRIX_BORDER

    sheet.column_dimensions["A"].width = 12.5
    for column_index in range(2, max_column + 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = 18 if dates[column_index - 2] in multiple_dates else 4.5
    sheet.row_dimensions[1].height = 18
    sheet.row_dimensions[2].height = 18
    for row_number in range(3, max_row + 1):
        sheet.row_dimensions[row_number].height = None if row_number in wrapped_rows else 22

    sheet.freeze_panes = "B3"
    sheet.print_title_rows = "1:2"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 0 if multiple_dates else 1
    sheet.page_setup.fitToHeight = 0 if multiple_dates else 1
    if multiple_dates:
        sheet.print_title_cols = "A:A"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3
    print_area_end_column = get_column_letter(max(max_column, 1))
    print_area_end_row = max(max_row, 3)
    sheet.print_area = f"A1:{print_area_end_column}{print_area_end_row}"
@app.put("/staff/{staff_id}", responses={200: {"model": StaffCommandResponse}})
def update_staff(staff_id: int, payload: StaffInput, session: Session = Depends(get_session)) -> dict:
    begin_editor_write(session)
    claim_collection(session, "staff", payload.collection_revision)
    claim_item(session, Staff, staff_id, payload.revision, "staff_revision")
    staff = session.get(Staff, staff_id)
    assert staff is not None
    duplicate = session.query(Staff).filter(Staff.name == payload.name, Staff.id != staff_id).first()
    if duplicate:
        raise HTTPException(status_code=409, detail="staff name already exists")
    if payload.half_day_shift_template_id is not None:
        template = session.get(ShiftTemplate, payload.half_day_shift_template_id)
        if template is None or template.shift_type != "half_day":
            raise HTTPException(status_code=422, detail="半日勤務用の勤務時間を選択してください")
    for field, value in payload.model_dump(exclude={"half_day_shift_template_id", "revision", "collection_revision"}).items():
        setattr(staff, field, value)
    _save_staff_half_day_template(session, staff.id, payload.half_day_shift_template_id)
    template = session.get(ShiftTemplate, payload.half_day_shift_template_id) if payload.half_day_shift_template_id is not None else None
    return finalize_command(session, "staff", lambda: _staff_payload(staff, payload.half_day_shift_template_id, {template.id: template} if template is not None else {}))


@app.delete("/staff/{staff_id}", responses={200: {"model": StaffDeleteResponse}})
def delete_staff(staff_id: int, revision: int | None = None, collection_revision: int | None = None, session: Session = Depends(get_session)) -> dict:
    begin_editor_write(session)
    claim_collection(session, "staff", collection_revision)
    claim_item(session, Staff, staff_id, revision, "staff_revision")
    staff = session.get(Staff, staff_id)
    assert staff is not None
    references = [
        (session.query(StaffSkill).filter(StaffSkill.staff_id == staff_id), "担当できる仕事"),
        (session.query(Availability).filter(Availability.staff_id == staff_id), "休み・勤務できない日"),
        (session.query(Assignment).filter(Assignment.staff_id == staff_id), "作成済みシフト"),
        (session.query(CandidateStaffStatus).filter(CandidateStaffStatus.staff_id == staff_id), "作成済み勤務状態"),
        (
            session.query(ReferenceAssignment).filter(ReferenceAssignment.staff_id == staff_id),
            "確認用の過去シフト",
        ),
        (session.query(PositionPriority).filter(PositionPriority.staff_id == staff_id), "持ち場の優先順"),
        (
            session.query(StaffPositionPreference).filter(StaffPositionPreference.staff_id == staff_id),
            "個別の担当希望",
        ),
        (
            session.query(StaffPositionRule).filter(StaffPositionRule.staff_id == staff_id),
            "持ち場の許可・禁止条件",
        ),
        (session.query(ScheduleException).filter(ScheduleException.staff_id == staff_id), "特別な予定"),
        (session.query(PositionAssignmentRule).filter(PositionAssignmentRule.staff_id == staff_id), "期間付き担当設定"),
        (session.query(Violation).filter(Violation.staff_id == staff_id), "作成時の確認結果"),
        (
            session.query(StaffPairRule).filter(
                (StaffPairRule.first_staff_id == staff_id) | (StaffPairRule.second_staff_id == staff_id)
            ),
            "人の組み合わせ条件",
        ),
    ]
    for query, label in references:
        if query.first():
            raise HTTPException(status_code=409, detail=f"{label}に使われている人員は削除できません")
    if any(staff_id in (position.required_staff_ids or []) for position in session.query(Position).all()):
        raise HTTPException(status_code=409, detail="持ち場の必須担当に設定されている人員は削除できません")
    if any(staff_id in policy.member_staff_ids for policy in session.query(RotationPolicy).all()):
        raise HTTPException(status_code=409, detail="交代順に入っている人員は削除できません")
    if any(staff_id in state.queue_staff_ids for state in session.query(RotationQueueState).all()):
        raise HTTPException(status_code=409, detail="確定済みの交代順に入っている人員は削除できません")
    settings = _active_settings(session)
    configured_ids = {
        person_id for key in ("weekday_off_exempt_staff_ids", "weekend_preferred_staff_ids", "weekend_comparison_staff_ids", "weekend_coverage_staff_ids")
        for person_id in settings.get(key, [])
    }
    configured_ids.update(row["staff_id"] for row in settings.get("staff_weekend_rules", []))
    configured_ids.update(person_id for row in settings.get("position_staffing_groups", []) for person_id in row["staff_ids"])
    configured_ids.update(person_id for row in settings.get("staff_day_category_staffing_ranges", []) for person_id in row["staff_ids"])
    if staff_id in configured_ids or any(staff_id in (row.full_day_staff_ids or []) for row in session.query(WorkdayRule).all()):
        raise HTTPException(status_code=409, detail="作り方や日別勤務に設定されている人員は、設定を外してから削除してください")
    _save_staff_half_day_template(session, staff_id, None)
    _remove_staff_from_display_order(session, staff_id)
    session.delete(staff)
    return finalize_command(session, "staff", lambda: {"id": staff_id, "revision": revision + 1, "deleted": True})


@app.get("/constraints", responses={200: {"model": ConstraintCollectionResponse}})
def list_constraints(session: Session = Depends(get_session), date_from: date | None = None, date_to: date | None = None) -> dict:
    if not isinstance(date_from, date) or not isinstance(date_to, date) or date_from >= date_to:
        raise HTTPException(422, "予定の期間を開始日・終了日（含まない）で指定してください")
    def items():
        query = session.query(Availability).options(joinedload(Availability.staff))
        if date_from:
            query = query.filter(Availability.work_date >= date_from)
        if date_to:
            query = query.filter(Availability.work_date < date_to)
        return [_constraint_payload(row) for row in query.order_by(Availability.work_date.desc(), Availability.id.desc())]
    return read_collection(session, "constraints", items)


@app.post("/constraints", status_code=201, responses={201: {"model": ConstraintCommandResponse}})
def create_constraint(payload: ConstraintInput, session: Session = Depends(get_session)) -> dict:
    begin_editor_write(session)
    claim_collection(session, "constraints", payload.collection_revision)
    if session.get(Staff, payload.staff_id) is None:
        raise HTTPException(status_code=404, detail="staff not found")
    if payload.shift_template_id is not None:
        _validate_constraint_template(session, payload.shift_template_id, payload.status, missing_status=404)
    constraint = Availability(**payload.model_dump(exclude={"revision", "collection_revision"}))
    session.add(constraint)
    return finalize_command(session, "constraints", lambda: _constraint_payload(constraint))


@app.put("/constraints/{constraint_id}", responses={200: {"model": ConstraintCommandResponse}})
def update_constraint(constraint_id: int, payload: ConstraintInput, session: Session = Depends(get_session)) -> dict:
    begin_editor_write(session)
    claim_collection(session, "constraints", payload.collection_revision)
    claim_item(session, Availability, constraint_id, payload.revision, "constraint_revision")
    constraint = session.get(Availability, constraint_id)
    assert constraint is not None
    if session.get(Staff, payload.staff_id) is None:
        raise HTTPException(status_code=404, detail="staff not found")
    if payload.shift_template_id is not None:
        _validate_constraint_template(session, payload.shift_template_id, payload.status, missing_status=404)
    for field, value in payload.model_dump(exclude={"revision", "collection_revision"}).items():
        setattr(constraint, field, value)
    return finalize_command(session, "constraints", lambda: _constraint_payload(constraint))


@app.delete("/constraints/{constraint_id}", responses={200: {"model": ConstraintDeleteResponse}})
def delete_constraint(constraint_id: int, revision: int | None = None, collection_revision: int | None = None, session: Session = Depends(get_session)) -> dict:
    begin_editor_write(session)
    claim_collection(session, "constraints", collection_revision)
    claim_item(session, Availability, constraint_id, revision, "constraint_revision")
    constraint = session.get(Availability, constraint_id)
    assert constraint is not None
    session.delete(constraint)
    return finalize_command(session, "constraints", lambda: {"id": constraint_id, "revision": revision + 1, "deleted": True})


def _validate_constraint_template(session: Session, template_id: int, status: str, *, missing_status: int = 422) -> ShiftTemplate:
    template = session.get(ShiftTemplate, template_id)
    if template is None:
        raise HTTPException(status_code=missing_status, detail="shift template not found" if missing_status == 404 else "勤務時間が見つかりません")
    if status == "half_day" and template.shift_type != "half_day":
        raise HTTPException(status_code=422, detail="半日勤務用の勤務時間を選択してください")
    return template


def _bulk_constraint_expected_existing(cell: ConstraintBulkCellInput) -> dict[int, int]:
    if "existing" not in cell.model_fields_set:
        raise HTTPException(status_code=428, detail={
            "code": "constraints_bulk_existing_required",
            "message": "各セルの読取時の既存制約一覧を指定してください。",
        })
    if cell.existing is None:
        raise HTTPException(status_code=422, detail={
            "code": "constraints_bulk_existing_invalid",
            "message": "既存制約一覧は配列で指定してください。",
        })
    expected: dict[int, int] = {}
    for item in cell.existing:
        if item.id <= 0:
            raise HTTPException(status_code=422, detail={
                "code": "constraints_bulk_existing_invalid",
                "message": "既存制約IDは1以上の整数で指定してください。",
            })
        if item.id in expected:
            raise HTTPException(status_code=422, detail={
                "code": "constraints_bulk_existing_invalid",
                "message": "同じ既存制約IDを重複して指定できません。",
            })
        expected[item.id] = require_revision(item.revision, "constraint_revision")
    return expected


@app.put("/constraints-bulk", responses={200: {"model": ConstraintBulkResponse}})
def save_constraints_bulk(
    payload: ConstraintBulkInput,
    session: Session = Depends(get_session),
) -> dict:
    """Atomically apply one exception value to selected staff/date cells."""
    begin_editor_write(session)
    keys = [(cell.staff_id, cell.work_date) for cell in payload.cells]
    if len(keys) != len(set(keys)):
        raise HTTPException(status_code=422, detail="同じ人員・日付を重複して指定できません")
    _ensure_staff_ids_exist(session, [cell.staff_id for cell in payload.cells])
    expected_by_key = {key: _bulk_constraint_expected_existing(cell) for key, cell in zip(keys, payload.cells, strict=True)}
    existing_rows = session.query(Availability).filter(
        tuple_(Availability.staff_id, Availability.work_date).in_(keys)
    ).all()
    existing_by_key: dict[tuple[int, date], list[Availability]] = {}
    for row in existing_rows:
        existing_by_key.setdefault((row.staff_id, row.work_date), []).append(row)
    for key, expected in expected_by_key.items():
        actual = {row.id: row.revision for row in existing_by_key.get(key, [])}
        if actual != expected:
            session.rollback()
            raise conflict("constraints_bulk_existing")

    try:
        operations = {"request_off", "unavailable", "half_day", "request_work", "clear"}
        if payload.operation not in operations:
            raise HTTPException(status_code=422, detail="一括設定の内容が正しくありません")
        shift_template_id = None
        if payload.operation == "half_day":
            if payload.shift_template_id is None:
                raise HTTPException(status_code=422, detail="半日勤務には勤務時間を選択してください")
            shift_template = _validate_constraint_template(session, payload.shift_template_id, "half_day")
            shift_template_id = shift_template.id
        elif payload.shift_template_id is not None:
            raise HTTPException(status_code=422, detail="勤務時間は半日勤務を設定するときだけ選択できます")

        claim_collection(session, "constraints", payload.collection_revision)
        for key in keys:
            for row in existing_by_key.get(key, []):
                claim_item(session, Availability, row.id, expected_by_key[key][row.id], "constraint_revision")

        changed = 0
        cleared = 0
        for key in keys:
            rows = existing_by_key.get(key, [])
            if payload.operation == "clear":
                for row in rows:
                    session.delete(row)
                    cleared += 1
                continue
            status = payload.operation
            strength = "hard" if status in {"unavailable", "half_day"} else payload.strength
            if len(rows) == 1:
                row = rows[0]
            else:
                for row in rows:
                    session.delete(row)
                row = Availability(staff_id=key[0], work_date=key[1])
                session.add(row)
            row.status = status
            row.strength = strength
            row.weight = payload.weight
            row.shift_template_id = shift_template_id
            row.reason = payload.reason
            changed += 1
        return finalize_command(session, "constraints", lambda: {
            "selected_count": len(keys), "changed_count": changed, "cleared_count": cleared,
            "items": [_constraint_payload(row) for row in session.query(Availability).filter(tuple_(Availability.staff_id, Availability.work_date).in_(keys)).order_by(Availability.id).all()],
        })
    except Exception:
        session.rollback()
        raise


@app.get("/optimization-settings", responses={200: {"model": OptimizationSettingsResponse}})
def get_optimization_settings(session: Session = Depends(get_session)) -> dict:
    return _active_settings(session)


def _store_optimization_settings(payload: OptimizationSettingsInput, session: Session) -> dict:
    if payload.settings_revision is None:
        raise HTTPException(status_code=428, detail={"code": "settings_revision_required", "message": "設定を再読み込みしてから保存してください。"})
    row = _locked_active_config(session)
    if payload.settings_revision != _settings_revision(row.config_json if row else {}):
        raise HTTPException(status_code=409, detail={"code": "settings_revision_conflict", "message": "別の画面で設定が更新されています。変更は保存していません。設定を再読み込みしてからやり直してください。"})
    validated = _validated_optimization_settings_payload(payload)
    _ensure_staff_ids_exist(
        session,
        validated["weekday_off_exempt_staff_ids"]
        + validated["weekend_preferred_staff_ids"]
        + validated["weekend_comparison_staff_ids"]
        + validated["weekend_coverage_staff_ids"]
        + [row["staff_id"] for row in validated["staff_weekend_rules"]]
        + [staff_id for key in ("position_staffing_groups", "staff_day_category_staffing_ranges") for group in validated[key] for staff_id in group["staff_ids"]],
    )
    position_ids = [row["position_id"] for key in ("position_day_category_prohibitions", "position_staffing_groups") for row in validated[key]]
    existing_position_ids = {
        row[0] for row in session.query(Position.id).filter(Position.id.in_(position_ids)).all()
    } if position_ids else set()
    missing_position_ids = sorted(set(position_ids) - existing_position_ids)
    if missing_position_ids:
        raise HTTPException(status_code=422, detail=f"担当が見つかりません: {missing_position_ids}")
    validated["staff_half_day_shift_template_ids"] = {
        str(staff_id): template_id
        for staff_id, template_id in _staff_half_day_template_ids(session).items()
    }
    validated["staff_display_order_ids"] = _staff_display_order_ids(session)
    if row is None:
        row = RuleConfig(name="active", config_json={})
        session.add(row)
    row.config_json = validated
    return validated


@app.put("/optimization-settings", responses={200: {"model": OptimizationSettingsResponse}})
def save_optimization_settings(payload: OptimizationSettingsInput, session: Session = Depends(get_session)) -> dict:
    _store_optimization_settings(payload, session)
    session.commit()
    return _active_settings(session)


@app.post("/dev/seed")
def seed(session: Session = Depends(get_session)) -> dict[str, str]:
    if os.getenv("ENABLE_DEV_SEED", "false").strip().lower() not in {"1", "true", "yes"}:
        raise HTTPException(status_code=404, detail="Not Found")
    _require_dev_seed_schema(session)
    seed_sample_data(session)
    return {"status": "seeded"}


@app.post("/dev/seed/pdf-page5")
def seed_pdf_page5(payload: PdfPage5SeedInput, session: Session = Depends(get_session)) -> dict:
    from .solver import solve_period

    if os.getenv("ENABLE_DEV_SEED", "false").strip().lower() not in {"1", "true", "yes"}:
        raise HTTPException(status_code=404, detail="Not Found")
    if payload.confirmation != "RESET_SHIFT_STG_FOR_PDF_PAGE5":
        raise HTTPException(status_code=422, detail="確認文字列が一致しません")

    _require_dev_seed_schema(session)

    if payload.reset_data:
        with dev_seed_write_snapshot(session):
            for table in reversed(Base.metadata.sorted_tables):
                if table is not EditorRevision.__table__:
                    session.execute(table.delete())
            seed_pdf_page5_full_manual_data(session)

    try:
        # Finish schema/seed setup before the solver owns its input transaction.
        session.commit()
        solve_run = solve_period(
            session=session,
            label="pdf-2026-07-21_2026-08-20-full-page5",
            start_date=PDF_PAGE5_START,
            end_date=PDF_PAGE5_END_EXCLUSIVE,
            strategy="balanced",
            reference_source=None,
            execution_request={"producer": "seed_pdf_page5", "request": payload.model_dump(mode="json")},
        )
    except SQLAlchemyError:
        session.rollback()
        logger.exception("PDF page 5 solve failed after canonical fixture seeding")
        raise HTTPException(status_code=500, detail="page5の計算中にDBエラーが発生しました")
    candidate = (
        session.query(CandidateSchedule)
        .filter(CandidateSchedule.solve_run_id == solve_run.id)
        .one_or_none()
    )
    mismatches = compare_reference_cells(
        session=session,
        candidate_id=candidate.id if candidate else None,
        source_name=PDF_REFERENCE_SOURCE_PAGE5_FULL,
        staff_names=PDF_STAFF_NAMES,
        start_date=PDF_PAGE5_START,
        end_date=PDF_PAGE5_END_EXCLUSIVE,
    )
    return {
        "status": "seeded_and_solved",
        "solve_run_id": solve_run.id,
        "solve_status": solve_run.status,
        "solver_status": solve_run.solver_status,
        "period_start": PDF_PAGE5_START.isoformat(),
        "period_end": (PDF_PAGE5_END_EXCLUSIVE - date.resolution).isoformat(),
        "staff_count": len(PDF_STAFF_NAMES),
        "reference_status_count": reference_status_count(
            session, PDF_REFERENCE_SOURCE_PAGE5_FULL
        ),
        "cell_count": len(PDF_STAFF_NAMES) * (PDF_PAGE5_END_EXCLUSIVE - PDF_PAGE5_START).days,
        "cell_mismatch_count": len(mismatches),
        "mismatches": mismatches,
    }


@app.post("/solve", responses={200: {"model": SolveResponse}})
def solve(request: SolveRequest, session: Session = Depends(get_session)) -> dict:
    from .solve_delivery_dispatch_service import configured_transport
    try:
        transport = configured_transport(os.environ)
    except (ValueError, TypeError) as error:
        raise HTTPException(503, "shift_transport_configuration_required") from error
    try:
        initialize_initial_profile(session)
        session.commit()
    except ValueError as error:
        session.rollback()
        raise HTTPException(
            status_code=409,
            detail={"code": "rule_profile_unresolved", "message": str(error)},
        ) from error
    begin_solver_input_snapshot(session)
    try:
        settings = effective_rule_config(session)
    except ValueError as error:
        session.rollback()
        raise HTTPException(status_code=409, detail={
            "code": "rule_profile_unresolved", "message": str(error),
        }) from error
    try:
        resolved_period = resolve_period_request(
            target_month=request.target_month,
            start_date=request.start_date,
            end_date_exclusive=request.end_date_exclusive,
            settings_config=settings,
        )
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    preflight = analyze_rule_feasibility(
        start_date=resolved_period.start_date,
        end_date_exclusive=resolved_period.end_date_exclusive,
        settings=settings,
        staffs=session.query(Staff).all(),
        demands=session.query(Demand).filter(
            Demand.work_date >= resolved_period.start_date,
            Demand.work_date < resolved_period.end_date_exclusive,
        ).all(),
        positions=session.query(Position).all(),
    )
    _add_rotation_preflight(preflight, session, resolved_period.start_date, resolved_period.end_date_exclusive)
    if preflight["blocking"]:
        raise HTTPException(
            status_code=409,
            detail={"message": "作成前に条件の調整が必要です", "preflight": preflight},
        )
    prior_statuses = tuple(
        (row.staff_id, row.work_date, row.status) for row in request.prior_statuses
    )
    from .operational_constraints import history_days
    prior_time_history = [row.model_dump(mode="json") for row in request.prior_time_history]
    requires_history = bool(settings.get("require_prior_period_continuity") or history_days(settings, session.query(Staff).all()))
    has_prior_staff = any(
        (staff.active_from is None or staff.active_from < resolved_period.start_date)
        and (staff.active_to is None or staff.active_to >= resolved_period.start_date)
        for staff in session.query(Staff).all()
    )
    if not prior_statuses and requires_history and has_prior_staff:
        prior_statuses = _automatic_prior_statuses(
            session,
            start_date=resolved_period.start_date,
            settings=settings,
            time_history=prior_time_history,
        )
        if not prior_statuses:
            raise HTTPException(
                status_code=409,
                detail={"code": "prior_history_required", "message": "前期間を含めて勤務・休暇を確認するため、直前期間の確定済みまたは有効なシフトが必要です。"},
            )
    from .solve_delivery_service import enqueue_configured_solve
    solve_run = enqueue_configured_solve(
        session,
        SolveJobRequest(
            target_month=request.target_month,
            start_date=request.start_date,
            end_date_exclusive=request.end_date_exclusive,
            strategy=request.strategy,
            locked_assignment_ids=tuple(request.locked_assignment_ids),
            candidate_count=request.candidate_count,
            prior_statuses=prior_statuses,
            prior_time_history=tuple(prior_time_history),
            effective_settings=settings,
            original_request=request.model_dump(mode="json"),
        ), config=transport,
    )
    if solve_run.status == "queued":
        _dispatch_saved_solve(session, solve_run.id)
    return _solve_response(session, solve_run)


def _solve_response(session, solve_run):
    if solve_run is None:  # pragma: no cover - the job contract guarantees this row
        raise HTTPException(status_code=503, detail="計算結果を読み込めませんでした")
    candidates = (
        session.query(CandidateSchedule)
        .filter(CandidateSchedule.solve_run_id == solve_run.id)
        .order_by(CandidateSchedule.candidate_no)
        .all()
    )
    selected_candidate = candidates[0] if candidates else None
    assignments_count = (
        session.query(Assignment).filter(Assignment.candidate_id == selected_candidate.id).count()
        if selected_candidate
        else 0
    )
    violations_count = session.query(Violation).filter(Violation.solve_run_id == solve_run.id).count()
    return {
        "solve_run_id": solve_run.id,
        **_solve_run_period_payload(solve_run),
        "job_status": serialize_job_status(session, solve_run),
        "status": solve_run.status,
        "solver_status": solve_run.solver_status,
        "candidates": [
            {
                "id": candidate.id,
                "strategy": candidate.strategy,
                "total_score": candidate.total_score,
                "score_breakdown": candidate.score_breakdown_json,
            }
            for candidate in candidates
        ],
        "assignments_count": assignments_count,
        "violations_count": violations_count,
    }


def _dispatch_saved_solve(session, run_id):
    from . import agent_trial_transport as local
    if local.enabled():
        local.dispatch(session, run_id)
        return
    from google.cloud import tasks_v2
    from .models import SolveDelivery
    from .solve_delivery_dispatch_service import DispatchError, dispatch_delivery, pending_delivery
    from .solve_delivery_service import DeliveryFenced
    delivery = session.query(SolveDelivery).filter_by(solve_run_id=run_id).one_or_none()
    if delivery is None:
        raise HTTPException(409, "shift_task_pending_intent_required")
    try:
        pending_delivery(session, delivery.id)
    except (ValueError, DeliveryFenced, DispatchError) as error:
        session.rollback()
        raise HTTPException(409, "shift_delivery_pending_capture_required") from error
    except SQLAlchemyError as error:
        session.rollback()
        raise HTTPException(409, "shift_delivery_busy") from error
    acknowledged = session.query(Violation).filter_by(solve_run_id=run_id, rule_name="shift_task_dispatched").first()
    if acknowledged is not None:
        raise HTTPException(409, "shift_task_already_dispatched")
    try:
        client = tasks_v2.CloudTasksClient()
    except Exception:
        _record_dispatch_warning(session, run_id)
        return
    try:
        try:
            dispatch_delivery(session, delivery.id, client=client)
        finally:
            try:
                client.transport.close()
            except Exception:
                logger.warning("shift_task_client_cleanup_failed")
    except (ValueError, DeliveryFenced) as error:
        session.rollback()
        raise HTTPException(409, "shift_delivery_capture_invalid") from error
    except DispatchError as error:
        if str(error) not in {"shift_task_dispatch_failed", "shift_task_lookup_failed"}:
            session.rollback()
            raise HTTPException(409, "shift_task_contract_unresolved") from error
        _record_dispatch_warning(session, run_id)
        return
    except SQLAlchemyError as error:
        session.rollback()
        raise HTTPException(409, "shift_delivery_busy") from error
    session.query(Violation).filter_by(solve_run_id=run_id, rule_name="shift_task_dispatch_pending").delete()
    session.add(Violation(solve_run_id=run_id, severity="info", rule_name="shift_task_dispatched",
                          message="保存済みの計算を配送しました。計算完了ではありません。", penalty=0))
    session.commit()


def _record_dispatch_warning(session, run_id):
    if session.query(Violation).filter_by(solve_run_id=run_id, rule_name="shift_task_dispatch_pending").first() is None:
        session.add(Violation(solve_run_id=run_id, severity="warning", rule_name="shift_task_dispatch_pending",
                              message="計算は受付済みですが配送を確認できません。同じ作成履歴から再配送してください。", penalty=0))
    session.commit()


@app.post("/solve-runs/{run_id}/dispatch", responses={200: {"model": SolveResponse}})
def redispatch_solve(run_id: int, session: Session = Depends(get_session)) -> dict:
    run = session.get(SolveRun, run_id)
    if run is None:
        raise HTTPException(404, "solve_run_not_found")
    _dispatch_saved_solve(session, run_id)
    session.refresh(run)
    return _solve_response(session, run)


def _add_rotation_preflight(preflight: dict, session: Session, start: date, end: date) -> None:
    policies = session.query(RotationPolicy).filter(RotationPolicy.enabled.is_(True)).all()
    errors = validate_rotation_period(session, policies, start, end)
    if not errors:
        from .solver import dated_assignment_preflight
        try:
            errors.extend(dated_assignment_preflight(session, start, end, effective_rule_config(session)))
        except ValueError as error:
            errors.append(str(error))
    preflight["issues"].extend(
        {"code": f"rotation_period_conflict_{index}", "severity": "error", "message": message, "suggestions": []}
        for index, message in enumerate(errors)
    )
    preflight["blocking"] = preflight["blocking"] or bool(errors)


@app.post("/rule-preflight", responses={200: {"model": RulePreflightResponse}})
def rule_preflight(request: RulePreflightRequest, session: Session = Depends(get_session)) -> dict:
    try:
        resolved_period = resolve_period_request(
            target_month=request.target_month,
            start_date=request.start_date,
            end_date_exclusive=request.end_date_exclusive,
            settings_config=_active_settings(session),
        )
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    preflight = analyze_rule_feasibility(
        start_date=resolved_period.start_date,
        end_date_exclusive=resolved_period.end_date_exclusive,
        settings=_active_settings(session),
        staffs=session.query(Staff).all(),
        demands=session.query(Demand).filter(
            Demand.work_date >= resolved_period.start_date,
            Demand.work_date < resolved_period.end_date_exclusive,
        ).all(),
        positions=session.query(Position).all(),
    )
    _add_rotation_preflight(preflight, session, resolved_period.start_date, resolved_period.end_date_exclusive)
    return preflight


@app.post("/rule-assistant/candidates", responses={200: {"model": RuleAssistantCandidateResponse}})
def rule_assistant_candidates(
    payload: RuleAssistantInput,
    session: Session = Depends(get_session),
) -> dict:
    tokens, started_here = begin_read_snapshot(session, ("staff", "pair_rules"))
    try:
        settings_row = (
            session.query(RuleConfig)
            .filter(RuleConfig.name == "active")
            .with_for_update(read=True)
            .one_or_none()
        )
        staffs = session.query(Staff).order_by(Staff.id).all()
        pair_rules = session.query(StaffPairRule).order_by(StaffPairRule.id).all()
        config = settings_row.config_json if settings_row else {}
        result = build_rule_candidates(
            text=payload.text,
            staffs=staffs,
            settings={**build_active_settings(config), "settings_revision": _settings_revision(config)},
        )
        result["input_versions"] = _rule_assistant_versions(
            session,
            staff_collection_revision=tokens["staff"],
            pair_rules_collection_revision=tokens["pair_rules"],
            staffs=staffs,
            pair_rules=pair_rules,
            config=config,
        )
    except Exception:
        end_read_snapshot(session, started_here, failed=True)
        raise
    end_read_snapshot(session, started_here)
    return result


def _rule_assistant_versions(
    session: Session,
    *,
    staff_collection_revision: int | None = None,
    pair_rules_collection_revision: int | None = None,
    staffs: list[Staff] | None = None,
    pair_rules: list[StaffPairRule] | None = None,
    config: dict | None = None,
) -> dict:
    if staff_collection_revision is None:
        staff_collection_revision = technical_revision(
            session.get(EditorRevision, "staff").revision,
            "staff_collection_revision",
        )
    if pair_rules_collection_revision is None:
        pair_rules_collection_revision = technical_revision(
            session.get(EditorRevision, "pair_rules").revision,
            "pair_rules_collection_revision",
        )
    if staffs is None:
        staffs = session.query(Staff).order_by(Staff.id).all()
    if pair_rules is None:
        pair_rules = session.query(StaffPairRule).order_by(StaffPairRule.id).all()
    if config is None:
        row = session.query(RuleConfig).filter(RuleConfig.name == "active").one_or_none()
        config = row.config_json if row else {}
    return {
        "staff": {
            "collection_revision": staff_collection_revision,
            "items": [{"id": row.id, "revision": row.revision} for row in staffs],
        },
        "pair_rules": {
            "collection_revision": pair_rules_collection_revision,
            "items": [{"id": row.id, "revision": row.revision} for row in pair_rules],
        },
        "settings_revision": _settings_revision(config),
    }


def _require_rule_assistant_collection_snapshot(
    supplied: RuleAssistantCollectionVersionsInput | None,
    actual_rows: list[Staff] | list[StaffPairRule],
    resource: str,
) -> dict[int, int]:
    if supplied is None or supplied.items is None:
        raise HTTPException(
            status_code=428,
            detail={
                "code": f"rule_assistant_{resource}_items_required",
                "message": "候補を再読み込みしてから適用してください。",
            },
        )
    expected: dict[int, int] = {}
    for item in supplied.items:
        if item.id in expected:
            raise HTTPException(
                status_code=422,
                detail={
                    "code": f"rule_assistant_{resource}_items_invalid",
                    "message": "候補の更新対象が重複しています。",
                },
            )
        expected[item.id] = require_revision(item.revision, f"{resource}_revision")
    actual = {row.id: row.revision for row in actual_rows}
    if actual != expected:
        raise conflict(f"rule_assistant_{resource}_snapshot")
    return expected


def _require_rule_assistant_input_versions(
    payload: RuleAssistantApplyInput,
) -> RuleAssistantInputVersionsInput:
    if payload.input_versions is None:
        raise HTTPException(
            status_code=428,
            detail={"code": "rule_assistant_input_versions_required", "message": "候補を再読み込みしてから適用してください。"},
        )
    versions = payload.input_versions
    if versions.staff is None or versions.pair_rules is None or versions.settings_revision is None:
        raise HTTPException(
            status_code=428,
            detail={"code": "rule_assistant_input_versions_required", "message": "候補を再読み込みしてから適用してください。"},
        )
    if not isinstance(versions.settings_revision, str) or not versions.settings_revision:
        raise HTTPException(
            status_code=422,
            detail={"code": "rule_assistant_settings_revision_invalid", "message": "設定の更新番号が不正です。"},
        )
    require_revision(versions.staff.collection_revision, "staff_collection_revision")
    require_revision(versions.pair_rules.collection_revision, "pair_rules_collection_revision")
    return versions


@app.post("/rule-assistant/apply", responses={200: {"model": RuleAssistantApplyResponse}})
def apply_rule_assistant_candidate(
    payload: RuleAssistantApplyInput,
    session: Session = Depends(get_session),
) -> dict:
    versions = _require_rule_assistant_input_versions(payload)
    staff_update_ids = [row.staff_id for row in payload.staff_updates]
    if len(set(staff_update_ids)) != len(staff_update_ids):
        raise HTTPException(status_code=422, detail="人員更新を重複指定できません")
    pair_keys: list[tuple[int, int, str]] = []
    for candidate in payload.pair_rule_candidates:
        if candidate.first_staff_id == candidate.second_staff_id:
            raise HTTPException(status_code=422, detail="同時休み禁止には異なる2人を指定してください")
        first_staff_id, second_staff_id = sorted(
            [candidate.first_staff_id, candidate.second_staff_id]
        )
        pair_keys.append((first_staff_id, second_staff_id, candidate.kind))
    if len(set(pair_keys)) != len(pair_keys):
        raise HTTPException(status_code=422, detail="同時休み禁止を重複指定できません")

    if payload.settings.settings_revision != versions.settings_revision:
        raise conflict("rule_assistant_settings_snapshot")
    collection_revisions = begin_multi_resource_write_snapshot(
        session,
        {
            "staff": versions.staff.collection_revision,
            "pair_rules": versions.pair_rules.collection_revision,
        },
        {
            resource
            for resource, changed in (
                ("staff", bool(payload.staff_updates)),
                ("pair_rules", bool(payload.pair_rule_candidates)),
            )
            if changed
        },
    )
    try:
        _store_optimization_settings(payload.settings, session)

        staffs = session.query(Staff).order_by(Staff.id).all()
        _require_rule_assistant_collection_snapshot(versions.staff, staffs, "staff")

        pair_rules = session.query(StaffPairRule).order_by(StaffPairRule.id).all()
        _require_rule_assistant_collection_snapshot(versions.pair_rules, pair_rules, "pair_rules")

        _ensure_staff_ids_exist(
            session,
            staff_update_ids + [staff_id for key in pair_keys for staff_id in key[:2]],
        )
        for update in payload.staff_updates:
            claim_item(session, Staff, update.staff_id, next(
                item.revision for item in versions.staff.items if item.id == update.staff_id
            ), "staff_revision")
            staff = session.get(Staff, update.staff_id)
            assert staff is not None
            staff.paid_leave_eligible = update.paid_leave_eligible
        pair_rule_by_key = {
            (row.first_staff_id, row.second_staff_id, row.kind): row for row in pair_rules
        }
        for candidate, (first_staff_id, second_staff_id, kind) in zip(
            payload.pair_rule_candidates,
            pair_keys,
            strict=True,
        ):
            row = pair_rule_by_key.get((first_staff_id, second_staff_id, kind))
            if row is None:
                row = StaffPairRule(
                    first_staff_id=first_staff_id,
                    second_staff_id=second_staff_id,
                    kind=kind,
                )
                session.add(row)
            else:
                claim_item(session, StaffPairRule, row.id, next(
                    item.revision for item in versions.pair_rules.items if item.id == row.id
                ), "pair_rule_revision")
            row.strength = candidate.strength
            row.weight = candidate.weight
            row.reason = candidate.reason

        session.flush()
        response = {
            "settings": _active_settings(session),
            "staff_updates_applied": len(payload.staff_updates),
            "pair_rules_applied": len(payload.pair_rule_candidates),
            "input_versions": _rule_assistant_versions(
                session,
                staff_collection_revision=collection_revisions["staff"],
                pair_rules_collection_revision=collection_revisions["pair_rules"],
            ),
        }
        session.commit()
        return response
    except Exception:
        session.rollback()
        raise


@app.get("/solve-runs/{solve_run_id}", responses={200: {"model": SolveRunDetail}})
def get_solve_run(
    solve_run_id: int,
    session: Session = Depends(get_session),
    candidate_id: int | None = None,
    prior_for_start_date: date | None = None,
) -> dict:
    with candidate_snapshot(session):
        return _get_solve_run_snapshot(solve_run_id, session, candidate_id, prior_for_start_date)


def _get_solve_run_snapshot(solve_run_id, session, candidate_id, prior_for_start_date):
    solve_run = session.get(SolveRun, solve_run_id)
    if solve_run is None:
        raise HTTPException(status_code=404, detail="solve run not found")
    candidates = (
        session.query(CandidateSchedule)
        .filter(CandidateSchedule.solve_run_id == solve_run.id)
        .order_by(CandidateSchedule.candidate_no)
        .all()
    )
    selected_candidate = (
        next((row for row in candidates if row.id == candidate_id), None)
        if candidate_id is not None
        else (confirmed_candidate(session, solve_run) or (candidates[0] if candidates else None))
    )
    if candidate_id is not None and selected_candidate is None:
        raise HTTPException(status_code=409, detail="solve run has no candidate")
    if solve_run.status in {"failed", "error"}:
        selected_candidate = None
    assignments = (
        session.query(Assignment)
        .options(joinedload(Assignment.staff), joinedload(Assignment.shift_template),
                 joinedload(Assignment.position))
        .filter(Assignment.candidate_id == selected_candidate.id)
        .order_by(Assignment.work_date, Assignment.staff_id)
        .all()
    ) if selected_candidate else []
    statuses = (
        session.query(CandidateStaffStatus)
        .options(joinedload(CandidateStaffStatus.staff), joinedload(CandidateStaffStatus.shift_template))
        .filter(CandidateStaffStatus.candidate_id == selected_candidate.id)
        .order_by(CandidateStaffStatus.work_date, CandidateStaffStatus.staff_id)
        .all()
    ) if selected_candidate else []
    violations = session.query(Violation).filter(Violation.solve_run_id == solve_run.id).order_by(Violation.id).all()
    status_reasons = {row.status_id: row for row in session.query(CandidateStaffStatusExplanation).filter(
        CandidateStaffStatusExplanation.status_id.in_([row.id for row in statuses]),
    ).all()}
    explanation_by_assignment_id = {
        row.assignment_id: row
        for row in session.query(AssignmentExplanation)
        .filter(AssignmentExplanation.assignment_id.in_([assignment.id for assignment in assignments]))
        .all()
    }
    confirmation = (
        session.query(ScheduleConfirmation)
        .filter(ScheduleConfirmation.solve_run_id == solve_run.id)
        .one_or_none()
    )
    run_period = _solve_run_period_payload(solve_run)
    precondition = None
    precondition_error = None
    if selected_candidate:
        try:
            precondition = candidate_precondition(session, selected_candidate, prior_for_start_date=prior_for_start_date)
        except HTTPException as error:
            if not isinstance(error.detail, dict) or error.detail.get("code") != "candidate_rule_source_unresolved":
                raise
            # Historical display/reprint remains available; no command token
            # is minted from absent provenance or today's replacement rules.
            precondition_error = error.detail
    return {
        "id": solve_run.id,
        **run_period,
        "job_status": serialize_job_status(session, solve_run),
        "status": solve_run.status,
        "solver_status": solve_run.solver_status,
        "selected_candidate_id": selected_candidate.id if selected_candidate else None,
        "command_precondition": precondition,
        "command_precondition_error": precondition_error,
        "candidates": [
            {
                "id": candidate.id,
                "strategy": candidate.strategy,
                "total_score": candidate.total_score,
                "score_breakdown": candidate.score_breakdown_json,
            }
            for candidate in candidates
        ],
        "assignments": [
            {
                "id": assignment.id,
                "staff_id": assignment.staff_id,
                "position_id": assignment.position_id,
                "shift_template_id": assignment.shift_template_id,
                "staff_name": assignment.staff.name,
                "work_date": assignment.work_date.isoformat(),
                "shift_template": _template_display_name(assignment.shift_template),
                "position": _position_display_name(assignment.position),
                "source": assignment.source,
                "is_locked": assignment.is_locked,
                "reason": explanation_by_assignment_id[assignment.id].message
                if assignment.id in explanation_by_assignment_id
                else "必要人数と担当可能条件を満たすように選ばれました",
            }
            for assignment in assignments
        ],
        "statuses": [
            {
                "staff_name": status.staff.name,
                "id": status.id,
                "staff_id": status.staff_id,
                "shift_template_id": status.shift_template_id,
                "source": status_reasons[status.id].category if status.id in status_reasons else None,
                "reason": status_reasons[status.id].message if status.id in status_reasons else None,
                "work_date": status.work_date.isoformat(),
                "status": status.status,
                "work_kind": status.work_kind,
                "shift_template": _template_display_name(status.shift_template) or None,
            }
            for status in statuses
        ],
        "violations": [
            {"severity": violation.severity, "rule_name": violation.rule_name, "message": violation.message}
            for violation in violations
        ],
        "confirmed": confirmation is not None,
        "confirmed_at": confirmation.stored_timestamp(session) if confirmation else None,
        "candidate_validation": (
            _historical_candidate_validation(
                selected_candidate, session, prior_for_start_date
            )
            if selected_candidate and prior_for_start_date is not None
            else _validate_candidate(selected_candidate.id, session) if selected_candidate else None
        ),
    }


@app.get("/solve-run-periods")
def list_solve_run_periods(session: Session = Depends(get_session), initial_context: bool = False,
                          end_date: date | None = None, limit: int = Query(default=100, ge=1, le=200),
                          before_id: int | None = Query(default=None, ge=1)) -> list[dict]:
    from sqlalchemy import select
    if session.query(SolveRun.id).filter(SolveRun.period_display_error.is_(True)).first():
        raise HTTPException(409, "保存済みのシフト期間を確認できません。")
    confirmed = select(1).where(ScheduleConfirmation.solve_run_id == SolveRun.id).exists()
    query = select(SolveRun.id, SolveRun.status, SolveRun.period_summary_json,
                   SolveRun.started_at, SolveRun.finished_at, confirmed.label('confirmed'))
    if initial_context:
        unknown = session.query(SolveRun.id).filter(
            SolveRun.status == 'succeeded', SolveRun.period_known.is_(False), confirmed,
        ).order_by(SolveRun.id.desc()).first()
        query = query.where(SolveRun.id == unknown[0]) if unknown else query.where(
            SolveRun.status == 'succeeded', confirmed,
        ).order_by(SolveRun.display_period_end_key.desc(), SolveRun.id.desc())
        query = query.limit(1)
    else:
        if end_date:
            query = query.where(SolveRun.display_period_end_key == end_date.isoformat())
        if isinstance(before_id, int):
            query = query.where(SolveRun.id < before_id)
        query = query.order_by(SolveRun.id.desc()).limit(limit if isinstance(limit, int) else 100)
    result = []
    for row in session.execute(query):
        if not isinstance(row.period_summary_json, dict):
            raise HTTPException(409, "保存済みのシフト期間を確認できません。")
        result.append({'id': row.id, 'status': row.status, **row.period_summary_json,
                       'confirmed': row.confirmed,
                       'started_at': row.started_at.isoformat() if row.started_at else None,
                       'finished_at': row.finished_at.isoformat() if row.finished_at else None})
    return result


@app.get("/solve-runs", responses={200: {"model": list[SolveRunSummary]}})
def list_solve_runs(session: Session = Depends(get_session), limit: int | None = Query(default=50, ge=1, le=200), before_id: int | None = Query(default=None, ge=1), include_job_status: bool = False) -> list[dict]:
    from sqlalchemy import func
    from .jobs import load_job_status_sources, serialize_job_status_from_sources

    query = session.query(SolveRun)
    if not include_job_status:
        query = query.options(load_only(SolveRun.id, SolveRun.target_month,
            SolveRun.period_summary_json, SolveRun.period_display_error, SolveRun.status,
            SolveRun.solver_status, SolveRun.started_at, SolveRun.finished_at))
    if isinstance(before_id, int):
        query = query.filter(SolveRun.id < before_id)
    query = query.order_by(SolveRun.id.desc())
    query = query.limit(limit if isinstance(limit, int) else 50)
    rows = query.all()
    if not rows:
        return []
    if not include_job_status and any(row.period_display_error or not isinstance(row.period_summary_json, dict) for row in rows):
        raise HTTPException(409, "保存済みのシフト期間を確認できません。")
    job_status_sources = load_job_status_sources(session, rows) if include_job_status else None
    run_ids = [row.id for row in rows]
    assignment_counts = dict(
        session.query(CandidateSchedule.solve_run_id, func.count(Assignment.id))
        .join(Assignment, Assignment.candidate_id == CandidateSchedule.id)
        .filter(CandidateSchedule.candidate_no == 1)
        .filter(CandidateSchedule.solve_run_id.in_(run_ids))
        .group_by(CandidateSchedule.solve_run_id).all()
    )
    confirmed_ids = {run_id for (run_id,) in session.query(ScheduleConfirmation.solve_run_id).filter(
        ScheduleConfirmation.solve_run_id.in_(run_ids)).all()}
    return [
        {
            "id": row.id,
            **(_solve_run_period_payload(row) if include_job_status else row.period_summary_json),
            "job_status": serialize_job_status_from_sources(row, job_status_sources) if include_job_status else None,
            "status": row.status,
            "solver_status": row.solver_status,
            "started_at": row.started_at.isoformat() if row.started_at else None,
            "finished_at": row.finished_at.isoformat() if row.finished_at else None,
            "assignments_count": assignment_counts.get(row.id, 0),
            "confirmed": row.id in confirmed_ids,
        }
        for row in rows
    ]


@app.put("/assignments/{assignment_id}/lock", responses={200: {"model": AssignmentLockResponse}})
def set_assignment_lock(
    assignment_id: int,
    payload: AssignmentLockInput,
    session: Session = Depends(get_session),
) -> dict:
    return execute_candidate_command(
        session, payload.command_precondition, None,
        lambda candidate: _set_assignment_lock(assignment_id, payload, session, candidate.id),
    )


def _set_assignment_lock(assignment_id, payload, session, candidate_id):
    assignment = session.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="assignment not found")
    if assignment.candidate_id != candidate_id:
        raise HTTPException(409, detail={"code": "candidate_precondition_conflict"})
    run = session.get(SolveRun, assignment.candidate.solve_run_id)
    if run is None or run.status != "succeeded":
        raise HTTPException(status_code=409, detail="作成に成功したシフトだけ固定できます")
    if run.status == "succeeded" and session.query(ScheduleConfirmation).filter(
        ScheduleConfirmation.solve_run_id == run.id
    ).first():
        raise HTTPException(status_code=409, detail="確定済みのシフトは固定内容を変更できません")
    assignment.is_locked = payload.locked
    session.flush()
    return {"assignment_id": assignment.id, "locked": assignment.is_locked,
            "validation": _validate_candidate(candidate_id, session)}


@app.put("/candidates/{candidate_id}/cells/{staff_id}/{work_date}", responses={200: {"model": ScheduleCellResponse}})
def update_schedule_cell(
    candidate_id: int,
    staff_id: int,
    work_date: date,
    payload: ScheduleCellInput,
    session: Session = Depends(get_session),
) -> dict:
    return execute_candidate_command(
        session, payload.command_precondition, candidate_id,
        lambda candidate: _update_schedule_cell(candidate.id, staff_id, work_date, payload, session),
    )


def _update_schedule_cell(candidate_id, staff_id, work_date, payload, session):
    result = _apply_schedule_cell(candidate_id, staff_id, work_date, payload, session)
    result["validation"] = _validate_candidate(candidate_id, session)
    return result


def _apply_schedule_cell(
    candidate_id: int,
    staff_id: int,
    work_date: date,
    payload: ScheduleCellInput,
    session: Session,
) -> dict:
    candidate = session.get(CandidateSchedule, candidate_id)
    if candidate is None:
        raise HTTPException(status_code=404, detail="candidate not found")
    run = session.get(SolveRun, candidate.solve_run_id)
    if run is None or run.status != "succeeded":
        raise HTTPException(status_code=409, detail="作成に成功したシフトだけ手動修正できます")
    if session.query(ScheduleConfirmation).filter(
        ScheduleConfirmation.solve_run_id == run.id
    ).first():
        raise HTTPException(status_code=409, detail="確定済みのシフトは手動修正できません")
    staff = session.get(Staff, staff_id)
    if staff is None:
        raise HTTPException(status_code=404, detail="staff not found")
    if payload.work_kind not in {"position", "regular", "half_day", "off", PAID_LEAVE_HALF_DAY_WORK_KIND}:
        raise HTTPException(status_code=422, detail="勤務区分を選択してください")

    shift_template = None
    if payload.work_kind != "off":
        if payload.shift_template_id is None:
            raise HTTPException(status_code=422, detail="勤務時間を選択してください")
        shift_template = session.get(ShiftTemplate, payload.shift_template_id)
        if shift_template is None:
            raise HTTPException(status_code=404, detail="shift template not found")
        if payload.work_kind in {"half_day", PAID_LEAVE_HALF_DAY_WORK_KIND} and shift_template.shift_type != "half_day":
            raise HTTPException(status_code=422, detail="半日勤務用の勤務時間を選択してください")
    if payload.work_kind == PAID_LEAVE_HALF_DAY_WORK_KIND and not staff.paid_leave_eligible:
        raise HTTPException(status_code=422, detail="この人員には有休0.5を設定できません")

    position = None
    if payload.work_kind == "position":
        if payload.position_id is None:
            raise HTTPException(status_code=422, detail="持ち場を選択してください")
        position = session.get(Position, payload.position_id)
        if position is None:
            raise HTTPException(status_code=404, detail="position not found")

    assignments = (
        session.query(Assignment)
        .filter(
            Assignment.candidate_id == candidate.id,
            Assignment.staff_id == staff.id,
            Assignment.work_date == work_date,
        )
        .all()
    )
    status = (
        session.query(CandidateStaffStatus)
        .filter(
            CandidateStaffStatus.candidate_id == candidate.id,
            CandidateStaffStatus.staff_id == staff.id,
            CandidateStaffStatus.work_date == work_date,
        )
        .one_or_none()
    )
    if status is None and not assignments:
        raise HTTPException(status_code=404, detail="この候補に対象の人員・日付がありません")
    if status is None:
        status = CandidateStaffStatus(
            candidate_id=candidate.id,
            staff_id=staff.id,
            work_date=work_date,
            status="off" if payload.work_kind == "off" else "working",
            work_kind=payload.work_kind,
            shift_template_id=shift_template.id if shift_template else None,
        )
        session.add(status)
    else:
        status.status = "off" if payload.work_kind == "off" else "working"
        status.work_kind = payload.work_kind
        status.shift_template_id = shift_template.id if shift_template else None

    if payload.reason is not None:
        session.flush()
        reason = session.query(CandidateStaffStatusExplanation).filter_by(status_id=status.id).one_or_none()
        if reason is None:
            session.add(CandidateStaffStatusExplanation(status_id=status.id, category="manual", message=payload.reason))
        else:
            reason.category = "manual"
            reason.message = payload.reason
    assignment = assignments[0] if assignments else None
    if payload.work_kind == "position":
        for duplicate in assignments[1:]:
            session.query(AssignmentExplanation).filter(
                AssignmentExplanation.assignment_id == duplicate.id
            ).delete(synchronize_session=False)
            session.delete(duplicate)
        if assignment is None:
            assignment = Assignment(
                candidate_id=candidate.id,
                staff_id=staff.id,
                work_date=work_date,
                shift_template_id=shift_template.id,
                position_id=position.id,
                is_locked=True,
                source="manual",
            )
            session.add(assignment)
            session.flush()
        else:
            assignment.shift_template_id = shift_template.id
            assignment.position_id = position.id
            assignment.is_locked = True
            assignment.source = "manual"
        explanation = (
            session.query(AssignmentExplanation)
            .filter(AssignmentExplanation.assignment_id == assignment.id)
            .one_or_none()
        )
        if explanation is None:
            session.add(AssignmentExplanation(
                assignment_id=assignment.id,
                category="manual",
                message=payload.reason if payload.reason is not None else "確認画面で手動修正されました",
            ))
        else:
            explanation.category = "manual"
            explanation.message = payload.reason if payload.reason is not None else "確認画面で手動修正されました"
    else:
        assignment_ids = [row.id for row in assignments]
        if assignment_ids:
            session.query(AssignmentExplanation).filter(
                AssignmentExplanation.assignment_id.in_(assignment_ids)
            ).delete(synchronize_session=False)
        for row in assignments:
            session.delete(row)
        assignment = None

    session.flush()
    return {
        "candidate_id": candidate.id,
        "staff_id": staff.id,
        "work_date": work_date.isoformat(),
        "work_kind": status.work_kind,
        "assignment_id": assignment.id if assignment else None,
    }


@app.put("/candidates/{candidate_id}/cells-bulk", responses={200: {"model": ScheduleBulkResponse}})
def update_schedule_cells_bulk(
    candidate_id: int,
    payload: ScheduleBulkInput,
    session: Session = Depends(get_session),
) -> dict:
    return execute_candidate_command(
        session, payload.command_precondition, candidate_id,
        lambda candidate: _update_schedule_cells_bulk(candidate.id, payload, session),
    )


def _update_schedule_cells_bulk(candidate_id, payload, session):
    """Apply all requested schedule changes in one database transaction."""
    keys = [(cell.staff_id, cell.work_date) for cell in payload.cells]
    if len(keys) != len(set(keys)):
        raise HTTPException(status_code=422, detail="同じ人員・日付を重複して指定できません")
    # Validate every referenced object before changing the first cell. This keeps
    # an invalid cell from leaving an earlier cell partially updated.
    candidate = session.get(CandidateSchedule, candidate_id)
    if candidate is None:
        raise HTTPException(status_code=404, detail="candidate not found")
    run = session.get(SolveRun, candidate.solve_run_id)
    if run is None or run.status != "succeeded":
        raise HTTPException(status_code=409, detail="作成に成功したシフトだけ手動修正できます")
    if session.query(ScheduleConfirmation).filter(ScheduleConfirmation.solve_run_id == run.id).first():
        raise HTTPException(status_code=409, detail="確定済みのシフトは手動修正できません")
    _ensure_staff_ids_exist(session, [cell.staff_id for cell in payload.cells])
    for cell in payload.cells:
        if cell.work_kind not in {"position", "regular", "half_day", "off", PAID_LEAVE_HALF_DAY_WORK_KIND}:
            raise HTTPException(status_code=422, detail="勤務区分を選択してください")
        if cell.work_kind != "off" and (
            cell.shift_template_id is None or session.get(ShiftTemplate, cell.shift_template_id) is None
        ):
            raise HTTPException(status_code=422, detail="勤務時間を選択してください")
        if cell.work_kind in {"half_day", PAID_LEAVE_HALF_DAY_WORK_KIND}:
            shift_template = session.get(ShiftTemplate, cell.shift_template_id)
            if shift_template is None or shift_template.shift_type != "half_day":
                raise HTTPException(status_code=422, detail="半日勤務用の勤務時間を選択してください")
        if cell.work_kind == PAID_LEAVE_HALF_DAY_WORK_KIND:
            staff = session.get(Staff, cell.staff_id)
            if staff is None or not staff.paid_leave_eligible:
                raise HTTPException(status_code=422, detail="この人員には有休0.5を設定できません")
        if cell.work_kind == "position" and (
            cell.position_id is None or session.get(Position, cell.position_id) is None
        ):
            raise HTTPException(status_code=422, detail="持ち場を選択してください")
        exists = session.query(CandidateStaffStatus.id).filter(
            CandidateStaffStatus.candidate_id == candidate_id,
            CandidateStaffStatus.staff_id == cell.staff_id,
            CandidateStaffStatus.work_date == cell.work_date,
        ).first() or session.query(Assignment.id).filter(
            Assignment.candidate_id == candidate_id,
            Assignment.staff_id == cell.staff_id,
            Assignment.work_date == cell.work_date,
        ).first()
        if not exists:
            raise HTTPException(status_code=404, detail="この候補に対象の人員・日付がありません")
    results = [
        _apply_schedule_cell(candidate_id, cell.staff_id, cell.work_date, cell, session)
        for cell in payload.cells
    ]
    session.flush()
    validation = _validate_candidate(candidate_id, session)
    return {"updated_count": len(results), "cells": results, "validation": validation}


def _validate_candidate(candidate_id: int, session: Session) -> dict:
    return validate_candidate_state(candidate_id, session)


def _historical_candidate_validation(
    candidate: CandidateSchedule, session: Session, prior_for_start_date: date
) -> dict:
    from .validation_snapshot import historical_candidate_validation
    from .schedule_selection import require_direct_predecessor

    require_direct_predecessor(session, candidate.solve_run, prior_for_start_date)
    validation = historical_candidate_validation(session, candidate)
    errors = prior_history_coverage_errors(session, candidate, prior_for_start_date)
    if not errors:
        return validation
    return {**validation, "blocking": True, "warnings": [*validation["warnings"], *(
        {"severity": "error", "code": "prior_history_incomplete", "message": message}
        for message in errors
    )]}


@app.get("/candidates/{candidate_id}/validation", responses={200: {"model": CandidateValidation}})
def validate_candidate(candidate_id: int, session: Session = Depends(get_session)) -> dict:
    return _validate_candidate(candidate_id, session)


@app.get("/solve-runs/{solve_run_id}/export")
def export_solve_run(
    solve_run_id: int,
    format: str = "csv",
    session: Session = Depends(get_session),
    candidate_id: int | None = None,
):
    with candidate_snapshot(session):
        return _export_solve_run_snapshot(solve_run_id, format, session, candidate_id)


def _export_solve_run_snapshot(solve_run_id, format, session, candidate_id):
    solve_run = session.get(SolveRun, solve_run_id)
    if solve_run is None:
        raise HTTPException(status_code=404, detail="solve run not found")
    candidate_query = session.query(CandidateSchedule).filter(CandidateSchedule.solve_run_id == solve_run_id)
    candidate = (
        candidate_query.filter(CandidateSchedule.id == candidate_id).one_or_none()
        if candidate_id is not None
        else (confirmed_candidate(session, solve_run) or (
            candidate_query.one_or_none() if candidate_query.count() == 1 else None
        ))
    )
    if candidate is None:
        if candidate_id is None and candidate_query.count() > 1:
            raise HTTPException(status_code=409, detail={"code": "export_candidate_required"})
        raise HTTPException(status_code=404, detail="candidate not found")
    rows = (
        session.query(Assignment)
        .filter(Assignment.candidate_id == candidate.id)
        .order_by(Assignment.work_date, Assignment.staff_id, Assignment.id)
        .all()
    )
    statuses = (
        session.query(CandidateStaffStatus)
        .filter(CandidateStaffStatus.candidate_id == candidate.id)
        .order_by(CandidateStaffStatus.work_date, CandidateStaffStatus.staff_id)
        .all()
    )
    assignment_by_staff_date = group_candidate_assignments(rows)
    if statuses:
        table = [["staff_name", "work_date", "status", "work_kind", "shift_template", "position"]]
        status_by_staff_date = {(row.staff_id, row.work_date): row for row in statuses}
        for key in sorted(status_by_staff_date.keys() | assignment_by_staff_date.keys(), key=lambda key: (key[1], key[0])):
            row = status_by_staff_date.get(key)
            cell_assignments = assignment_by_staff_date.get(key, [])
            if row is None:
                table.extend([
                    assignment.staff.name, assignment.work_date.isoformat(), "working", "position",
                    _template_display_name(assignment.shift_template), _position_display_name(assignment.position),
                ] for assignment in cell_assignments)
                continue
            for assignment in cell_assignments or [None]:
                table.append([
                    row.staff.name, row.work_date.isoformat(), row.status, row.work_kind,
                    _template_display_name(assignment.shift_template if len(cell_assignments) > 1 else row.shift_template),
                    _position_display_name(assignment.position) if assignment is not None else "",
                ])
    else:
        table = [["staff_name", "work_date", "status", "work_kind", "shift_template", "position"]] + [
            [
                row.staff.name,
                row.work_date.isoformat(),
                "working",
                "position",
                _template_display_name(row.shift_template),
                _position_display_name(row.position),
            ]
            for row in rows
        ]
    filename_base = f"shift-run-{solve_run_id}-candidate-{candidate.id}"
    if format == "csv":
        output = StringIO(newline="")
        csv.writer(output).writerows(
            [_neutralize_spreadsheet_cell(value) for value in row] for row in table
        )
        content = output.getvalue().encode("utf-8-sig")
        try:
            stored_object = store_bytes(
                f"exports/{filename_base}.csv", content, "text/csv; charset=utf-8"
            )
        except StorageError as error:
            raise HTTPException(status_code=503, detail=str(error)) from error
        headers = {
            "Content-Disposition": f'attachment; filename="{filename_base}.csv"',
            "X-Storage-Backend": stored_object.backend,
            "X-Storage-Key": stored_object.storage_key,
        }
        if stored_object.bucket_name:
            headers["X-GCS-Object"] = stored_object.storage_key
        return StreamingResponse(
            iter([content]),
            media_type="text/csv; charset=utf-8",
            headers=headers,
        )
    if format == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Shift schedule"
        _build_shift_matrix_sheet(sheet, solve_run, rows, statuses, session, assignment_by_staff_date)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        content = output.getvalue()
        try:
            stored_object = store_bytes(
                f"exports/{filename_base}.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except StorageError as error:
            raise HTTPException(status_code=503, detail=str(error)) from error
        headers = {
            "Content-Disposition": f'attachment; filename="{filename_base}.xlsx"',
            "X-Storage-Backend": stored_object.backend,
            "X-Storage-Key": stored_object.storage_key,
        }
        if stored_object.bucket_name:
            headers["X-GCS-Object"] = stored_object.storage_key
        return StreamingResponse(
            iter([content]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers=headers,
        )
    raise HTTPException(status_code=400, detail="format must be csv or xlsx")
