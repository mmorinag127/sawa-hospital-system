"""Real packaged APIs/owned databases; synthetic inputs and no model/network fees."""
import asyncio
from unittest.mock import AsyncMock
from contextlib import contextmanager, suppress
import hashlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
from io import BytesIO
import os
from pathlib import Path
import sys
import threading

import httpx
import pytest

from experiments.agent_assistant.app import create_app
from experiments.agent_assistant.gemini import command_identity
from experiments.agent_assistant.mcp_server import build_mcp
from experiments.agent_assistant.providers import Providers
from experiments.agent_assistant.sandbox import Sandbox
from experiments.agent_assistant.store import Store
from experiments.agent_assistant.tools import Tools
from experiments.agent_assistant.tests.test_gemini import interaction

ROOT=Path(__file__).resolve().parents[3]
pytestmark=pytest.mark.skipif(os.environ.get('SAWA_AGENT_NATIVE_TESTS')!='1',reason='Opt in to owned localhost native subprocesses')


@contextmanager
def portal():
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            valid=self.headers.get('Authorization')=='Bearer synthetic-native-operator'
            body=json.dumps({'account':'native-operator','role':'operator','systems':['shift','school-lunch'],'auth_disabled':False}).encode()
            self.send_response(200 if valid else 401)
            self.send_header('Content-Type','application/json');self.end_headers()
            self.wfile.write(body if valid else b'{}')
        def log_message(self,*_):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:yield 'http://127.0.0.1:'+str(server.server_port)+'/auth/me'
    finally:server.shutdown();server.server_close();thread.join()


def setup(root,domain,url):
    cfg={'worktree':str(ROOT/'agent_runtime/native'/domain),
         'source_manifest':str(ROOT/'agent_runtime/manifests'/(domain+'.json')),
         'baseline':str(ROOT/'agent_runtime/baselines'/domain),
         'python':os.environ.get('SAWA_AGENT_NATIVE_'+domain.upper().replace('-','_')+'_PYTHON',sys.executable),
         'env':{'AUTH_PROVIDER':'portal','PORTAL_AUTH_ME_URL':url}}
    store=Store(root);store.set_budget(1000)
    sandbox=Sandbox(root,{domain:cfg});tools=Tools(store,sandbox)
    return store,sandbox,tools


def test_real_registered_day_suggestions_do_not_need_calendar_input(tmp_path):
    with portal() as url:
        store, sandbox, tools = setup(tmp_path, 'school-lunch', url)
        async def scenario():
            selection = await tools.flow.options('lunch_daily_bundle', {'facility_id': 1}, 'native-operator', 'Bearer synthetic-native-operator')
            assert selection['work_days']
            proposal = selection['work_days'][0]
            targets = [item for item in selection['choices'] if item['data']['service_date'] == proposal['service_date']]
            assert len(targets) == proposal['menu_count']
            assert sorted(item['data']['name'] for item in targets) == sorted(proposal['menu_names'])
            context = tools.flow.selection('lunch_daily_bundle', 'native-operator', selection['selection_id'], [item['key'] for item in targets])
            assert context['inputs']['service_date'] == proposal['service_date']
            assert context['inputs']['month'] == proposal['service_date'][:7]
            assert all(item['data']['facility']['id'] == 1 for item in context['targets'])
            assert not sandbox.children
        asyncio.run(scenario())


@pytest.mark.parametrize('domain',['shift','school-lunch'])
def test_packaged_api_auth_and_extensions_with_latest_contracts(tmp_path,domain):
    with portal() as url:
        store,sandbox,tools=setup(tmp_path,domain,url)
        run,_=store.create('native-operator','native-smoke-'+domain,'synthetic',[domain],100)
        sandbox.prepare(run)
        tools.authorizations[run['id']]='Bearer synthetic-native-operator'
        async def scenario():
            try:
                await sandbox.start(run)
                path='/planning/period' if domain=='shift' else '/api/school-lunch/facilities'
                params={'month':'2026-11'} if domain=='shift' else {}
                async with httpx.AsyncClient(trust_env=False) as outsider:
                    denied=await outsider.get(sandbox.urls[(run['id'],domain)]+path,params=params,
                                               headers={'X-Trial-Token':'invalid','X-Sawa-Service-Token':'invalid'})
                assert denied.status_code in (401,403),denied.text
                response=await sandbox.request(run,domain,'GET',path,params=params,authorization='Bearer synthetic-native-operator')
                assert response.status_code==200,response.text
                schema=await sandbox.request(run,domain,'GET','/openapi.json',authorization='Bearer synthetic-native-operator')
                assert schema.status_code==200,schema.text
                from experiments.agent_assistant.catalog import OPERATIONS
                for name,op in OPERATIONS.items():
                    if op.domain==domain:
                        described=await tools.describe(run,{'operation':name})
                        assert described['ok'] and described['method']==op.method
                if domain=='shift':
                    assert response.json()['start_date']<'2026-12-01'
                else:
                    query={'facility_id':1,'template_id':'kitchen-rice'}
                    drawing=await sandbox.request(run,domain,'GET','/api/school-lunch/menu-executions/1/process/question-drawing',params=query,authorization='Bearer synthetic-native-operator')
                    assert drawing.status_code==200,drawing.text
                    import base64
                    value=drawing.json()
                    assert hashlib.sha256(base64.b64decode(value['svg_base64'])).hexdigest()==value['sha256']
                    options=await sandbox.request(run,domain,'GET','/api/school-lunch/menu-executions/1/process/route-options',params={'facility_id':1},authorization='Bearer synthetic-native-operator')
                    assert options.status_code==200,options.text
                    assert options.json()['templates']
            finally:await sandbox.stop(run['id'])
        asyncio.run(scenario())


def test_real_shift_approval_restart_confirmation_output_and_rollback(tmp_path,monkeypatch):
    for key,value in {'GEMINI_API_KEY':'synthetic','GEMINI_AGENT_MODEL':'gemini-test',
        'GEMINI_AGENT_INPUT_USD_PER_MILLION':'1','GEMINI_AGENT_OUTPUT_USD_PER_MILLION':'2',
        'SAWA_TRIAL_OPERATOR_ACCOUNT':'native-operator'}.items():monkeypatch.setenv(key,value)
    with portal() as url:
        store,sandbox,tools=setup(tmp_path,'shift',url)
        async def scenario():
            # The pinned synthetic fixture has demands and leave in August.
            selected=await tools.flow.options('shift_next_month',{'month':'2026-08'},'native-operator','Bearer synthetic-native-operator')
            context=tools.flow.selection('shift_next_month','native-operator',selected['selection_id'],['2026-08'])
            run,_=store.create('native-operator','native-shift-full','synthetic',['shift'],100,task_id='shift_next_month',task_inputs=context,provider='gemini')
            ident=run['id'];sandbox.prepare(run)
            tools.authorizations[ident]='Bearer synthetic-native-operator'
            try:
                await sandbox.start(run);store.update(ident,status='running',session_id='turn')
                review=await tools.execute(ident,0,'prepare_shift_review',{},'native-prepare')
                assert review.get('ok'),review
                saved=tools.flow.review(ident);assert saved['state']=='awaiting_approval'
                calls=[('wait','wait_for_task_decision',{}),('complete','complete_shift_work',{})]
                remote=interaction('turn',calls)
                store.event(ident,'gemini_dispatch',{'id':'dispatch'})
                store.event(ident,'gemini_response',{'id':'turn','dispatch':'dispatch'})
                store.event(ident,'gemini_state',{'interaction':'turn','calls':remote['steps'],'results':[],'index':0,'pending':None})
                waiting=asyncio.create_task(tools.execute(ident,0,'wait_for_task_decision',{},command_identity('turn','wait')))
                await asyncio.sleep(.05);waiting.cancel()
                with suppress(asyncio.CancelledError):await waiting
                await sandbox.stop(ident);store.recover()
                requests=[]
                def handler(request):
                    requests.append(request.method);assert request.method=='GET'
                    return httpx.Response(200,json=remote)
                client=httpx.AsyncClient(transport=httpx.MockTransport(handler))
                agent=Providers(store,sandbox,tools,build_mcp(tools),{},gemini_client=client)
                await agent.resume(ident)
                async with tools.lock(ident):
                    await tools.flow.decide(store.get(ident),saved['id'],True,'合成試験の勤務案を承認')
                await asyncio.wait_for(agent.tasks[ident],60)
                assert store.get(ident)['status']=='completed',store.events(ident)[-8:]
                assert store.get(ident)['reserved']==0
                with store.read() as db:artifacts=[dict(row) for row in db.execute('SELECT * FROM artifacts WHERE run_id=?',(ident,))]
                assert artifacts
                from openpyxl import load_workbook
                workbook=load_workbook(BytesIO(Path(artifacts[-1]['path']).read_bytes()))
                assert workbook.sheetnames
                assert any('試験職員' in str(cell.value) for sheet in workbook for row in sheet for cell in row)
                app=create_app({'domains':sandbox.configs},tmp_path,gemini_client=client)
                async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://testserver') as browser:
                    headers={'x-trial-token':app.state.token}
                    restored=await browser.post('/runs/'+ident+'/rollback',headers=headers)
                    assert restored.status_code==200,restored.text
                    assert restored.json()['status']=='rolled_back'
                    assert (await browser.get('/runs/'+ident+'/artifacts/'+artifacts[-1]['id'],headers=headers)).status_code==409
                assert requests and set(requests)=={'GET'}
                await client.aclose()
            finally:await sandbox.stop(ident)
        asyncio.run(scenario())


def test_real_lunch_questions_save_approval_restart_and_three_official_outputs(tmp_path,monkeypatch):
    for key,value in {'GEMINI_API_KEY':'synthetic','GEMINI_AGENT_MODEL':'gemini-test',
        'GEMINI_AGENT_INPUT_USD_PER_MILLION':'1','GEMINI_AGENT_OUTPUT_USD_PER_MILLION':'2'}.items():
        monkeypatch.setenv(key,value)
    with portal() as url:
        store,sandbox,tools=setup(tmp_path,'school-lunch',url)
        async def scenario():
            selection=await tools.flow.options('lunch_daily_bundle',{'facility_id':1,'service_date':'2026-11-02'},'native-operator','Bearer synthetic-native-operator')
            context=tools.flow.selection('lunch_daily_bundle','native-operator',selection['selection_id'],['1'],1)
            run,_=store.create('native-operator','native-lunch-full','synthetic',['school-lunch'],100,
                               task_id='lunch_daily_bundle',task_inputs=context,provider='gemini')
            ident=run['id'];sandbox.prepare(run);tools.authorizations[ident]='Bearer synthetic-native-operator'
            school=context['targets'][0]['data']['facility']['name']
            # Explicit synthetic operator inputs, never inferred actual costs/times.
            price={'facility_id':1,'implementation_date':'2026-11-02','source_menu_execution_ids':[1],
                   'corporation_name':'合成試験法人','representative_name':'合成試験代表','tax_rate':.08,
                   'schools':[{'name':school,'recipient':school,'menu_name':'野菜スープ（試験データ）',
                               'headcount':10,'headcount_confirmed':True,'allocation_headcount':10,'reserve_meals':0,
                               'quantity_items':[{'name':name,'unit':'g','amount_per_person':100,'cost_key':key}
                                                 for key,name in [('carrot','にんじん'),('potato','じゃがいも')]]}],
                   'cost_items':[{'key':key,'name':name,'status':'confirmed','unit_price_tax_included':1000,
                                  'price_unit':'kg','origin_or_maker':'合成試験産地','supplier':'合成試験業者'}
                                 for key,name in [('carrot','にんじん'),('potato','じゃがいも')]]}
            schedule={'plating_start':'10:00','plating_end':'10:20',
                      'timeline_ticks':['08:00','08:30','09:00','09:30','10:00','10:30','11:00','11:30'],
                      'delivery_assignees':['試験担当A']}
            responses={};phase=0
            def handler(request):
                nonlocal phase
                if request.method=='GET':return httpx.Response(200,json=responses[request.url.path.rsplit('/',1)[-1]])
                assert request.method=='POST'
                if phase==0:
                    state=tools.flow.questions.state(store.get(ident))
                    complete=state and not state['pending'] and all(len(d['answers'])==len(d['document']['steps']) for d in state['documents'])
                    name,args=('save_lunch_process_answers',{}) if complete else ('ask_lunch_process_question',{'restart':False})
                    if complete:phase=1
                else:
                    name,args={1:('lunch_save_price_source',{'parameters':{'query':{'facility_id':1,'month':'2026-11'},'body':price,'reason':'合成試験の明示入力：各100g・10食・1000円/kgを保存'}}),
                               2:('save_lunch_daily_schedule',schedule),3:('prepare_lunch_daily_review',{}),
                               4:('wait_for_task_decision',{}),5:('complete_lunch_daily_work',{})}[phase]
                    phase+=1
                key='turn-'+str(len(responses)+1)
                response=interaction(key,[('call',name,args)]);responses[key]=response
                return httpx.Response(200,json=response)
            client=httpx.AsyncClient(transport=httpx.MockTransport(handler))
            agent=Providers(store,sandbox,tools,build_mcp(tools),{},gemini_client=client)
            try:
                await agent.start(run)
                answered=set()
                for _ in range(2400):
                    current=store.get(ident)
                    state=tools.flow.questions.state(current)
                    pending=state['pending'] if state else None
                    if current['status']=='waiting_input' and pending and pending['question_id'] not in answered:
                        if pending['field']=='grouping':label='一緒に調理しました'
                        elif pending['field']=='equipment':
                            chosen='釜A' if pending['key']=='cook' else '野菜前処理作業台'
                            matches=[c['label'] for c in pending['choices'] if c['value'].get('name')==chosen]
                            assert len(matches)==1,(chosen,pending)
                            label=matches[0]
                        else:raise AssertionError(pending)
                        await agent.reply(ident,label,pending['question_id'])
                        answered.add(pending['question_id'])
                    review=tools.flow.review(ident)
                    if review and review['state']=='awaiting_approval':break
                    assert current['status'] not in ('stopped','failed','reconciliation_required'),store.events(ident)[-6:]
                    await asyncio.sleep(.05)
                else:raise AssertionError(store.events(ident)[-8:])
                assert len(answered)==4
                # Crash during the read-only approval wait; preserve all answers.
                for _ in range(100):
                    with store.read() as db:
                        waiting=db.execute("SELECT 1 FROM commands WHERE run_id=? AND status='executing'",(ident,)).fetchone()
                    if waiting:break
                    await asyncio.sleep(.02)
                agent.tasks[ident].cancel()
                with suppress(asyncio.CancelledError):await agent.tasks[ident]
                await sandbox.stop(ident);store.recover()
                recovered=Providers(store,sandbox,tools,build_mcp(tools),{},gemini_client=client)
                await recovered.resume(ident)
                async with tools.lock(ident):await tools.flow.decide(store.get(ident),review['id'],True,'合成試験の三帳票を承認')
                await asyncio.wait_for(recovered.tasks[ident],120)
                assert store.get(ident)['status']=='completed',store.events(ident)[-8:]
                assert store.get(ident)['reserved']==0
                with store.read() as db:artifacts=[dict(row) for row in db.execute('SELECT * FROM artifacts WHERE run_id=?',(ident,))]
                assert len(artifacts)==3
                from openpyxl import load_workbook
                for artifact in artifacts:
                    assert hashlib.sha256(Path(artifact['path']).read_bytes()).hexdigest()==artifact['sha256']
                    assert load_workbook(BytesIO(Path(artifact['path']).read_bytes())).sheetnames
                result=[e['data'] for e in tools.flow.review(ident)['evidence'] if e['operation']=='lunch_get_price_result'][0]
                assert float(result['preview']['schools'][0]['total_price'])==2000
                assert len(store.get(ident)['task_inputs']['process_questions']['documents'][0]['answers'])==3
            finally:
                await sandbox.stop(ident);await client.aclose()
        asyncio.run(scenario())


def test_real_shift_business_condition_question_parse_confirm_resolve_and_output(tmp_path,monkeypatch):
    """Actual packaged rule parser/apply/solver/export; synthetic owned DB and no model."""
    from experiments.agent_assistant.business_workflow import BusinessWorkflow
    from experiments.agent_assistant.conversation import pending_question
    from experiments.agent_assistant.tests.test_business_workflow import NoModel
    monkeypatch.setenv('SAWA_TRIAL_OPERATOR_ACCOUNT','native-operator')
    with portal() as url:
        store,sandbox,tools=setup(tmp_path,'shift',url)
        async def scenario():
            selection=await tools.flow.options('shift_next_month',{'month':'2026-08'},'native-operator','Bearer synthetic-native-operator')
            context=tools.flow.selection('shift_next_month','native-operator',selection['selection_id'],['2026-08'])
            run,_=store.create('native-operator','native-shift-question','synthetic',['shift'],100,task_id='shift_next_month',task_inputs=context)
            sandbox.prepare(run);tools.authorizations[run['id']]='Bearer synthetic-native-operator'
            engine=BusinessWorkflow(store,tools,NoModel());engine.kick=lambda *a,**kw:None
            engine.kick_many=AsyncMock()
            engine.initialize(run);store.update(run['id'],status='running');run=store.get(run['id'])
            try:
                await sandbox.start(run)
                await engine.shift_questions.ask(run,'native-description','変更したい勤務条件を教えてください。',[],kind='text')
                question=pending_question(store.events(run['id']))
                assert engine.shift_questions.answer(store.get(run['id']),'休日数を2日にする',question['question_id'],None)
                assert await engine.advance(run['id']) is False
                question=pending_question(store.events(run['id']))
                assert '2日' in question['text'] and 'この変更を反映して作り直す' in question['options']
                assert not any(entry['request']['name']=='shift_apply_rules' for entry in tools.flow.commands(run['id']))
                assert engine.shift_questions.answer(store.get(run['id']),'この変更を反映して作り直す',question['question_id'],'この変更を反映して作り直す')
                assert await engine.advance(run['id']) is True
                assert (await tools.flow.read(store.get(run['id']),'shift_get_settings'))['monthly_holiday_exact_days']==2
                assert await engine.advance(run['id']) is False
                # The owned fixture has max_consecutive_work_days=6: two days
                # off in August cannot satisfy it. Ask again, never fake a plan.
                assert tools.flow.review(run['id']) is None
                question=pending_question(store.events(run['id']))
                assert question and '変更したい勤務条件を伝える' in question['options']
                assert store.get(run['id'])['status']=='waiting_input'
                assert engine.shift_questions.answer(store.get(run['id']),'変更したい勤務条件を伝える',question['question_id'],'変更したい勤務条件を伝える')
                assert await engine.advance(run['id']) is False
                question=pending_question(store.events(run['id']))
                assert engine.shift_questions.answer(store.get(run['id']),'休日数を9日にする',question['question_id'],None)
                assert await engine.advance(run['id']) is False
                question=pending_question(store.events(run['id']))
                assert '9日' in question['text']
                assert engine.shift_questions.answer(store.get(run['id']),'この変更を反映して作り直す',question['question_id'],'この変更を反映して作り直す')
                assert await engine.advance(run['id']) is True
                assert (await tools.flow.read(store.get(run['id']),'shift_get_settings'))['monthly_holiday_exact_days']==9
                assert await engine.advance(run['id']) is False
                # Four staff and exactly three required each day imply 31 total
                # off-days: a uniform exact count cannot fit this fixture. The
                # user explicitly removes that rule; no solver rule is weakened
                # automatically and no fabricated schedule is accepted.
                assert tools.flow.review(run['id']) is None
                question=pending_question(store.events(run['id']))
                assert '休日数を全員同じにする指定を外す' in question['options']
                assert engine.shift_questions.answer(store.get(run['id']),'休日数を全員同じにする指定を外す',question['question_id'],'休日数を全員同じにする指定を外す')
                assert await engine.advance(run['id']) is False
                question=pending_question(store.events(run['id']))
                assert engine.shift_questions.answer(store.get(run['id']),'この変更を反映して作り直す',question['question_id'],'この変更を反映して作り直す')
                assert await engine.advance(run['id']) is True
                settings=await tools.flow.read(store.get(run['id']),'shift_get_settings')
                assert settings['monthly_holiday_exact_days'] is None
                assert settings['max_consecutive_work_days']==6
                assert await engine.advance(run['id']) is False
                review=tools.flow.review(run['id']);assert review and review['state']=='awaiting_approval'
                await tools.flow.decide(store.get(run['id']),review['id'],True,'合成試験の確認')
                assert await engine.advance(run['id']) is False
                assert store.get(run['id'])['status']=='completed'
                assert not engine.providers.started
                with store.read() as db:
                    artifacts=db.execute('SELECT path FROM artifacts WHERE run_id=?',(run['id'],)).fetchall()
                assert artifacts and all(Path(item['path']).is_file() for item in artifacts)
            finally:
                await engine.close();await sandbox.stop(run['id'])
        asyncio.run(scenario())
