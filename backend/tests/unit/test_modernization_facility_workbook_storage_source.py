"""Synthetic storage components, not native OCR, live UI, or real GCS proof."""

import asyncio
from hashlib import sha256
from io import BytesIO
import json

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook

from src.services import order_form_service, storage_service


ENV_KEYS = (
    "OCR_OUTPUT_BUCKET", "RAW_BUCKET", "OCR_ARTIFACT_BUCKET", "OCR_ARTIFACT_DIR",
    "FAX_SOURCE_TEMPLATE_UPLOAD_DIR",
)


class MemoryBlob:
    def __init__(self, transport, key):
        self.transport = transport
        self.key = key

    def upload_from_string(self, data, content_type=None):
        self.transport.events.append(("upload", self.key, bytes(data), content_type))
        if self.transport.failure is not None:
            raise self.transport.failure
        self.transport.objects[self.key] = bytes(data)

    def reload(self, *, retry):
        assert retry is None
        self.generation = self.transport.generations[self.key]
        self.transport.events.append(("reload", self.key, self.generation))

    def download_as_bytes(self, **kwargs):
        if kwargs:
            generation = self.transport.generations[self.key]
            assert kwargs == {"if_generation_match": int(generation), "retry": None}
            assert type(kwargs["if_generation_match"]) is int
            assert self.generation == generation
            data = self.transport.objects[self.key]
            self.transport.events.append(("download-pinned", self.key, generation, sha256(data).hexdigest()))
            return data
        self.transport.events.append(("download", self.key))
        return self.transport.objects[self.key]


class MemoryBucket:
    def __init__(self, transport, name):
        self.transport = transport
        self.name = name

    def blob(self, name):
        return MemoryBlob(self.transport, (self.name, name))


class MemoryClient:
    """Only the external GCS Client/Bucket/Blob transport is replaced."""

    def __init__(self):
        self.events = []
        self.objects = {}
        self.generations = {}
        self.failure = None

    def bucket(self, name):
        return MemoryBucket(self, name)


@pytest.fixture
def storage_component(tmp_path, monkeypatch):
    for name in ENV_KEYS:
        monkeypatch.delenv(name, raising=False)
    client = MemoryClient()
    monkeypatch.setattr(storage_service._GCS_CLIENT_LOCAL, "client", client, raising=False)
    materialized = tmp_path / "materialized"
    monkeypatch.setattr(order_form_service, "_SOURCE_WORKBOOK_MATERIALIZED_DIR", materialized)
    source = tmp_path / "authored.xlsx"
    workbook = Workbook()
    workbook.active.title = "Authored week"
    workbook.active["A1"] = "Synthetic workbook storage component"
    workbook.active["B2"] = 0
    workbook.create_sheet("Second sheet")["C3"] = 69
    workbook.save(source)
    workbook.close()
    return client, source.read_bytes(), materialized


@pytest.mark.parametrize(("environment", "expected"), [
    ({}, None),
    ({"OCR_OUTPUT_BUCKET": ""}, None),
    ({"OCR_OUTPUT_BUCKET": " \t "}, None),
    ({"RAW_BUCKET": "raw-only"}, None),
    ({"OCR_ARTIFACT_BUCKET": "artifact-only"}, None),
    ({"OCR_OUTPUT_BUCKET": "  ", "RAW_BUCKET": "raw"}, None),
    ({"OCR_OUTPUT_BUCKET": "explicit-output"}, "explicit-output"),
    ({"OCR_OUTPUT_BUCKET": " \texplicit-output \n"}, "explicit-output"),
    ({"OCR_OUTPUT_BUCKET": "explicit-output", "RAW_BUCKET": "raw",
      "OCR_ARTIFACT_BUCKET": "artifact", "OCR_ARTIFACT_DIR": "/unused"}, "explicit-output"),
])
def test_explicit_output_resolver(storage_component, monkeypatch, environment, expected):
    for name, value in environment.items():
        monkeypatch.setenv(name, value)
    assert storage_service.get_default_output_bucket() == expected


@pytest.mark.parametrize("environment", [
    {}, {"RAW_BUCKET": "raw-only"}, {"OCR_ARTIFACT_BUCKET": "artifact-only"},
    {"OCR_OUTPUT_BUCKET": " \t ", "RAW_BUCKET": "raw"},
])
def test_missing_bucket_stops_before_storage_or_mkdir(
    storage_component, monkeypatch, tmp_path, environment,
):
    client, data, materialized = storage_component
    local = tmp_path / "forbidden-local-upload"
    monkeypatch.setenv("FAX_SOURCE_TEMPLATE_UPLOAD_DIR", str(local))
    for name, value in environment.items():
        monkeypatch.setenv(name, value)
    before = sorted(str(path.relative_to(tmp_path)) for path in tmp_path.rglob("*"))
    with pytest.raises(ValueError, match="^source_workbook_gcs_storage_required$"):
        order_form_service.save_facility_source_workbook_upload(
            facility_id="synthetic-storage", filename="authored.xlsx", data=data,
        )
    assert client.events == []
    assert client.objects == {}
    assert not local.exists()
    assert not materialized.exists()
    assert sorted(str(path.relative_to(tmp_path)) for path in tmp_path.rglob("*")) == before


@pytest.mark.parametrize("output_bucket", [None, " \t "])
def test_cloud_run_registered_endpoint_retains_missing_storage_503(
    storage_component, monkeypatch, tmp_path, output_bucket,
):
    """Direct endpoint component with actual JSON snapshot, not HTTP/auth proof."""
    from src.api import facilities
    from src.services import facility_master_service

    client, data, materialized = storage_component
    master_path = tmp_path / "authored-facilities.json"
    master_path.write_text(json.dumps({"facilities": [{
        "facility_id": "synthetic-storage",
        "facility_name": "Synthetic workbook storage facility", "areas": [],
    }]}), encoding="utf-8")
    master_bytes = master_path.read_bytes()
    master_digest = sha256(master_bytes).hexdigest()
    master_key = ("componentcanonical", "facilities.json")
    master_uri = "gs://componentcanonical/facilities.json"
    master_generation = "123456789"
    client.objects[master_key] = master_bytes
    client.generations[master_key] = master_generation
    monkeypatch.setenv("FACILITY_MASTER_URI", master_uri)
    monkeypatch.delenv("FACILITY_MASTER_PATH", raising=False)
    monkeypatch.setenv("K_SERVICE", "synthetic-cloud-run-component")
    monkeypatch.setenv("RAW_BUCKET", "must-not-use-raw")
    if output_bucket is not None:
        monkeypatch.setenv("OCR_OUTPUT_BUCKET", output_bucket)
    snapshot = facility_master_service.get_snapshot()
    assert snapshot.source == master_uri
    assert snapshot.revision.startswith("gcs:")
    assert snapshot.revision.endswith(f":{master_generation}")
    assert snapshot.master == json.loads(master_bytes)
    assert snapshot.master["facilities"][0]["facility_id"] == "synthetic-storage"
    assert any(route.endpoint is facilities.upload_order_form_source_workbook
               for route in facilities.router.routes)
    upload = UploadFile(filename="authored.xlsx", file=BytesIO(data))
    try:
        with pytest.raises(HTTPException) as caught:
            asyncio.run(facilities.upload_order_form_source_workbook(
                "synthetic-storage", file=upload, month_id=None,
                expected_revision=snapshot.revision,
            ))
        assert caught.value.status_code == 503
        assert caught.value.detail == "source_workbook_gcs_storage_required"
        assert upload.file.tell() == 0
        assert not materialized.exists()
        assert facility_master_service.get_snapshot().revision == snapshot.revision
        assert client.objects == {master_key: master_bytes}
        assert client.generations == {master_key: master_generation}
        assert client.events == [
            ("reload", master_key, master_generation),
            ("download-pinned", master_key, master_generation, master_digest),
        ] * 3
    finally:
        upload.file.close()


def test_actual_upload_materialize_and_workbook_validation(storage_component, monkeypatch):
    client, data, materialized = storage_component
    monkeypatch.setenv("OCR_OUTPUT_BUCKET", " explicit-output ")
    monkeypatch.setenv("RAW_BUCKET", "must-not-use-raw")
    digest = sha256(data).hexdigest()
    object_name = f"artifacts/facility-template-synthetic-storage/authored_{digest[:12]}.xlsx"
    result = order_form_service.save_facility_source_workbook_upload(
        facility_id="synthetic-storage", filename="authored.xlsx", data=data,
    )
    assert result == {
        "uri": f"gs://explicit-output/{object_name}", "storage": "gcs",
        "filename": "authored.xlsx", "sha256": digest,
        "sheetnames": ["Authored week", "Second sheet"],
    }
    key = ("explicit-output", object_name)
    assert client.objects == {key: data}
    assert client.events == [
        ("upload", key, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("download", key),
    ]
    files = list(materialized.iterdir())
    assert files == [materialized / f"{digest}.xlsx"]
    assert files[0].read_bytes() == data
    workbook = load_workbook(files[0], read_only=True)
    try:
        assert workbook.sheetnames == ["Authored week", "Second sheet"]
        assert workbook["Authored week"]["B2"].value == 0
        assert workbook["Second sheet"]["C3"].value == 69
    finally:
        workbook.close()


def test_provider_failure_propagates_without_local_alternative(storage_component, monkeypatch, tmp_path):
    client, data, materialized = storage_component
    failure = OSError("synthetic provider upload failure")
    client.failure = failure
    local = tmp_path / "forbidden-local-upload"
    monkeypatch.setenv("OCR_OUTPUT_BUCKET", "explicit-output")
    monkeypatch.setenv("RAW_BUCKET", "must-not-use-raw")
    monkeypatch.setenv("FAX_SOURCE_TEMPLATE_UPLOAD_DIR", str(local))
    with pytest.raises(OSError) as caught:
        order_form_service.save_facility_source_workbook_upload(
            facility_id="synthetic-storage", filename="authored.xlsx", data=data,
        )
    assert caught.value is failure
    assert len(client.events) == 1
    assert client.events[0][0] == "upload"
    assert client.events[0][1][0] == "explicit-output"
    assert client.objects == {}
    assert not local.exists()
    assert not materialized.exists()
