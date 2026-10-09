"""Real SQLite/JSON/openpyxl and public HTTP unit-boundary regressions.

Authorization is synthetic. These definitions do not prove operator FAX/OCR,
PostgreSQL concurrency, staging, or live source availability.
"""

from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import datetime
from html.parser import HTMLParser
from io import BytesIO
from threading import Barrier

from openpyxl import load_workbook
import pytest

from modernization_saved_sheet_export_fixtures import SavedSheetFixture, WEEK, source_workbook
from src.models.facility_template_version import FacilityTemplateVersion
from src.models.order import Order
from src.models.order_ocr_evidence_run import OrderOcrEvidenceRun
from src.models.order_sheet_draft import OrderSheetDraft
from src.models.order_workflow_state import OrderWorkflowState
from src.services import order_form_service


class PreviewRows(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        elif tag == "td":
            self.cell = ""

    def handle_data(self, data):
        if self.cell is not None:
            self.cell += data

    def handle_endtag(self, tag):
        if tag == "td" and self.cell is not None:
            self.row.append(self.cell)
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None


@pytest.fixture
def saved(tmp_path, monkeypatch):
    fixture = SavedSheetFixture(tmp_path, monkeypatch)
    try:
        yield fixture
    finally:
        fixture.close()


def assert_excel(content, saved, *, marker="authored-target"):
    workbook = load_workbook(BytesIO(content))
    try:
        worksheet = workbook[WEEK]
        assert [worksheet[cell].value for cell in ("E11", "F11", "G11", "H11")] == [12, 3, 4, 0]
        assert worksheet["D11"].value == "選択献立"
        assert worksheet["A11"].value.date().isoformat() == "2026-03-01"
        assert worksheet["A4"].value == "保存済み施設"
        assert worksheet["J5"].value == marker
        metadata = {row[0]: row[1] for row in workbook["設定"].iter_rows(values_only=True)
                    if len(row) > 1 and row[0] is not None}
        assert metadata["order_id"] == saved.order_id
        assert metadata["saved_sheet_draft_id"] == saved.draft_id
        assert metadata["month_id"] == "2026-03"
        assert workbook.sheetnames == [WEEK, "設定"]
    finally:
        workbook.close()


def test_pre_confirmation_excel_uses_selected_sheet_and_preserves_raw_payload(saved):
    saved.verify_current()
    response = saved.request("excel")
    assert response.status_code == 200, response.text if response.status_code != 200 else ""
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "fax_read_sheet_excel_" in response.headers["content-disposition"]
    assert_excel(response.content, saved)
    saved.verify_current()


def test_pre_confirmation_public_preview_shows_same_selected_quantities(saved):
    response = saved.request("preview")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/html")
    assert "選択献立" in response.text
    assert "999" not in response.text and "777" not in response.text
    parsed = PreviewRows()
    parsed.feed(response.text)
    selected_rows = [row for row in parsed.rows if "選択献立" in row]
    assert len(selected_rows) == 1
    assert selected_rows[0][4:8] == ["12", "3", "4", "0"]
    files = list(saved.output.rglob("*.xlsx"))
    assert len(files) == 1
    assert_excel(files[0].read_bytes(), saved)
    saved.verify_current()


STOP_CASES = [
    ("missing-pointer", "saved sheet not found"),
    ("foreign-draft", "saved sheet order mismatch"),
    ("selected-evidence", "saved sheet selected OCR mismatch"),
    ("template", "saved sheet template mismatch"),
    ("embedded-evidence", "saved sheet OCR lineage mismatch"),
    ("facility-context", "saved sheet context mismatch"),
    ("date-context", "saved sheet context mismatch"),
    ("dated-version", "facility_version_date_unrecorded"),
    ("projection-proof", "canonical_lineage_selection_proof_missing"),
    ("candidate-source", "saved sheet candidate not adopted"),
    ("candidate-adopted-false", "saved sheet candidate not adopted"),
    ("row-date-not-period", "order week is unresolved"),
    ("month-not-period", "order week is unresolved"),
]


def invalidate(saved, case):
    with saved.session_scope() as session:
        workflow = session.get(OrderWorkflowState, saved.order_id)
        draft = session.get(OrderSheetDraft, saved.draft_id)
        evidence = session.get(OrderOcrEvidenceRun, saved.evidence_id)
        if case == "missing-pointer":
            workflow.draft_id = None
        elif case == "foreign-draft":
            workflow.draft_id = "DRAFT-FOREIGN"
        elif case == "selected-evidence":
            draft.base_evidence_run_id = None
        elif case == "template":
            evidence.template_version_id = None
        elif case in {"embedded-evidence", "candidate-source", "candidate-adopted-false"}:
            payload = deepcopy(draft.draft_sheet_json)
            if case == "embedded-evidence":
                payload["base_evidence_run_id"] = "OCR-UNSELECTED"
            elif case == "candidate-source":
                payload["source"] = "reparse_candidate"
            else:
                payload["adopted"] = False
            draft.draft_sheet_json = payload
        elif case in {"facility-context", "date-context"}:
            metadata = deepcopy(workflow.secondary_actions_json)
            metadata["workflow_v2"]["facility_id" if case == "facility-context" else "week_start"] = (
                "FOREIGN" if case == "facility-context" else "2026-03-02")
            workflow.secondary_actions_json = metadata
        elif case == "dated-version":
            session.get(Order, saved.order_id).week_code = "2026-04@2026-04-01~2026-04-07"
            workflow.secondary_actions_json = {}
        elif case == "projection-proof":
            projection = session.get(FacilityTemplateVersion, saved.projection_id)
            metadata = deepcopy(projection.validation_json)
            del metadata["creation_source_proof"]
            projection.validation_json = metadata
        elif case in {"row-date-not-period", "month-not-period"}:
            session.get(Order, saved.order_id).week_code = "" if case == "row-date-not-period" else "2026-03"
        else:
            raise AssertionError(case)


@pytest.mark.parametrize("case,detail", STOP_CASES, ids=[item[0] for item in STOP_CASES])
def test_both_public_boundaries_reject_retired_or_foreign_lineage_without_fallback(saved, case, detail):
    invalidate(saved, case)
    for boundary in ("excel", "preview"):
        response = saved.request(boundary)
        assert response.status_code == 400, (boundary, response.text)
        assert response.json()["detail"] == detail
        assert list(saved.output.iterdir()) == []


@pytest.mark.parametrize("case", ["missing-source", "missing-week", "different-alias-bytes"])
def test_both_public_boundaries_never_substitute_configured_artifact(saved, case):
    if case == "missing-source":
        saved.source.unlink()
        source_workbook(saved.root / "earlier-month.xlsx", sheet="2月1日～2月7日")
    elif case == "missing-week":
        source_workbook(saved.source, sheet="2月1日～2月7日")
    else:
        other = saved.root / "different.xlsx"
        source_workbook(other, marker="foreign-bytes")
        saved.authored["order_form_source_workbook_uri"] = other.as_uri()
        saved.republish()
    for boundary in ("excel", "preview"):
        response = saved.request(boundary)
        assert response.status_code == (500 if case == "missing-source" else 400), response.text
        assert ("No such file" if case == "missing-source" else
                "saved sheet week not found in configured workbook" if case == "missing-week" else
                "saved sheet source workbook ambiguous") in response.json()["detail"]
        assert list(saved.output.iterdir()) == []
    if case == "missing-source":
        del saved.authored["order_form_source_workbook"]
        saved.republish()
        for boundary in ("excel", "preview"):
            response = saved.request(boundary)
            assert response.status_code == 400, response.text
            assert response.json()["detail"] == "saved sheet source workbook not configured"
            assert list(saved.output.iterdir()) == []


def test_registered_filename_and_native_file_uri_are_same_artifact_aliases(saved):
    saved.authored["order_form_source_workbook_uri"] = saved.source.as_uri()
    saved.republish()
    for boundary in ("excel", "preview"):
        response = saved.request(boundary)
        assert response.status_code == 200, response.text if response.status_code != 200 else ""
    for path in saved.output.rglob("*.xlsx"):
        assert_excel(path.read_bytes(), saved)


def test_explicit_month_mapping_not_general_workbook_is_authoritative(saved):
    target = saved.root / "explicit-march.xlsx"
    source_workbook(target, marker="explicit-month")
    source_workbook(saved.source, sheet="2月1日～2月7日", marker="base-must-not-win")
    saved.authored["order_form_month_sources"] = {"2026-03": target.name}
    saved.authored["order_form_month_source_uris"] = {"2026-03": target.as_uri()}
    saved.republish()
    response = saved.request("excel")
    assert response.status_code == 200, response.text if response.status_code != 200 else ""
    assert_excel(response.content, saved, marker="explicit-month")
    response = saved.request("preview")
    assert response.status_code == 200 and "explicit-month" in response.text


@pytest.mark.parametrize("retirement", ["sheet", "canonical-json", "source-bytes"])
def test_final_capture_rejects_input_retired_after_real_row_write(saved, monkeypatch, retirement):
    actual_write = order_form_service._write_saved_sheet_rows_to_order_form
    calls = []

    def write_then_retire(*args, **kwargs):
        result = actual_write(*args, **kwargs)
        assert result == 1
        calls.append(retirement)
        if retirement == "sheet":
            with saved.session_scope() as session:
                workflow = session.get(OrderWorkflowState, saved.order_id)
                workflow.draft_id = ("DRAFT-LATEST-UNSELECTED"
                                     if workflow.draft_id == saved.draft_id else saved.draft_id)
        elif retirement == "canonical-json":
            saved.authored["facility_name"] = "Changed authored configuration"
            saved.write_master()
        else:
            source_workbook(saved.source, marker=f"replaced-after-render-{len(calls)}")
        return result

    # Schedule actual external retirement, forwarding the entire real quantity writer.
    monkeypatch.setattr(order_form_service, "_write_saved_sheet_rows_to_order_form", write_then_retire)
    for boundary in ("excel", "preview"):
        response = saved.request(boundary)
        assert response.status_code == 400, response.text
        assert response.json()["detail"] == (
            "saved sheet dated template mismatch" if retirement == "canonical-json" else "saved sheet output input changed")
        assert list(saved.output.iterdir()) == []
    assert len(calls) == (1 if retirement == "canonical-json" else 2)


def test_same_clock_parallel_http_generations_own_distinct_intact_artifacts(saved, monkeypatch):
    class SameClock(datetime):
        @classmethod
        def utcnow(cls):
            return cls(2026, 3, 1, 10, 0, 0)

    monkeypatch.setattr(order_form_service, "datetime", SameClock)
    actual_write = order_form_service._write_saved_sheet_rows_to_order_form
    barrier = Barrier(2)

    def synchronized_actual_write(*args, **kwargs):
        result = actual_write(*args, **kwargs)
        barrier.wait(timeout=10)
        return result

    monkeypatch.setattr(order_form_service, "_write_saved_sheet_rows_to_order_form", synchronized_actual_write)
    with ThreadPoolExecutor(max_workers=2) as executor:
        replies = list(executor.map(lambda _: saved.request("excel"), range(2)))
    assert [reply.status_code for reply in replies] == [200, 200]
    paths = list(saved.output.rglob("*.xlsx"))
    assert len(paths) == 2 and len({path.parent for path in paths}) == 2
    assert paths[0].name == paths[1].name
    for response in replies:
        assert "20260301_100000.xlsx" in response.headers["content-disposition"]
        assert_excel(response.content, saved)
    for path in paths:
        assert_excel(path.read_bytes(), saved)
    saved.verify_current()
