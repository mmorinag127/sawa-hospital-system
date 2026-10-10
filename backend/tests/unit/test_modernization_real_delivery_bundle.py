"""Real confirmed synthetic saved-sheet receipts -> mounted XLSX/ZIP HTTP.

AUTH_DISABLED, owned SQLite, no startup, manually seeded OCR. Not native/auth/
PG/live proof. Only OUTPUT_DIR is rebound. The single invoice uses an explicitly
authored disposable workbook; daily embedding uses the existing non-template
invoice contract because templated daily embedding is explicitly prohibited.
"""

import csv
from datetime import date, timedelta
from hashlib import sha256
from io import BytesIO, StringIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
import pytest

from test_modernization_real_bagging_confirmation import owned_bagging as owned_bagging
from test_modernization_workflow_surface_parity import surface as surface
from test_modernization_real_daily_delivery_html import (
    confirm, preservation, assert_preserved,
)
from src.models.document import OrderDocument
from src.models.order import Order
from src.services import order_service as orders, output_builder as builder


MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def bundle_order(owned_bagging, surface):
    fixture = owned_bagging
    # The non-template writer consumes row keys, with optional header labels.
    # Explicit source-key names keep date/menu identity in the real workbook.
    fixture.authored["invoice_template"] = {
        "include_menu_name": True,
        "columns": [{"name": "date", "header": "Date", "source": "date"},
                    {"name": "menu_name", "header": "Menu", "source": "menu_name"},
                    {"name": "Quantity", "source": "quantity", "diet_type": "regular", "area_id": "X"}],
    }
    fixture.republish()
    with fixture.session_scope() as session:
        orders._record_current_order_version(session, session.get(Order, fixture.order_id),
                                             session.get(OrderDocument, "UNIT-DOC"))
    return fixture, surface[1], confirm(fixture)


def receive(response, fixture, name, mime, capsys):
    assert response.status_code == 200, response.text
    body = response.content
    assert body and int(response.headers["content-length"]) == len(body)
    assert response.headers["content-type"] == mime
    assert response.headers["content-disposition"].startswith("attachment;")
    assert name.rsplit(".", 1)[1] in response.headers["content-disposition"]
    path = fixture.root / "http-received" / name
    path.parent.mkdir(exist_ok=True)
    path.write_bytes(body)
    assert path.read_bytes() == body
    with capsys.disabled():
        print(f"http-received artifact={path} sha256={sha256(body).hexdigest()}")
    return body


def workbook_rows(body):
    with ZipFile(BytesIO(body)) as archive:
        assert archive.testzip() is None
    workbook = load_workbook(BytesIO(body), data_only=False)
    try:
        return {sheet.title: list(sheet.values) for sheet in workbook.worksheets}
    finally:
        workbook.close()


def label_rows(rows, *, zero, empty_header=None):
    assert rows[0][6:10] == ("商品名１", "商品名２", "内容量", "内容詳細")
    assert rows[0][10] == empty_header
    values = {row[6]: (row[8], row[9], row[10]) for row in rows[1:]}
    assert values == ({"選択献立0": ("700g", "100g", "7人前"),
                       "選択献立3": ("0g", "100g", "0人前")} if zero else
                      {"選択献立0": ("700g", "100g", "7人前")})


def test_real_single_delivery_workbook_uses_captured_name_and_quantities(
    owned_bagging, surface, monkeypatch, capsys,
):
    fixture = owned_bagging
    template = fixture.root / "synthetic-delivery-template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic invoice"
    sheet["E1"] = "施設名"
    sheet.append(["Date", "Menu", "Quantity"])
    workbook.save(template)
    workbook.close()
    template_hash = sha256(template.read_bytes()).hexdigest()
    fixture.authored["invoice_template"] = {
        "include_menu_name": True, "template_uri": str(template), "sheet_name": "Synthetic invoice",
        "columns": [{"name": "Date", "source": "date"}, {"name": "Menu", "source": "menu_name"},
                    {"name": "Quantity", "source": "quantity", "diet_type": "regular", "area_id": "X"}],
    }
    fixture.republish()
    with fixture.session_scope() as session:
        orders._record_current_order_version(session, session.get(Order, fixture.order_id),
                                             session.get(OrderDocument, "UNIT-DOC"))
    candidate = confirm(fixture)
    monkeypatch.setattr(builder, "OUTPUT_DIR", fixture.output)
    before = preservation(fixture)
    body = receive(surface[1].get("/outputs/delivery-notes", params={"order_id": fixture.order_id}),
                   fixture, "single-delivery.xlsx", MIME, capsys)
    rows = workbook_rows(body)["Synthetic invoice"]
    assert rows[1][4] == candidate["configured_input"]["facility"]["facility_name"]
    values = {row[1]: row[2] for row in rows[2:] if row[1]}
    assert values["選択献立0"] == 7
    assert values["選択献立3"] is None  # Existing zero_as_empty rendering.
    assert all(row[0].date().isoformat() == str(candidate["lines"][0]["date"])[:10]
               for row in rows[2:] if row[1])
    assert sha256(template.read_bytes()).hexdigest() == template_hash
    assert_preserved(fixture, before)


@pytest.mark.parametrize("mode", ["labels", "labels_csv", "delivery", "both"])
def test_real_daily_bundle_modes_receive_exact_members_and_accepted_identity(
    bundle_order, mode, monkeypatch, capsys,
):
    fixture, client, candidate = bundle_order
    monkeypatch.setattr(builder, "OUTPUT_DIR", fixture.output)
    day = date.fromisoformat(str(candidate["lines"][0]["date"])[:10])
    before = preservation(fixture)
    response = client.get("/outputs/daily-bundle", params={"date": day.isoformat(), "bundle_type": mode})
    zipped = mode in {"labels", "labels_csv"}
    body = receive(response, fixture, f"daily-{mode}.{'zip' if zipped else 'xlsx'}",
                   "application/zip" if zipped else MIME, capsys)
    for key, expected in {"total-orders": "1", "success-orders": "1", "empty-orders": "0",
                          "error-orders": "0", "type": mode}.items():
        assert response.headers["x-daily-bundle-" + key] == expected
    facility_name = candidate["configured_input"]["facility"]["facility_name"]
    if zipped:
        with ZipFile(BytesIO(body)) as archive:
            assert archive.testzip() is None
            assert len(archive.namelist()) == 1
            name = archive.namelist()[0]
            assert day.isoformat() in name and facility_name in name
            member = archive.read(name)
        if mode == "labels_csv":
            rows = list(csv.reader(StringIO(member.decode("cp932"))))
            label_rows([tuple(row) for row in rows], zero=True, empty_header="")
        else:
            sheets = workbook_rows(member)
            assert len(sheets) == 1
            label_rows(next(iter(sheets.values())), zero=True)
    else:
        sheets = workbook_rows(body)
        delivery = [rows for name, rows in sheets.items()
                    if name == (facility_name if mode == "delivery" else "納品書_" + facility_name)]
        assert len(delivery) == 1
        assert delivery[0][0] == ("Date", "Menu", "Quantity")
        assert all(row[0].date() == day for row in delivery[0][1:])
        quantities = {row[1]: row[2] for row in delivery[0][1:]}
        assert quantities == ({"選択献立0": 7} if mode == "delivery" else
                              {"選択献立0": 7, "選択献立1": None,
                               "選択献立2": None, "選択献立3": None})
        assert all(facility_name in name for name in sheets)
        if mode == "both":
            labels = [rows for name, rows in sheets.items() if name == "ラベル_" + facility_name]
            assert len(labels) == 1
            label_rows(labels[0], zero=False)
        else:
            assert len(sheets) == 1
    assert_preserved(fixture, before)
    absent = client.get("/outputs/daily-bundle", params={
        "date": (day + timedelta(days=7)).isoformat(), "bundle_type": mode,
    })
    assert absent.status_code == 400
    assert absent.json() == {"detail": "対象日の出力対象がありません"}
    assert_preserved(fixture, before)
