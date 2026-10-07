from contextlib import contextmanager
from datetime import date
from hashlib import sha256
import importlib.util
from io import BytesIO
from pathlib import Path

from alembic.migration import MigrationContext
from alembic.operations import Operations
from openpyxl import Workbook
import pytest
import sqlalchemy as sa
from sqlalchemy.orm import Session

from src.models.facility import Facility, FacilityConfig
from src.models.menu import (BaseMenuCycleItem, MenuEditRevision, MenuFacilityOverride,
    MenuMaster, MenuRule, MonthlyMenu, MonthlyMenuEntry, MonthlyMenuImport, MonthlyMenuItem)
from src.models.user import AuditLog
from src.services import menu_service as service
from src.services import menu_edit_revision_service as revisions
from src.services.menu_upload_archive_service import SavedMenuUpload


ROOT = Path(__file__).resolve().parents[3]


def migration(connection, number):
    path = next((ROOT / "backend/migrations").glob(f"{number}_*.py"))
    spec = importlib.util.spec_from_file_location(f"monthly_import_{number}", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    with Operations.context(MigrationContext.configure(connection)):
        module.upgrade()


@pytest.fixture
def store(tmp_path, monkeypatch):
    engine = sa.create_engine(sa.URL.create("sqlite", database=str(tmp_path / "menu.sqlite3")))
    metadata = sa.MetaData()
    for model in (MonthlyMenu, MonthlyMenuItem, MonthlyMenuEntry, MenuMaster, MenuFacilityOverride,
                  MenuRule, BaseMenuCycleItem, AuditLog, Facility, FacilityConfig):
        model.__table__.to_metadata(metadata)
    metadata.create_all(engine)
    with engine.begin() as connection:
        migration(connection, "0030")
        migration(connection, "0031")

    @contextmanager
    def scope():
        with Session(engine, autoflush=False, expire_on_commit=True) as session:
            try:
                yield session
                session.commit()
            except Exception:
                session.rollback()
                raise

    monkeypatch.setattr(service, "session_scope", scope)
    yield scope, engine
    engine.dispose()


def archived(tmp_path, content=None):
    if content is None:
        workbook = Workbook()
        workbook.active.title = "source"
        workbook.active.append(["2026年10月"])
        stream = BytesIO()
        workbook.save(stream)
        content = stream.getvalue()
    path = tmp_path / "source.xlsx"
    path.write_bytes(content)
    return content, SavedMenuUpload(str(path), path.name, sha256(content).hexdigest(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def parsed(monkeypatch, *, diet="regular", month=10, slots=4):
    items = [{"name": f"Fish {slot}", "daypart": meal, "category": "主菜", "diet_type": diet}
             for meal in ("朝食", "昼食", "夕食") for slot in range(slots)]
    entries = [{**item, "menu_date": date(2026, month, 1), "slot_index": index % slots}
               for index, item in enumerate(items)]
    monkeypatch.setattr(service, "_parse_monthly_menu", lambda *args: (date(2026, month, 1), diet, items, entries))


def masters(scope, *, quantity=1):
    with scope() as session:
        for slot in range(4):
            session.add(MenuMaster(id=f"fish-{slot}", name=f"Fish {slot}", normalized_name=f"fish{slot}",
                unit_type="cut", qty_per_serving=quantity, bag_max_qty=20, bag_max_unit="cut"))


def save(tmp_path, **kwargs):
    content, source = archived(tmp_path)
    return service.save_menu_import("2026-10", content, source, sheet_name=None, scope_override=None, **kwargs)


def test_saved_receipt_does_not_apply_and_survives_commit_close(store, tmp_path, monkeypatch):
    scope, engine = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    assert receipt["source"]["content_sha256"] == sha256((tmp_path / "source.xlsx").read_bytes()).hexdigest()
    assert receipt["source"]["sheet_name"] == "source"
    assert receipt["review"]["item_count"] == 12
    assert receipt["state"] == "review"
    with scope() as session:
        assert session.query(MonthlyMenuItem).count() == 0
        assert session.query(MonthlyMenuEntry).count() == 0
        assert session.get(MenuEditRevision, "monthly-menu:2026-10").revision == 1
    assert service.read_menu_import("2026-10", receipt["id"]) == receipt
    own = service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True)
    assert own["receipt"]["state"] == "applied"
    assert len(own["snapshot"]["items"]) == 12
    assert len(own["snapshot"]["entries"]) == 12
    assert all(item["qty_per_serving"] == 1 and item["unit_type"] == "cut" for item in own["snapshot"]["items"])
    assert {entry["daypart"] for entry in own["snapshot"]["entries"]} == {"朝食", "昼食", "夕食"}
    for meal in ("朝食", "昼食", "夕食"):
        assert len([entry for entry in own["snapshot"]["entries"] if entry["daypart"] == meal]) == 4
    assert own["snapshot"]["revision"] != receipt["revision"]
    with pytest.raises(revisions.RevisionConflict, match="already applied"):
        service.review_menu_import("2026-10", receipt["id"], own["receipt"]["revision"], None, apply=True)


@pytest.mark.parametrize("token", [None, "malformed", "stale"])
def test_missing_invalid_stale_apply_leaves_receipt_and_business_rows(store, tmp_path, monkeypatch, token):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    error = revisions.RevisionPreconditionRequired if token is None else revisions.RevisionConflict
    with pytest.raises(error):
        service.review_menu_import("2026-10", receipt["id"], token, None, apply=True)
    assert service.read_menu_import("2026-10", receipt["id"]) == receipt
    with scope() as session:
        assert session.query(MonthlyMenuItem).count() == 0
        assert session.get(MenuEditRevision, "monthly-menu:2026-10").revision == 1


def test_reload_never_rebases_saved_review_to_new_master_source(store, tmp_path, monkeypatch):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    with scope() as session:
        session.get(MenuMaster, "fish-0").qty_per_serving = 9
    refreshed = service.read_menu_import("2026-10", receipt["id"])
    assert refreshed["review"] == receipt["review"]
    assert refreshed["revision"] == receipt["revision"]
    assert refreshed["input_revision"] == receipt["input_revision"]
    assert refreshed["current_input_revision"] != receipt["input_revision"]
    with pytest.raises(revisions.RevisionConflict):
        service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True)
    with scope() as session:
        assert session.query(MonthlyMenuItem).count() == 0
        assert session.get(MenuMaster, "fish-0").qty_per_serving == 9
    reviewed = service.review_menu_import("2026-10", receipt["id"], refreshed["revision"], [],
        input_expected=refreshed["current_input_revision"])
    assert reviewed["input_revision"] == refreshed["current_input_revision"]
    own = service.review_menu_import("2026-10", receipt["id"], reviewed["revision"], None, apply=True)
    assert own["receipt"]["state"] == "applied"


def test_archive_changed_or_missing_never_applies(store, tmp_path, monkeypatch):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    (tmp_path / "source.xlsx").write_bytes(b"changed source")
    with pytest.raises(ValueError, match="digest mismatch"):
        service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True)
    (tmp_path / "source.xlsx").unlink()
    with pytest.raises(FileNotFoundError):
        service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True)
    assert service.list_menu_imports("2026-10")[0]["state"] == "review"
    with scope() as session:
        assert session.query(MonthlyMenuEntry).count() == 0


@pytest.mark.parametrize("diet", ["regular", "soft", "mixer", None])
def test_repeat_import_targets_recorded_diet_only_and_preserves_other_scope(store, tmp_path, monkeypatch, diet):
    scope, _ = store
    parsed(monkeypatch, diet=diet)
    masters(scope, quantity=0)
    for _ in range(2):
        receipt = save(tmp_path)
        own = service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True)
        with scope() as session:
            if session.get(MonthlyMenuItem, "unrelated") is None:
                session.add(MonthlyMenuItem(id="unrelated", monthly_menu_id="2026-10", name="Other diet", diet_type="tea", qty_per_serving=7))
                session.add(MonthlyMenuEntry(id="unrelated-entry", monthly_menu_id="2026-10", name="Other diet", diet_type="tea", menu_date=date(2026, 10, 1), daypart="昼食", slot_index=0))
                session.add(MonthlyMenuItem(id="facility-item", monthly_menu_id="2026-10", name="Facility exception", facility_override="F1", qty_per_serving=3))
    with scope() as session:
        assert session.query(MonthlyMenuItem).filter(MonthlyMenuItem.name.like("Fish%"), MonthlyMenuItem.facility_override.is_(None)).count() == 12
        assert session.query(MenuMaster).count() == 4
        assert session.get(MonthlyMenuItem, "unrelated").qty_per_serving == 7
        assert session.get(MonthlyMenuItem, "facility-item").qty_per_serving == 3
    assert all(item["qty_per_serving"] == 0 for item in own["snapshot"]["items"] if item["name"].startswith("Fish"))


def test_failure_after_materialization_rolls_back_rows_receipt_and_cas(store, tmp_path, monkeypatch):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    actual = revisions.compare_and_advance

    def fail(session, *args):
        actual(session, *args)
        raise RuntimeError("after CAS")

    monkeypatch.setattr(revisions, "compare_and_advance", fail)
    with pytest.raises(RuntimeError, match="after CAS"):
        service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True)
    assert service.read_menu_import("2026-10", receipt["id"]) == receipt
    with scope() as session:
        assert session.query(MonthlyMenuItem).count() == 0
        assert session.get(MenuEditRevision, "monthly-menu:2026-10").revision == 1


def test_unresolved_master_is_saved_then_explicit_review_and_apply(store, tmp_path, monkeypatch):
    scope, _ = store
    parsed(monkeypatch)
    receipt = save(tmp_path)
    assert len(receipt["review"]["issues"]) == 4
    with pytest.raises(service.MenuMasterResolutionRequired):
        service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True)
    resolutions = [{"source_name": f"Fish {slot}", "action": "create", "unit_type": "count", "qty_per_serving": 0,
        "bag_max_qty": 20, "bag_max_unit": "count"} for slot in range(4)]
    reviewed = service.review_menu_import("2026-10", receipt["id"], receipt["revision"], resolutions,
        input_expected=receipt["input_revision"])
    assert reviewed["review"]["issues"] == []
    assert reviewed["revision"] != receipt["revision"]
    assert service.read_menu_import("2026-10", receipt["id"]) == reviewed
    with pytest.raises(revisions.RevisionConflict):
        service.review_menu_import("2026-10", receipt["id"], reviewed["revision"], [], apply=True)
    own = service.review_menu_import("2026-10", receipt["id"], reviewed["revision"], None, apply=True)
    assert own["item_count"] == 12
    with scope() as session:
        assert session.query(MenuMaster).count() == 4
        assert all(master.qty_per_serving == 0 for master in session.query(MenuMaster))


def test_adjacent_month_read_and_command_use_explicit_recorded_source(store, tmp_path, monkeypatch):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    own = service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True)
    with scope() as session:
        row = session.query(MonthlyMenuEntry).first()
        row.menu_date = date(2026, 11, 1)
    read = service.get_monthly_editor_snapshot("2026-11")
    assert read["menu"]["id"] == "2026-11"
    assert read["source_month_id"] == "2026-10"
    assert read["entries"][0]["source_month_id"] == "2026-10"
    item = read["items"][0]
    saved = service.monthly_editor_command("2026-11", item["id"], {"name": item["name"], "qty_per_serving": 2}, read["revision"], action="item")
    assert saved["snapshot"]["source_month_id"] == "2026-10"
    assert saved["snapshot"]["revision"] != read["revision"]
    with pytest.raises(revisions.RevisionConflict):
        service.monthly_editor_command("2026-12", item["id"], {"name": "Wrong month"}, saved["snapshot"]["revision"], action="item")
    with scope() as session:
        assert session.get(MonthlyMenuItem, item["id"]).monthly_menu_id == "2026-10"
        assert session.get(MenuMaster, session.get(MonthlyMenuItem, item["id"]).menu_master_id).qty_per_serving == 2


def test_month_mismatch_and_exact_import_target_stop(store, tmp_path, monkeypatch):
    scope, _ = store
    parsed(monkeypatch, month=11)
    masters(scope)
    with pytest.raises(ValueError, match="menu_month_mismatch"):
        save(tmp_path)
    with scope() as session:
        assert session.query(MonthlyMenuImport).count() == 0
        assert session.get(MenuEditRevision, "monthly-menu:2026-10") is None
    parsed(monkeypatch)
    receipt = save(tmp_path)
    with scope() as session:
        session.add(MonthlyMenu(id="2026-11", month_start=date(2026, 11, 1)))
        session.add(MenuEditRevision(scope="monthly-menu:2026-11", revision=1))
    with pytest.raises(ValueError, match="not found"):
        service.read_menu_import("2026-11", receipt["id"])


def test_two_commands_from_same_receipt_have_one_winner(store, tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    barrier = Barrier(2)

    def apply():
        barrier.wait()
        try:
            return service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True)
        except revisions.RevisionConflict:
            return None

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: apply(), range(2)))
    assert sum(result is not None for result in results) == 1
    with scope() as session:
        assert session.query(MonthlyMenuItem).count() == 12
        assert session.get(MenuEditRevision, "monthly-menu:2026-10").revision == 2


def test_existing_month_missing_formal_authority_stops_without_repair(store, tmp_path, monkeypatch):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    with scope() as session:
        session.add(MonthlyMenu(id="2026-10", month_start=date(2026, 10, 1)))
    with pytest.raises(revisions.RevisionSchemaNotMigrated):
        save(tmp_path)
    with scope() as session:
        assert session.query(MonthlyMenuImport).count() == 0
        assert session.get(MenuEditRevision, "monthly-menu:2026-10") is None


@pytest.mark.parametrize("flushed", [False, True])
def test_snapshot_rejects_caller_transaction_without_discarding_edits(store, flushed):
    scope, engine = store
    with Session(engine) as session:
        row = MonthlyMenu(id="2026-10")
        session.add(row)
        if flushed:
            session.flush()
        with pytest.raises(revisions.RevisionConflict, match="owned transaction"):
            revisions.begin_monthly_snapshot(session, write=True)
        assert session.in_transaction()
        assert row in session
        with engine.connect() as connection:
            assert connection.execute(sa.select(MonthlyMenu.id)).all() == []
        session.rollback()


def test_real_workbook_parser_keeps_four_plus_per_meal_and_final_week():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "掲示用（料理・食品）"
    sheet.append(["2026年5月"])
    sheet.append(["", "日", "月", "火", "水", "木", "金", "土"])
    sheet.append(["", 31, "", "", "", "", "", ""])
    for meal in ("朝食", "昼食", "夕食"):
        for slot, name in enumerate(("魚の塩焼き", "南瓜の煮物", "胡瓜の和え物", "豆サラダ", "さつま芋の煮物")):
            sheet.append([meal if slot == 0 else "", name])
    content = BytesIO()
    workbook.save(content)
    _, _, _, entries = service._parse_monthly_menu(content.getvalue(), "menu.xlsx", None, "2026-05")
    for meal in ("朝食", "昼食", "夕食"):
        assert [entry["name"] for entry in entries if entry["menu_date"] == date(2026, 5, 31) and entry["daypart"] == meal] == [
            "魚の塩焼き", "南瓜の煮物", "胡瓜の和え物", "豆サラダ", "さつま芋の煮物"]


@pytest.mark.parametrize("diet,label", [(None, None), ("regular", "常食"), ("soft", "軟菜"), ("mixer", "ミキサー")])
def test_real_workbook_save_review_apply_passes_nested_producer_contract(store, tmp_path, diet, label):
    scope, _ = store
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "保存原本"
    sheet.append(["2026年10月"])
    if label is not None:
        sheet.append([f"献立種類 {label}"])
    sheet.append(["", "日", "月", "火", "水", "木", "金", "土"])
    sheet.append(["", "", "", "", "", "", "", 31])
    names = ["魚の塩焼き", "南瓜の煮物", "胡瓜の和え物", "豆サラダ", "さつま芋の煮物"]
    for meal in ("朝食", "昼食", "夕食"):
        for slot, name in enumerate(names):
            sheet.append([meal if slot == 0 else "", "", "", "", "", "", "", name])
    stream = BytesIO()
    workbook.save(stream)
    content, original = archived(tmp_path, stream.getvalue())
    selected = None if diet is None else [diet]
    receipt = service.save_menu_import("2026-10", content, original, sheet_name=sheet.title,
        scope_override=None, diet_types=selected)
    assert receipt["source"]["diet_types"] == selected and receipt["source"]["scope_override"] is None
    assert receipt["source"]["content_sha256"] == sha256(content).hexdigest()
    assert receipt["review"]["item_count"] == len(receipt["review"]["items"]) == 15
    assert len(receipt["review"]["entries"]) == 15
    assert all(entry["diet_type"] == diet for entry in receipt["review"]["entries"])
    with scope() as session:
        service._validate_import_receipt(session.get(MonthlyMenuImport, receipt["id"]))
        assert session.query(MonthlyMenuItem).count() == 0
    choices = [{"source_name": name, "action": "create", "unit_type": "count", "qty_per_serving": 0,
                "bag_max_qty": 20, "bag_max_unit": "count"} for name in names]
    own = service.review_menu_import("2026-10", receipt["id"], receipt["revision"], choices,
        input_expected=receipt["current_input_revision"])
    assert own["source"] == receipt["source"] and own["review"]["issues"] == []
    assert len(own["review"]["plans"]) == 5 and own["review"]["resolutions"] == choices
    with scope() as session:
        service._validate_import_receipt(session.get(MonthlyMenuImport, receipt["id"]))
    applied = service.review_menu_import("2026-10", receipt["id"], own["revision"], None, apply=True)
    assert applied["created"] is True and applied["item_count"] == 15
    assert applied["receipt"]["state"] == "applied" and applied["receipt"]["source"] == receipt["source"]
    for meal in ("朝食", "昼食", "夕食"):
        entries = sorted((entry for entry in applied["snapshot"]["entries"] if entry["daypart"] == meal), key=lambda entry: entry["slot_index"])
        assert [entry["name"] for entry in entries] == names
        assert [entry["slot_index"] for entry in entries] == list(range(5))
        assert all(entry["menu_date"] == "2026-10-31" and entry["diet_type"] == diet for entry in entries)
    assert all(item["qty_per_serving"] == 0 and item["unit_type"] == "count" for item in applied["snapshot"]["items"])
    with scope() as session:
        service._validate_import_receipt(session.get(MonthlyMenuImport, receipt["id"]))
        assert session.query(MenuMaster).count() == 5 and session.query(MonthlyMenuItem).count() == 15


def test_formal_migration_keeps_existing_month_and_creates_only_its_authority(tmp_path):
    engine = sa.create_engine(sa.URL.create("sqlite", database=str(tmp_path / "migration.sqlite3")))
    metadata = sa.MetaData()
    for model in (MonthlyMenu, MenuRule, BaseMenuCycleItem):
        model.__table__.to_metadata(metadata)
    metadata.create_all(engine)
    with engine.begin() as connection:
        connection.execute(sa.insert(MonthlyMenu).values(id="2026-10", filename="original.xlsx", month_start=date(2026, 10, 1)))
        migration(connection, "0030")
        migration(connection, "0031")
        assert connection.execute(sa.select(MonthlyMenu.id, MonthlyMenu.filename)).all() == [("2026-10", "original.xlsx")]
        rows = connection.execute(sa.select(MenuEditRevision.scope, MenuEditRevision.revision)).all()
        assert sorted(row for row in rows if row.scope.startswith("monthly-menu:")) == [("monthly-menu:2026-10", 1)]
        assert connection.execute(sa.select(MonthlyMenuImport.id)).all() == []
        with pytest.raises(RuntimeError, match="unknown ownership"):
            migration(connection, "0031")
    engine.dispose()


@pytest.mark.parametrize("input_expected", [None, "old-input", 0, {}])
def test_review_requires_exact_current_read_input_without_saving_partial_review(store, tmp_path, monkeypatch, input_expected):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    exception = revisions.RevisionPreconditionRequired if input_expected is None else revisions.RevisionConflict
    with pytest.raises(exception):
        service.review_menu_import("2026-10", receipt["id"], receipt["revision"], [], input_expected=input_expected)
    with scope() as session:
        row = session.get(MonthlyMenuImport, receipt["id"])
        assert row.revision == 1
        assert row.state == "review"
        assert row.review_json["resolutions"] == []
        assert session.query(MonthlyMenuItem).count() == 0


def test_exact_saved_source_download_survives_closed_owner_and_rejects_changed_bytes(store, tmp_path, monkeypatch):
    parsed(monkeypatch)
    masters(store[0])
    receipt = save(tmp_path)
    own = service.get_menu_import_download("2026-10", receipt["id"])
    assert sha256(own["bytes"]).hexdigest() == receipt["source"]["content_sha256"]
    assert own["filename"] == "source.xlsx"
    (tmp_path / "source.xlsx").write_bytes(b"changed original")
    with pytest.raises(ValueError, match="digest mismatch"):
        service.get_menu_import_download("2026-10", receipt["id"])


@pytest.mark.parametrize("field", ["sheet_name", "content_sha256", "file_uri"])
def test_missing_saved_source_identity_never_uses_current_or_first_source(store, tmp_path, monkeypatch, field):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    with scope() as session:
        row = session.get(MonthlyMenuImport, receipt["id"])
        row.source_json = {**row.source_json, field: None}
    with pytest.raises(ValueError, match="unresolved"):
        service.read_menu_import("2026-10", receipt["id"])
    with scope() as session:
        assert session.query(MonthlyMenuItem).count() == 0
        assert session.get(MenuEditRevision, "monthly-menu:2026-10").revision == 1


@pytest.mark.parametrize("autoflush", [False, True])
@pytest.mark.parametrize("pending", ["new", "dirty", "deleted"])
def test_snapshot_pending_guard_runs_before_sql_and_preserves_operator_state(store, autoflush, pending):
    scope, engine = store
    with scope() as session:
        session.add(MonthlyMenu(id="2026-10", filename="saved.xlsx"))
    with Session(engine, autoflush=autoflush) as session:
        row = session.get(MonthlyMenu, "2026-10")
        if pending == "new":
            edit = MonthlyMenu(id="2026-11", filename="operator.xlsx")
            session.add(edit)
        elif pending == "dirty":
            edit = row
            row.filename = "operator.xlsx"
        else:
            edit = row
            session.delete(row)
        statements = []
        def record(_connection, _cursor, statement, _parameters, _context, _many):
            statements.append(statement)
        sa.event.listen(engine, "before_cursor_execute", record)
        try:
            with pytest.raises(revisions.RevisionConflict):
                revisions.begin_monthly_snapshot(session, write=True)
            assert statements == []
        finally:
            sa.event.remove(engine, "before_cursor_execute", record)
        assert edit in getattr(session, {"new": "new", "dirty": "dirty", "deleted": "deleted"}[pending])
        assert session.in_transaction()
        with engine.connect() as connection:
            assert connection.execute(sa.select(MonthlyMenu.id, MonthlyMenu.filename)).all() == [("2026-10", "saved.xlsx")]
        session.rollback()


@pytest.mark.parametrize("field", ["file_uri", "filename", "content_sha256", "sheet_name", "scope_override", "diet_types", "media_type"])
@pytest.mark.parametrize("operation", ["read", "list", "review", "apply", "download"])
def test_missing_persisted_source_keys_stop_every_caller_before_archive_or_business_write(store, tmp_path, monkeypatch, field, operation):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    with scope() as session:
        row = session.get(MonthlyMenuImport, receipt["id"])
        row.source_json = {key: value for key, value in row.source_json.items() if key != field}
        version = row.revision + 1
    def unexpected(*args, **kwargs):
        pytest.fail("invalid persisted source must stop before loading or applying any archive")
    monkeypatch.setattr(service, "load_bytes_from_uri", unexpected)
    calls = {
        "read": lambda: service.read_menu_import("2026-10", receipt["id"]),
        "list": lambda: service.list_menu_imports("2026-10"),
        "review": lambda: service.review_menu_import("2026-10", receipt["id"], receipt["revision"], [], input_expected=receipt["input_revision"]),
        "apply": lambda: service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True),
        "download": lambda: service.get_menu_import_download("2026-10", receipt["id"]),
    }
    with pytest.raises(ValueError, match="unresolved"):
        calls[operation]()
    with scope() as session:
        row = session.get(MonthlyMenuImport, receipt["id"])
        assert field not in row.source_json
        assert row.revision == version and row.state == "review"
        assert session.query(MonthlyMenuItem).count() == 0
        assert session.get(MenuEditRevision, "monthly-menu:2026-10").revision == 1


@pytest.mark.parametrize("field,value", [("scope_override", ""), ("scope_override", 0), ("scope_override", " TAG: "),
    ("diet_types", []), ("diet_types", "regular"), ("diet_types", [None]), ("diet_types", [0]),
    ("file_uri", 0), ("filename", False), ("sheet_name", None), ("media_type", None), ("content_sha256", 0)])
def test_malformed_saved_source_values_are_not_coerced(store, tmp_path, monkeypatch, field, value):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    with scope() as session:
        row = session.get(MonthlyMenuImport, receipt["id"])
        row.source_json = {**row.source_json, field: value}
    with pytest.raises(ValueError, match="unresolved"):
        service.read_menu_import("2026-10", receipt["id"])


@pytest.mark.parametrize("field", ["issues", "plans", "item_count", "master_diffs", "entries", "items", "scope_override", "resolutions", "input_snapshot"])
def test_missing_review_keys_stop_even_apply_before_input_hash_or_archive(store, tmp_path, monkeypatch, field):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    with scope() as session:
        row = session.get(MonthlyMenuImport, receipt["id"])
        row.review_json = {key: value for key, value in row.review_json.items() if key != field}
    def unexpected(*args, **kwargs):
        pytest.fail("invalid review must stop before constructing a current input token")
    monkeypatch.setattr(service, "_monthly_input_digest", unexpected)
    with pytest.raises(ValueError, match="unresolved"):
        service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True)


def test_explicit_null_source_semantics_zero_and_unknown_ownership_are_retained(store, tmp_path, monkeypatch):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope, quantity=0)
    receipt = save(tmp_path)
    with scope() as session:
        row = session.get(MonthlyMenuImport, receipt["id"])
        row.source_json = {**row.source_json, "operator_extension": {"owner": "recorded-owner", "value": None}}
        row.review_json = {**row.review_json, "review_extension": {"zero": 0},
            "input_snapshot": {**row.review_json["input_snapshot"], "schema_owner": "recorded-owner"}}
    read = service.read_menu_import("2026-10", receipt["id"])
    assert read["source"]["scope_override"] is None and read["source"]["diet_types"] is None
    own = service.review_menu_import("2026-10", receipt["id"], read["revision"], [], input_expected=read["current_input_revision"])
    assert own["source"]["operator_extension"] == {"owner": "recorded-owner", "value": None}
    assert own["review"]["review_extension"] == {"zero": 0}
    assert own["review"]["input_snapshot"]["schema_owner"] == "recorded-owner"
    applied = service.review_menu_import("2026-10", receipt["id"], own["revision"], None, apply=True)
    assert all(item["qty_per_serving"] == 0 for item in applied["snapshot"]["items"])
    assert applied["receipt"]["source"] == own["source"]


@pytest.mark.parametrize("patch", [{"resolutions": None}, {"items": [None]}, {"plans": []},
    {"item_count": None}, {"item_count": True}, {"scope_override": "unrelated-facility"},
    {"input_snapshot": None}, {"input_snapshot": {"revision": 1}},
    {"input_snapshot": {"revision": 0, "digest": "a" * 64}},
    {"input_snapshot": {"revision": True, "digest": "a" * 64}},
    {"input_snapshot": {"revision": 1, "digest": None}}])
def test_malformed_recorded_review_is_not_reconstructed_from_current_input(store, tmp_path, monkeypatch, patch):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    with scope() as session:
        row = session.get(MonthlyMenuImport, receipt["id"])
        row.review_json = {**row.review_json, **patch}
    with pytest.raises(ValueError, match="unresolved"):
        service.read_menu_import("2026-10", receipt["id"])
    with scope() as session:
        assert session.query(MonthlyMenuItem).count() == 0


@pytest.mark.parametrize("operation", ["read", "list", "review", "apply", "download"])
@pytest.mark.parametrize("patch", [{"items": [{}]}, {"items": [{"name": "Fish"}]}, {"entries": [{}]}, {"plans": {}}, {"plans": {"Fish": {}}},
    {"plans": {"Fish": {"action": "existing", "seed_fields": {}}}},
    {"plans": {"Fish": {"action": "unknown", "seed_fields": {}, "name": "Fish"}}},
    {"resolutions": [{}]}, {"resolutions": [{"source_name": "Fish", "action": "unknown"}]},
    {"resolutions": [{"source_name": "Fish", "action": "create", "unit_type": "count"}]},
    {"issues": [{}]}, {"master_diffs": [{}]},
    {"entries": [{"menu_date": "2026-10-01", "daypart": "朝食", "name": "Fish", "category": None, "diet_type": None, "slot_index": True}]}])
def test_corrupt_nested_saved_review_stops_all_read_and_command_paths(store, tmp_path, monkeypatch, patch, operation):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    receipt = save(tmp_path)
    with scope() as session:
        row = session.get(MonthlyMenuImport, receipt["id"])
        row.review_json = {**row.review_json, **patch}
    monkeypatch.setattr(service, "load_bytes_from_uri", lambda *_: pytest.fail("corrupt receipt cannot be repaired from source"))
    calls = {
        "read": lambda: service.read_menu_import("2026-10", receipt["id"]),
        "list": lambda: service.list_menu_imports("2026-10"),
        "review": lambda: service.review_menu_import("2026-10", receipt["id"], receipt["revision"], [], input_expected=receipt["input_revision"]),
        "apply": lambda: service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None, apply=True),
        "download": lambda: service.get_menu_import_download("2026-10", receipt["id"]),
    }
    with pytest.raises(ValueError, match="unresolved"):
        calls[operation]()
    with scope() as session:
        assert session.query(MonthlyMenuItem).count() == 0
        assert session.get(MenuEditRevision, "monthly-menu:2026-10").revision == 1


@pytest.mark.parametrize("second", ["Fish 0", "Fish0"])
@pytest.mark.parametrize("mode", ["preview", "review", "apply"])
def test_duplicate_resolution_aliases_never_select_last_master(store, tmp_path, monkeypatch, second, mode):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope)
    choices = [{"source_name": "Fish 0", "action": "existing", "menu_master_id": "fish-0"},
               {"source_name": second, "action": "existing", "menu_master_id": "fish-1"}]
    receipt = save(tmp_path)
    if mode == "apply":
        with scope() as session:
            row = session.get(MonthlyMenuImport, receipt["id"])
            row.review_json = {**row.review_json, "resolutions": choices}
    with pytest.raises(ValueError, match="ambiguous"):
        if mode == "preview":
            content, source = archived(tmp_path)
            service.create_menu("2026-10", content, source.original_filename, "source", preview=True, menu_master_resolutions=choices)
        else:
            service.review_menu_import("2026-10", receipt["id"], receipt["revision"], None if mode == "apply" else choices,
                apply=mode == "apply", input_expected=receipt["input_revision"])
    with scope() as session:
        assert session.query(MonthlyMenuItem).count() == 0
        assert session.get(MenuMaster, "fish-0").qty_per_serving == 1


def test_nested_authored_null_zero_and_unknown_fields_are_preserved(store, tmp_path, monkeypatch):
    scope, _ = store
    parsed(monkeypatch)
    masters(scope, quantity=0)
    receipt = save(tmp_path)
    with scope() as session:
        row = session.get(MonthlyMenuImport, receipt["id"])
        review = {**row.review_json}
        review["items"] = [{**item, "qty_per_serving": 0, "category": None, "extension": {"owner": "original"}} for item in review["items"]]
        review["entries"] = [{**entry, "diet_type": None, "category": None, "extension": 0} for entry in review["entries"]]
        row.review_json = review
    read = service.read_menu_import("2026-10", receipt["id"])
    assert all(item["qty_per_serving"] == 0 and item["extension"] == {"owner": "original"} for item in read["review"]["items"])
    assert all(entry["diet_type"] is None and entry["extension"] == 0 for entry in read["review"]["entries"])


def test_resolution_unique_aliases_preserve_one_explicit_choice_and_duplicate_identical_rows_stop():
    choice = {"source_name": "Fish 0", "issue_key": "fish0", "normalized_name": "fish0", "action": "existing", "menu_master_id": "fish-0"}
    indexed = service._index_menu_master_resolutions([choice])
    assert indexed["Fish 0"] is choice and indexed["fish0"] is choice
    with pytest.raises(ValueError, match="ambiguous"):
        service._index_menu_master_resolutions([choice, dict(choice)])
    with pytest.raises(ValueError, match="ambiguous"):
        service._index_menu_master_resolutions([choice, {"issue_key": "Fish 0", "source_name": "Other", "action": "existing", "menu_master_id": "fish-1"}])
    authored = {"source_name": "Fish 0", "action": " CREATE ", "unit_type": "count", "qty_per_serving": "0", "extension": None}
    assert service._index_menu_master_resolutions([authored])["fish0"] is authored
    assert authored["qty_per_serving"] == "0" and authored["extension"] is None
