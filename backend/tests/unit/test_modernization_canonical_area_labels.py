"""Captured-config label component definitions; no native/PG/live proof.

Real owned fixture provides dated input and 7/0 lines. Area cases explicitly
author copied component inputs, not purported operator/OCR observations.
Real bag producer, zero-label producer and XLSX/CSV writers are exercised.
"""

from copy import deepcopy
from datetime import date
from io import BytesIO

from openpyxl import load_workbook
import pytest

from test_modernization_real_bagging_confirmation import owned_bagging
from test_modernization_real_daily_delivery_html import confirm
from test_modernization_workflow_surface_parity import database_snapshot, surface
from src.models.document import OrderDocument
from src.models.order import Order
from src.models.order_sheet_draft import OrderSheetDraft
from src.services import order_service as orders
from src.services.canonical_facility_store import CanonicalStoreError
from src.services import output_builder as output


DAY = date(2026, 3, 1)


@pytest.fixture
def captured(owned_bagging, surface):
    fixture = owned_bagging
    confirm(fixture)
    with fixture.session_scope() as session:
        adoption = output._capture_output_adoption(fixture.order_id, session)
    order = deepcopy(adoption["order"])
    candidate = adoption["candidate"]
    configured = deepcopy(candidate["configured_input"])
    # Explicit synthetic area names differ from both old hardcoded policies.
    configured["facility"]["areas"] = [{"id": "2F", "name": "正本東棟"}, {"id": "3F", "name": "正本西棟"}]
    lines = output.build_order_lines_for_outputs(order, accepted_candidate=candidate)
    assert [line["quantity_original"] for line in lines] == [7, 0]
    lines = deepcopy(lines)
    for line in lines:
        line["area_id"] = "2F"
    return fixture, order, configured, lines


def produce(order, configured, lines):
    bags = output.build_bag_rows_for_materialization(order, configured_input=configured, order_lines=lines)
    context = {"order_for_outputs": order, "order_lines": lines, "facility_config": configured["facility"]}
    return output._append_zero_quantity_label_bags_for_order_context(bags, context, DAY)


def render(bags, facility):
    group = {"bags": bags, "label_profile": {}, "facility_config": facility,
             "facility_code": facility["facility_id"], "facility_name": facility["facility_name"]}
    xlsx, _ = output._create_daily_label_workbook_bytes(target_date=DAY, group=group)
    csv, _ = output._create_daily_label_csv_bytes(target_date=DAY, group=group)
    workbook = load_workbook(BytesIO(xlsx))
    try:
        cells = [cell.value for sheet in workbook for row in sheet for cell in row if cell.value is not None]
    finally:
        workbook.close()
    return cells, csv.decode("cp932")


@pytest.mark.parametrize("facility_id", ["F", "FAC00008", "FAC00009", "FAC00010"])
def test_canonical_name_reaches_real_positive_zero_xlsx_and_csv_not_fac_branch(captured, facility_id):
    fixture, order, configured, lines = captured
    order["facility"] = facility_id
    configured["facility"]["facility_id"] = facility_id
    before = database_snapshot()
    bags = produce(order, configured, lines)
    assert sorted(bag["quantity"] for bag in bags) == [0, 7]
    assert {bag["area_name"] for bag in bags} == {"正本東棟"}
    assert {bag["area_id"] for bag in bags} == {"2F"}
    cells, csv = render(bags, configured["facility"])
    assert "朝正本東棟" in cells and "朝正本東棟" in csv
    assert "選択献立0" in cells and "選択献立3" in cells
    assert "朝花" not in csv and "朝２階" not in csv
    assert database_snapshot() == before


@pytest.fixture
def dated_area_order(owned_bagging, surface, monkeypatch):
    """Publish actual synthetic March input before capture/confirmation.

    Only owned authored JSON and manually seeded draft field identities change;
    PDF identity, source cells, dates and recognized 7/0 values stay unchanged.
    """
    fixture = owned_bagging
    fixture.authored["areas"] = [{"id": "2F", "name": "正本東棟"}, {"id": "3F", "name": "正本西棟"}]
    template = fixture.authored["fax_template"]
    for column in template["columns"]:
        if column.get("role") == "quantity":
            column["area_id"] = "2F"
    template["body_merge_policy"]["columns"] = ["qty.regular_2f"]
    with fixture.session_scope() as session:
        draft = session.get(OrderSheetDraft, fixture.draft_id)
        payload = deepcopy(draft.draft_sheet_json)
        payload["fields"] = [field[:-2] + "_2f" if field.startswith("qty.") and field.endswith("_x") else field
                             for field in payload["fields"]]
        assert [row[4] for row in payload["rows"]] == ["7", "", "", "0"]
        draft.draft_sheet_json = payload
    fixture.republish()
    with fixture.session_scope() as session:
        orders._record_current_order_version(session, session.get(Order, fixture.order_id),
                                             session.get(OrderDocument, "UNIT-DOC"))
    confirm(fixture)
    monkeypatch.setattr(output, "OUTPUT_DIR", fixture.output)
    with fixture.session_scope() as session:
        captured_input = output._capture_output_adoption(fixture.order_id, session)
    candidate = captured_input["candidate"]
    assert candidate["configured_input"]["facility"]["areas"] == fixture.authored["areas"]
    assert [line["area_id"] for line in candidate["output_lines"]] == ["2F", "2F"]
    assert [line["quantity_original"] for line in candidate["output_lines"]] == [7, 0]
    return fixture, captured_input


def test_captured_version_name_is_not_replaced_by_newer_unselected_facility(dated_area_order):
    fixture, captured_input = dated_area_order
    march = deepcopy(fixture.master["facilities"][0]["versions"][0])
    fixture.master["facilities"][0]["versions"].append({
        "version_id": "synthetic-April-unselected", "valid_from": "2026-04-01", "valid_to": "2026-04-30",
        "facility": {**deepcopy(fixture.authored), "areas": [{"id": "2F", "name": "最新版別名"}]},
    })
    fixture.write_master()
    before = database_snapshot()
    # No caller-supplied config/candidate: both public producers capture the order.
    xlsx, _ = output.build_label_workbook_for_order(fixture.order_id, target_date=DAY)
    csv_path, _ = output.build_label_csv_for_order(fixture.order_id, target_date=DAY)
    workbook = load_workbook(xlsx)
    try:
        cells = [cell.value for sheet in workbook for row in sheet for cell in row if cell.value is not None]
    finally:
        workbook.close()
    csv = csv_path.read_bytes().decode("cp932")
    assert "朝正本東棟" in cells and "最新版別名" not in csv
    assert "朝正本東棟" in csv and "選択献立0" in csv and "選択献立3" in csv
    with fixture.session_scope() as session:
        actual = output._capture_output_adoption(fixture.order_id, session)
    assert actual["candidate"]["configured_input"] == captured_input["candidate"]["configured_input"]
    bags = output.build_bag_rows_for_outputs(actual["order"], accepted_candidate=actual["candidate"])
    context = {"order_for_outputs": actual["order"], "order_lines": actual["candidate"]["output_lines"],
               "facility_config": actual["candidate"]["configured_input"]["facility"]}
    bags = output._append_zero_quantity_label_bags_for_order_context(bags, context, DAY)
    assert sorted(bag["quantity"] for bag in bags) == [0, 7]
    assert fixture.master["facilities"][0]["versions"][0] == march
    assert database_snapshot() == before


@pytest.mark.parametrize("producer", ["build_label_workbook_for_order", "build_label_csv_for_order"])
def test_changed_order_applicable_canonical_version_blocks_not_latest_fallback(dated_area_order, producer):
    fixture, captured_input = dated_area_order
    fixture.authored["areas"][0]["name"] = "適用版改変名"
    fixture.write_master()
    before = database_snapshot()
    files = {path: path.read_bytes() for path in fixture.output.rglob("*") if path.is_file()}
    with pytest.raises((ValueError, CanonicalStoreError)):
        getattr(output, producer)(fixture.order_id, target_date=DAY)
    assert database_snapshot() == before
    assert {path: path.read_bytes() for path in fixture.output.rglob("*") if path.is_file()} == files
    assert captured_input["candidate"]["configured_input"]["facility"]["areas"][0]["name"] == "正本東棟"


@pytest.mark.parametrize("fault", ["duplicate", "missing-name", "unregistered", "alias-collision"])
def test_invalid_canonical_area_stops_before_render_and_preserves_db(captured, fault):
    fixture, order, configured, lines = captured
    if fault == "duplicate":
        configured["facility"]["areas"].append({"id": "2F", "name": "別名"})
    elif fault == "missing-name":
        configured["facility"]["areas"][0].pop("name")
    elif fault == "unregistered":
        lines[0]["area_id"] = "UNREGISTERED"
    else:
        configured["facility"]["areas"][1]["name"] = "2F"
    before = database_snapshot()
    with pytest.raises(ValueError):
        produce(order, configured, lines)
    assert database_snapshot() == before


@pytest.mark.parametrize("area", [None, "", "X"])
def test_no_area_retains_empty_suffix_and_unchanged_quantities(captured, area):
    fixture, order, configured, lines = captured
    configured["facility"]["areas"] = None
    for line in lines:
        line["area_id"] = area
    bags = produce(order, configured, lines)
    assert {bag["area_name"] for bag in bags} == {""}
    assert sorted(bag["quantity"] for bag in bags) == [0, 7]
    cells, csv = render(bags, configured["facility"])
    assert "朝" in cells and "正本東棟" not in csv


def test_legacy_label_details_and_serialized_bags_carry_canonical_name(captured):
    fixture, order, configured, lines = captured
    bags = produce(order, configured, lines)
    serialized = output._serialize_bag_payload_rows(bags)
    assert {row["area_name"] for row in serialized} == {"正本東棟"}
    labels, _, kind = output._build_label_rows(bags, {"label_fields": ["product_name", "details", "quantity"]}, "合成施設")
    assert kind == "legacy"
    assert len(labels) == 2
    assert all("正本東棟" in label["details"] for label in labels)
    assert sorted(label["quantity"] for label in labels) == [0, 7]
