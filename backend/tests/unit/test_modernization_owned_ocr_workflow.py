"""Synthetic native/local service proof; NOT operator FAX, HTTP bagging, PG or live.

No recognizer, geometry, ownership, assignment or projection doubles. Missing
native dependencies/model cache fail normally, never skip. Only session/engine,
local canonical JSON/output and offline environment wiring are substituted.
Receipts, expected targets, execution IDs and evidence manifests are produced by
the real capture/rerun/persistence services, not fixture-authored dictionaries.
"""

from copy import deepcopy
from datetime import date, datetime, timedelta
from hashlib import sha256
import json
import shutil

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
import pytest
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker

from src import db, models  # noqa: F401 -- register real FK targets
from src.models.document import OrderDocument
from src.models.facility import Facility
from src.models.menu import MonthlyMenu, MonthlyMenuEntry
from src.models.ocr_job import OcrJob
from src.models.order import Order
from src.models.order_ocr_evidence_run import OrderOcrEvidenceRun
from src.models.order_sheet_draft import OrderSheetDraft
from src.models.order_workflow_state import OrderWorkflowState
from src.models.uploaded_pdf import UploadedPdf
from src.services import canonical_facility_lineage as lineage
from src.services import config_service, facility_template_version_service
from src.services import hakodate_cell_ocr_batch_service as batch
from src.services import master_order_form_template_service as master
from src.services import menu_service, ocr_job_service, order_service
from src.services import order_workflow_v2_service as workflow
from src.services.storage_service import load_selected_pdf_bytes
from src.services.workbook_pdf_renderer import render_workbook_path_to_pdf
from src.services.workflow_attempt_guard_service import capture_selected_pdf_receipt


START = date(2026, 3, 1)
END = START + timedelta(days=6)
WEEK = "2026-03@2026-03-01~2026-03-07"
FACILITY = "UNIT-NATIVE-F"
ORDER = "UNIT-NATIVE-ORDER"
OTHER = "UNIT-NATIVE-OTHER"
TEMPLATE = "fax_layout_regular_forbidden_v1"
JOB = "UNIT-NATIVE-JOB"


class OwnedProducer:
    def __init__(self, root, monkeypatch):
        self.root = root
        self.engine = create_engine(f"sqlite:///{root / 'owned.sqlite'}")

        @event.listens_for(self.engine, "connect")
        def foreign_keys(connection, _record):
            connection.execute("PRAGMA foreign_keys=ON")

        db.Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine, expire_on_commit=False)
        monkeypatch.setattr(db, "SessionLocal", self.sessions)
        monkeypatch.setattr(db, "engine", self.engine)
        monkeypatch.setattr(menu_service, "engine", self.engine)
        monkeypatch.setattr(facility_template_version_service, "engine", self.engine)
        for name in ("FACILITY_MASTER_URI", "OCR_OUTPUT_BUCKET", "RAW_BUCKET", "K_SERVICE"):
            monkeypatch.delenv(name, raising=False)
        monkeypatch.setenv("HF_HUB_OFFLINE", "1")
        monkeypatch.setenv("TRANSFORMERS_OFFLINE", "1")
        monkeypatch.setenv("HF_DATASETS_OFFLINE", "1")
        monkeypatch.setenv("TMPDIR", str(root))
        self.master_json = root / "authored-facilities.json"
        monkeypatch.setenv("FACILITY_MASTER_PATH", str(self.master_json))
        policy = root / "ingest-policy.json"
        shutil.copyfile(config_service.INGEST_POLICY_PATH, policy)
        monkeypatch.setattr(config_service, "INGEST_POLICY_PATH", policy)
        self.locked_hash = sha256(master.MASTER_TEMPLATE_PATH.read_bytes()).hexdigest()
        self.master_copy = root / "copied-master.xlsx"
        shutil.copyfile(master.MASTER_TEMPLATE_PATH, self.master_copy)
        columns = [
            {"index": 0, "source_index": 0, "role": "date", "header": "日付"},
            {"index": 1, "source_index": 1, "role": "daypart", "header": "区分"},
            {"index": 2, "source_index": 2, "role": "aux", "name": "aux.weekday", "header": "曜日"},
            {"index": 3, "source_index": 3, "role": "menu_name", "header": "献立"},
            {"index": 4, "source_index": 4, "role": "quantity", "header": "常食",
             "diet_type": "regular", "area_id": "X"},
        ]
        self.facility = {
            "facility_id": FACILITY, "facility_name": "Synthetic native unit facility",
            "areas": None, "fax_template_id": TEMPLATE,
            "fax_template": {"template_id": TEMPLATE, "page": 1,
                             "units": "normalized", "table_box": [0, 0, 1, 1],
                             "header_rows": 0, "fill_forward_roles": ["date", "daypart"],
                             "columns": columns},
        }
        # Both the authored root and the explicitly dated snapshot are identical;
        # the actual capture selects/verifies the dated snapshot, never merges it.
        authored = {"version_id": "native-March", "valid_from": "2026-03-01",
                    "valid_to": "2026-03-31", "facility": deepcopy(self.facility)}
        self.master_json.write_text(json.dumps({"facilities": [
            {**deepcopy(self.facility), "versions": [authored]},
        ]}, ensure_ascii=False), encoding="utf-8")
        with db.session_scope() as session:
            session.add(Facility(id=FACILITY, name=self.facility["facility_name"]))
            session.add(MonthlyMenu(id="2026-03", month_start=START))
            session.flush()
            for offset in range(7):
                day = START + timedelta(days=offset)
                for meal, count in (("朝", 2), ("昼", 3), ("夕", 3)):
                    for slot in range(count):
                        session.add(MonthlyMenuEntry(
                            id=f"NATIVE-{offset}-{meal}-{slot}", monthly_menu_id="2026-03",
                            menu_date=day, daypart=meal, name=f"献立{offset}{meal}{slot}",
                            category="主", diet_type="regular", slot_index=slot,
                        ))
            session.flush()
            selected, _ = lineage.resolve(session, FACILITY, START)
            _, projection = lineage.materialize(session, selected)
            self.projection_id = projection.id
            self.lineage = lineage.verify_projection(projection)
        entries = order_service._build_position_menu_entries_safe(WEEK, FACILITY)
        assert len(entries) == 56, "actual monthly SQL must resolve all authored physical rows"
        configured = config_service.capture_authored_processing_config(
            FACILITY, selected_authored_version=authored,
        )
        self.plan = batch.capture_original_target_plan(
            configured_input=configured, physical_menu_entries=entries, week_code=WEEK,
        )
        assert len(self.plan) == 56
        self.expected = {region["sheet_cell"]: "" for region in self.plan}
        self.expected[self.plan[0]["sheet_cell"]] = "7"
        self.expected[self.plan[1]["sheet_cell"]] = "0"
        fax_xlsx = master.build_facility_template_xlsx(
            facility_config=self.facility, master_template_path=self.master_copy,
            output_path=root / "synthetic-fax.xlsx", week_value=WEEK, week_menu_entries=entries,
        )
        workbook = load_workbook(fax_xlsx)
        try:
            worksheet = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
            for cell, value in self.expected.items():
                worksheet[cell] = value if value != "" else None
                worksheet[cell].font = Font(name="DejaVu Sans", size=18, color="000000")
                worksheet[cell].alignment = Alignment(horizontal="center", vertical="center")
            workbook.save(fax_xlsx)
        finally:
            workbook.close()
        self.pdf = render_workbook_path_to_pdf(
            fax_xlsx, output_path=root / "synthetic-fax.pdf",
            sheet_name=master.FACILITY_TEMPLATE_SHEET_NAME,
        )
        with db.session_scope() as session:
            for identifier in (ORDER, OTHER):
                session.add(Order(id=identifier, facility_code=FACILITY, week_code=WEEK,
                                  document_uri=str(self.pdf), message_id=identifier,
                                  received_at=datetime(2026, 3, 1, 9),
                                  template_version_id=self.projection_id))
            session.flush()
            self.add_document(session, "DOC-NATIVE", "PDF-NATIVE", self.pdf)
        initial, error = workflow.get_workflow_readonly(ORDER)
        assert error is None, error
        self.context, error = workflow.confirm_context(
            order_id=ORDER, facility_id=FACILITY, week_start=START.isoformat(),
            week_end=END.isoformat(), template_id=TEMPLATE, expected_revision=initial["revision"],
        )
        assert error is None, (error, self.context)

    def add_document(self, session, document_id, uploaded_id, path):
        session.add(OrderDocument(
            id=document_id, facility_code=FACILITY, week_code=WEEK, storage_uri=str(path),
            source_email_id=ORDER, received_at=datetime(2026, 3, 1, 9), status="done",
        ))
        session.flush()
        session.add(UploadedPdf(
            id=uploaded_id, message_id=ORDER, storage_uri=str(path), original_filename=path.name,
            content_sha256=sha256(path.read_bytes()).hexdigest(), page_count=1,
            received_at=datetime(2026, 3, 1, 9), current_document_id=document_id,
            current_order_id=ORDER,
        ))
        session.flush()
        order = session.get(Order, ORDER)
        order.current_document_id = document_id
        order.document_uri = str(path)
        session.flush()
        order_service._record_current_order_version(session, order, session.get(OrderDocument, document_id))
        self.receipt = capture_selected_pdf_receipt(session, order)

    def produce(self):
        assert load_selected_pdf_bytes(str(self.pdf), self.receipt) == self.pdf.read_bytes()
        _, created = ocr_job_service.create_job(
            JOB, str(self.pdf), order_id=ORDER, uploaded_pdf_id=self.receipt["uploaded_pdf_id"],
            order_document_id=self.receipt["document_id"],
            input_artifact_digest=self.receipt["sha256"], template_version_id=self.projection_id,
        )
        assert created
        evidence, error = order_service.rerun_ocr_evidence_only(
            ORDER, job_id=JOB, project_sheet=False, refresh_workflow=False,
            selected_document_id=self.receipt["document_id"],
        )
        assert error is None, (error, evidence)
        assert evidence["terminal_receipt"]["status"] == "done"
        self.evidence = evidence
        return evidence

    def select(self):
        with db.session_scope() as session:
            revision = session.get(OrderWorkflowState, ORDER).revision
        selected, error = workflow.select_ocr_result(
            ORDER, self.evidence["id"], expected_revision=revision,
        )
        assert error is None, (error, selected)
        return selected


@pytest.fixture
def owned_producer(tmp_path, monkeypatch):
    fixture = OwnedProducer(tmp_path, monkeypatch)
    try:
        yield fixture
    finally:
        fixture.engine.dispose()
        assert sha256(master.MASTER_TEMPLATE_PATH.read_bytes()).hexdigest() == fixture.locked_hash


def test_native_producer_real_lineage_terminal_select_projection_and_save(owned_producer):
    f = owned_producer
    evidence = f.produce()
    payload = evidence["payload_json"]
    execution = payload["_ocr_execution"]
    source = payload["hakodate_repair_source_receipt"]
    assert execution["job_id"] == JOB
    assert execution["evidence_run_id"] == evidence["id"]
    assert execution["nonworkflow_execution_id"]
    assert source["execution_id"] == execution["nonworkflow_execution_id"]
    assert source["selected_pdf_receipt"] == f.receipt
    assert evidence["artifact_manifest_json"]["original_hakodate_source"] == source
    assert execution["processing_start_input"]["original_template_lineage"] == f.lineage
    records = payload["hakodate_ocr_evidence_records"]
    by_cell = {record["engine_metadata"]["sheet_cell"]: record for record in records}
    assert set(by_cell) == set(f.expected)
    for cell, expected in f.expected.items():
        assert by_cell[cell]["raw_text"] == expected, (cell, by_cell[cell])
        assert by_cell[cell]["normalized_value"] == expected
    with db.session_scope() as session:
        stored = session.get(OrderOcrEvidenceRun, evidence["id"])
        job = session.get(OcrJob, JOB)
        assert stored.payload_json == payload
        assert job.status == "done"
        assert job.metrics["rerun_terminal_receipt"] == evidence["terminal_receipt"]
        assert session.get(OrderWorkflowState, ORDER).evidence_run_id is None
    f.select()
    projected, error = workflow.build_sheet_from_selected_ocr(ORDER)
    assert error is None, error
    assert projected["selected_ocr_result_id"] == evidence["id"]
    assert projected["blockers"] == []
    sheet = projected["sheet"]
    qty = sheet["fields"].index("qty.regular_x")
    assert [row[qty] for row in sheet["rows"]] == list(f.expected.values())
    original = deepcopy(sheet)
    with db.session_scope() as session:
        revision = session.get(OrderWorkflowState, ORDER).revision
    # Public service contract, not HTTP pre-save-check/auth proof.
    saved, error = workflow.save_sheet(
        order_id=ORDER, sheet=sheet, edited_by="synthetic-native-unit", expected_revision=revision,
    )
    assert error is None, (error, saved)
    assert sheet == original
    with db.session_scope() as session:
        state = session.get(OrderWorkflowState, ORDER)
        draft = session.get(OrderSheetDraft, state.draft_id)
        assert draft.draft_sheet_json == original
        assert draft.base_evidence_run_id == evidence["id"]
        assert draft.template_version_id == f.projection_id
        assert state.revision == revision + 1
        assert session.get(OrderOcrEvidenceRun, evidence["id"]).payload_json == payload


def test_real_generated_evidence_cannot_be_selected_for_foreign_order(owned_producer):
    f = owned_producer
    evidence = f.produce()
    initial, error = workflow.get_workflow_readonly(OTHER)
    assert error is None, error
    result, error = workflow.select_ocr_result(
        OTHER, evidence["id"], expected_revision=initial["revision"],
    )
    assert result is None
    assert error == "ocr_result_not_found"
    with db.session_scope() as session:
        assert session.get(OrderWorkflowState, OTHER) is None
        assert session.get(OrderOcrEvidenceRun, evidence["id"]).order_id == ORDER
        assert session.execute(select(OrderSheetDraft)).scalars().all() == []


def test_selected_real_artifacts_stop_after_owned_pdf_document_retirement(owned_producer):
    f = owned_producer
    evidence = f.produce()
    f.select()
    before, error = workflow.build_sheet_from_selected_ocr(ORDER)
    assert error is None and before["blockers"] == [], (error, before)
    old_receipt = deepcopy(f.receipt)
    alternate = f.root / "owned-successor.pdf"
    shutil.copyfile(f.pdf, alternate)
    with db.session_scope() as session:
        f.add_document(session, "DOC-SUCCESSOR", "PDF-SUCCESSOR", alternate)
    assert f.receipt != old_receipt
    assert f.receipt["sha256"] == old_receipt["sha256"], "same bytes do not authorize another document"
    result, error = workflow.build_sheet_from_selected_ocr(ORDER)
    assert result is None
    assert error == "hakodate_artifact_source_mismatch"
    with db.session_scope() as session:
        assert session.get(OrderWorkflowState, ORDER).evidence_run_id == evidence["id"]
        assert session.get(OrderOcrEvidenceRun, evidence["id"]).payload_json == evidence["payload_json"]
        assert session.execute(select(OrderSheetDraft)).scalars().all() == []


def _queue_and_claim_owned_workflow_attempt(f, *, force=False):
    from src.services import workflow_ocr_delivery_service as delivery

    with db.session_scope() as session:
        order = session.get(Order, ORDER)
        input_reference, selected = order_service._resolve_order_document_selection(
            session, order, f.receipt["document_id"],
        )
        revision = session.get(OrderWorkflowState, ORDER).revision
    assert input_reference == str(f.pdf)
    assert selected["selected_version_input"]["document_id"] == f.receipt["document_id"]
    queued, error = workflow.mark_ocr_run_queued(
        ORDER, JOB, input_reference=input_reference, request_mode="ocr_rerun",
        job_metrics={"selected_version_input": selected["selected_version_input"]},
        delivery_selected_document_id=f.receipt["document_id"],
        expected_revision=revision, force=force,
    )
    assert error is None, (error, queued)
    assert queued["workflow"]["state"] == "ocr_running"
    claim = delivery.claim_delivery(
        queued["delivery"]["id"], queued["delivery"]["generation"],
        lease_expires_at=datetime.utcnow() + timedelta(minutes=5),
    )
    assert not claim.get("duplicate"), claim
    assert claim["order_id"] == ORDER
    assert claim["job_id"] == JOB
    assert claim["workflow_revision"] == queued["workflow"]["revision"]
    assert claim["selected_pdf_receipt"] == f.receipt
    with db.session_scope() as session:
        job = session.get(OcrJob, JOB)
        assert job.status == "running"
        assert job.metrics["workflow_attempt_id"] == claim["workflow_attempt_id"]
        assert job.metrics["workflow_completion_revision"] == claim["workflow_revision"]
        assert job.metrics["selected_version_input"] == selected["selected_version_input"]
        assert job.metrics["selected_pdf_receipt"] == f.receipt
        assert job.metrics.get("rerun_terminal_receipt") is None
    return claim


def _owned_workflow_database_snapshot():
    # Include deliveries, cache, drafts, evidence and source/version rows, not
    # just the reusable job whose attempt ID changed.
    with db.session_scope() as session:
        return {
            table.name: deepcopy([
                dict(row) for row in session.execute(
                    select(table).order_by(*table.primary_key.columns),
                ).mappings()
            ])
            for table in db.Base.metadata.sorted_tables
        }


def _persist_owned_workflow_failure(claim):
    from src.services import workflow_ocr_delivery_service as delivery

    with delivery.claim_context(claim), order_service.workflow_ocr_write_context(
        claim["workflow_attempt_id"], claim["workflow_revision"], JOB, ORDER,
    ):
        return order_service._persist_ocr_rerun_completion(
            ORDER, JOB, own_payload=None, error="quad_estimation_failed",
        )


def _assert_owned_workflow_failure_receipt(claim, result, error):
    from src.models.workflow_ocr_delivery import WorkflowOcrDelivery

    assert error == "quad_estimation_failed"
    returned = result["terminal_receipt"]
    assert returned["order_id"] == ORDER
    assert returned["job_id"] == JOB
    assert returned["workflow_attempt_id"] == claim["workflow_attempt_id"]
    assert returned["workflow_revision"] == claim["workflow_revision"]
    assert returned["status"] == "failed"
    assert returned["error"] == error
    assert returned["execution_id"] is None
    assert returned["evidence_run_id"] is None
    with db.session_scope() as session:
        state = session.get(OrderWorkflowState, ORDER)
        job = session.get(OcrJob, JOB)
        delivery = session.get(WorkflowOcrDelivery, claim["delivery_id"])
        assert state.state == "ocr_failed"
        assert state.primary_action == "run_ocr"
        assert state.blockers_json == [error]
        assert state.revision == claim["workflow_revision"] + 1
        assert state.evidence_run_id is None
        assert state.draft_id is None
        assert job.status == "failed"
        assert job.error_message == error
        assert job.metrics["workflow_attempt_id"] == claim["workflow_attempt_id"]
        assert job.metrics["workflow_result_revision"] == state.revision
        persisted = job.metrics["rerun_terminal_receipt"]
        assert persisted == delivery.result_receipt
        assert {key: persisted[key] for key in returned} == returned
        assert persisted["workflow_delivery_owner"] == {
            key: claim[key] for key in ("delivery_id", "generation", "claim_token")
        }
        assert session.execute(select(OrderOcrEvidenceRun)).scalars().all() == []
        assert session.execute(select(OrderSheetDraft)).scalars().all() == []


def test_old_failed_workflow_attempt_cannot_overwrite_force_queued_successor(owned_producer):
    """Local failure handoff proof only; no native OCR result is asserted."""
    f = owned_producer
    pdf_before = f.pdf.read_bytes()
    expected_before = deepcopy(f.expected)
    claim_a = _queue_and_claim_owned_workflow_attempt(f)
    claim_b = _queue_and_claim_owned_workflow_attempt(f, force=True)
    assert claim_a["job_id"] == claim_b["job_id"]
    assert claim_a["workflow_attempt_id"] != claim_b["workflow_attempt_id"]
    assert claim_a["delivery_id"] != claim_b["delivery_id"]
    assert claim_b["workflow_revision"] > claim_a["workflow_revision"]
    before = _owned_workflow_database_snapshot()
    with pytest.raises(ValueError, match="^reparse_attempt_superseded$"):
        _persist_owned_workflow_failure(claim_a)
    assert _owned_workflow_database_snapshot() == before

    result, error = _persist_owned_workflow_failure(claim_b)
    _assert_owned_workflow_failure_receipt(claim_b, result, error)
    from src.models.workflow_ocr_delivery import WorkflowOcrDelivery
    with db.session_scope() as session:
        assert session.get(WorkflowOcrDelivery, claim_a["delivery_id"]).result_receipt is None
    assert f.pdf.read_bytes() == pdf_before
    assert f.expected == expected_before


def test_completed_current_workflow_failure_cannot_be_replayed(owned_producer):
    """Replay the real completed claim, without rewriting its CAS revision."""
    claim = _queue_and_claim_owned_workflow_attempt(owned_producer)
    result, error = _persist_owned_workflow_failure(claim)
    _assert_owned_workflow_failure_receipt(claim, result, error)
    before = _owned_workflow_database_snapshot()
    with pytest.raises(ValueError, match="^reparse_attempt_superseded$"):
        _persist_owned_workflow_failure(claim)
    assert _owned_workflow_database_snapshot() == before
