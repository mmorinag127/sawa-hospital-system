"""Owned SQLite/JSON and synthetic source/public export units, not native FAX/stg.

Only the CAS negative case schedules a real revision change at the external XLSX
load boundary; the real loader, writer and both capture guards still execute.
"""

from copy import deepcopy
from datetime import date, datetime, timedelta
import hashlib
from html.parser import HTMLParser
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font
from openpyxl.worksheet.print_settings import PrintArea
from PIL import Image
import pytest
from sqlalchemy import select

from modernization_saved_sheet_export_fixtures import SavedSheetFixture, WEEK
from src.db import Base
from src.models.order_sheet_draft import OrderSheetDraft
from src.models.order_workflow_state import OrderWorkflowState
from src.services import master_order_form_template_service as master
from src.services import order_form_service as forms


@pytest.fixture
def saved(tmp_path, monkeypatch):
    original_master = hashlib.sha256(master.MASTER_TEMPLATE_PATH.read_bytes()).hexdigest()
    fixture = SavedSheetFixture(tmp_path, monkeypatch)
    try:
        yield fixture
    finally:
        fixture.close()
        assert hashlib.sha256(master.MASTER_TEMPLATE_PATH.read_bytes()).hexdigest() == original_master


def database_snapshot(saved):
    with saved.engine.connect() as connection:
        return {table.name: sorted(repr(dict(row)) for row in connection.execute(select(table)).mappings())
                for table in Base.metadata.sorted_tables}


def publish_rows(saved, values):
    payload = deepcopy(saved.payload)
    payload["rows"] = [["3/1", "朝", "主", f"operator-menu-{index}", value, "3", "4", "0"]
                       for index, value in enumerate(values)]
    saved.payload = payload
    with saved.session_scope() as session:
        session.get(OrderSheetDraft, saved.draft_id).draft_sheet_json = deepcopy(payload)


class PreviewRows(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self.row, self.cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        elif tag == "td":
            self.cell = ""

    def handle_data(self, data):
        if self.cell is not None:
            self.cell += data

    def handle_endtag(self, tag):
        if tag == "td" and self.cell is not None:
            self.row.append(self.cell)
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)


def assert_excel_rows(content, values, *, expected_columns=None):
    workbook = load_workbook(BytesIO(content))
    try:
        ws = workbook[WEEK]
        assert isinstance(ws["A11"].value, datetime)
        assert ws["A11"].value.date() == date(2026, 3, 1)
        assert ws["J5"].value == "authored-target"
        assert [ws.cell(11 + index, 4).value for index in range(len(values))] == [
            f"operator-menu-{index}" for index in range(len(values))]
        assert [ws.cell(11 + index, 5).value for index in range(len(values))] == values
        if expected_columns is not None:
            assert [ws.cell(11, column).value for column in (5, 6, 7, 8)] == expected_columns
        metadata = {row[0]: row[1] for row in workbook["設定"].iter_rows(values_only=True) if row[0] is not None}
        assert metadata["saved_sheet_draft_id"] == SavedSheetFixture.draft_id
        assert metadata["order_id"] == SavedSheetFixture.order_id
        assert metadata["saved_sheet_row_count"] == len(values)
        assert metadata["physical_row_plan"]
        assert ws.cell(11 + len(values), 4).value is None
    finally:
        workbook.close()


def test_quantity_only_json_blank_source_public_excel_and_preview_preserve_zero_blank_unknown_and_selection(saved):
    publish_rows(saved, [0, "0", "", "?", 7])
    source_before = saved.source.read_bytes()
    before = database_snapshot(saved)
    response = saved.request("excel")
    assert response.status_code == 200, response.text if response.status_code != 200 else ""
    assert_excel_rows(response.content, [0, 0, None, "?", 7], expected_columns=[0, 3, 4, 0])
    response = saved.request("preview")
    assert response.status_code == 200, response.text
    parser = PreviewRows()
    parser.feed(response.text)
    for index, expected in enumerate(["0", "0", "", "?", "7"]):
        row = next(row for row in parser.rows if f"operator-menu-{index}" in row)
        menu_index = row.index(f"operator-menu-{index}")
        assert row[menu_index + 1:menu_index + 5] == [expected, "3", "4", "0"]
    assert "999" not in response.text and "777" not in response.text
    assert database_snapshot(saved) == before
    assert saved.source.read_bytes() == source_before
    saved.verify_current()


def test_public_saved_sheet_above_57_keeps_every_row_and_raw_source_index(saved):
    import json

    values = [7, 0, "", "?"] * 15
    publish_rows(saved, values)
    before = database_snapshot(saved)
    response = saved.request("excel")
    assert response.status_code == 200, response.text if response.status_code != 200 else ""
    assert_excel_rows(response.content, [7, 0, None, "?"] * 15)
    workbook = load_workbook(BytesIO(response.content))
    try:
        ws = workbook[WEEK]
        assert ws["D70"].value == "operator-menu-59"
        assert all(ws.cell(70, column).border.bottom.style == "thick" for column in range(1, 11))
        assert [region.bounds for region in PrintArea.from_string(ws.print_area).ranges] == [(1, 1, 12, 72)]
        metadata = {row[0]: row[1] for row in workbook["設定"].iter_rows(values_only=True) if row[0] is not None}
        assert [row["source_index"] for row in json.loads(metadata["physical_row_plan"])] == list(range(60))
    finally:
        workbook.close()
    assert database_snapshot(saved) == before


@pytest.mark.parametrize("swap", [False, True], ids=["authored-areas", "authored-area-destinations-swapped"])
def test_same_diet_distinct_areas_use_exact_canonical_source_indexes(saved, swap):
    publish_rows(saved, [7])
    columns = saved.authored["fax_template"]["columns"]
    for column, area, source in zip(columns[:2], ("2F", "3F"), ((5, 4) if swap else (4, 5))):
        column.update(diet_type="regular", area_id=area, source_index=source)
    saved.payload["fields"][4:6] = ["qty.regular_2f", "qty.regular_3f"]
    saved.payload["rows"][0][5] = 0
    with saved.session_scope() as session:
        session.get(OrderSheetDraft, saved.draft_id).draft_sheet_json = deepcopy(saved.payload)
    workbook = load_workbook(saved.source)
    workbook[WEEK]["F7"] = "常食"
    workbook.save(saved.source)
    workbook.close()
    saved.republish()
    before = database_snapshot(saved)
    response = saved.request("excel")
    assert response.status_code == 200, response.text if response.status_code != 200 else ""
    expected = [0, 7, 4, 0] if swap else [7, 0, 4, 0]
    assert_excel_rows(response.content, [expected[0]], expected_columns=expected)
    assert database_snapshot(saved) == before


@pytest.mark.parametrize("case,detail", [
    ("wrong-index", "source workbook quantity header mismatch:qty.regular_x"),
    ("unknown-field", "saved sheet canonical field unmapped"),
    ("no-fried-alias", "saved sheet canonical field unmapped"),
])
def test_public_source_mapping_missing_duplicate_unknown_or_other_diet_never_falls_back(saved, case, detail):
    publish_rows(saved, [7])
    columns = saved.authored["fax_template"]["columns"]
    if case == "wrong-index":
        columns[0]["source_index"] = 8
    elif case == "unknown-field":
        saved.payload["fields"].append("qty.unowned_x")
        saved.payload["rows"][0].append("99")
    else:
        columns[1]["diet_type"] = "forbidden_other"
        saved.payload["fields"][5] = "qty.no_fried_x"
        workbook = load_workbook(saved.source)
        workbook[WEEK]["F7"] = "肉卵魚禁"
        workbook.save(saved.source)
        workbook.close()
    with saved.session_scope() as session:
        session.get(OrderSheetDraft, saved.draft_id).draft_sheet_json = deepcopy(saved.payload)
    saved.republish()
    before = database_snapshot(saved)
    for boundary in ("excel", "preview"):
        response = saved.request(boundary)
        assert response.status_code == 400, response.text
        assert response.json()["detail"] == detail
    assert database_snapshot(saved) == before
    assert not list(saved.output.rglob("*.xlsx"))


@pytest.mark.parametrize("case,validation_error", [
    ("duplicate-index", "template_source_index_duplicate"),
    ("missing-index", "template_source_index_missing"),
])
def test_authored_column_tamper_is_rejected_without_republishing_projection(saved, case, validation_error):
    from src.services.canonical_facility_store import CanonicalStoreError, validate_authored_snapshot

    before = database_snapshot(saved)
    columns = saved.authored["fax_template"]["columns"]
    if case == "duplicate-index":
        columns[1]["source_index"] = 4
    else:
        del columns[0]["source_index"]
    # Tamper only the authored JSON; retain the real previously valid projection.
    saved.write_master()
    with pytest.raises(CanonicalStoreError) as raised:
        validate_authored_snapshot(saved.authored, "authored-March")
    assert str(raised.value) == "facility_version_snapshot_incomplete"
    assert validation_error in raised.value.detail["validation_errors"]
    for boundary in ("excel", "preview"):
        response = saved.request(boundary)
        assert response.status_code == 400, response.text
        assert response.json()["detail"] == "facility_version_snapshot_incomplete"
    assert database_snapshot(saved) == before
    assert not list(saved.output.rglob("*.xlsx"))


@pytest.mark.parametrize("second,success", [(7, True), (0, False), ("", False)], ids=["equal-seven", "seven-zero-conflict", "seven-blank-conflict"])
def test_explicit_quantity_merge_policy_preserves_equal_cells_but_blocks_information_loss(saved, second, success):
    publish_rows(saved, [7, second])
    saved.authored["fax_template"]["body_merge_policy"] = {
        "mode": "daypart", "columns": ["qty.regular_x"], "required": True}
    saved.republish()
    before = database_snapshot(saved)
    response = saved.request("excel")
    if success:
        assert response.status_code == 200, response.text if response.status_code != 200 else ""
        workbook = load_workbook(BytesIO(response.content))
        try:
            assert "E11:E12" in {str(region) for region in workbook[WEEK].merged_cells.ranges}
            assert workbook[WEEK]["E11"].value == 7
            assert [workbook[WEEK][cell].value for cell in ("D11", "D12")] == ["operator-menu-0", "operator-menu-1"]
        finally:
            workbook.close()
    else:
        assert response.status_code == 400, response.text
        assert response.json()["detail"] == "saved sheet source merge quantity conflict"
        assert not list(saved.output.rglob("*.xlsx"))
    assert database_snapshot(saved) == before


def test_real_revision_change_after_source_load_is_rejected_by_final_capture(saved, monkeypatch):
    publish_rows(saved, [7, 0])
    actual_load = forms.load_workbook
    calls = []
    with saved.session_scope() as session:
        old_revision = session.get(OrderWorkflowState, saved.order_id).revision

    def load_then_external_revision_change(*args, **kwargs):
        workbook = actual_load(*args, **kwargs)
        with saved.session_scope() as session:
            session.get(OrderWorkflowState, saved.order_id).revision += 1
        calls.append(True)
        return workbook

    monkeypatch.setattr(forms, "load_workbook", load_then_external_revision_change)
    response = saved.request("excel")
    assert response.status_code == 400, response.text
    assert response.json()["detail"] == "saved sheet output input changed"
    assert calls == [True]
    assert not list(saved.output.rglob("*.xlsx"))
    with saved.session_scope() as session:
        assert session.get(OrderWorkflowState, saved.order_id).revision == old_revision + 1
        assert session.get(OrderSheetDraft, saved.draft_id).draft_sheet_json == saved.payload


def monthly_entries():
    result = []
    for day in range(31):
        for meal in ("朝", "昼", "夕"):
            for slot in range(3):
                result.append({"menu_date": (date(2026, 3, 1) + timedelta(days=day)).isoformat(),
                               "_menu_date_obj": date(2026, 3, 1) + timedelta(days=day),
                               "_source_index": len(result), "daypart": meal, "category": "主",
                               "name": f"monthly-{day}-{meal}-{slot}"})
    return result


def test_public_monthly_build_resolves_full_dated_facility_and_keeps_all_menu_rows(saved, monkeypatch):
    from src.models.menu import MonthlyMenu, MonthlyMenuEntry, MonthlyMenuItem
    from src.services import menu_service

    # Only database infrastructure is redirected; real schema guards/readers remain active.
    monkeypatch.setattr(menu_service, "engine", saved.engine)
    monkeypatch.setattr(menu_service, "session_scope", saved.session_scope)
    monkeypatch.setattr(menu_service, "_MENU_SCHEMA_INITIALIZED", False)
    entries = monthly_entries()
    with saved.session_scope() as session:
        session.add(MonthlyMenu(id="2026-03", month_start=date(2026, 3, 1)))
        session.flush()
        for index, entry in enumerate(entries):
            session.add(MonthlyMenuEntry(
                id=f"public-entry-{index}", monthly_menu_id="2026-03",
                menu_date=entry["_menu_date_obj"], daypart=entry["daypart"],
                name=entry["name"], category=entry["category"], diet_type="regular",
                slot_index=index, facility_override="F"))
            session.add(MonthlyMenuItem(
                id=f"public-item-{index}", monthly_menu_id="2026-03",
                name=entry["name"], daypart=entry["daypart"], category=entry["category"],
                diet_type="regular", facility_override="F", unit_type="g",
                qty_per_serving=100, master_resolution_mode="month_only"))
    before = database_snapshot(saved)
    source_before = saved.source.read_bytes()
    assert saved.master["facilities"][0]["facility_name"] == "Reference only"
    output = forms.build_order_form_excel(facility_id="F", month_id="2026-03")
    workbook = load_workbook(output)
    try:
        actual_names = []
        for start, end in forms._build_week_ranges_for_month("2026-03"):
            ws = workbook[forms._format_week_sheet_name(start, end)]
            expected = [entry["name"] for entry in entries if start <= entry["_menu_date_obj"] <= end]
            names = [ws.cell(11 + index, 4).value for index in range(len(expected))]
            assert names == expected
            assert ws["J5"].value == "authored-target"
            actual_names.extend(names)
        assert actual_names == [entry["name"] for entry in entries]
        assert len(actual_names) == 279
    finally:
        workbook.close()
    assert database_snapshot(saved) == before
    assert saved.source.read_bytes() == source_before


@pytest.mark.parametrize("body_end,footer,styled_tail", [(66, 67, False), (68, 69, False), (66, 67, True)],
                         ids=["standard-footer", "footer69-no-collision", "outside-print-empty-style86"])
def test_actual_monthly_source_layout_preserves_footer_image_unknown_content_and_all_rows(saved, tmp_path, body_end, footer, styled_tail):
    workbook = load_workbook(saved.source)
    ws = workbook[WEEK]
    ws.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=3)
    ws.merge_cells(start_row=footer, start_column=5, end_row=footer, end_column=12)
    ws.cell(footer, 1).value = "authored-footer"
    ws.row_dimensions[footer].height = 34.5
    ws.print_area = f"A1:L{body_end}"
    image_path = tmp_path / "owned-footer-image.png"
    Image.new("RGB", (8, 8), "red").save(image_path)
    ws.add_image(ExcelImage(image_path), f"E{footer}")
    if styled_tail:
        ws["A86"].font = Font(italic=True)
        ws.row_dimensions[86].height = 18
    workbook.save(saved.source)
    workbook.close()
    source_before = saved.source.read_bytes()
    saved.republish()
    before = database_snapshot(saved)
    entries = monthly_entries()
    generated = forms._build_monthly_fax_order_form_workbook(
        facility=deepcopy(saved.authored), month_id="2026-03", entries=entries, pattern={"pattern_id": "synthetic"})
    output = tmp_path / "owned-monthly-output.xlsx"
    generated.save(output)
    generated.close()
    reloaded = load_workbook(output)
    try:
        for first, last in ((1, 7), (8, 14), (15, 21), (22, 28), (29, 31)):
            title = f"3月{first}日～3月{last}日"
            ws = reloaded[title]
            expected = [entry for entry in entries if first <= entry["_menu_date_obj"].day <= last]
            end = 10 + len(expected)
            delta = end - body_end
            assert [ws.cell(11 + index, 4).value for index in range(len(expected))] == [entry["name"] for entry in expected]
            assert ws.cell(end, 4).value == f"monthly-{last - 1}-夕-2"
            assert ws.cell(footer + delta, 1).value == "authored-footer"
            assert ws["J5"].value == "authored-target"
            assert len(ws._images) == 1
            assert ws._images[0].anchor._from.row == footer + delta - 1
            marker = end + 3
            assert ws.cell(marker, 1).fill.fgColor.rgb in {"00000000", "FF000000"}
            assert f"B{marker}:K{marker}" in {str(region) for region in ws.merged_cells.ranges}
            assert [region.bounds for region in PrintArea.from_string(ws.print_area).ranges] == [(1, 1, 12, marker)]
            assert all(ws.cell(end, column).border.bottom.style == "thick" for column in range(1, 13))
            if styled_tail:
                assert ws.cell(86 + delta, 1).font.italic is True
                assert ws.row_dimensions[86 + delta].height == 18
                assert marker < 86 + delta
    finally:
        reloaded.close()
    assert saved.source.read_bytes() == source_before
    assert database_snapshot(saved) == before


@pytest.mark.parametrize("case,detail", [("missing-json-source", "saved sheet source workbook not configured"),
    ("different-alias-bytes", "saved sheet source workbook ambiguous"),
    ("empty-week", "week menu source unavailable")])
def test_monthly_source_missing_ambiguous_or_empty_week_stops_without_family_or_old_month_fallback(saved, tmp_path, case, detail):
    facility_config = deepcopy(saved.authored)
    entries = monthly_entries()
    if case == "missing-json-source":
        del facility_config["order_form_source_workbook"]
        facility_config["order_form_month_sources"] = {"2026-02": saved.source.name}
    elif case == "different-alias-bytes":
        other = tmp_path / "different-owned-source.xlsx"
        workbook = load_workbook(saved.source)
        workbook[WEEK]["J5"] = "different-content"
        workbook.save(other)
        workbook.close()
        facility_config["order_form_source_workbook_uri"] = other.as_uri()
    else:
        entries = [entry for entry in entries if entry["_menu_date_obj"].day > 7]
    original = saved.source.read_bytes()
    before = database_snapshot(saved)
    with pytest.raises(ValueError, match=f"^{detail}$"):
        forms._build_monthly_fax_order_form_workbook(facility=facility_config, month_id="2026-03", entries=entries, pattern={})
    assert saved.source.read_bytes() == original
    assert database_snapshot(saved) == before
    assert not list(saved.output.rglob("*.xlsx"))
