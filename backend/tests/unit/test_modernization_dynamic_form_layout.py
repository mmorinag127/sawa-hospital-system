"""Synthetic copied-master/layout/geometry proofs, not native FAX or staging proof."""

from copy import copy
from datetime import date, timedelta
import json
from shutil import copyfile

from openpyxl import load_workbook
from openpyxl.worksheet.print_settings import PrintArea
import pytest

from src.services import hakodate_assignment_service as assignment
from src.services import hakodate_step_review_pipeline_service as review
from src.services import master_order_form_template_service as master


START = date(2026, 9, 27)
WEEK = "2026-09@2026-09-27~2026-10-03"
MEALS = ("朝", "昼", "夕")


def facility():
    return {
        "facility_id": "DYNAMIC-SYNTHETIC", "facility_name": "合成行計画試験",
        "fax_template": {
            "columns": [
                {"index": 0, "role": "date", "name": "date_mmdd", "header": "日付"},
                {"index": 1, "role": "daypart", "name": "daypart", "header": "区分"},
                {"index": 2, "role": "menu_name", "name": "menu", "header": "献立"},
                {"index": 3, "role": "quantity", "name": "qty.regular_x", "header": "常食",
                 "diet_type": "regular", "area_id": "X"},
                {"index": 4, "role": "quantity", "name": "qty.no_meat_x", "header": "肉禁",
                 "diet_type": "no_meat", "area_id": "X"},
                {"index": 5, "role": "note", "name": "remarks", "header": "備考"},
            ],
            "body_merge_policy": {"mode": "daypart", "columns": ["qty.regular_x"], "required": True},
        },
    }


def entries_for(days, sizes):
    result = []
    for offset in days:
        for meal, count in zip(MEALS, sizes):
            for slot in range(count):
                result.append({"menu_date": (START + timedelta(days=offset)).isoformat(),
                               "daypart": meal, "category": "主" if slot == 0 else "",
                               "name": f"physical-{offset}-{meal}-{slot}", "sort_order": slot + 1})
    return result


@pytest.fixture
def copied_master(tmp_path):
    original = master.MASTER_TEMPLATE_PATH.read_bytes()
    path = tmp_path / "copied-locked-master.xlsx"
    copyfile(master.MASTER_TEMPLATE_PATH, path)
    yield path
    assert master.MASTER_TEMPLATE_PATH.read_bytes() == original
    assert path.read_bytes() == original


def build_reload(tmp_path, copied_master, entries, *, week=WEEK):
    output = tmp_path / "generated.xlsx"
    master.build_facility_template_xlsx(
        facility_config=facility(), output_path=output, master_template_path=copied_master,
        week_value=week, week_menu_entries=entries,
    )
    return load_workbook(output)


CASES = [("56", range(7), (2, 3, 3), 0), ("57", range(7), (2, 3, 3), 1),
         ("58", range(7), (2, 3, 3), 2), ("four-per-meal", range(7), (4, 4, 4), 0),
         ("five-per-meal", range(7), (5, 5, 5), 0),
         ("September-four-days", range(4), (4, 5, 4), 0),
         ("October-three-days", range(4, 7), (5, 4, 5), 0)]


@pytest.mark.parametrize("label,days,sizes,extra", CASES, ids=[case[0] for case in CASES])
def test_actual_saved_reloaded_layout_and_all_consumers_share_every_physical_row(
    tmp_path, copied_master, label, days, sizes, extra,
):
    entries = entries_for(days, sizes)
    for index in range(extra):
        entries.append({**entries[-1], "name": f"last-extra-{index}", "sort_order": 4 + index})
    expected_count = len(entries)
    workbook = build_reload(tmp_path, copied_master, entries)
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        plan = master.validated_body_row_plan(ws)
        end = 10 + expected_count
        assert len(plan) == expected_count
        assert [row["source_index"] for row in plan] == list(range(expected_count))
        assert [(row["date"], row["daypart"], row["category"], row["name"]) for row in plan] == [
            (entry["menu_date"], entry["daypart"], entry["category"], entry["name"]) for entry in entries
        ]
        assert ws.cell(end, 4).value == entries[-1]["name"]
        assert master.validated_body_end_row(ws) == end
        assert [region.bounds for region in PrintArea.from_string(ws.print_area).ranges] == [(1, 1, 7, end + 2)]
        assert all(ws.cell(end, column).border.bottom.style == "thick" for column in range(1, 8))
        assert ws.cell(end + 1, 4).value is None
        assert ws.row_dimensions[end].height == ws.row_dimensions[11].height
        spans = []
        for offset, entry in enumerate(entries):
            key = (entry["menu_date"], entry["daypart"])
            if spans and spans[-1][2] == key:
                spans[-1] = (spans[-1][0], 11 + offset, key)
            else:
                spans.append((11 + offset, 11 + offset, key))
        assert {(region.min_row, region.max_row) for region in ws.merged_cells.ranges
                if region.min_col == region.max_col == 5 and region.min_row >= 11} == {
                    (first, last) for first, last, _key in spans}
        assert {(region.min_row, region.max_row) for region in ws.merged_cells.ranges
                if region.min_col == region.max_col == 2 and region.min_row >= 11} == {
                    (first, last) for first, last, _key in spans}
        for row in range(11, end):
            assert ws.cell(row, 6).border.bottom.style == "thin"
            assert ws.cell(row, 6).border.right.style == "thin"

        draft = {"fields": ["date_mmdd", "daypart", "menu", "qty.regular_x", "qty.no_meat_x", "remarks"],
                 "rows": [[entry["menu_date"], entry["daypart"], entry["name"], "7", "0", ""] for entry in entries]}
        logical = review.canonical_logical_target_plan(worksheet=ws, fax_template=facility()["fax_template"], draft_sheet=draft)
        covered = [cell for region in logical for cell in region["covered_sheet_cells"]]
        assert len(covered) == expected_count * 3
        assert set(covered) == {f"{column}{row}" for row in range(11, end + 1) for column in "EFG"}
        grid = assignment._structure_grid_from_rendered_worksheet(ws, dpi=144)
        assert grid["detected_row_count"] == expected_count + 2
        assert len(grid["row_edges"]) == expected_count + 3
        merge_regions = assignment._worksheet_merge_regions_for_grid(
            ws, row_edges=grid["row_edges"], column_edges=grid["column_edges"],
            quantity_columns={5, 6},
        )
        quantity_regions = [region for region in merge_regions
                            if region["is_quantity"] and region["min_row"] >= 11]
        assert len(quantity_regions) == len(spans)
        assert {region["range"] for region in quantity_regions} == {
            f"E{first}:E{last}" for first, last, _key in spans
        }
        lines = assignment._merge_aware_grid_line_segments(
            row_edges=grid["row_edges"], column_edges=grid["column_edges"],
            merge_regions=merge_regions,
        )
        x0, x1, adjacent_end = grid["column_edges"][4:7]
        for first, last, _key in spans:
            start_index, end_index = first - 9, last - 9
            y0, y1 = grid["row_edges"][start_index], grid["row_edges"][end_index + 1]
            region = next(region for region in quantity_regions if region["range"] == f"E{first}:E{last}")
            assert region == {
                "range": f"E{first}:E{last}",
                "min_col": 5, "max_col": 5, "min_row": first, "max_row": last,
                "row_span": last - first + 1, "col_span": 1,
                "start_col_index": 4, "end_col_index": 4,
                "start_row_index": start_index, "end_row_index": end_index,
                "bbox": pytest.approx([x0, y0, x1, y1]), "is_quantity": True,
            }
            for edge_index in range(start_index + 1, end_index + 1):
                y = grid["row_edges"][edge_index]
                horizontal = [line for line in lines if line["orientation"] == "horizontal"
                              and line["start"][1] == line["end"][1] == y]
                assert all(line["end"][0] <= x0 or line["start"][0] >= x1 for line in horizontal)
                assert any(line["start"][0] <= x1 and line["end"][0] >= adjacent_end
                           for line in horizontal)
            for x in (x0, x1):
                assert any(line["orientation"] == "vertical"
                           and line["start"][0] == line["end"][0] == x
                           and line["start"][1] <= y0 and line["end"][1] >= y1
                           for line in lines)
        maps = [assignment._workbook_physical_row_map(ws, row_count=expected_count + 2),
                review._step_review_physical_row_map(ws, row_count=expected_count + 2)]
        for row_map in maps:
            assert list(row_map) == list(range(2, expected_count + 2))
            assert [row["worksheet_row"] for row in row_map.values()] == list(range(11, end + 1))
            assert [row["menu_name"] for row in row_map.values()] == [entry["name"] for entry in entries]
        assert review._step_review_worksheet_row_to_grid_index(end, body_end_row=end) == expected_count + 1
        assert review._step_review_worksheet_row_to_grid_index(end + 1, body_end_row=end) is None
        regions, metadata = review._post_menu_target_regions(
            worksheet=ws, column_edges=grid["column_edges"], row_edges=grid["row_edges"],
            fax_template=facility()["fax_template"], draft_sheet=draft,
            menu_body_row_indexes=list(range(expected_count)),
        )
        assert metadata["logical_target_count"] == expected_count * 3
        assert {cell for region in regions for cell in region["covered_sheet_cells"]} == set(covered)
        for region in regions:
            indexes = [target["grid_row_index"] for target in region["logical_targets"]]
            column = region["worksheet_col"] - 1
            assert region["bbox"] == pytest.approx([
                grid["column_edges"][column], grid["row_edges"][min(indexes)],
                grid["column_edges"][column + 1], grid["row_edges"][max(indexes) + 1],
            ])
    finally:
        workbook.close()


def schema_cell(workbook, key):
    return next(row[1] for row in workbook["generated_template_schema"].iter_rows() if row[0].value == key)


@pytest.mark.parametrize("tamper,code", [
    ("schema", "generated_body_schema_missing"), ("rowcount", "generated_body_row_plan_changed"),
    ("menu", "generated_body_row_identity_changed"), ("date", "generated_body_row_identity_changed"),
    ("meal", "generated_body_row_identity_changed"), ("category", "generated_body_row_identity_changed"),
    ("source-index", "generated_body_row_identity_invalid"), ("merge", "generated_body_merge_plan_changed"),
    ("border", "generated_body_end_border_changed"), ("print", "generated_body_row_plan_changed"),
])
def test_saved_workbook_tamper_stops_shared_validation(tmp_path, copied_master, tamper, code):
    workbook = build_reload(tmp_path, copied_master, entries_for(range(4), (4, 4, 4)))
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        if tamper == "schema":
            del workbook["generated_template_schema"]
        elif tamper == "rowcount":
            schema_cell(workbook, "week_menu_rows").value = 47
        elif tamper in {"menu", "date", "meal", "category"}:
            ws.cell(11, {"date": 1, "meal": 2, "category": 3, "menu": 4}[tamper]).value = "changed"
        elif tamper == "source-index":
            cell = schema_cell(workbook, "week_menu_physical_row_plan")
            plan = json.loads(cell.value)
            plan[1]["source_index"] = plan[0]["source_index"]
            cell.value = json.dumps(plan)
        elif tamper == "merge":
            ws.unmerge_cells("E11:E14")
        elif tamper == "border":
            border = copy(ws.cell(58, 7).border)
            border.bottom = None
            ws.cell(58, 7).border = border
        else:
            ws.print_area = "A1:G59"
        output = tmp_path / "tampered.xlsx"
        workbook.save(output)
        reloaded = load_workbook(output)
        try:
            with pytest.raises(master.FacilityTemplateBuildError, match=f"^{code}$"):
                master.validated_body_row_plan(reloaded[master.FACILITY_TEMPLATE_SHEET_NAME])
        finally:
            reloaded.close()
    finally:
        workbook.close()


@pytest.mark.parametrize("field,value", [("date_mmdd", "2026-09-28"), ("daypart", "昼"), ("menu", "wrong-menu")])
def test_logical_plan_rejects_same_count_draft_identity_substitution(tmp_path, copied_master, field, value):
    entries = entries_for(range(3), (4, 4, 4))
    workbook = build_reload(tmp_path, copied_master, entries)
    try:
        fields = ["date_mmdd", "daypart", "menu"]
        draft = {"fields": fields, "rows": [[entry["menu_date"], entry["daypart"], entry["name"]] for entry in entries]}
        draft["rows"][0][fields.index(field)] = value
        with pytest.raises(ValueError, match="^hakodate_physical_row_changed$"):
            review.canonical_logical_target_plan(worksheet=workbook[master.FACILITY_TEMPLATE_SHEET_NAME],
                                                 fax_template=facility()["fax_template"], draft_sheet=draft)
    finally:
        workbook.close()


@pytest.mark.parametrize("week,entries,code", [(WEEK, [], "week_menu_entries_missing"),
    ("not-a-week", [], "week_menu_week_invalid"),
    (WEEK, [{"menu_date": "?", "daypart": "朝", "name": "invalid"}], "week_menu_entry_identity_unresolved")])
def test_explicit_week_unresolved_inputs_never_generate_static_body(tmp_path, copied_master, week, entries, code):
    with pytest.raises(master.FacilityTemplateBuildError, match=f"^{code}$"):
        build_reload(tmp_path, copied_master, entries, week=week)
    assert not (tmp_path / "generated.xlsx").exists()


def test_weekless_prototype_keeps_explicit_static_contract(tmp_path, copied_master):
    workbook = build_reload(tmp_path, copied_master, None, week=None)
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        assert master.validated_body_row_plan(ws, allow_static=True) is None
        assert master.validated_body_end_row(ws, allow_static=True) == 67
        assert [region.bounds for region in PrintArea.from_string(ws.print_area).ranges] == [(1, 1, 7, 69)]
        assert ws["A4"].value == facility()["facility_name"]
        assert schema_cell(workbook, "week_menu_rows").value == 0
        assert {"E11:E12", "E13:E15", "E16:E18"}.issubset({str(region) for region in ws.merged_cells.ranges})
    finally:
        workbook.close()


@pytest.mark.parametrize("sizes,extra,expected_count", [
    ((2, 3, 3), 0, 56), ((2, 3, 3), 2, 58), ((5, 5, 5), 0, 105),
], ids=["render-56", "render-58", "render-five-per-meal"])
def test_actual_synthetic_png_pixels_match_dynamic_body_geometry(tmp_path, copied_master, sizes, extra, expected_count):
    """Copied-master synthetic rendering only, not recognizer/native FAX/stg proof."""
    from PIL import Image
    from src.services import workbook_pdf_renderer as renderer

    entries = entries_for(range(7), sizes)
    for index in range(extra):
        entries.append({**entries[-1], "name": f"render-last-extra-{index}", "sort_order": 4 + index})
    workbook = build_reload(tmp_path, copied_master, entries)
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        plan = master.validated_body_row_plan(ws)
        assert len(plan) == expected_count
        end = 10 + expected_count
        assert plan[-1]["worksheet_row"] == end
        assert ws.cell(end, 4).value == entries[-1]["name"]
        geometry = renderer.worksheet_render_geometry(ws, dpi=144)
        grid = assignment._structure_grid_from_rendered_worksheet(ws, dpi=144)
        image = renderer.render_worksheet_to_image(ws, dpi=144)
        png = tmp_path / f"synthetic-body-{expected_count}.png"
        try:
            image.save(png, format="PNG")
        finally:
            image.close()
        with Image.open(png) as actual:
            assert actual.format == "PNG"
            assert actual.size == (geometry["image_width"], geometry["image_height"])
            assert geometry["print_range"] == [1, 1, 7, end + 2]
            width, height = actual.size
            assert grid["detected_row_count"] == expected_count + 2
            assert round(grid["row_edges"][-2] * height) == geometry["y_positions"][end]
            bottom = round(grid["row_edges"][-1] * height)
            assert bottom == geometry["y_positions"][end + 1]
            assert bottom < height - renderer._DEFAULT_MARGIN_PX
            left, right = [round(grid["column_edges"][index] * width) for index in (5, 6)]
            assert (left, right) == (geometry["x_positions"][6], geometry["x_positions"][7])
            center = (left + right) // 2
            pixels = actual.convert("RGB")
            try:
                # F is an empty, unmerged quantity column: text cannot mimic this border.
                assert all(pixels.getpixel((x, bottom)) == (0, 0, 0) for x in range(left + 4, right - 3))
                assert [pixels.getpixel((center, bottom + dy)) for dy in (-1, 0, 1)] == [(0, 0, 0)] * 3
                assert pixels.getpixel((center, bottom - 5)) == (255, 255, 255)
                last_row_center = (geometry["y_positions"][end] + bottom) // 2
                assert pixels.getpixel((right, last_row_center)) == (0, 0, 0)
            finally:
                pixels.close()
    finally:
        workbook.close()
