"""Business environment snapshots; no external service or live operator data."""
from copy import deepcopy
from contextvars import copy_context
from concurrent.futures import ThreadPoolExecutor

import pytest
from src import staging_data as data
from src.services import gemini_ocr_service as gemini


def snapshot(monkeypatch, **settings):
    identity = 'replay-runtime-settings'
    value = {'id': identity, 'runtime_settings_provenance': 'production_environment', 'components': {
        feature: {'runtime_settings': {key: None for key in keys}} for feature, keys in data.RUNTIME_SETTING_KEYS.items()}}
    value['components']['order']['runtime_settings'].update(settings)
    monkeypatch.setenv('K_SERVICE', 'worker-stg')
    monkeypatch.setenv('APP_ENV', 'stg')
    monkeypatch.setattr(data, 'manifest', lambda: value)
    monkeypatch.setattr(data, 'manifest_for', lambda selected: deepcopy(value))
    return identity, value


def test_model_and_limits_use_copy_values_and_restore_test_even_across_thread(monkeypatch):
    monkeypatch.setenv('GEMINI_OCR_MODEL', 'gemini-test-env')
    identity, _ = snapshot(monkeypatch, GEMINI_OCR_MODEL='gemini-production-copy', OPENAI_OCR_MODEL='gpt-copy', GEMINI_OCR_MAX_TOKENS='321')
    assert gemini._get_model({}) == 'gemini-test-env'
    with data.dataset_context(identity):
        assert gemini._get_model({}) == 'gemini-production-copy'
        assert data.order_runtime_setting('OPENAI_OCR_MODEL') == 'gpt-copy'
        context = copy_context()
        with ThreadPoolExecutor(max_workers=1) as executor:
            assert executor.submit(context.run, gemini._get_model, {}).result() == 'gemini-production-copy'
    assert gemini._get_model({}) == 'gemini-test-env'


def test_explicit_source_unset_uses_callsite_default_and_never_test_env(monkeypatch):
    monkeypatch.setenv('GEMINI_OCR_MODEL', 'gemini-test-env')
    identity, _ = snapshot(monkeypatch)
    with data.dataset_context(identity):
        assert data.order_runtime_setting('GEMINI_OCR_MODEL', 'code-default') == 'code-default'
        assert gemini._get_model({}) == 'gemini-2.5-flash'
        assert data.order_runtime_setting('OCR_REPARSE_OCR_PIPELINE_TIMEOUT_SECONDS', '45') == '45'


def test_missing_setting_and_secret_names_fail_without_reading_test_env(monkeypatch):
    identity, value = snapshot(monkeypatch)
    del value['components']['order']['runtime_settings']['GEMINI_OCR_MODEL']
    with data.dataset_context(identity):
        with pytest.raises(data.DatasetUnavailable, match='setting_required'):
            gemini._get_model({})
        with pytest.raises(data.DatasetUnavailable, match='setting_not_allowed'):
            data.order_runtime_setting('GEMINI_API_KEY')
        with pytest.raises(data.DatasetUnavailable, match='setting_not_allowed'):
            data.order_runtime_setting('OCR_REPARSE_UNKNOWN_SECRET_TIMEOUT_SECONDS')


def test_explicit_legacy_provenance_preserves_existing_stg_settings(monkeypatch):
    monkeypatch.setenv('GEMINI_OCR_MODEL', 'gemini-existing-stg')
    identity, value = snapshot(monkeypatch)
    value['runtime_settings_provenance'] = 'legacy_stg_runtime'
    with data.dataset_context(identity):
        assert gemini._get_model({}) == 'gemini-existing-stg'


def test_production_manifest_cannot_claim_legacy_provenance_to_use_test_env(monkeypatch):
    identity, value = snapshot(monkeypatch)
    value['runtime_settings_provenance'] = 'legacy_stg_runtime'
    monkeypatch.setenv('SAWA_STG_REPLAY_MANIFEST_URI', 'gs://fixture-stg-raw/stg-replay/replay-original/manifest.json')
    monkeypatch.setenv('SAWA_STG_REPLAY_MANIFEST_SHA256', 'a' * 64)
    monkeypatch.setattr(data, '_read_manifest', lambda *args: value)
    monkeypatch.setattr(data, 'legacy_manifest', lambda: {'id': 'replay-original'})
    with pytest.raises(data.DatasetUnavailable, match='runtime_settings_provenance'):
        data._runtime_manifest({'uri': 'gs://fixture-stg-raw/stg-replay/' + identity + '/manifest.json', 'sha256': 'b' * 64})



def test_shipping_uses_verified_copy_template_and_leaves_test_template_unchanged(monkeypatch, tmp_path):
    from io import BytesIO
    from openpyxl import Workbook, load_workbook
    from src.services import shipping_service as shipping
    original = '/production/templates/shipping.xlsx'
    identity, _ = snapshot(monkeypatch, SHIPPING_TEMPLATE_PATH=original)
    test_path = tmp_path / 'test.xlsx'
    test = Workbook(); test.active['A1'] = 'test-template'; test.save(test_path)
    before = test_path.read_bytes()
    copied = Workbook(); copied.active['A1'] = 'copy-template'; content = BytesIO(); copied.save(content)
    calls = []
    def load(reference):
        calls.append(reference)
        assert reference == original
        return content.getvalue()
    monkeypatch.setattr(shipping, 'load_bytes_from_uri', load)
    monkeypatch.setattr(shipping, 'SHIPPING_TEMPLATE_PATH', test_path)
    monkeypatch.setattr(shipping, 'OUTPUT_DIR', tmp_path)
    assert load_workbook(shipping.build_shipping_excel([])).active['A1'].value == 'test-template'
    with data.dataset_context(identity):
        assert load_workbook(shipping.build_shipping_excel([])).active['A1'].value == 'copy-template'
    assert load_workbook(shipping.build_shipping_excel([])).active['A1'].value == 'test-template'
    assert calls == [original] and test_path.read_bytes() == before


def test_execution_slot_limits_and_producer_setting_do_not_mix(monkeypatch):
    from src.services import ocr_execution_lock_service as slots
    monkeypatch.setenv('WORKFLOW_V2_OCR_MAX_CONCURRENT', '2')
    monkeypatch.setenv('WORKFLOW_V2_OCR_SLOT_POLL_SECONDS', '3')
    monkeypatch.setenv('WORKFLOW_V2_OCR_SLOT_TIMEOUT_SECONDS', '60')
    identity, value = snapshot(monkeypatch, WORKFLOW_V2_OCR_MAX_CONCURRENT='4', WORKFLOW_V2_OCR_SLOT_POLL_SECONDS='7', WORKFLOW_V2_OCR_SLOT_TIMEOUT_SECONDS='120', OCR_TEMPLATE_MIN_CONFIDENCE='0.7')
    value['components']['order']['producer_runtime_settings'] = {key: None for key in data.RUNTIME_PRODUCER_KEYS}
    value['components']['order']['producer_runtime_settings']['OCR_TEMPLATE_MIN_CONFIDENCE'] = '0.8'
    assert (slots._max_slots(), slots._poll_seconds(), slots._timeout_seconds()) == (2, 3, 60)
    with data.dataset_context(identity):
        assert (slots._max_slots(), slots._poll_seconds(), slots._timeout_seconds()) == (4, 7, 120)
        assert data.order_runtime_setting('OCR_TEMPLATE_MIN_CONFIDENCE') == '0.7'
        assert data.producer_runtime_setting('OCR_TEMPLATE_MIN_CONFIDENCE') == '0.8'
    assert (slots._max_slots(), slots._poll_seconds(), slots._timeout_seconds()) == (2, 3, 60)
