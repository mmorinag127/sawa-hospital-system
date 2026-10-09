"""Synthetic workbook/render component definitions, not native, live or PG proof."""

from datetime import date, datetime

import numpy as np
from openpyxl import load_workbook
from PIL import ImageDraw
import pytest

from test_modernization_dynamic_form_layout import copied_master
from test_modernization_serialized_form_geometry import narrow_config
from src.services import hakodate_assignment_service as assignment
from src.services import hakodate_fixed_quad_registration_service as quad
from src.services import master_order_form_template_service as master
from src.services import workbook_pdf_renderer as renderer


DATES = (date(2026, 12, 31), date(2027, 1, 1))
WEEK = "2026-12@2026-12-27~2027-01-02"
FORMAT = "m/d\n(aaa)"


def build(tmp_path, copied_master, quantity_count, days):
    entries = [
        {"menu_date": day.isoformat(), "daypart": "朝", "category": "主",
         "name": f"synthetic-{index}", "sort_order": index + 1}
        for index, day in enumerate(days)
    ]
    path = master.build_facility_template_xlsx(
        facility_config=narrow_config(quantity_count), master_template_path=copied_master,
        output_path=tmp_path / "dates.xlsx", week_value=WEEK, week_menu_entries=entries,
    )
    return path, entries


@pytest.mark.parametrize("quantity_count", [1, 2])
@pytest.mark.parametrize("days", [(DATES[0],), DATES])
def test_singleton_saved_dates_text_bounds_and_all_geometry_agree(
    tmp_path, copied_master, quantity_count, days,
):
    path, entries = build(tmp_path, copied_master, quantity_count, days)
    workbook = load_workbook(path)
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        plan = master.validated_body_row_plan(ws)
        assert [item["date"] for item in plan] == [day.isoformat() for day in days]
        assert [item["name"] for item in plan] == [item["name"] for item in entries]
        assert ws.max_column == 4 + quantity_count
        geometry = renderer.worksheet_render_geometry(ws, dpi=144)
        xs, ys = quad.canonical_template_axes_from_workbook(
            path, sheet_name=ws.title, canvas_width=geometry["image_width"],
            canvas_height=geometry["image_height"],
        )
        edge_rows = master.validated_template_row_edges(ws)
        assert ys == [geometry["y_positions"][row] for row in edge_rows]
        assert xs == [geometry["x_positions"][col] for col in range(1, ws.max_column + 2)]
        grid = assignment._structure_grid_from_rendered_worksheet(ws, dpi=144)
        assert [round(edge * geometry["image_height"]) for edge in grid["row_edges"]] == ys
        image = renderer.render_worksheet_to_image(ws, dpi=144)
        try:
            image.save(tmp_path / "synthetic-singleton-dates.png")
            pixels = np.asarray(image.convert("RGB"))
            draw = ImageDraw.Draw(image)
            for row, day in enumerate(days, 11):
                cell = ws.cell(row, 1)
                assert isinstance(cell.value, (date, datetime))
                assert (cell.value.year, cell.value.month, cell.value.day) == (day.year, day.month, day.day)
                assert cell.number_format == FORMAT
                assert cell.alignment.wrap_text is True
                assert cell.alignment.horizontal == cell.alignment.vertical == "center"
                text = renderer._cell_text(cell.value, number_format=cell.number_format)
                assert text == f"{day.month}/{day.day}\n({'月火水木金土日'[day.weekday()]})"
                font = renderer._load_font(round(float(cell.font.sz or 11) * 144 / 72), bold=bool(cell.font.bold))
                width, height = renderer._text_bbox(draw, text, font)
                left, right = geometry["x_positions"][1:3]
                top, bottom = geometry["y_positions"][row:row + 2]
                text_x = left + max(4, ((right - left) - width) // 2)
                text_y = top + max(3, ((bottom - top) - height) // 2)
                bbox = draw.multiline_textbbox((text_x, text_y), text, font=font, spacing=2)
                assert left < bbox[0] < bbox[2] < right
                assert top < bbox[1] < bbox[3] < bottom
                assert ws.row_dimensions[row].height >= 2 * float(cell.font.sz or 11) + 8
                assert np.any(pixels[top + 3:bottom - 3, left + 3:right - 3] < 255)
                assert [ws.cell(row, col).value for col in range(5, ws.max_column + 1)] == [None] * quantity_count
        finally:
            image.close()
    finally:
        workbook.close()


def test_multiday_identity_keeps_existing_date_and_weekday_layout(tmp_path, copied_master):
    path, _entries = build(tmp_path, copied_master, 1, (DATES[0], DATES[0], DATES[0]))
    workbook = load_workbook(path)
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        assert "A11:A12" in ws.merged_cells
        assert (ws["A11"].value.year, ws["A11"].value.month, ws["A11"].value.day) == (2026, 12, 31)
        assert ws["A13"].value == "木"
        assert ws["A11"].number_format != FORMAT
        assert [row["date"] for row in master.validated_body_row_plan(ws)] == ["2026-12-31"] * 3
    finally:
        workbook.close()


def test_preserve_date_cells_branch_retains_single_line_format(tmp_path, copied_master):
    path, entries = build(tmp_path, copied_master, 1, (DATES[0],))
    workbook = load_workbook(path)
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        height = ws.row_dimensions[11].height
        alignment = ws["A11"].alignment.copy()
        normalized = [{**entry, "_menu_date_obj": DATES[0], "_source_index": index}
                      for index, entry in enumerate(entries)]
        master._write_week_menu_identity(ws, normalized, source_body_end_row=11, preserve_date_cells=True)
        assert ws["A11"].value == DATES[0]
        assert ws["A11"].number_format == "m/d (aaa)"
        assert renderer._cell_text(ws["A11"].value, number_format=ws["A11"].number_format) == "12/31"
        assert ws.row_dimensions[11].height == height
        assert ws["A11"].alignment == alignment
        workbook.save(path)
    finally:
        workbook.close()
    reloaded = load_workbook(path)
    try:
        cell = reloaded[master.FACILITY_TEMPLATE_SHEET_NAME]["A11"]
        assert (cell.value.year, cell.value.month, cell.value.day) == (2026, 12, 31)
        assert cell.number_format == "m/d (aaa)"
    finally:
        reloaded.close()


@pytest.mark.parametrize("value,number_format,expected", [
    (None, FORMAT, ""), (DATES[0], None, "12/31"),
    (datetime(2027, 1, 1, 12, 30), None, "1/1"),
    (DATES[0], "yyyy-mm-dd", "12/31"),
    (datetime(2027, 1, 1), FORMAT, "1/1\n(金)"),
    (0, FORMAT, "0"), ("", FORMAT, ""), ("7", None, "7"),
])
def test_renderer_other_values_and_formats_remain_unchanged(value, number_format, expected):
    assert renderer._cell_text(value, number_format=number_format) == expected
