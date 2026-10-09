"""Synthetic serialized workbook/render/component boundaries, not native/live proof."""

import ast
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

import cv2
import numpy as np
from openpyxl import load_workbook
from openpyxl.worksheet.print_settings import PrintArea
import pytest

from test_modernization_dynamic_form_layout import copied_master, entries_for, facility, WEEK
from src.services import hakodate_assignment_service as assignment
from src.services import hakodate_fixed_quad_registration_service as quad
from src.services import hakodate_step_review_pipeline_service as review
from src.services import master_order_form_template_service as master
from src.services import workbook_pdf_renderer as renderer


def narrow_config(quantity_count):
    config = deepcopy(facility())
    template = config["fax_template"]
    template["columns"] = template["columns"][:3] + template["columns"][3:3 + quantity_count]
    template.pop("body_merge_policy")
    return config


def physical_entries(count):
    if count in (84, 105):
        return entries_for(range(7), (count // 21,) * 3)
    rows = entries_for(range(7), (2, 3, 3))
    if count == 1:
        return rows[:1]
    for index in range(count - 56):
        rows.append({**rows[-1], "name": f"extra-{index}", "sort_order": 4 + index})
    return rows


def generated(tmp_path, copied_master, quantity_count=1, count=56):
    config = narrow_config(quantity_count)
    entries = physical_entries(count)
    path = master.build_facility_template_xlsx(
        facility_config=config, master_template_path=copied_master,
        output_path=tmp_path / "serialized.xlsx", week_value=WEEK, week_menu_entries=entries,
    )
    return path, config, entries


def media_bytes(path):
    with ZipFile(path) as archive:
        return sorted(archive.read(name) for name in archive.namelist() if name.startswith("xl/media/"))


@pytest.mark.parametrize("quantity_count", [1, 2])
@pytest.mark.parametrize("count", [1, 56, 57, 58, 84, 105])
def test_serialized_narrow_workbook_all_consumers_and_render_share_authored_geometry(
    tmp_path, copied_master, quantity_count, count,
):
    path, config, entries = generated(tmp_path, copied_master, quantity_count, count)
    assert media_bytes(path) == media_bytes(copied_master)
    assert media_bytes(path)
    workbook = load_workbook(path)
    original = load_workbook(copied_master)
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        source = original.active
        end, max_col = 10 + count, 4 + quantity_count
        plan = master.validated_body_row_plan(ws)
        schema = {row[0]: row[1] for row in workbook["generated_template_schema"].iter_rows(values_only=True)
                  if len(row) >= 2 and isinstance(row[0], str)}
        assert schema["generated_end_col"] == max_col
        assert schema["generated_column_count"] == quantity_count
        assert schema["week_menu_rows"] == count
        assert schema["body_end_row"] == end
        assert len(plan) == count
        assert [row["source_index"] for row in plan] == list(range(count))
        assert [row["name"] for row in plan] == [entry["name"] for entry in entries]
        assert ws.max_column == max_col
        assert [region.bounds for region in PrintArea.from_string(ws.print_area).ranges] == [(1, 1, max_col, end + 2)]
        assert all(region.max_col <= max_col for region in ws.merged_cells.ranges)
        assert ws["A4"].value == config["facility_name"]
        assert "A4:D5" in ws.merged_cells
        assert [ws.cell(row, 5).value for row in range(2, 7)] == [source.cell(row, 7).value for row in range(2, 7)]
        assert ws.cell(7, 5).value == "常食"
        if quantity_count == 2:
            assert ws.cell(7, 6).value == "肉禁"
        assert len(ws._images) == len(source._images)
        source_geometry = renderer.worksheet_render_geometry(source, dpi=144)
        geometry = renderer.worksheet_render_geometry(ws, dpi=144)
        logo_rectangles = []
        for drawing, original_drawing in zip(ws._images, source._images, strict=True):
            anchor = drawing.anchor
            original_anchor = original_drawing.anchor
            def original_point(marker):
                return (
                    source_geometry["x_positions"][marker.col + 1] + round(marker.colOff / 9525 * 144 / 96),
                    source_geometry["y_positions"][marker.row + 1] + round(marker.rowOff / 9525 * 144 / 96),
                )
            old_left, old_top = original_point(original_anchor._from)
            old_right, old_bottom = original_point(original_anchor.to)
            # OneCell extents use integral 96-dpi pixels; allow only that exact quantization.
            assert anchor.ext.cx == round((old_right - old_left) * 96 / 144) * 9525
            assert anchor.ext.cy == round((old_bottom - old_top) * 96 / 144) * 9525
            left = geometry["x_positions"][anchor._from.col + 1] + round(anchor._from.colOff / 9525 * 144 / 96)
            top = geometry["y_positions"][anchor._from.row + 1] + round(anchor._from.rowOff / 9525 * 144 / 96)
            right = left + round(anchor.ext.cx / 9525 * 144 / 96)
            bottom = top + round(anchor.ext.cy / 9525 * 144 / 96)
            assert abs((right - left) - (old_right - old_left)) <= 1
            assert abs((bottom - top) - (old_bottom - old_top)) <= 1
            assert geometry["x_positions"][1] <= left < right <= geometry["x_positions"][max_col + 1]
            assert geometry["y_positions"][1] <= top < bottom <= geometry["y_positions"][end + 3]
            assert right < geometry["image_width"] and bottom < geometry["image_height"]
            content_boxes = [
                (geometry["x_positions"][1], geometry["y_positions"][4],
                 geometry["x_positions"][5], geometry["y_positions"][6]),
                (geometry["x_positions"][5], geometry["y_positions"][2],
                 geometry["x_positions"][max_col + 1], geometry["y_positions"][7]),
            ]
            for box_left, box_top, box_right, box_bottom in content_boxes:
                assert right <= box_left or left >= box_right or bottom <= box_top or top >= box_bottom
            logo_rectangles.append((left, top, right, bottom))
        edge_rows = (7, 9) + tuple(range(11, end + 2))
        assert master.validated_template_row_edges(ws) == edge_rows
        assert [master.template_row_to_grid_index(row, body_end_row=end) for row in (7, 9, *range(11, end + 1))] == list(range(count + 2))
        assert review._step_review_worksheet_row_to_grid_index(end, body_end_row=end) == count + 1
        assert quad.template_axis_edge_counts_from_workbook(path, sheet_name=ws.title) == (max_col + 1, count + 3)
        xs, ys = quad.canonical_template_axes_from_workbook(
            path, sheet_name=ws.title, canvas_width=geometry["image_width"], canvas_height=geometry["image_height"],
        )
        assert xs == [geometry["x_positions"][index] for index in range(1, max_col + 2)]
        assert ys == [geometry["y_positions"][row] for row in edge_rows]
        grid = assignment._structure_grid_from_rendered_worksheet(ws, dpi=144)
        assert grid["detected_row_count"] == count + 2
        assert [round(value * geometry["image_height"]) for value in grid["row_edges"]] == ys
        for row_map in (assignment._workbook_physical_row_map(ws, row_count=count + 2),
                        review._step_review_physical_row_map(ws, row_count=count + 2)):
            assert [row["worksheet_row"] for row in row_map.values()] == list(range(11, end + 1))
        fields = ["date_mmdd", "daypart", "menu"] + [column["name"] for column in config["fax_template"]["columns"][3:]]
        draft = {"fields": fields, "rows": [[entry["menu_date"], entry["daypart"], entry["name"]] + [""] * quantity_count for entry in entries]}
        logical = review.canonical_logical_target_plan(worksheet=ws, fax_template=config["fax_template"], draft_sheet=draft)
        assert {cell for target in logical for cell in target["covered_sheet_cells"]} == {
            f"{chr(64 + col)}{row}" for row in range(11, end + 1) for col in range(5, max_col + 1)}
        image = renderer.render_worksheet_to_image(ws, dpi=144)
        try:
            image.save(tmp_path / f"synthetic-{quantity_count}-{count}.png")
            assert image.size == (geometry["image_width"], geometry["image_height"])
            pixels = np.asarray(image.convert("RGB"))
            bottom = geometry["y_positions"][end + 1]
            left, right = geometry["x_positions"][max_col], geometry["x_positions"][max_col + 1]
            assert np.all(pixels[bottom, left + 4:right - 3] == 0)
            assert np.any(pixels[:geometry["y_positions"][7]] < 255)
            assert np.any(pixels[geometry["y_positions"][end]:bottom] < 255)
            for left, top, right, logo_bottom in logo_rectangles:
                assert np.any(pixels[top:logo_bottom, left:right] < 255)
            facility_rect = pixels[geometry["y_positions"][4] + 4:geometry["y_positions"][6] - 4,
                                   geometry["x_positions"][1] + 4:geometry["x_positions"][5] - 4]
            assert np.any(facility_rect < 255)
            for row in range(2, 7):
                if ws.cell(row, 5).value:
                    assert np.any(pixels[geometry["y_positions"][row] + 3:geometry["y_positions"][row + 1] - 3,
                                         geometry["x_positions"][5] + 3:geometry["x_positions"][max_col + 1] - 3] < 255)
        finally:
            image.close()
    finally:
        original.close()
        workbook.close()


@pytest.mark.parametrize("axis", ["x", "y"])
@pytest.mark.parametrize("tamper", ["short", "extra", "duplicate", "nonmonotone"])
def test_explicit_axes_cannot_self_approve_invalid_counts_or_order(tmp_path, copied_master, axis, tamper):
    path, _config, _entries = generated(tmp_path, copied_master)
    xs, ys = quad.canonical_template_axes_from_workbook(path, sheet_name=master.FACILITY_TEMPLATE_SHEET_NAME, canvas_width=1200, canvas_height=2400)
    axes = {"x": list(xs), "y": list(ys)}
    values = axes[axis]
    if tamper == "short":
        values.pop()
    elif tamper == "extra":
        values.append(values[-1] + 1)
    elif tamper == "duplicate":
        values[1] = values[0]
    else:
        values[0], values[1] = values[1], values[0]
    reason = f"canonical template {axis} axes " + ("incomplete" if tamper in {"short", "extra"} else "are not strictly increasing")
    with pytest.raises(ValueError, match=reason):
        quad.resolve_template_axes_from_manifest_or_image(
            item={"template_xlsx": path, "template_sheet_name": master.FACILITY_TEMPLATE_SHEET_NAME,
                  "template_axes_x": axes["x"], "template_axes_y": axes["y"]},
            template_image=np.full((2400, 1200, 3), 255, np.uint8), manifest_template_bbox=[0, 0, 1199, 2399],
        )


def test_canonical_axes_missing_schema_and_bbox_crop_stop(tmp_path, copied_master):
    with pytest.raises(ValueError, match="canonical_template_workbook_required"):
        quad.template_axis_edge_counts_from_workbook(None, sheet_name=None)
    path, _config, _entries = generated(tmp_path, copied_master)
    with pytest.raises(ValueError, match="canonical template x axes incomplete"):
        quad.canonical_template_axes_from_workbook(path, sheet_name=master.FACILITY_TEMPLATE_SHEET_NAME,
                                                 canvas_width=1200, canvas_height=2400, table_bbox=[500, 0, 1199, 2399])
    workbook = load_workbook(path)
    try:
        del workbook["generated_template_schema"]
        workbook.save(path)
    finally:
        workbook.close()
    with pytest.raises(master.FacilityTemplateBuildError):
        quad.canonical_template_axes_from_workbook(path, sheet_name=master.FACILITY_TEMPLATE_SHEET_NAME,
                                                 canvas_width=1200, canvas_height=2400)


@pytest.mark.parametrize("key,value", [("week_menu_physical_row_plan", "not-json"), ("week_menu_rows", 55), ("generated_end_col", 6)])
def test_serialized_corrupt_schema_cannot_supply_axis_counts(tmp_path, copied_master, key, value):
    path, _config, _entries = generated(tmp_path, copied_master)
    workbook = load_workbook(path)
    try:
        cell = next(row[1] for row in workbook["generated_template_schema"].iter_rows() if row[0].value == key)
        cell.value = value
        workbook.save(path)
    finally:
        workbook.close()
    with pytest.raises(master.FacilityTemplateBuildError):
        quad.template_axis_edge_counts_from_workbook(path, sheet_name=master.FACILITY_TEMPLATE_SHEET_NAME)


@pytest.mark.parametrize("missing", ["x", "y"])
def test_explicit_axes_missing_one_axis_cannot_use_image_fallback(tmp_path, copied_master, missing):
    path, _config, _entries = generated(tmp_path, copied_master)
    xs, ys = quad.canonical_template_axes_from_workbook(path, sheet_name=master.FACILITY_TEMPLATE_SHEET_NAME,
                                                      canvas_width=1200, canvas_height=2400)
    item = {"template_xlsx": path, "template_sheet_name": master.FACILITY_TEMPLATE_SHEET_NAME,
            "template_axes_x": xs, "template_axes_y": ys}
    del item[f"template_axes_{missing}"]
    with pytest.raises(ValueError, match="canonical template axes require both x and y"):
        quad.resolve_template_axes_from_manifest_or_image(item=item, template_image=np.full((2400, 1200, 3), 255, np.uint8),
                                                         manifest_template_bbox=[0, 0, 1199, 2399])


def test_real_component_threshold_and_score_retained_but_nongrid_cannot_fallback():
    mask = np.zeros((1000, 1000), np.uint8)
    cv2.rectangle(mask, (100, 100), (800, 800), 255, 3)
    chosen, stats, source = quad._edge_locked_choose_component(mask)
    assert source == "large_grid_component"
    assert stats["score"] == stats["w"] * stats["h"] + stats["area"] * 8
    assert np.array_equal(chosen > 0, mask > 0)
    nongrid = np.zeros_like(mask)
    cv2.rectangle(nongrid, (50, 50), (950, 200), 255, -1)
    with pytest.raises(quad.QuadEstimationError, match="edge locked grid component not found"):
        quad._edge_locked_choose_component(nongrid)


def test_actual_nested_failure_classifier_requires_typed_quad_not_message_lookalike():
    # Extract the production nested pure classifier, not a copied implementation or producer double.
    source = Path(__file__).resolve().parents[2] / "src/services/order_service.py"
    tree = ast.parse(source.read_text())
    node, = [node for node in ast.walk(tree) if isinstance(node, ast.FunctionDef) and node.name == "_hakodate_rerun_failure_payload"]
    namespace = {"Any": object}
    exec(compile(ast.Module(body=[node], type_ignores=[]), str(source), "exec"), namespace)
    classify = namespace[node.name]
    failure = quad.QuadEstimationError("synthetic estimator refusal")
    typed = classify(failure)
    assert typed["error"] == "quad_estimation_failed"
    assert typed["recovery_action"] == "review_or_edit_quad"
    assert typed["error_detail"] == str(failure)
    assert "edge locked quad estimation failed" in typed["error_detail"]
    lookalike = classify(ValueError(str(failure)))
    assert lookalike["error"] == "hakodate_live_rerun_failed"
    assert lookalike["recovery_action"] == "review_ocr_error"
