"""Synthetic-native/manual_confirmation_service definitions, not UI/stg proof.

One real production supplies this separate manual-review scenario. The original
blank-native and numeric automatic-transfer failures remain unchanged. Human
confirmation is simulated from the authored workbook, never from OCR values.
"""

from copy import deepcopy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
import pytest

from test_modernization_owned_numeric_fax import numeric_fax
from test_modernization_owned_ocr_workflow import FACILITY, ORDER, WEEK
from src import db
from src.models.order_ocr_evidence_run import OrderOcrEvidenceRun
from src.models.order_sheet_draft import OrderSheetDraft
from src.models.order_workflow_state import OrderWorkflowState
from src.services import canonical_ocr_artifact_service as artifacts
from src.services import master_order_form_template_service as master
from src.services import ocr_evidence_service, order_service
from src.services import order_workflow_v2_service as workflow


def test_synthetic_native_manual_confirmation_save_and_packet_rejection(numeric_fax):
    f, identity = numeric_fax
    workbook = load_workbook(f.root / "synthetic-numeric-fax.xlsx")
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        authored = [str(ws[target["sheet_cell"]].value) for target in f.plan]
    finally:
        workbook.close()
    assert len(authored) == len(f.plan) == 56
    assert authored == [str((7, 0, 6, 9)[index % 4]) for index in range(56)]

    evidence = f.produce()
    payload = deepcopy(evidence["payload_json"])
    records = payload["hakodate_ocr_evidence_records"]
    by_cell = {record["engine_metadata"]["sheet_cell"]: record for record in records}
    assert len(records) == len(by_cell) == 56
    assert set(by_cell) == {target["sheet_cell"] for target in f.plan}
    for target, expected in zip(f.plan, authored, strict=True):
        assert by_cell[target["sheet_cell"]]["raw_text"] == expected
        assert by_cell[target["sheet_cell"]]["normalized_value"] == expected
    artifacts.validate_payload(payload)
    assert len(payload["hakodate_canonical_pipeline"]["artifact_packet"]["outputs"]) == 8
    f.select()
    projected, error = workflow.build_sheet_from_selected_ocr(ORDER)
    assert error is None, (error, projected)
    assert projected["selected_ocr_result_id"] == evidence["id"]
    assert "hakodate_sheet_projection_incomplete" in projected["blockers"]
    sheet = projected["sheet"]
    assert sheet["review_state"] == "blocked"
    assert len(sheet["rows"]) == sheet["physical_menu_row_count"] == 56
    assert (sheet["order_id"], sheet["facility_id"], sheet["week_id"]) == (ORDER, FACILITY, WEEK)
    assert sheet["physical_menu_row_ids"] == identity["physical_menu_row_ids"]
    identity_columns = [sheet["fields"].index(field) for field in ("date_mmdd", "daypart", "menu")]
    public_dayparts = {"朝食": "朝", "昼食": "昼", "夕食": "夕"}
    expected_identity = [
        [date.fromisoformat(day).strftime("%m/%d"), public_dayparts[meal], menu]
        for day, meal, menu in identity["rows"]
    ]
    assert [[row[index] for index in identity_columns] for row in sheet["rows"]] == expected_identity
    quantity = sheet["fields"].index("qty.regular_x")
    deferred = sheet["hakodate_evidence_projection"]["deferred"]
    assert deferred, "This scenario must retain actual low-confidence candidates"
    for candidate in deferred:
        assert candidate["field"] == "qty.regular_x"
        assert candidate["defer_reason"] == "hakodate_ocr_confidence_below_strict_threshold"
        assert candidate["assignment_confidence"] < 0.90
        assert candidate["classification"] != "accepted"
        assert candidate["value"] == authored[candidate["row_index"]]
        assert sheet["rows"][candidate["row_index"]][quantity] == ""
    assert len(sheet["ocr_numeric_cell_items"]) == 56

    captured = order_service._capture_synchronous_cache_write_input(ORDER)
    assert captured is not None
    order_service._validate_selected_hakodate_artifacts(evidence, captured)
    uri = Path(payload["hakodate_canonical_pipeline"]["outputs"]["records"])
    original_bytes = uri.read_bytes()
    for change, blocker in (
        ("missing_packet", "canonical_artifact_packet_required"),
        ("tampered_bytes", "canonical_artifact_bytes_changed"),
        ("deleted_bytes", "canonical_artifact_bytes_unavailable"),
        ("changed_source", "canonical_artifact_source_changed"),
    ):
        changed = deepcopy(evidence)
        pipeline = changed["payload_json"]["hakodate_canonical_pipeline"]
        with db.session_scope() as session:
            state = session.get(OrderWorkflowState, ORDER)
            before = (state.revision, state.evidence_run_id, state.draft_id)
        try:
            if change == "missing_packet":
                pipeline.pop("artifact_packet")
            elif change == "tampered_bytes":
                uri.write_bytes(b"[]")
            elif change == "deleted_bytes":
                uri.unlink()
            else:
                source = pipeline["artifact_packet"]["source_receipt"]
                digest = source["configured_input_digest"]
                source["configured_input_digest"] = ("1" if digest[0] != "1" else "0") + digest[1:]
            assessment = ocr_evidence_service.classify_evidence_payload(changed["payload_json"])
            assert assessment["persistable"] is False, change
            assert blocker in assessment["message"], (change, assessment)
            with db.session_scope() as session:
                session.get(OrderOcrEvidenceRun, evidence["id"]).payload_json = deepcopy(changed["payload_json"])
            with db.session_scope() as session:
                assert ocr_evidence_service.evidence_run_is_selectable(
                    session.get(OrderOcrEvidenceRun, evidence["id"])
                ) is False, change
            with pytest.raises(ValueError, match=f"^{blocker}$"):
                order_service._validate_selected_hakodate_artifacts(changed, captured)
            rejected, error = workflow.build_sheet_from_selected_ocr(ORDER)
            assert rejected is None, (change, rejected)
            assert error == blocker, (change, error)
            with db.session_scope() as session:
                state = session.get(OrderWorkflowState, ORDER)
                assert (state.revision, state.evidence_run_id, state.draft_id) == before
        finally:
            uri.write_bytes(original_bytes)
            with db.session_scope() as session:
                session.get(OrderOcrEvidenceRun, evidence["id"]).payload_json = deepcopy(payload)
        artifacts.validate_payload(payload)
        with db.session_scope() as session:
            assert ocr_evidence_service.evidence_run_is_selectable(
                session.get(OrderOcrEvidenceRun, evidence["id"])
            ) is True
        order_service._validate_selected_hakodate_artifacts(evidence, captured)

    # Explicit simulated human entry changes quantities only, not OCR provenance.
    confirmed = deepcopy(sheet)
    for row, expected in zip(confirmed["rows"], authored, strict=True):
        row[quantity] = expected
    assert confirmed["hakodate_evidence_projection"] == sheet["hakodate_evidence_projection"]
    assert confirmed["ocr_numeric_cell_items"] == sheet["ocr_numeric_cell_items"]
    assert [[row[index] for index in identity_columns] for row in confirmed["rows"]] == expected_identity
    with db.session_scope() as session:
        state = session.get(OrderWorkflowState, ORDER)
        revision, previous_draft = state.revision, state.draft_id
    rejected, error = workflow.save_sheet(
        order_id=ORDER, sheet=confirmed, edited_by="synthetic-manual-confirmation",
        expected_revision=revision - 1,
    )
    assert error == workflow.WORKFLOW_REVISION_CONFLICT, (error, rejected)
    with db.session_scope() as session:
        state = session.get(OrderWorkflowState, ORDER)
        assert (state.revision, state.draft_id, state.evidence_run_id) == (
            revision, previous_draft, evidence["id"],
        )
    submitted = deepcopy(confirmed)
    saved, error = workflow.save_sheet(
        order_id=ORDER, sheet=confirmed, edited_by="synthetic-manual-confirmation",
        expected_revision=revision,
    )
    assert error is None, (error, saved)
    assert confirmed == submitted
    with db.session_scope() as session:
        state = session.get(OrderWorkflowState, ORDER)
        draft = session.get(OrderSheetDraft, state.draft_id)
        assert state.revision == revision + 1
        assert state.evidence_run_id == draft.base_evidence_run_id == evidence["id"]
        assert draft.template_version_id == f.projection_id
        assert draft.draft_sheet_json == submitted
        assert [row[quantity] for row in draft.draft_sheet_json["rows"]] == authored
        assert session.get(OrderOcrEvidenceRun, evidence["id"]).payload_json == payload
