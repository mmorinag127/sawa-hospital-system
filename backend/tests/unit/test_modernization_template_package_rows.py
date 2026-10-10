"""Synthetic real Workbook/package patch/reload component definitions.

No cell writer, style serializer or package writer doubles. ZIP/XML mutations
are explicit component faults, not public operator actions. The tracked
reference is read/copied only; these are not HTTP/auth/native/live proofs.
"""

from copy import copy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import pytest

from src.services import output_builder as builder


NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
SHEET = "xl/worksheets/sheet1.xml"


def parts(content):
    with ZipFile(BytesIO(content)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def package(values):
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        for name, content in values.items():
            archive.writestr(name, content)
    return stream.getvalue()


def original():
    workbook = Workbook()
    workbook.active.title = "Synthetic"
    workbook.active["A1"] = "original"
    workbook.active["C1"] = "keep"
    workbook.active["A3"] = "last"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def patch(workbook, content, tmp_path):
    output = tmp_path / "patched.xlsx"
    builder._patch_template_package_with_workbook_values(
        workbook, template_bytes=content, output_path=output,
    )
    return output, parts(output.read_bytes())


@pytest.mark.parametrize("case", [
    "new-row", "multiple-rows", "middle-row", "empty-sheet-data",
    "self-closed-row", "missing-cell", "existing-cell",
])
def test_real_patch_inserts_ordered_rows_cells_and_dimensions(tmp_path, case):
    source = parts(original())
    root = ET.fromstring(source[SHEET])
    data = root.find(NS + "sheetData")
    targets = {"A5": 7}
    if case == "multiple-rows":
        targets = {"A9": 9, "A5": 7, "B7": 0}
    elif case == "middle-row":
        targets = {"A2": 7}
    elif case == "empty-sheet-data":
        data.clear()
        targets = {"A2": 7, "C4": 0}
    elif case == "self-closed-row":
        data.find(NS + "row").clear()
        data[0].set("r", "1")
        data[0].set("ht", "31")
        data[0].set("customHeight", "1")
        targets = {"B1": 7}
    elif case == "missing-cell":
        targets = {"B1": 7}
    elif case == "existing-cell":
        targets = {"A1": 7}
    source[SHEET] = ET.tostring(root)
    content = package(source)
    workbook = load_workbook(BytesIO(content))
    try:
        for coordinate, value in targets.items():
            workbook.active[coordinate] = value
        output, result = patch(workbook, content, tmp_path)
        with ZipFile(output) as archive:
            assert archive.testzip() is None
        reloaded = load_workbook(output)
        try:
            assert {coordinate: reloaded.active[coordinate].value for coordinate in targets} == targets
            if case != "empty-sheet-data":
                assert reloaded.active["A3"].value == "last"
            if case == "self-closed-row":
                assert reloaded.active.row_dimensions[1].height == 31
        finally:
            reloaded.close()
        tree = ET.fromstring(result[SHEET])
        rows = list(tree.find(NS + "sheetData"))
        indices = [int(row.attrib["r"]) for row in rows]
        assert indices == sorted(set(indices))
        for row in rows:
            cells = list(row.findall(NS + "c"))
            columns = [builder.column_index_from_string(
                "".join(char for char in cell.attrib["r"] if char.isalpha())) for cell in cells]
            assert columns == sorted(set(columns))
        # Preserve original A1:C3, union all generated cells (empty data: A2/C4).
        expected_dimension = {"new-row": "A1:C5", "multiple-rows": "A1:C9",
                              "empty-sheet-data": "A1:C4"}.get(case, "A1:C3")
        assert tree.find(NS + "dimension").attrib["ref"] == expected_dimension
        assert parts(content) == source
    finally:
        workbook.close()


@pytest.mark.parametrize("value,kind,number_format", [
    (7, "n", "General"), (0, "n", "General"), (0.0, "n", "General"),
    (False, "b", "General"), (True, "b", "General"),
    ("#DIV/0!", "e", "General"), ("=SUM(A1:A2)", "f", "General"),
    ("authored text", "s", "General"),
    (datetime(2026, 10, 4, 12, 30), "d", "yyyy-mm-dd hh:mm"),
])
def test_real_cell_writer_types_dates_zero_and_appended_styles(tmp_path, value, kind, number_format):
    content = original()
    workbook = load_workbook(BytesIO(content))
    try:
        cell = workbook.active["D5"]
        cell.value = value
        cell.font = Font(name="Synthetic appended font", bold=True, color="FF123456")
        cell.number_format = number_format
        workbook.active.row_dimensions[5].height = 33
        workbook.active.row_dimensions[5].hidden = True
        output, result = patch(workbook, content, tmp_path)
        reloaded = load_workbook(output, data_only=False)
        try:
            actual = reloaded.active["D5"]
            assert actual.value == value
            assert actual.data_type == kind
            assert copy(actual.font) == copy(cell.font)
            assert actual.number_format == number_format
            assert reloaded.active.row_dimensions[5].height == 33
            assert reloaded.active.row_dimensions[5].hidden is True
        finally:
            reloaded.close()
        old = ET.fromstring(parts(content)["xl/styles.xml"])
        new = ET.fromstring(result["xl/styles.xml"])
        for group in old:
            if group.tag in {NS + name for name in (
                "numFmts", "fonts", "fills", "borders", "cellStyleXfs", "cellXfs", "cellStyles", "dxfs",
            )}:
                assert [ET.tostring(node) for node in new.find(group.tag)][:len(group)] == [
                    ET.tostring(node) for node in group]
        assert len(new.find(NS + "fonts")) > len(old.find(NS + "fonts"))
    finally:
        workbook.close()


def test_raw_unknown_xml_graphic_parts_rename_and_existing_calc_chain_cleanup(tmp_path):
    content = original()
    workbook = load_workbook(BytesIO(content))
    source = parts(content)
    workbook.active.title = "Renamed & authored"
    workbook.active["B2"] = "=A1"
    source["xl/workbook.xml"] = source["xl/workbook.xml"].replace(
        b"<sheets>", b"<!--raw-workbook-marker--><sheets>")
    source[SHEET] = source[SHEET].replace(b"</worksheet>", b"<!--raw-sheet-marker--></worksheet>")
    source[SHEET] = source[SHEET].replace(b'<row r="1"', b'<row customFormat="0" r="1"')
    source[SHEET] = source[SHEET].replace(
        b'<c r="A1"', b'<c xmlns:opaque="urn:synthetic" opaque:marker="preserved" r="A1"')
    extension = b'<extLst><ext uri="urn:synthetic"><opaque xmlns="urn:synthetic"/></ext></extLst>'
    source[SHEET] = source[SHEET].replace(b"</c>", extension + b"</c>", 1)
    source["xl/drawings/owned-unknown.xml"] = b'<drawing xmlns="urn:synthetic"><opaque/></drawing>'
    source["xl/worksheets/_rels/sheet1.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="synthetic" Type="urn:synthetic" Target="../drawings/owned-unknown.xml"/>'
        b'</Relationships>')
    source["custom/unknown.bin"] = b"\x00opaque\xff"
    source["xl/calcChain.xml"] = (
        b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="B2" i="1"/></calcChain>')
    rels = ET.fromstring(source["xl/_rels/workbook.xml.rels"])
    ET.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {
        "Id": "calc-test", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain",
        "Target": "calcChain.xml",
    })
    source["xl/_rels/workbook.xml.rels"] = ET.tostring(rels)
    types = ET.fromstring(source["[Content_Types].xml"])
    ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
        "PartName": "/xl/calcChain.xml", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml",
    })
    source["[Content_Types].xml"] = ET.tostring(types)
    try:
        output, result = patch(workbook, package(source), tmp_path)
        expected = dict(source)
        builder._drop_calc_chain(expected)
        assert result.keys() == expected.keys()
        for name in source.keys() - {SHEET, "xl/workbook.xml", "xl/styles.xml", "xl/calcChain.xml"}:
            assert result[name] == expected[name]
        assert "xl/calcChain.xml" not in result
        assert all(not str(node.attrib.get("Target", "")).endswith("calcChain.xml")
                   for node in ET.fromstring(result["xl/_rels/workbook.xml.rels"]))
        assert all(node.attrib.get("PartName") != "/xl/calcChain.xml"
                   for node in ET.fromstring(result["[Content_Types].xml"]))
        assert b"<!--raw-sheet-marker-->" in result[SHEET]
        assert b'opaque:marker="preserved"' in result[SHEET]
        assert extension in result[SHEET]
        assert b'<row customFormat="0" r="1"' in result[SHEET]
        before_xml, after_xml = source["xl/workbook.xml"], result["xl/workbook.xml"]
        before_tree, after_tree = ET.fromstring(before_xml), ET.fromstring(after_xml)
        before_sheet = before_tree.find(NS + "sheets")[0]
        after_sheet = after_tree.find(NS + "sheets")[0]
        assert before_sheet.attrib["name"] == "Synthetic"
        assert after_sheet.attrib["name"] == "Renamed & authored"
        assert {key: value for key, value in after_sheet.attrib.items() if key != "name"} == {
            key: value for key, value in before_sheet.attrib.items() if key != "name"}
        assert list(ET.iterparse(BytesIO(after_xml), events=("start-ns",))) == list(
            ET.iterparse(BytesIO(before_xml), events=("start-ns",)))
        spans = []
        for content in (before_xml, after_xml):
            root = builder._template_xml_spans(content)
            sheets = next(node for node in root["children"] if node["tag"] == NS + "sheets")
            spans.append(next(node for node in sheets["children"] if node["tag"] == NS + "sheet"))
        before_span, after_span = spans
        assert after_xml[:after_span["start"]] == before_xml[:before_span["start"]]
        assert after_xml[after_span["open_end"]:] == before_xml[before_span["open_end"]:]
        reloaded = load_workbook(output)
        try:
            assert reloaded.sheetnames == ["Renamed & authored"]
            assert reloaded.active["B2"].value == "=A1"
        finally:
            reloaded.close()
    finally:
        workbook.close()


@pytest.mark.parametrize("fault,error", [
    ("row-order", "template_row_order_invalid"),
    ("cell-order", "template_cell_order_invalid"),
    ("cell-foreign-row", "template_cell_order_invalid"),
    ("missing-sheet-data", "template_sheet_data_required"),
    ("missing-styles", "template_styles_required"),
    ("style-count", "template_style_count_invalid:fonts"),
    ("style-prefix", "template_style_prefix_changed:fonts"),
    ("style-reference", "template_style_reference_invalid"),
])
def test_invalid_structure_explicitly_stops_before_any_output(tmp_path, fault, error):
    content = original()
    workbook = load_workbook(BytesIO(content))
    source = parts(content)
    root = ET.fromstring(source[SHEET])
    data = root.find(NS + "sheetData")
    if fault == "row-order":
        data.insert(0, data[-1])
    elif fault == "cell-order":
        data[0].insert(0, data[0][-1])
    elif fault == "cell-foreign-row":
        data[0][0].set("r", "A2")
    elif fault == "missing-sheet-data":
        root.remove(data)
    elif fault == "missing-styles":
        del source["xl/styles.xml"]
    elif fault in {"style-count", "style-prefix"}:
        styles = ET.fromstring(source["xl/styles.xml"])
        fonts = styles.find(NS + "fonts")
        if fault == "style-count":
            fonts.set("count", "999")
        else:
            fonts[0].find(NS + "name").set("val", "Invalid original prefix")
        source["xl/styles.xml"] = ET.tostring(styles)
    else:
        workbook.active["A1"]._style.fontId = 999
    source[SHEET] = ET.tostring(root)
    before = {path: path.read_bytes() for path in tmp_path.rglob("*") if path.is_file()}
    try:
        with pytest.raises(ValueError, match="^" + error + "$"):
            patch(workbook, package(source), tmp_path)
        assert {path: path.read_bytes() for path in tmp_path.rglob("*") if path.is_file()} == before
        assert not (tmp_path / "patched.xlsx").exists()
    finally:
        workbook.close()


def test_tracked_reference_readonly_copy_preserves_all_unrelated_parts(tmp_path):
    reference = Path(builder.DAILY_DELIVERY_REFERENCE_TEMPLATE)
    content = reference.read_bytes()
    digest = sha256(content).hexdigest()
    copied = tmp_path / "reference-copy.xlsx"
    copied.write_bytes(content)
    workbook = load_workbook(copied)
    try:
        sheet = workbook.worksheets[0]
        coordinate = f"A{sheet.max_row + 2}"
        sheet[coordinate] = 7
        output, result = patch(workbook, content, tmp_path)
        reloaded = load_workbook(output)
        try:
            assert reloaded.worksheets[0][coordinate].value == 7
        finally:
            reloaded.close()
        source = parts(content)
        changed_sheets = set(builder._worksheet_paths_by_name(source).values())
        exempt = changed_sheets | {"xl/styles.xml", "xl/calcChain.xml",
                                  "xl/_rels/workbook.xml.rels", "[Content_Types].xml"}
        for name in source.keys() - exempt:
            assert result[name] == source[name]
        assert copied.read_bytes() == content
        assert sha256(reference.read_bytes()).hexdigest() == digest
    finally:
        workbook.close()
