from types import SimpleNamespace
from unittest.mock import patch

import pytest

from src.services import output_builder
from src.services.order_workflow_v2_service import _canonical_payload_digest


def _save_owned_receipt(monkeypatch, tmp_path, *, quantity=0, invoice_template=None):
    from copy import deepcopy
    from src.db import session_scope
    from src.models.order_workflow_state import OrderWorkflowState
    from src.models.order_sheet_draft import OrderSheetDraft
    from src.services import order_output_artifact_service as artifacts
    from src.services import order_service
    from tests.helpers.newnarrowbaggingfixture import accepted_bagging_order

    order = accepted_bagging_order(monkeypatch, tmp_path, [{
        "line_id": "acceptance-line", "date": "2026-10-07", "daypart": "昼",
        "menu_name": "own", "diet_type": "regular", "area_id": "2F",
        "bag_type": "standard", "quantity_original": quantity,
    }], facility_id="fixture", facility_name="Synthetic acceptance fixture",
        invoice_template=invoice_template)
    # The helper registers its authored original, selects real owned evidence,
    # formally saves the sheet and runs the actual first bagging producer.
    with session_scope() as session:
        row = session.get(OrderWorkflowState, order["id"])
        saved = session.get(OrderSheetDraft, row.draft_id)
        payload = artifacts.load_bagging_result_payload(session,
            row.secondary_actions_json["workflow_v2"]["bagging_result_id"], expected_order_id=order["id"])
        assert payload is not None
        assert payload["source_ocr_result_id"] == row.evidence_run_id == saved.base_evidence_run_id
        state = SimpleNamespace(state=row.state, secondary_actions_json=deepcopy(row.secondary_actions_json))
        draft = SimpleNamespace(id=saved.id, template_version_id=saved.template_version_id)
    return order_service.get_order_for_output_read(order["id"]), state, draft, payload




def _read_component_receipt(order, state, draft):
    from src.db import session_scope
    from src.db import get_engine
    from tests.unit.test_disabled_facility_columns import _snapshot
    before = _snapshot(get_engine())
    try:
        with session_scope() as session:
            return output_builder._workflow_v2_materialization_candidate(session, order, state, draft)
    finally:
        assert _snapshot(get_engine()) == before




def test_first_bagging_producer_then_consumer_uses_same_captured_output(isolated_formal_schema, monkeypatch, tmp_path):
    from copy import deepcopy
    from src.db import session_scope
    from src.models.order_output_artifact import OrderBaggingResult
    from src.services import order_workflow_v2_service as workflow

    producer = workflow._build_bagging_result_payload
    calls = []

    def observed_producer(**inputs):
        with session_scope() as session:
            assert session.query(OrderBaggingResult).filter(
                OrderBaggingResult.order_id == inputs["order"].id).count() == 0
        configured = deepcopy(inputs["materialization_candidate"]["configured_input"])
        with patch.object(output_builder, "_accepted_output_candidate", side_effect=AssertionError("producer prior artifact read")), patch("src.services.config_service.get_facility_config", side_effect=AssertionError("producer late facility read")), patch("src.services.config_service.load_ingest_policy", side_effect=AssertionError("producer late policy read")):
            result = producer(**inputs)
        assert inputs["materialization_candidate"]["configured_input"] == configured
        assert result["materialization_candidate"]["lines"][0]["quantity_original"] == 0
        calls.append(inputs["order"].id)
        return result

    monkeypatch.setattr(workflow, "_build_bagging_result_payload", observed_producer)
    owned_order, state, saved_draft, artifact = _save_owned_receipt(monkeypatch, tmp_path)
    assert calls == [owned_order["id"]]
    candidate = deepcopy(artifact["materialization_candidate"])
    configured = deepcopy(candidate["configured_input"])
    order_id = owned_order["id"]
    assert artifact["materialization_digest"] == _canonical_payload_digest(candidate)
    assert candidate["configured_input"] == configured
    assert candidate["lines"][0]["quantity_original"] == 0
    accepted = _read_component_receipt(owned_order, state, saved_draft)
    configured["facility"]["bagging_exceptions"] = [{"type": "quantity_multiplier", "multiplier": 99}]
    expected_bags = output_builder.build_bag_rows_for_outputs(owned_order, accepted_candidate=accepted)
    with patch("src.services.config_service.get_facility_config", side_effect=AssertionError("changed master read")):
        assert output_builder.build_order_lines_for_outputs(owned_order, accepted_candidate=accepted) == candidate["output_lines"]
        with patch.object(output_builder, "_collect_current_menu_items_for_week", side_effect=AssertionError("late monthly read")), patch.object(output_builder, "_resolve_current_menu_defaults", side_effect=AssertionError("late master read")):
            assert output_builder.build_bag_rows_for_outputs(owned_order, accepted_candidate=accepted) == expected_bags
    foreign_lines = [{**candidate["output_lines"][0], "quantity_original": 99}]
    with pytest.raises(ValueError, match="accepted output lines mismatch"):
        output_builder.build_bag_rows_for_outputs({"id": order_id}, order_lines=foreign_lines, accepted_candidate=accepted)


def test_output_consumes_exact_accepted_candidate_without_config_recapture(isolated_formal_schema, monkeypatch, tmp_path):
    order, workflow, draft, artifact = _save_owned_receipt(monkeypatch, tmp_path)
    candidate = artifact["materialization_candidate"]
    with patch("src.services.config_service.capture_authored_processing_config", side_effect=AssertionError("late config read")):
        result = _read_component_receipt(order, workflow, draft)
    assert result == candidate
    assert result is not candidate
    assert result["lines"][0]["quantity_original"] == 0
    assert result["lines"][0]["quantity_corrected"] == candidate["lines"][0]["quantity_corrected"]


@pytest.mark.parametrize("change", ["missing", "wrong_sheet", "wrong_digest", "missing_policy", "missing_template"])
def test_output_missing_or_inconsistent_acceptance_receipt_stops(isolated_formal_schema, monkeypatch, tmp_path, change):
    order, workflow, draft, artifact = _save_owned_receipt(monkeypatch, tmp_path)
    candidate = artifact["materialization_candidate"]
    artifact_id = artifact["bagging_result_id"]
    if change == "missing":
        artifact = None
    elif change == "wrong_sheet":
        artifact["source_saved_sheet_id"] = "other"
    elif change == "wrong_digest":
        artifact["materialization_digest"] = "other"
    elif change == "missing_template":
        del artifact["template_version_id"]
    else:
        del candidate["configured_input"]["policy"]
        artifact["materialization_digest"] = _canonical_payload_digest(candidate)
    from src.db import session_scope
    from src.models.order_output_artifact import OrderBaggingResult
    from src.services import order_output_artifact_service as artifacts
    with session_scope() as session:
        row = session.get(OrderBaggingResult, artifact_id)
        if artifact is None:
            session.delete(row)
        else:
            row.payload_json = artifact
            row.payload_digest = artifacts.payload_digest(artifact)
    with pytest.raises(ValueError):
        _read_component_receipt(order, workflow, draft)


def test_actual_delivery_authoritative_empty_does_not_read_current(monkeypatch):
    monkeypatch.setattr(output_builder, "_build_ocr_menu_meta", lambda *args, **kwargs: pytest.fail("empty receipt replaced by current OCR"))
    monkeypatch.setattr(output_builder, "_resolve_current_menu_defaults", lambda *args: pytest.fail("current menu defaults"))
    rows = output_builder._build_delivery_rows({"id": "own", "facility": "fixture", "lines": []}, {"columns": []}, {"zero_as_empty": False}, {}, {"entries": [], "by_menu": {}, "entry_defaults": {}})
    assert rows == []


@pytest.mark.parametrize("area", ["2F", None, "X"])
@pytest.mark.parametrize("quantity", [7, 0, None])
def test_actual_nonempty_delivery_columns_use_only_authored_mapping(area, quantity, tmp_path):
    from tests.helpers.authored_invoice import declare_invoice
    column = {"source": "quantity", "name": "肉禁3F misleading display", "diet_type": "regular", "area_id": area}
    line = {"date": "2026-10-07", "daypart": "lunch", "menu_name": "own", "menu_category": "主菜", "diet_type": "regular", "area_id": area, "quantity_original": quantity}
    configured = {"areas": [{"area_id": area, "name": "explicit authored area"}]}
    rows = output_builder._build_delivery_rows({"facility": "fixture", "lines": [line]}, {"columns": [column]}, {"zero_as_empty": False}, configured, {"entries": [], "entry_defaults": {}})
    if quantity is None:
        assert rows == []
    else:
        assert len(rows) == 1
        assert rows[0][column["name"]] == quantity
    signature = output_builder._invoice_quantity_signature(column, {})
    assert signature == ("regular", None if area is None else str(area))
    declared = declare_invoice(tmp_path, {"columns": [column]})
    invoice = output_builder._resolve_captured_invoice_template({"invoice_template": declared})
    quantity_columns = [col for col in invoice["columns"] if col.get("source") == "quantity"]
    assert len(quantity_columns) == 1
    assert quantity_columns[0]["area_id"] == (None if area is None else str(area))


@pytest.mark.parametrize("change", ["missing_diet", "null_diet", "invalid_area", "duplicate", "duplicate_name"])
def test_actual_delivery_rejects_unresolved_or_ambiguous_nonempty_mapping(change):
    from copy import deepcopy
    columns = [{"source": "quantity", "name": "常食2F", "diet_type": "regular", "area_id": "2F"}]
    if change == "missing_diet":
        del columns[0]["diet_type"]
    elif change == "null_diet":
        columns[0]["diet_type"] = None
    elif change == "invalid_area":
        columns[0]["area_id"] = []
    else:
        columns.append(deepcopy(columns[0]))
        if change == "duplicate":
            columns[1]["name"] = "another name"
        else:
            columns[1]["area_id"] = "3F"
    line = {"date": "2026-10-07", "menu_name": "own", "diet_type": "regular", "area_id": "2F", "quantity": 7}
    with pytest.raises(ValueError, match="delivery_quantity"):
        output_builder._build_delivery_rows({"lines": [line]}, {"columns": columns}, {"zero_as_empty": False}, {}, {"entries": [], "entry_defaults": {}})
    if not change.startswith("duplicate"):
        with pytest.raises(ValueError, match="delivery_quantity"):
            output_builder._invoice_quantity_signature(columns[0], {})
        with pytest.raises(ValueError, match="delivery_quantity"):
            output_builder._fax_quantity_signature({**columns[0], "role": "quantity"}, {})


@pytest.mark.parametrize("area_definition", [{}, {"area_id": None}, {"area_id": ""}])
def test_authored_optional_area_aggregates_without_display_name_inference(area_definition):
    column = {"source": "quantity", "name": "軟菜3F misleading", "diet_type": "soft", **area_definition}
    lines = [{"date": "2026-10-07", "menu_name": "own", "diet_type": "soft", "area_id": area, "quantity": qty} for area, qty in [("2F", 3), ("3F", 5), (None, 2), ("X", 4)]]
    rows = output_builder._build_delivery_rows({"lines": lines}, {"columns": [column]}, {"zero_as_empty": False}, {}, {"entries": [], "entry_defaults": {}})
    assert len(rows) == 1
    assert rows[0][column["name"]] == 14
    assert output_builder._invoice_quantity_signature(column, {}) == ("soft", None)


def test_generated_invoice_aggregate_and_explicit_x_preserve_both_quantities(monkeypatch, tmp_path):
    from openpyxl import load_workbook
    from tests.helpers.authored_invoice import declare_invoice
    facility = {"invoice_template": {"columns": [{"name": "soft total", "header": "soft total", "source": "quantity", "diet_type": "soft"}, {"name": "soft X", "header": "soft X", "source": "quantity", "diet_type": "soft", "area_id": "X"}]}}
    facility["invoice_template"] = declare_invoice(tmp_path, facility["invoice_template"])
    invoice = output_builder._resolve_captured_invoice_template(facility)
    columns = [col for col in invoice["columns"] if col.get("source") == "quantity"]
    assert len(columns) == 2 and columns[0]["name"] != columns[1]["name"]
    lines = [{"date": "2026-10-07", "menu_name": "own", "diet_type": "soft", "area_id": area, "quantity": qty} for area, qty in [("2F", 7), ("X", 3)]]
    rows = output_builder._build_delivery_rows({"lines": lines}, invoice, {"zero_as_empty": False}, facility, {"entries": [], "entry_defaults": {}})
    assert rows[0][columns[0]["name"]] == 10
    assert rows[0][columns[1]["name"]] == 3
    destination = tmp_path / "delivery.xlsx"
    output_builder._write_delivery_note(destination, rows, columns, invoice["template_uri"], False,
                                        invoice["sheet_name"])
    workbook = load_workbook(destination)
    try:
        sheet = workbook[invoice["sheet_name"]]
        assert [sheet.cell(2, column).value for column in (1, 2)] == [10, 3]
        assert [sheet.cell(1, column).value for column in (1, 2)] == [col["header"] for col in columns]
        assert all(sheet.cell(2, column).border.left.style == "thin" for column in (1, 2))
    finally:
        workbook.close()


def test_requested_delivery_template_failure_never_exports_alternative(monkeypatch, tmp_path):
    def unavailable(uri):
        raise OSError("requested template unavailable")
    monkeypatch.setattr(output_builder, "load_bytes_from_uri", unavailable)
    monkeypatch.setattr(output_builder.pd.DataFrame, "to_excel", lambda *args, **kwargs: pytest.fail("alternative delivery write"))
    with pytest.raises(ValueError, match="delivery_template_load_failed"):
        output_builder._write_delivery_note(tmp_path / "delivery.xlsx", [{"own": 0}], [{"name": "own"}], "gs://own/template.xlsx", False)


@pytest.mark.parametrize("columns", [[], [{"source": "menu", "name": "menu"}]])
@pytest.mark.parametrize("quantity", [7, 0])
def test_nonnull_delivery_quantity_without_mapping_stops(columns, quantity):
    with pytest.raises(ValueError, match="mapping_unresolved"):
        output_builder._build_delivery_rows({"lines": [{"menu_name": "own", "diet_type": "regular", "quantity": quantity}]}, {"columns": columns}, {"zero_as_empty": False}, {}, {"entries": [], "entry_defaults": {}})


def test_existing_no_fried_forbidden_other_business_aggregate_is_retained():
    column = {"source": "quantity", "name": "authored combined", "diet_type": "forbidden_other"}
    lines = [{"menu_name": "own", "diet_type": diet, "area_id": "X", "quantity": qty} for diet, qty in [("no_fried", 7), ("forbidden_other", 3)]]
    rows = output_builder._build_delivery_rows({"lines": lines}, {"columns": [column]}, {"zero_as_empty": False}, {}, {"entries": [], "entry_defaults": {}})
    assert rows[0]["authored combined"] == 10


@pytest.mark.parametrize("consumer", ["bundle", "preview", "delivery_preview", "outputs"])
def test_actual_delivery_consumers_empty_quantity_mapping_stop_before_output_write(isolated_formal_schema, monkeypatch, tmp_path, consumer):
    from tests.helpers.newnarrowbaggingfixture import accepted_bagging_order
    from src.db import get_engine, session_scope
    from tests.unit.test_disabled_facility_columns import _snapshot
    order = accepted_bagging_order(monkeypatch, tmp_path, [{"line_id": "1", "date": "2026-10-07", "daypart": "昼", "menu_name": "own", "diet_type": "regular", "area_id": "2F", "bag_type": "standard", "quantity_original": 7}], invoice_template={"columns": [{"name": "menu_name", "source": "menu_name"}]}, confirm=True)
    monkeypatch.setattr(output_builder, "OUTPUT_DIR", tmp_path / "outputs")
    monkeypatch.setattr(output_builder, "_write_label_csv", lambda *args: pytest.fail("output before quantity gate"))
    output_before = {path: path.read_bytes() for path in (tmp_path / "outputs").rglob("*") if path.is_file()}
    before = _snapshot(get_engine())
    try:
        with session_scope() as session:
            original = output_builder._capture_output_adoption(order["id"], session)
            with pytest.raises(ValueError, match="mapping_unresolved"):
                if consumer == "bundle":
                    ctx = output_builder._prepare_output_context(order["id"], include_bags=False, original_adoption=original)
                    output_builder._build_delivery_rows_for_bundle(
                        ctx["order_for_outputs"], ctx["invoice_template"], ctx["quantity_rules"],
                        ctx["facility_config"], ctx["ocr_menu_meta"], allow_ocr_menu_meta=False)
                elif consumer == "preview":
                    output_builder.build_output_preview(order["id"], "delivery")
                elif consumer == "delivery_preview":
                    output_builder.build_delivery_preview(order["id"], include_diagnostics=False)
                else:
                    output_builder.build_outputs(order["id"], write_session=session, original_adoption=original)
    finally:
        assert _snapshot(get_engine()) == before
    assert {path: path.read_bytes() for path in (tmp_path / "outputs").rglob("*") if path.is_file()} == output_before

def test_actual_delivery_header_and_reference_mapper_never_pick_name_alias():
    from openpyxl import Workbook
    workbook = Workbook()
    ws = workbook.active
    ws.cell(1, 5, "常食2F")
    column = {"source": "quantity", "name": "own soft aggregate", "header": "常食2F", "diet_type": "soft"}
    assert output_builder._build_delivery_column_map(ws, 1, [column]) == {"own soft aggregate": 5}
    # Reference header text itself is deliberately misleading; authored diet
    # and absent-area aggregate semantics still win.
    ws.cell(9, 5, "常食2F")
    mapped = output_builder._daily_delivery_column_meta(ws, authored_columns=[column])
    quantity = [item for item in mapped if item.get("source") == "quantity"]
    assert len(quantity) == 1
    assert output_builder._invoice_quantity_signature(quantity[0], {}) == ("soft", None)
    ws.cell(1, 6, "常食2F extra")
    assert output_builder._build_delivery_column_map(ws, 1, [column]) == {"own soft aggregate": 5}
    ws.cell(1, 6, "常食2F")
    with pytest.raises(ValueError, match="header_mapping_unresolved"):
        output_builder._build_delivery_column_map(ws, 1, [column])
    with pytest.raises(ValueError, match="header_mapping_unresolved"):
        output_builder._daily_delivery_column_meta(ws, authored_columns=[])
    workbook.close()


def test_existing_authored_floor_invoice_and_aggregate_exact_coexistence():
    import json
    from pathlib import Path
    from copy import deepcopy
    from src.services import config_validator
    authored = json.loads((Path(__file__).parents[2] / "src/data/facility_master.template.json").read_text())
    facilities = [facility for facility in authored["facilities"] if facility["facility_id"] == "FAC00003"]
    assert len(facilities) == 1
    template = deepcopy(facilities[0]["invoice_template"])
    captured = {"invoice_template": template, "fax_template": {"columns": []}}
    assert output_builder._resolve_captured_invoice_template(captured) == template
    assert output_builder._resolve_captured_invoice_template(captured) is not template
    errors = []
    config_validator._validate_invoice_columns(template["columns"], "invoice_template.columns", errors)
    assert errors == []
    quantities = [("regular", "2F", 3), ("regular", "3F", 5), ("soft", "2F", 7), ("soft", "3F", 11), ("mixer", "2F", 0), ("mixer", "3F", 13), ("soft", None, 2), ("soft", "X", 4), ("soft", "2F", None)]
    lines = [{"date": "2026-10-07", "daypart": "lunch", "menu_name": "own", "diet_type": diet, "area_id": area, "quantity": qty} for diet, area, qty in quantities]
    rows = output_builder._build_delivery_rows({"lines": lines}, template, {"zero_as_empty": False}, {}, {"entries": [], "entry_defaults": {}})
    assert len(rows) == 1
    assert {name: rows[0][name] for name in ["常食2F", "常食3F", "軟菜", "ミキサー2F", "ミキサー3F"]} == {"常食2F": 3, "常食3F": 5, "軟菜": 24, "ミキサー2F": 0, "ミキサー3F": 13}
    template["columns"].append({"source": "quantity", "name": "soft exact2F", "diet_type": "soft", "area_id": "2F"})
    rows = output_builder._build_delivery_rows({"lines": lines}, template, {"zero_as_empty": False}, {}, {"entries": [], "entry_defaults": {}})
    assert rows[0]["軟菜"] == 24
    assert rows[0]["soft exact2F"] == 7
    del template["columns"][-1]["diet_type"]
    with pytest.raises(ValueError, match="delivery_quantity_diet_required"):
        output_builder._build_delivery_rows({"lines": lines}, template, {"zero_as_empty": False}, {}, {"entries": [], "entry_defaults": {}})


def test_label_missing_amount_does_not_substitute_servings_or_diet_category():
    bag = {"date": "2026-10-07", "menu_name": "own", "diet_type": "regular", "quantity": 7, "menu_qty_per_serving": None, "menu_unit_type": None, "menu_category": None}
    label = output_builder._label_payload_jp(bag)
    assert label["内容量"] == ""
    assert label["内容詳細"] == ""
    assert label["メニュー"] == ""
    legacy = output_builder._label_payload_legacy(bag, {}, "fixture")
    assert legacy["menu_category"] is None
    assert legacy["quantity"] == 7


def test_required_numeric_label_amount_without_unit_stops():
    with pytest.raises(ValueError, match="accepted_label_unit_required"):
        output_builder._label_payload_jp({"quantity": 7, "menu_qty_per_serving": 0, "menu_unit_type": None})


@pytest.mark.parametrize("omit_invoice", [False, True])
def test_actual_delivery_context_duplicate_preserves_producer_first(isolated_formal_schema, monkeypatch, tmp_path, omit_invoice):
    from copy import deepcopy
    from pathlib import Path
    import os
    from src.db import get_engine
    from tests.helpers.newnarrowbaggingfixture import accepted_bagging_order
    from tests.unit.test_disabled_facility_columns import _snapshot

    # Same date/name, but two explicitly authored physical meal slots. The
    # second quantity must not replace the first zero through a name-only key.
    lines = [{"line_id": str(index), "date": "2026-10-07", "daypart": daypart,
              "menu_name": "own", "diet_type": "regular", "area_id": "2F",
              "bag_type": "standard", "quantity_original": quantity}
             for index, (daypart, quantity) in enumerate([("朝", 0), ("昼", 99)])]
    order = accepted_bagging_order(monkeypatch, tmp_path, lines, omit_invoice=omit_invoice,
        invoice_template={"include_menu_name": True, "columns": [
            {"name": "menu_name", "source": "menu_name"},
            {"name": "qty", "source": "quantity", "diet_type": "regular", "area_id": "2F"}]})
    accepted = output_builder._accepted_output_candidate(order)
    captured = deepcopy(accepted)
    source = Path(os.environ["FACILITY_MASTER_PATH"])
    source_bytes = source.read_bytes()
    before = _snapshot(get_engine())
    monkeypatch.setattr(output_builder, "_resolve_current_menu_defaults",
                        lambda *args: pytest.fail("consumer current master"))
    monkeypatch.setattr(output_builder, "_build_ocr_menu_meta",
                        lambda *args, **kwargs: pytest.fail("consumer current OCR"))

    context = output_builder._prepare_output_context(order["id"], include_bags=False,
                                                     require_invoice=not omit_invoice)
    assert output_builder._accepted_output_candidate(order) == captured
    assert [(line["menu_name"], line["quantity_original"]) for line in context["order_lines"]] == [
        ("own", 0), ("own", 99)]
    assert [output_builder._normalize_delivery_daypart(line["daypart"])
            for line in context["order_lines"]] == ["朝", "昼"]
    assert all(type(line["source_row_index"]) is int for line in context["order_lines"])
    # Current physical slots are authoritative; the legacy name-only reader
    # is not used even though both captured menus have the same name/date.
    assert context["ocr_menu_meta"] == {"entries": [], "by_menu": {}, "entry_defaults": {}}
    if omit_invoice:
        assert context["invoice_template"] is None
        with pytest.raises(ValueError, match="accepted_invoice_template_required"):
            output_builder._prepare_output_context(order["id"], include_bags=False)
    else:
        delivery = output_builder._build_delivery_rows_for_bundle(
            context["order_for_outputs"], context["invoice_template"], context["quantity_rules"],
            context["facility_config"], context["ocr_menu_meta"], allow_ocr_menu_meta=False)
        assert [(row["menu_name"], row["qty"]) for row in delivery] == [("own", 99)]
        assert output_builder._build_delivery_rows(
            {"lines": []}, context["invoice_template"], context["quantity_rules"],
            context["facility_config"], {"entries": [], "by_menu": {}, "entry_defaults": {}}) == []
    assert accepted == captured
    assert source.read_bytes() == source_bytes
    assert _snapshot(get_engine()) == before


@pytest.mark.parametrize("template", [None, {}, {"columns": []}, {"columns": "bad"}, {"columns": [{"name": "qty", "source": None}]}, {"columns": [{"name": "qty", "source": "quantity"}], "template_uri": False}, {"columns": [{"name": "qty", "source": "quantity"}], "include_menu_name": 1}])
def test_captured_invoice_invalid_stops_without_fax_replacement(template):
    with pytest.raises(ValueError, match="accepted_invoice_|delivery_quantity_"):
        output_builder._resolve_captured_invoice_template({"invoice_template": template, "fax_template": {"columns": [{"name": "regular2F", "source": "quantity"}]}})


def test_captured_invoice_retains_template_and_slot_metadata():
    from copy import deepcopy
    template = {"template_uri": "file:///owned/invoice.xlsx", "sheet_name": "Template", "columns": [{"name": "soft aggregate", "source": "quantity", "diet_type": "soft", "area_id": None, "column_index": 4}], "slot_metadata": {"raw": [0, None, False]}, "include_menu_name": False}
    original = deepcopy(template)
    resolved = output_builder._resolve_captured_invoice_template({"invoice_template": template})
    assert resolved == original
    resolved["slot_metadata"]["raw"].append(9)
    assert template == original


@pytest.mark.parametrize("sheet_name", ["missing", None])
def test_delivery_template_worksheet_selection_stops_before_file(tmp_path, sheet_name):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.create_sheet("second")
    source = tmp_path / "source.xlsx"
    workbook.save(source)
    workbook.close()
    target = tmp_path / "output.xlsx"
    with pytest.raises(ValueError, match="delivery_worksheet_"):
        output_builder._write_delivery_note(target, [], [{"name": "menu_name", "source": "menu_name"}], source.as_uri(), False, sheet_name)
    assert not target.exists()


@pytest.mark.parametrize("category,daypart,slot,expected", [("主A", "朝", 0, "主A"), ("主Ａ", "夕", 0, "主Ａ"), ("主A", "朝", 1, "主A"), ("主菜", "夕", 0, "主菜"), ("副菜1", "昼", 2, "副菜1")])
def test_delivery_authored_category_precedes_slot_inference(category, daypart, slot, expected):
    assert output_builder._menu_entry_category_label({"category": category, "daypart": daypart, "slot_index": slot}) == expected


@pytest.mark.parametrize("category", [None, "", "generic", "主B", "主Ｂ"])
def test_delivery_missing_or_unresolved_category_stops(category):
    with pytest.raises(ValueError, match="delivery_menu_category_unresolved"):
        output_builder._menu_entry_category_label({"category": category, "daypart": "朝", "slot_index": 0})


@pytest.mark.parametrize("daypart,count", [("朝", 2), ("昼", 3), ("夕", 3), ("昼", 4)])
def test_canonical_generic_side_slots_ignore_diet_variants(daypart, count):
    lines = [{"date": "2026-10-07", "daypart": daypart, "menu_name": f"side{slot}", "menu_category": "副菜", "diet_type": diet, "source_row_index": slot, "_monthly_entry_raw_category": "副菜", "_monthly_entry_override_applied": True, "_monthly_entry_slot_index": slot} for slot in range(count) for diet in ("regular", "soft")]
    entries = [{"menu_date": "2026-10-07", "daypart": daypart, "slot_index": slot, "category": "副菜", "name": f"side{slot}", "diet_type": "regular"} for slot in range(count)]
    result = output_builder._apply_label_meal_slot_categories(output_builder._apply_menu_entry_overrides(lines, entries))
    assert [line["menu_category"] for line in result] == [f"副菜{slot + 1}" for slot in range(count) for _ in range(2)]
    assert [line["diet_type"] for line in result] == [line["diet_type"] for line in lines]


@pytest.mark.parametrize("excluded_first", [False, True])
def test_complete_captured_sides_survive_partial_output(excluded_first):
    entries = [{"menu_date": "2026-10-07", "daypart": "昼", "slot_index": slot, "category": "副菜1" if slot == 1 else "副菜", "name": f"side{slot}", "diet_type": "regular"} for slot in (1, 2)]
    lines = [{"date": "2026-10-07", "daypart": "昼", "menu_name": "side2", "diet_type": "regular"}]
    if excluded_first:
        lines.insert(0, {"date": "2026-10-07", "daypart": "昼", "menu_name": "side1", "diet_type": "regular", "_exclude_from_outputs": True})
    result = output_builder._apply_label_meal_slot_categories(output_builder._apply_menu_entry_overrides(lines, entries))
    assert result[-1]["menu_category"] == "副菜2"


def test_contradictory_captured_slot_stops():
    entry = {"menu_date": "2026-10-07", "daypart": "昼", "slot_index": 1, "category": "副菜", "name": "own", "diet_type": "regular"}
    with pytest.raises(ValueError, match="delivery_monthly_slot_ambiguous"):
        output_builder._captured_monthly_categories([entry, {**entry, "category": "主菜", "diet_type": "soft"}])


@pytest.mark.parametrize("categories,expected", [(["副菜2"], ["副菜2"]), (["副菜1", "副菜3"], ["副菜1", "副菜3"]), (["副菜", "副菜2"], ["副菜1", "副菜2"])])
def test_captured_explicit_side_numbers_preserve_gaps(categories, expected):
    entries = [{"menu_date": "2026-10-07", "daypart": "昼", "slot_index": slot, "category": category, "name": f"side{slot}", "diet_type": diet} for slot, category in enumerate(categories) for diet in ("regular", "soft")]
    resolved = output_builder._captured_monthly_categories(entries)
    assert [resolved[("2026-10-07", "昼", slot)] for slot in range(len(categories))] == expected
    lines = [{"date": "2026-10-07", "daypart": "昼", "menu_name": f"side{slot}", "diet_type": "regular", "source_row_index": slot} for slot in range(len(categories))]
    result = output_builder._apply_label_meal_slot_categories(output_builder._apply_menu_entry_overrides(lines, entries))
    assert [line["menu_category"] for line in result] == expected


@pytest.mark.parametrize("categories", [["副菜", "副菜1"]])
def test_captured_side_number_conflicts_stop_without_shift(categories):
    entries = [{"menu_date": "2026-10-07", "daypart": "昼", "slot_index": slot, "category": category, "name": f"side{slot}", "diet_type": "regular"} for slot, category in enumerate(categories)]
    with pytest.raises(ValueError, match="delivery_monthly_slot_ambiguous"):
        output_builder._captured_monthly_categories(entries)


def test_repeated_explicit_side_numbers_preserved_through_shared_paths_and_writer(tmp_path):
    from openpyxl import load_workbook
    from tests.helpers.authored_invoice import declare_invoice
    entries = [{"menu_date": "2026-10-07", "daypart": "昼", "slot_index": slot, "category": "副菜2", "name": f"own{slot}", "diet_type": "regular"} for slot in (0, 1)]
    lines = [{"date": "2026-10-07", "daypart": "昼", "menu_name": f"own{slot}", "diet_type": "regular", "source_row_index": slot, "quantity_original": quantity} for slot, quantity in [(0, 5), (1, 7)]]
    enriched = output_builder._apply_label_meal_slot_categories(output_builder._apply_menu_entry_overrides(lines, entries))
    assert [line["menu_category"] for line in enriched] == ["副菜2", "副菜2"]
    assert [line["menu_name"] for line in enriched] == ["own0", "own1"]
    columns = [{"name": "menu_category", "source": "menu_category"}, {"name": "menu_name", "source": "menu_name"}, {"name": "qty", "source": "quantity", "diet_type": "regular", "area_id": None}]
    rows = output_builder._build_delivery_rows({"facility": "fixture", "lines": enriched}, {"columns": columns}, {"zero_as_empty": False}, {}, {"entries": [], "entry_defaults": {}})
    assert len(rows) == 2
    target = tmp_path / "repeated-categories.xlsx"
    template = declare_invoice(tmp_path, {"columns": columns})
    output_builder._write_delivery_note(target, rows, columns, template["template_uri"], True,
                                      template["sheet_name"])
    workbook = load_workbook(target, read_only=True)
    try:
        assert list(workbook.active.values) == [("menu_category", "menu_name", "qty"), ("副②", "own0", 5), ("副②", "own1", 7)]
    finally:
        workbook.close()


def test_canonical_explicit_categories_survive_shared_label_path():
    categories = ["主A", "主Ａ", "主菜", "副菜1", "副菜2", "副菜3"]
    lines = [{"date": "2026-10-07", "daypart": "朝", "menu_name": f"own{index}", "menu_category": raw, "_monthly_entry_override_applied": True, "_monthly_entry_raw_category": raw} for index, raw in enumerate(categories)]
    assert [line["menu_category"] for line in output_builder._apply_label_meal_slot_categories(lines)] == categories


@pytest.mark.parametrize("slot", [None, -1, True, "1"])
def test_generic_category_without_canonical_slot_stops(slot):
    with pytest.raises(ValueError, match="delivery_monthly_slot_required"):
        output_builder._apply_label_meal_slot_categories([{"date": "2026-10-07", "daypart": "朝", "menu_category": "副菜", "_monthly_entry_raw_category": "副菜", "_monthly_entry_override_applied": True, "_monthly_entry_slot_index": slot}])


@pytest.mark.parametrize("daypart,count", [("朝", 2), ("昼", 3), ("夕", 3), ("昼", 4)])
def test_actual_monthly_generic_categories_reach_owned_bagging(isolated_formal_schema, monkeypatch, tmp_path, daypart, count):
    from tests.helpers.newnarrowbaggingfixture import accepted_bagging_order
    from src.db import session_scope
    lines = [{"line_id": str(slot), "date": "2026-10-07", "daypart": daypart, "menu_name": f"side{slot}", "diet_type": "regular", "area_id": "2F", "bag_type": "standard", "quantity_original": slot + 1} for slot in range(count)]
    order = accepted_bagging_order(monkeypatch, tmp_path, lines, menu_category="副菜", confirm=True)
    with session_scope() as session:
        original = output_builder._capture_output_adoption(order["id"], session)
        context = output_builder._prepare_output_context(order["id"], original_adoption=original, require_invoice=True)
        assert [line["menu_category"] for line in context["order_lines"]] == [f"副菜{slot + 1}" for slot in range(count)]
        assert [line["quantity_original"] for line in context["order_lines"]] == [slot + 1 for slot in range(count)]


def test_actual_invoice_free_label_consumers_and_delivery_stop(isolated_formal_schema, monkeypatch, tmp_path):
    from datetime import date
    from tests.helpers.newnarrowbaggingfixture import accepted_bagging_order
    order = accepted_bagging_order(monkeypatch, tmp_path, [{"line_id": "1", "date": "2026-10-07", "daypart": "昼", "menu_name": "own", "diet_type": "regular", "area_id": "2F", "bag_type": "standard", "quantity_original": 7}], confirm=True, omit_invoice=True)
    target = date(2026, 10, 7)
    directory = tmp_path / "outputs"
    directory.mkdir()
    monkeypatch.setattr(output_builder, "OUTPUT_DIR", directory)
    for consumer in (output_builder.build_label_workbook_for_order, output_builder.build_label_csv_for_order):
        path, _ = consumer(order["id"])
        assert path.exists() and path.stat().st_size > 0
    for mode in ("labels", "labels_csv"):
        path, manifest = output_builder.build_daily_output_bundle(target, bundle_type=mode)
        assert path.exists() and path.stat().st_size > 0
    assert output_builder._weekly_weight_collect_rows(target)
    with pytest.raises(ValueError, match="accepted_invoice_template_required"):
        output_builder.build_delivery_preview(order["id"])


@pytest.mark.parametrize("per_quantity", [0, 7, None])
def test_final_label_formatter_preserves_frozen_values_without_menu_defaults(monkeypatch, per_quantity):
    from src.services import menu_service
    monkeypatch.setattr(menu_service, "resolve_menu_defaults", lambda *args: pytest.fail("late menu master read"))
    bag = {"menu_name": "煮込みハンバーグ", "date": "2026-10-07", "quantity": 5, "menu_qty_per_serving": per_quantity, "menu_unit_type": "g", "menu_temp_type": None, "menu_category": None, "daypart": None, "bag_max_qty": 9}
    row = output_builder._label_payload_jp(bag)
    assert row["内容詳細"] == ("" if per_quantity is None else f"{per_quantity}g")
    assert row["温・冷"] == ""
    assert row["時間"] == ""
    assert row["メニュー"] == ""
    assert bag["quantity"] == 5
    assert bag["bag_max_qty"] == 9
    assert bag["menu_qty_per_serving"] is per_quantity
