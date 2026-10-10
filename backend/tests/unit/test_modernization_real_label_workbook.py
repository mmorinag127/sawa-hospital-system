"""Real accepted SQL producer to CSV/XLSX, with manually seeded synthetic OCR.

No recognition/validator/business-return doubles. This component proof is not
native automatic adoption, UI/auth/live/formal PG/stg proof and waives no old
failure. The native receipt sibling separately uses explicit operator input.
"""

from copy import deepcopy
import csv
from hashlib import sha256

from openpyxl import load_workbook

from test_modernization_real_bagging_confirmation import (
    actual_bagging, owned_bagging as owned_bagging,
)
from test_modernization_workflow_surface_parity import (
    database_snapshot, surface as surface,
)
from src.services import master_order_form_template_service as master
from src.services import output_builder as output


def test_real_accepted_label_workbook_matches_csv_and_preserves_owned_inputs(
    owned_bagging, monkeypatch,
):
    fixture = owned_bagging
    bagged = actual_bagging(fixture)
    bagged_before = deepcopy(bagged)
    candidate = bagged["bagging_result"]["materialization_candidate"]
    assert [(line["menu_name"], line["quantity_original"]) for line in candidate["lines"]] == [
        ("選択献立0", 7), ("選択献立3", 0),
    ]
    days = {output._ensure_date(line["date"]) for line in candidate["lines"]}
    assert len(days) == 1 and None not in days
    target = days.pop()
    saved = deepcopy(fixture.saved())
    assert [row[4] for row in saved["rows"]] == ["7", "", "", "0"]
    original_hashes = {
        path: sha256(path.read_bytes()).hexdigest()
        for path in (fixture.pdf, fixture.source, fixture.master_path, master.MASTER_TEMPLATE_PATH)
    }
    before = database_snapshot()
    # Only artifact destination is rebound; all accepted reads/rendering are real.
    monkeypatch.setattr(output, "OUTPUT_DIR", fixture.output)
    csv_path, _ = output.build_label_csv_for_order(fixture.order_id, target_date=target)
    xlsx_path, _ = output.build_label_workbook_for_order(fixture.order_id, target_date=target)
    with csv_path.open(encoding="cp932", newline="") as stream:
        reader = csv.DictReader(stream)
        csv_rows = list(reader)
        csv_headers = reader.fieldnames
    assert len(csv_rows) == 2
    assert {row["商品名１"]: (row[""], row["内容量"], row["内容詳細"])
            for row in csv_rows} == {
        "選択献立0": ("7人前", "700g", "100g"),
        "選択献立3": ("0人前", "0g", "100g"),
    }
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        assert len(workbook.worksheets) == 1
        sheet = workbook.worksheets[0]
        assert sheet.max_row == 3
        assert sheet.max_column == 11
        assert [sheet.cell(1, col).value for col in (7, 9, 10, 11)] == [
            "商品名１", "内容量", "内容詳細", None,
        ]
        assert [sheet.cell(2, col).value for col in (7, 9, 10, 11)] == [
            "選択献立0", "700g", "100g", "7人前",
        ]
        assert [sheet.cell(3, col).value for col in (7, 9, 10, 11)] == [
            "選択献立3", "0g", "100g", "0人前",
        ]
        # XLSX serializes an empty header/cell as None; CSV serializes it as
        # an empty string. Normalize only representation for full-row parity.
        headers = ["" if cell.value is None else str(cell.value) for cell in sheet[1]]
        assert headers == csv_headers
        xlsx_rows = [
            dict(zip(headers, ("" if value is None else str(value) for value in row), strict=True))
            for row in sheet.iter_rows(min_row=2, values_only=True)
        ]
        assert xlsx_rows == csv_rows
    finally:
        workbook.close()
    assert database_snapshot() == before
    assert fixture.saved() == saved
    assert bagged == bagged_before
    assert all(sha256(path.read_bytes()).hexdigest() == digest
               for path, digest in original_hashes.items())
