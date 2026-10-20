from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from src.services import output_builder as outputs
from src.services import output_file_validation_service as validation

COLUMNS = [
    {'name': '日付', 'header': '提供日', 'source': 'date'},
    {'name': '区分', 'source': 'daypart'},
    {'name': '献立', 'source': 'menu_display'},
    {'name': '常食2F', 'source': 'quantity'},
]
ROW = {'date': date(2026, 1, 8), 'daypart': '昼食', 'menu_display': 'Menu A', '常食2F': 2}


@pytest.fixture
def invoice_grid(tmp_path):
    path = tmp_path / 'declared-template.xlsx'
    book = Workbook()
    book.active.append(['提供日', '区分', '献立', '常食2F'])
    book.save(path)
    book.close()
    return str(path)


def test_delivery_requires_template_instead_of_plain_export(tmp_path):
    with pytest.raises(ValueError, match='delivery_template_uri_required'):
        outputs._write_delivery_note(tmp_path/'output.xlsx', [ROW], COLUMNS, None, True)
    assert not (tmp_path/'output.xlsx').exists()


def test_template_delivery_resolves_sources_and_preserves_zero(tmp_path, invoice_grid):
    path = tmp_path / 'delivery.xlsx'
    outputs._write_delivery_note(path, [ROW, {**ROW, '常食2F': 0}], COLUMNS, invoice_grid, True)
    book = load_workbook(path, data_only=True)
    assert [cell.value for cell in book.active[2]][1:] == ['昼', 'Menu A', 2]
    assert book.active['A2'].value.date() == ROW['date']
    assert book.active['D3'].value == 0
    book.close()


@pytest.mark.parametrize('source', ['date', 'daypart', 'menu_display'])
def test_missing_required_source_rejected_with_template(tmp_path, invoice_grid, source):
    row = {**ROW, source: None}
    row['日付'] = row['区分'] = row['献立'] = 'not authoritative'
    with pytest.raises(ValueError, match='delivery_required_value_missing'):
        outputs._write_delivery_note(tmp_path/'missing.xlsx', [row], COLUMNS, invoice_grid, True)


def workbook_bytes(values):
    book = Workbook()
    book.active.append(['date', 'menu', 'quantity'])
    book.active.append(values)
    result = BytesIO()
    book.save(result)
    book.close()
    return result.getvalue()


@pytest.fixture
def accepted(tmp_path, monkeypatch):
    identity = {'state': 'confirmed', 'revision': 9, 'confirmed_snapshot_id': 'owned-snapshot'}
    monkeypatch.setattr(validation.order_workflow_v2_service, 'get_workflow_readonly', lambda order: (identity, None))
    content = workbook_bytes(['2026-01-08', 'Menu A', 2])
    path = tmp_path / 'accepted.xlsx'
    path.write_bytes(content)
    monkeypatch.setattr(outputs, 'build_output_preview', lambda *a: {'delivery_note': path})
    return identity, content


def test_exact_business_content_validates_against_accepted_snapshot(accepted):
    identity, content = accepted
    result = validation.validate_received_output('order', 'delivery', content)
    assert result['valid'] and result['revision'] == 9
    assert result['confirmed_snapshot_id'] == identity['confirmed_snapshot_id']


@pytest.mark.parametrize('values', [[None, 'Menu A', 2], ['2026-01-08', None, 2], ['2026-01-08', 'Menu A', 99]])
def test_blank_or_wrong_received_fields_are_rejected(accepted, values):
    with pytest.raises(ValueError, match='output_business_content_mismatch'):
        validation.validate_received_output('order', 'delivery', workbook_bytes(values))


def test_unreadable_or_unconfirmed_output_cannot_validate(accepted):
    identity, content = accepted
    with pytest.raises(ValueError, match='output_file_unreadable'):
        validation.validate_received_output('order', 'delivery', b'not a workbook')
    identity['state'] = 'sheet_review'
    with pytest.raises(ValueError, match='output_confirmed_order_required'):
        validation.validate_received_output('order', 'delivery', content)


def test_order_change_during_validation_is_rejected(accepted, monkeypatch):
    identity, content = accepted
    original = outputs.build_output_preview
    def regenerate(*a):
        identity['revision'] += 1
        return original(*a)
    monkeypatch.setattr(outputs, 'build_output_preview', regenerate)
    with pytest.raises(ValueError, match='output_order_changed'):
        validation.validate_received_output('order', 'delivery', content)


def test_validation_endpoint_requires_registered_operator(isolated_formal_schema, hospital_operator):
    from fastapi.testclient import TestClient
    from src.main import app
    client = TestClient(app)
    payload = {'order_id':'owned', 'kind':'delivery'}
    files = {'artifact_file':('received.xlsx',b'not a workbook','application/octet-stream')}
    assert client.post('/outputs/validate-received',data=payload,files=files).status_code == 401
    response = client.post('/outputs/validate-received',data=payload,files=files,headers=hospital_operator.headers)
    assert response.status_code == 409
    assert response.json()['detail'] == 'output_confirmed_order_required'


from tests.contract.auth_support import hospital_operator  # noqa: E402,F401


def test_slot_body_clears_unmapped_history_but_retains_template_axes_and_styles():
    book = Workbook(); ws = book.active
    ws.merge_cells('A2:C2'); ws['A2']='【納品書】'
    ws.merge_cells('A4:D5'); ws['A4']='試験施設'
    ws['F10']='肉禁'; ws['G10']='魚禁'
    ws['A14']=date(2026,1,8);ws['B14']='昼';ws['C14']='主Ａ'
    ws['D14']=27;ws['F14']='過去の記入例';ws['G14']='過去の別列';ws['I14']='過去の備考'
    before_merge=list(ws.merged_cells.ranges);before_style=ws['G14'].style_id
    outputs._clear_delivery_dynamic_rows(ws,[14],3)
    assert all(ws.cell(14,c).value is None for c in range(3,10))
    assert ws['A14'].value==date(2026,1,8) and ws['B14'].value=='昼'
    assert ws['A2'].value=='【納品書】' and ws['A4'].value=='試験施設'
    assert ws['F10'].value=='肉禁' and ws['G10'].value=='魚禁'
    assert list(ws.merged_cells.ranges)==before_merge and ws['G14'].style_id==before_style
    book.close()


def test_slot_menu_text_fits_original_width_without_changing_borders():
    book=Workbook(); ws=book.active
    ws.column_dimensions['C'].width=8
    before=ws['C14'].border.copy()
    outputs._write_delivery_slot_row(ws,14,{'menu_display':'主Ａ 鶏肉の照り焼き'},[{'name':'献立','source':'menu_display'}],{'献立':3},False)
    assert ws['C14'].value=='主Ａ 鶏肉の照り焼き'
    assert ws['C14'].alignment.shrinkToFit is True
    assert ws.column_dimensions['C'].width==8
    assert ws['C14'].border==before
    book.close()


def test_delivery_output_rejects_unmapped_authored_quantity_source():
    from src.services.output_builder import _resolve_delivery_cell
    import pytest
    with pytest.raises(ValueError,match='delivery_column_source_unmapped'):
        _resolve_delivery_cell({'常食':20},{'source':'quantity','name':'別の数量'})
    assert _resolve_delivery_cell({'常食':0},{'source':'quantity','name':'常食'})==0
    assert _resolve_delivery_cell({'常食':None},{'source':'quantity','name':'常食'}) is None


@pytest.mark.parametrize('zero_as_empty',[True,False])
def test_delivery_rows_explicitly_materialize_declared_blank_columns(zero_as_empty):
    from src.services.output_builder import _build_delivery_rows,_resolve_delivery_cell
    columns=[{'source':'quantity','name':'常食','diet_type':'regular','area_id':'X'},
             {'source':'quantity','name':'小口','diet_type':'small','area_id':'X'}]
    order={'facility':'synthetic','lines':[{'date':'2026-01-08','daypart':'昼','menu_name':'合成献立',
            'quantity_original':5,'diet_type':'regular','area_id':'X'}]}
    result=_build_delivery_rows(order,{'columns':columns},{'zero_as_empty':zero_as_empty},allow_ocr_menu_meta=False)
    assert len(result)==1
    assert result[0]['常食']==5
    assert '小口' in result[0] and result[0]['小口'] is None
    assert _resolve_delivery_cell(result[0],columns[1]) is None
    with pytest.raises(ValueError,match='delivery_column_source_unmapped'):
        _resolve_delivery_cell(result[0],{'source':'quantity','name':'未定義'})
