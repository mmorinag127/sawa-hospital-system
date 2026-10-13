"""Generated presentation metadata; original authored values remain untouched."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from src.services import hakodate_assignment_service as assignment
from src.services import hakodate_step_review_pipeline_service as review
from src.services import master_order_form_template_service as master
from test_modernization_dynamic_form_layout import copied_master, entries_for, facility, WEEK


def test_generated_metadata_preserves_prod_labels_note_and_nfkc(tmp_path, copied_master):
    config = facility()
    column = config["fax_template"]["columns"][3]
    column.update(header="花", header_group="常食", area_id="2F", name="qty.regular_2f")
    config["fax_template"].pop("body_merge_policy")
    entries = entries_for(range(7), (2, 3, 3))
    entries[0]["name"] = "ＡＢＣ ｶﾚｰ"
    original = deepcopy((config, entries))
    output = tmp_path / "generated.xlsx"
    master.build_facility_template_xlsx(
        facility_config=config, output_path=output, master_template_path=copied_master,
        week_value=WEEK, week_menu_entries=entries,
    )
    draft = {"fields": ["date", "daypart", "menu"], "rows": [
        [entry["menu_date"], entry["daypart"], entry["name"]] for entry in entries
    ]}
    original_draft = deepcopy(draft)
    workbook = load_workbook(output)
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        plan = review.canonical_logical_target_plan(
            worksheet=ws, fax_template=config["fax_template"], draft_sheet=draft)
        assert plan[0]["field_label"] == "常食花"
        assert plan[0]["field"] == "qty.regular_2f"
        assert plan[0]["menu_name"] == "ABC カレー"
        assert {r["field"] for r in plan if r["role"] == "note"} == {"note"}
        regions, _ = review._post_menu_target_regions(
            worksheet=ws, fax_template=config["fax_template"], draft_sheet=draft,
            column_edges=list(range(ws.max_column + 1)),
            row_edges=list(range(review.STEP_REVIEW_HEADER_BANDS + len(entries) + 1)),
            menu_body_row_indexes=list(range(len(entries))))
        assert review.canonical_target_semantics(plan) == review.canonical_target_semantics(regions)
        assert (config, entries) == original
        assert draft == original_draft
    finally:
        workbook.close()


@pytest.mark.parametrize("fault,code", [
    ("role", "role_unresolved"),
    ("duplicate", "menu_role_unresolved"),
    ("diet", "quantity_diet_required"),
    ("count", "generated_columns_changed"),
])
def test_generated_metadata_keeps_authored_identity_stops(fault, code):
    template = facility()["fax_template"]
    count = 7
    if fault == "role":
        template["columns"][-1]["role"] = "unknown_authored_role"
    elif fault == "duplicate":
        template["columns"][-1]["index"] = template["columns"][0]["index"]
    elif fault == "diet":
        template["columns"][3].pop("diet_type")
    else:
        count += 1
    before = deepcopy(template)
    with pytest.raises(ValueError, match=code):
        assignment.canonical_generated_column_slots(template, col_count=count)
    assert template == before
