"""Copied-master synthetic PNG/OpenCV component proofs, not native/operator proof.

All FAX pixels originate in the saved/reloaded generated worksheet. Explicit
blank-month and damaged-line variants below are synthetic input stimuli, not
recognizer output, production FAX evidence, or a replacement for native3.
"""

from copy import copy, deepcopy
from types import SimpleNamespace

import cv2
import numpy as np
from openpyxl import load_workbook
import pytest

from test_modernization_dynamic_form_layout import copied_master, facility, WEEK
from test_modernization_serialized_form_geometry import physical_entries
from src.services import hakodate_cell_ocr_batch_service as batch
from src.services import hakodate_physical_menu_row_service as physical
from src.services import hakodate_step_review_pipeline_service as review
from src.services import master_order_form_template_service as master
from src.services import workbook_pdf_renderer as renderer


def _config(kind):
    config = deepcopy(facility())
    columns = config["fax_template"]["columns"][:3]
    specifications = {
        "one": [("regular", "X", "常食", "")],
        "distinct": [("regular", "X", "常食", ""), ("no_meat", "X", "肉禁", "")],
        "areas": [("regular", "X", "X棟", "常食"), ("regular", "Y", "Y棟", "常食")],
        "partial": [("regular", "X", "常食", ""), ("no_meat", "X", "肉禁", ""),
                    ("no_fish", "X", "X棟", "魚禁"), ("no_fish", "Y", "Y棟", "魚禁")],
    }[kind]
    for index, (diet, area, label, group) in enumerate(specifications, start=3):
        columns.append({"index": index, "role": "quantity", "name": f"qty.{diet}_{area.lower()}",
                        "header": label, "diet_type": diet, "area_id": area, "header_group": group})
    config["fax_template"]["columns"] = columns
    config["fax_template"].pop("body_merge_policy")
    return config


@pytest.fixture
def form_factory(tmp_path, copied_master):
    opened = []

    def build(kind="one", count=56):
        config = _config(kind)
        entries = physical_entries(count)
        path = tmp_path / f"{kind}-{count}.xlsx"
        master.build_facility_template_xlsx(
            facility_config=config, master_template_path=copied_master, output_path=path,
            week_value=WEEK, week_menu_entries=entries,
        )
        workbook = load_workbook(path)
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        if kind == "partial":
            # Explicit printable canonical border variant: retain the group
            # ruling G9:H9 but omit its internal vertical leaf separator.
            # Its only intersections are columns 6 and 8, outside menu support.
            for row in (9, 10):
                for column, side in ((7, "right"), (8, "left")):
                    border = copy(ws.cell(row, column).border)
                    setattr(border, side, None)
                    ws.cell(row, column).border = border
        values = [7, 0, None, "?"]
        quantities = []
        for offset in range(count):
            row_values = []
            for column in range(5, ws.max_column + 1):
                value = values[(offset + column - 5) % len(values)]
                ws.cell(11 + offset, column).value = value
                row_values.append(value)
            quantities.append(row_values)
        workbook.save(path)
        workbook.close()
        reloaded = load_workbook(path)
        opened.append(reloaded)
        ws = reloaded[master.FACILITY_TEMPLATE_SHEET_NAME]
        geometry = renderer.worksheet_render_geometry(ws, dpi=144)
        edge_rows = (7, 9, *range(11, count + 12))
        image = renderer.render_worksheet_to_image(ws, dpi=144)
        png = tmp_path / f"{kind}-{count}.png"
        try:
            image.save(png)
            assert image.size == (geometry["image_width"], geometry["image_height"])
            bgr = cv2.cvtColor(np.asarray(image.convert("RGB")), cv2.COLOR_RGB2BGR)
        finally:
            image.close()
        return SimpleNamespace(
            ws=ws, workbook=reloaded, entries=entries, config=config, quantities=quantities,
            geometry=geometry, image=bgr, png=png, path=path,
            xs=[float(geometry["x_positions"][col]) for col in range(1, ws.max_column + 2)],
            ys=[geometry["y_positions"][row] for row in edge_rows],
        )

    yield build
    for workbook in opened:
        workbook.close()


def _observe(form, *, image=None, xs=None, ys=None):
    return review._row_intersection_correct_ys(
        rectified_fax=form.image if image is None else image,
        corrected_xs=form.xs if xs is None else xs,
        template_ys=form.ys if ys is None else ys, worksheet=form.ws,
    )


def _sheet(entries, week):
    ids = physical.physical_row_ids_from_entries([
        {"menu_date": entry["menu_date"], "daypart": entry["daypart"],
         "slot_index": entry["sort_order"] - 1} for entry in entries
    ])
    return {"week_id": week, "physical_menu_row_ids": ids, "physical_menu_row_count": len(ids)}


def _assert_body(form, corrected, evidence):
    count = len(form.entries)
    assert corrected is not None
    assert evidence["used"] is True
    assert evidence["structural_match"]["detected_body_band_count"] == count
    assert len(corrected) == count + 3
    body_y = evidence["structural_match"]["body_start_cluster_index"]
    raw = evidence["raw_fax_y_clusters"]
    observed_start = raw[body_y]["value"]
    body = [cluster["value"] for cluster in evidence["fax_y_clusters"] if cluster["value"] > observed_start]
    assert corrected[2:] == [observed_start, *body]
    assert len(corrected[2:]) == count + 1
    # Independent worksheet pixel positions bound each actually detected line;
    # this does not replace exact observed-coordinate equality above.
    assert corrected[2:] == pytest.approx(form.ys[2:], abs=3)
    assert all(right > left for left, right in zip(corrected, corrected[1:]))
    plan = master.validated_body_row_plan(form.ws)
    assert len(plan) == count
    assert [row["source_index"] for row in plan] == list(range(count))
    assert [(row["date"], row["daypart"], row["name"]) for row in plan] == [
        (entry["menu_date"], entry["daypart"], entry["name"]) for entry in form.entries
    ]
    assert [[form.ws.cell(row, col).value for col in range(5, form.ws.max_column + 1)]
            for row in range(11, count + 11)] == form.quantities
    assert form.png.is_file()


@pytest.mark.parametrize("kind,count,printed_middle", [
    ("one", 56, False), ("one", 57, False), ("one", 58, False),
    ("distinct", 56, False), ("areas", 56, True),
], ids=["qty1-56", "qty1-57", "qty1-58", "qty2-distinct", "same-diet-two-areas"])
def test_actual_printed_header_preserves_logical_axes_and_every_observed_body_row(
    form_factory, kind, count, printed_middle,
):
    form = form_factory(kind, count)
    corrected, evidence = _observe(form)
    _assert_body(form, corrected, evidence)
    header = evidence["structural_match"]["header_boundaries"]
    assert [entry["worksheet_row"] for entry in header] == [7, 9]
    assert header[0]["source"] == "observed_printed_header_intersections"
    if printed_middle:
        assert header[1]["source"] == "observed_printed_header_intersections"
        assert header[1]["y"] in [cluster["value"] for cluster in evidence["raw_fax_y_clusters"]]
    else:
        assert header[1]["source"] == "canonical_relative_position_between_observed_header_boundaries"
        positions = form.geometry["y_positions"]
        ratio = (positions[9] - positions[7]) / (positions[11] - positions[7])
        assert corrected[1] == pytest.approx(corrected[0] + ratio * (corrected[2] - corrected[0]))
        assert header[1]["y"] not in [cluster["value"] for cluster in evidence["raw_fax_y_clusters"]]
    assert batch._validate_menu_day_boundaries(_sheet(form.entries, WEEK), evidence) == list(range(count))


def test_partial_two_vertex_printed_header_survives_generic_body_filter(form_factory):
    form = form_factory("partial")
    corrected, evidence = _observe(form)
    _assert_body(form, corrected, evidence)
    middle = evidence["structural_match"]["header_boundaries"][1]
    assert middle["source"] == "observed_printed_header_intersections"
    cluster = next(cluster for cluster in evidence["raw_fax_y_clusters"] if cluster["value"] == middle["y"])
    assert cluster["column_votes"] == [6, 8]
    assert review._filter_structural_y_clusters([cluster]) == []
    assert corrected[1] == cluster["value"]
    assert batch._validate_menu_day_boundaries(_sheet(form.entries, WEEK), evidence) == list(range(56))


@pytest.mark.parametrize("first,last,week,expected", [
    (0, 32, "2026-09@2026-09-27~2026-09-30", list(range(32))),
    (32, 56, "2026-10@2026-10-01~2026-10-03", list(range(32, 56))),
], ids=["trailing-blank-October", "leading-blank-September"])
def test_month_target_mapping_keeps_all_known_physical_fax_blank_day_bands(
    form_factory, tmp_path, first, last, week, expected,
):
    form = form_factory()
    image = form.image.copy()
    positions = form.geometry
    # Author a blank-month FAX from this exact weekly source: erase only cell
    # interiors outside the selected month, retaining date/meal and all rulings.
    for offset in range(56):
        if first <= offset < last:
            continue
        for column in range(4, form.ws.max_column + 1):
            left, right = positions["x_positions"][column], positions["x_positions"][column + 1]
            top, bottom = positions["y_positions"][11 + offset], positions["y_positions"][12 + offset]
            image[top + 4:bottom - 4, left + 4:right - 4] = 255
    assert cv2.imwrite(str(tmp_path / "synthetic-month-blank.png"), image)
    corrected, evidence = _observe(form, image=image)
    _assert_body(form, corrected, evidence)
    sheet = _sheet(form.entries[first:last], week)
    before = deepcopy(sheet)
    assert batch._validate_menu_day_boundaries(sheet, evidence) == expected
    assert sheet == before
    assert evidence["structural_match"]["body_day_boundary_indexes"] == list(range(0, 57, 8))
    assert corrected[2:] == [cluster["value"] for cluster in evidence["fax_y_clusters"]
                             if cluster["value"] >= corrected[2]]


@pytest.mark.parametrize("change,code", [
    ("schema", "generated_body_schema_missing"),
    ("menu", "generated_body_row_identity_changed"),
    ("date", "generated_body_row_identity_changed"),
])
def test_actual_source_tamper_stops_at_existing_canonical_validator(form_factory, change, code):
    form = form_factory()
    if change == "schema":
        del form.workbook["generated_template_schema"]
    else:
        form.ws.cell(11, 4 if change == "menu" else 1).value = "not-canonical"
    with pytest.raises(master.FacilityTemplateBuildError, match=f"^{code}$"):
        _observe(form)


def test_missing_worksheet_is_not_a_static_row_plan_fallback(form_factory):
    form = form_factory()
    with pytest.raises(ValueError, match="^fax_row_worksheet_required$"):
        review._row_intersection_correct_ys(
            rectified_fax=form.image, corrected_xs=form.xs, template_ys=form.ys, worksheet=None,
        )


@pytest.mark.parametrize("axis", ["x", "y"])
def test_nonmonotone_actual_canonical_axes_stop_without_sorting_or_duplicate_repair(form_factory, axis):
    form = form_factory()
    values = list(form.xs if axis == "x" else form.ys)
    values[1], values[2] = values[2], values[1]
    corrected, evidence = _observe(form, **{axis + "s": values})
    assert corrected is None
    assert evidence == {"used": False, "reason": "fax_row_canonical_axes_mismatch"}


@pytest.mark.parametrize("damage", ["required-outer", "required-group", "ambiguous-top"])
def test_actual_printed_required_line_loss_or_ambiguous_candidate_stops(form_factory, tmp_path, damage):
    form = form_factory("areas" if damage == "required-group" else "one")
    contract = renderer.worksheet_header_boundary_contract(
        form.ws, header_edge_rows=(7, 9), body_start_row=11,
    )
    required_boundary = contract[1] if damage == "required-group" else contract[0]
    assert required_boundary["horizontal_segments"]
    assert required_boundary["column_votes"]
    image = form.image.copy()
    if damage == "ambiguous-top":
        y = round(form.ys[0] + (form.ys[1] - form.ys[0]) * 0.4)
        cv2.line(image, (round(form.xs[0]), y), (round(form.xs[-1]), y), (0, 0, 0), 3)
    else:
        y = form.ys[1] if damage == "required-group" else form.ys[0]
        image[y - 4:y + 5, round(form.xs[0]) - 4:round(form.xs[-1]) + 5] = 255
    assert cv2.imwrite(str(tmp_path / f"synthetic-{damage}.png"), image)
    corrected, evidence = _observe(form, image=image)
    assert corrected is None
    assert evidence["used"] is False
    assert evidence["reason"] == "fax_body_header_boundary_unresolved"
    if damage == "ambiguous-top":
        assert len(evidence["header_candidates"]) == 2
    else:
        assert evidence["header_candidates"] == []
