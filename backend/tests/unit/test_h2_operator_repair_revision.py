from copy import deepcopy
from contextlib import nullcontext
from datetime import date, datetime
from hashlib import sha256
import json
import os
from pathlib import Path
from uuid import uuid4

import pytest
from fastapi import HTTPException

from src.api import orders
from src.db import session_scope
from src.models.order import Order
from src.models.document import OrderDocument
from src.models.uploaded_pdf import UploadedPdf
from src.models.order_ocr_evidence_run import OrderOcrEvidenceRun
from src.models.order_workflow_state import OrderWorkflowState
from src.models.order_ocr_cache import OrderOcrCache
from src.models.facility import Facility
from src.models.menu import MonthlyMenu, MonthlyMenuEntry
from src.services import config_service, order_service
from src.services import canonical_facility_lineage as lineage


@pytest.mark.parametrize("period", [None, "", "2026-13", "2026-10@2026-10-11~2026-10-05"])
def test_original_week_unresolved_stops_sheet_cache_and_weekly_write_before_sources(repair_order, monkeypatch, period):
    order_id, _revision, _configured = repair_order
    with session_scope() as session:
        session.get(Order, order_id).week_code = period
    monkeypatch.setattr(config_service, "capture_authored_processing_config", lambda *a, **kw: pytest.fail("unresolved period must stop before config/menu capture"))
    monkeypatch.setattr(order_service, "_build_position_menu_entries_safe", lambda *a, **kw: pytest.fail("no neighboring month menu lookup"))
    from src.services import order_current_state_service
    monkeypatch.setattr(order_current_state_service, "get_current_state_payload", lambda *a, **kw: pytest.fail("unresolved period must not adopt historical cached state"))
    monkeypatch.setattr(order_current_state_service, "persist_current_state", lambda *a, **kw: pytest.fail("unresolved period must not write state"))
    monkeypatch.setattr(order_current_state_service, "delete_current_state", lambda *a, **kw: pytest.fail("unresolved period must not delete state"))
    assert order_service.get_ocr_sheet(order_id) == (None, "week_unresolved")
    assert order_service.build_recoverable_ocr_sheet_payload(order_id, "week_unresolved") == (None, "week_unresolved")
    assert order_service.force_overwrite_current_sheet_with_weekly_menu(order_id) == (None, "week_unresolved")
    for read in (order_service.get_current_sheet_context, order_service.refresh_current_sheet_context):
        blocked = read(order_id)
        assert blocked["blockers"] == ["week_unresolved"]
        assert blocked["rows"] == []
        assert blocked["resolved_week_id"] is None
        visible = order_service._merge_order_review_summary_with_sheet_context({}, current_sheet_context=blocked)
        assert visible["ocr_apply_blockers"] == ["week_unresolved"]
        assert visible["ocr_confirm_blockers"] == ["week_unresolved"]
        assert visible["ocr_can_apply_draft"] is False
        assert visible["ocr_can_confirm"] is False


@pytest.mark.parametrize("period", ["2026-10@2026-10-01~2026-10-03", "2026-09@2026-09-29~2026-10-01"])
def test_explicit_partial_period_is_not_replaced_by_received_month_or_neighbor(period):
    assert order_service._resolve_sheet_week_id(
        current_week_id=period, received_at=datetime(2026, 11, 30),
        order_lines=[{"date": "2026-11-30"}], ocr_payload={"dates": ["2026-11-30"]},
        facility_id="fixture", week_hints=["2026-11", "2026-12"],
    ) == period
    assert order_service._resolve_current_sheet_context_week_id(
        order_id="explicit", order_payload={"stored_week_value": period, "week": "2026-11"},
        draft_payload={"week_id": "2026-12"}, facility_id="fixture",
    ) == period


def _owned_artifact_fixture(source, value=0):
    from src.services import hakodate_ocr_evidence_service as evidence
    regions = [{**deepcopy(region), "bbox": [0, 0, 10, 10]} for region in source["expected_target_semantics"]]
    targets = evidence.target_cells_from_regions(regions)
    for target in targets:
        target.update(source_receipt=deepcopy(source), value=value)
    records = [{"evidence_id": str(uuid4()), "source_scope": "hakodate_cell_crop_batch", "raw_text": value, "normalized_value": "0" if value == 0 else "", "center": [5, 5], "engine_metadata": {"sheet_cell": region["sheet_cell"], "source_receipt": deepcopy(source), "recognizer_output": deepcopy(region)}} for region in regions]
    return targets, records


def _write_repair_policy(configured):
    config_service.INGEST_POLICY_PATH.write_text(
        json.dumps(configured["policy"], ensure_ascii=False), encoding="utf-8",
    )


def _provider_output_files(workspace, *, records, regions, sheet_values):
    import fitz
    from PIL import Image

    workspace.mkdir(parents=True, exist_ok=True)
    outputs = {}
    for key in ("overlay", "overlay_png", "sheet_review_base", "contact_sheet"):
        path = workspace / f"{key}.png"
        with Image.new("RGB", (32, 32), "white") as image:
            image.save(path)
        outputs[key] = str(path)
    with fitz.open() as document:
        document.new_page(width=32, height=32)
        path = workspace / "output.pdf"
        document.save(path)
        outputs["pdf"] = str(path)
    for key, value in (("records", records), ("ocr_regions", regions), ("sheet_values", sheet_values)):
        path = workspace / f"{key}.json"
        path.write_text(json.dumps(order_service._json_snapshot(value), ensure_ascii=False), encoding="utf-8")
        outputs[key] = str(path)
    return outputs


def _attach_provider_artifact_packet(payload):
    from src.services import canonical_ocr_artifact_service as artifacts

    source = payload["hakodate_repair_source_receipt"]
    pipeline = payload["hakodate_canonical_pipeline"]
    records = payload["hakodate_ocr_evidence"]["records"]
    regions = payload["hakodate_preprocessing"]["target_cell_map"]
    workspace = Path(os.environ["OCR_ARTIFACT_DIR"]).parent / "provider-workspaces" / uuid4().hex
    outputs = _provider_output_files(
        workspace, records=records, regions=regions,
        sheet_values={region["sheet_cell"]: region.get("value") for region in regions},
    )
    packet = artifacts.persist_outputs(
        order_id=source["selected_pdf_receipt"]["order_id"], outputs=outputs,
        source_receipt=source, expected_sheet_cells=pipeline["expected_sheet_cells"],
        producer_version=pipeline["version"], workspace=workspace,
    )
    outputs = {key: receipt["uri"] for key, receipt in packet["outputs"].items()}
    pipeline.update(requested_order_id=packet["order_id"], outputs=outputs, artifact_packet=packet)
    payload["hakodate_overlay"] = {"uri": outputs["overlay"], "sheet_review_base_uri": outputs["sheet_review_base"]}
    for record in records:
        record["raw_payload_ref"] = outputs["records"]
    payload["hakodate_ocr_evidence_records"] = deepcopy(records)
    artifacts.validate_payload(payload)


def _publish_repair_facility(configured, *, select_for_order=None):
    master_path = Path(os.environ["FACILITY_MASTER_PATH"])
    master = json.loads(master_path.read_text(encoding="utf-8"))
    facility = deepcopy(configured["facility"])
    envelope = next(item for item in master["facilities"] if item["facility_id"] == facility["facility_id"])
    version = next(item for item in envelope["versions"] if item["version_id"] == configured["selected_authored_version"]["version_id"])
    version["facility"] = facility
    envelope.update(facility)
    master_path.write_text(json.dumps(master, ensure_ascii=False), encoding="utf-8")
    if select_for_order is not None:
        # Explicit setup selection before capture, not repair of an old receipt.
        with session_scope() as session:
            selected, _ = lineage.resolve(session, facility["facility_id"], date(2026, 10, 5))
            _, projection = lineage.materialize(session, selected)
            lineage.verify_projection(projection)
            order = session.get(Order, select_for_order)
            state = session.get(OrderWorkflowState, select_for_order)
            order.template_version_id = projection.id
            state.template_version_id = projection.id
            session.get(OrderOcrEvidenceRun, state.evidence_run_id).template_version_id = projection.id
        current = config_service.capture_authored_processing_config(
            facility["facility_id"], selected_authored_version=version,
        )
        configured.clear()
        configured.update(current)


@pytest.fixture
def repair_order(isolated_formal_schema, monkeypatch, tmp_path):
    order_id, document_id, evidence_id, message = [str(uuid4()) for _ in range(4)]
    import fitz
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((40, 40), "fixture / 2026-10 / explicit own repair source")
        body = pdf.tobytes()
    path = tmp_path / "repair.pdf"
    path.write_bytes(body)
    configured = {"facility": {"facility_id": "fixture", "fax_template_id": "own", "fax_template": {"template_id": "own", "columns": [{"column_id": "date", "index": 0, "source_index": 0, "role": "date"}, {"column_id": "daypart", "index": 1, "source_index": 1, "role": "daypart"}, {"column_id": "menu", "index": 2, "source_index": 2, "role": "menu"}, {"column_id": "quantity", "index": 3, "source_index": 3, "role": "quantity", "diet_type": "regular", "area_id": "x"}]}}, "policy": {"quantity_rules": {"zero_as_empty": False}}}
    configured["facility"]["fax_template"]["page"] = 1
    configured["facility"].update(facility_name="Synthetic October repair fixture", areas=None)
    config_service._validate_authored_facility_config(configured["facility"], "own")
    authored = {"version_id": "repair-October", "valid_from": "2026-10-01",
                "valid_to": "2026-10-31", "facility": deepcopy(configured["facility"])}
    master_path = tmp_path / "repair-facilities.json"
    master_path.write_text(json.dumps({"facilities": [
        {**deepcopy(configured["facility"]), "versions": [authored]},
    ]}, ensure_ascii=False), encoding="utf-8")
    monkeypatch.delenv("FACILITY_MASTER_URI", raising=False)
    monkeypatch.delenv("K_SERVICE", raising=False)
    monkeypatch.setenv("FACILITY_MASTER_PATH", str(master_path))
    monkeypatch.delenv("OCR_OUTPUT_BUCKET", raising=False)
    monkeypatch.setenv("OCR_ARTIFACT_DIR", str(tmp_path / "owned-artifact-storage"))
    monkeypatch.setattr(config_service, "INGEST_POLICY_PATH", tmp_path / "repair-policy.json")
    _write_repair_policy(configured)
    with session_scope() as session:
        session.add(Facility(id="fixture", name=configured["facility"]["facility_name"]))
        session.add(MonthlyMenu(id="2026-10", month_start=date(2026, 10, 1)))
        session.flush()
        session.add(MonthlyMenuEntry(
            id=str(uuid4()), monthly_menu_id="2026-10", menu_date=date(2026, 10, 7),
            daypart="昼", slot_index=0, name="own", category="主菜", diet_type="regular",
        ))
    with session_scope() as session:
        selected, _ = lineage.resolve(session, "fixture", date(2026, 10, 5))
        _, projection = lineage.materialize(session, selected)
        receipt = lineage.verify_projection(projection)
        assert receipt["authored_version"] == authored
        projection_id = projection.id
    configured = config_service.capture_authored_processing_config(
        "fixture", selected_authored_version=authored,
    )
    with session_scope() as session:
        session.add(Order(id=order_id, facility_code="fixture", week_code="2026-10@2026-10-05~2026-10-11", current_document_id=document_id, document_uri=str(path), message_id=message, received_at=datetime(2026, 10, 7), template_version_id=projection_id))
        session.add(OrderDocument(id=document_id, storage_uri=str(path), source_email_id=message))
        session.flush()
        session.add(UploadedPdf(id=str(uuid4()), message_id=message, storage_uri=str(path), original_filename="repair.pdf", content_sha256=sha256(body).hexdigest(), received_at=datetime(2026, 10, 7)))
        session.add(OrderOcrEvidenceRun(id=evidence_id, order_id=order_id, template_version_id=projection_id, schema_version="fixture", status="ready", payload_json={"own": 0}, artifact_digest="own"))
        session.flush()
        state = OrderWorkflowState(order_id=order_id, evidence_run_id=evidence_id, template_version_id=projection_id, state="ocr_review")
        session.add(state)
        session.flush()
        revision = state.revision
    return order_id, revision, configured


@pytest.mark.parametrize("endpoint", [orders.build_hakodate_assignment, orders.compare_quantity_assignment, orders.build_hakodate_projected_sheet])
def test_actual_endpoint_service_missing_and_stale_no_provider(repair_order, monkeypatch, endpoint):
    order_id, revision, _ = repair_order
    monkeypatch.setattr(order_service, "_augment_hakodate_ocr_payload_artifacts", lambda **_: pytest.fail("provider must not start"))
    with pytest.raises(HTTPException) as missing:
        orders._require_workflow_if_match(None)
    assert missing.value.status_code == 428
    with pytest.raises(HTTPException) as service_missing:
        endpoint(order_id, {}, expected_revision=None)
    assert service_missing.value.status_code == 428
    with pytest.raises(HTTPException) as stale:
        endpoint(order_id, {}, expected_revision=revision + 1)
    assert stale.value.status_code == 409
    with session_scope() as session:
        assert session.get(OrderOcrCache, order_id) is None


@pytest.mark.parametrize("endpoint", [orders.build_hakodate_assignment, orders.compare_quantity_assignment, orders.build_hakodate_projected_sheet])
@pytest.mark.parametrize("outcome", ["own", "provider_failed", "changed_revision", "changed_config", "partial_artifact", "both_partial", "wrong_page", "wrong_evidence", "wrong_config"])
def test_current_repair_actual_cache_write_gate(repair_order, monkeypatch, outcome, endpoint):
    order_id, revision, configured = repair_order
    base = {"fields": ["date", "daypart", "menu", "qty.regular_x"], "rows": [["2026-10-07", "lunch", "own", 0]], "week_id": "2026-10"}
    monkeypatch.setattr(order_service, "_build_hakodate_weekly_menu_base_sheet", lambda *args, **kw: (deepcopy(base), None))
    monkeypatch.setattr(order_service, "get_ocr_sheet", lambda *args, **kw: (deepcopy(base), None))
    monkeypatch.setattr(order_service, "_resolve_llm_review_baseline", lambda **_: {})
    monkeypatch.setattr(order_service, "_build_hakodate_evidence_assignment_from_payload", lambda **_: {"quantity": 0})
    seen = []
    def provider(**kw):
        seen.append(kw)
        assert kw["configured_input"] == configured
        assert kw["pdf_bytes_override"].selected_pdf_receipt["sha256"]
        if outcome == "provider_failed":
            return {"hakodate_canonical_pipeline": {"status": "blocked"}}
        if outcome == "changed_revision":
            with session_scope() as session:
                state = session.get(OrderWorkflowState, order_id)
                state.headline = "new operator edit"
        if outcome == "changed_config":
            configured["policy"]["quantity_rules"]["zero_as_empty"] = True
            _write_repair_policy(configured)
        proof = order_service._hakodate_owned_source(kw["configured_input"], kw["pdf_bytes_override"].selected_pdf_receipt, kw["synchronous_write_input"])
        cell = proof["expected_sheet_cells"][0]
        assert proof["expected_sheet_cells"] == ["E11"]
        result = {"own": 0, "hakodate_repair_source_receipt": deepcopy(proof), "hakodate_canonical_pipeline": {"expected_sheet_cells": [cell]}, "hakodate_preprocessing": {"target_cell_map": [{"row": 0, "sheet_cell": cell, "value": 0, "source_receipt": deepcopy(proof)}]}, "hakodate_ocr_evidence": {"records": [{"evidence_id": "own-provider-record", "source_scope": "hakodate_cell_crop_batch", "value": 0, "engine_metadata": {"sheet_cell": cell, "source_receipt": deepcopy(proof)}}]}}
        result["hakodate_canonical_pipeline"].update(status="ready", version=order_service.HAKODATE_CANONICAL_PIPELINE_VERSION, expected_sheet_cells=deepcopy(proof["expected_sheet_cells"]))
        result["hakodate_preprocessing"]["target_cell_map"][0].update(worksheet_row=11, worksheet_col=5, bbox=[0, 0, 10, 10])
        result["hakodate_preprocessing"]["target_cell_map"], result["hakodate_ocr_evidence"]["records"] = _owned_artifact_fixture(proof)
        _attach_provider_artifact_packet(result)
        if outcome == "both_partial":
            proof["expected_sheet_cells"].append("E12")
            result["hakodate_repair_source_receipt"] = deepcopy(proof)
            result["hakodate_canonical_pipeline"]["expected_sheet_cells"] = deepcopy(proof["expected_sheet_cells"])
            for target in result["hakodate_preprocessing"]["target_cell_map"]:
                target["source_receipt"] = deepcopy(proof)
            for record in result["hakodate_ocr_evidence"]["records"]:
                record["engine_metadata"]["source_receipt"] = deepcopy(proof)
            result["hakodate_ocr_evidence_records"] = deepcopy(result["hakodate_ocr_evidence"]["records"])
        if outcome == "partial_artifact":
            del result["hakodate_ocr_evidence"]
        elif outcome == "wrong_page":
            result["hakodate_repair_source_receipt"]["page"] = 2
        elif outcome == "wrong_evidence":
            result["hakodate_ocr_evidence"]["records"][0]["engine_metadata"]["source_receipt"]["evidence_run_id"] = "foreign"
        elif outcome == "wrong_config":
            result["hakodate_preprocessing"]["target_cell_map"][0]["source_receipt"]["configured_input_digest"] = "foreign"
        return result
    monkeypatch.setattr(order_service, "_hakodate_canonical_payload_from_manifest_item", provider)
    if outcome == "own":
        result = endpoint(order_id, {"strategy": "hakodate"}, expected_revision=revision)
        if endpoint is orders.compare_quantity_assignment:
            assert result["hakodate"]["summary"]["available"] is True
        else:
            assignment = result if endpoint is orders.build_hakodate_assignment else result["assignment"]
            assert assignment["quantity"] == 0
    else:
        with pytest.raises(HTTPException) as error:
            endpoint(order_id, {"strategy": "hakodate"}, expected_revision=revision)
        assert error.value.status_code == (409 if outcome in {"changed_revision", "changed_config"} else 500)
    assert len(seen) == 1
    with session_scope() as session:
        cache = session.get(OrderOcrCache, order_id)
        if outcome == "own":
            assert cache.payload["own"] == 0
        else:
            assert cache is None


@pytest.mark.parametrize("axis", ["header", "row"])
def test_axis_actual_save_and_consumption_original_source(repair_order, axis):
    from src.services import order_workflow_v2_service as workflow
    order_id, revision, configured = repair_order
    with session_scope() as session:
        order = session.get(Order, order_id)
        state = session.get(OrderWorkflowState, order_id)
        receipt, captured = workflow._axis_source_receipt(session, order, state)
    coordinates = [0, 50, 100]
    save = workflow.save_header_axis_review_decision if axis == "header" else workflow.save_row_axis_review_decision
    key = "corrected_xs" if axis == "header" else "corrected_ys"
    result, error = save(order_id, **{key: coordinates}, coordinate_space={"width": 100, "height": 100}, source_receipt=receipt, expected_revision=revision)
    assert error is None
    with session_scope() as session:
        state = session.get(OrderWorkflowState, order_id)
        saved = state.secondary_actions_json["workflow_v2"][axis + "_axis_override"]
    normalize = order_service._normalize_operator_header_axis_override if axis == "header" else order_service._normalize_operator_row_axis_override
    inputs = {"selected_pdf_receipt": receipt["selected_pdf_receipt"], "configured_input": configured, "order_input_identity": receipt["order_input_identity"]}
    normalized = normalize(saved, **inputs)
    assert normalized[key] == coordinates
    assert normalized["source_receipt"] == receipt
    with session_scope() as session:
        updated_receipt, _ = workflow._axis_source_receipt(session, session.get(Order, order_id), session.get(OrderWorkflowState, order_id))
    assert updated_receipt["workflow_revision"] != receipt["workflow_revision"]
    assert updated_receipt["input_digest"] != receipt["input_digest"]
    assert updated_receipt["geometry_source_digest"] == receipt["geometry_source_digest"]
    replaced = deepcopy(inputs)
    replaced["selected_pdf_receipt"]["sha256"] = "0" * 64
    with pytest.raises(ValueError, match="axis_source_input_changed"):
        normalize(saved, **replaced)
    missing = deepcopy(saved)
    del missing["source_receipt"]["geometry_source_digest"]
    with pytest.raises(ValueError, match="geometry_source_proof_required"):
        normalize(missing, **inputs)


@pytest.mark.parametrize("change", ["pdf", "config", "week", "physical_row"])
def test_geometry_source_rejects_changed_authority(repair_order, monkeypatch, change):
    from src.services import order_workflow_v2_service as workflow
    order_id, _, configured = repair_order
    with session_scope() as session:
        receipt, _ = workflow._axis_source_receipt(session, session.get(Order, order_id), session.get(OrderWorkflowState, order_id))
    pdf = deepcopy(receipt["selected_pdf_receipt"])
    identity = deepcopy(receipt["order_input_identity"])
    config = deepcopy(configured)
    if change == "pdf":
        pdf["sha256"] = "0" * 64
    elif change == "config":
        config["facility"]["fax_template"]["columns"][0]["index"] = 99
    elif change == "week":
        identity["week_code"] = "2026-11"
    else:
        monkeypatch.setattr(order_service, "_build_position_menu_entries_safe", lambda *args, **kwargs: [{"menu_date": "2026-10-07", "daypart_key": "lunch", "slot_index": 1, "menu_name": "own"}])
    with pytest.raises(ValueError, match="geometry_source_input_changed"):
        order_service._validate_geometry_source(receipt, pdf, config, identity)


def test_actual_workbook_generation_uses_original_a_during_b_then_a(repair_order, monkeypatch):
    from src.services import order_workflow_v2_service as workflow
    from src.services import master_order_form_template_service as master
    order_id, _, _ = repair_order
    original = [{"menu_date": "2026-10-07", "daypart": "lunch", "daypart_key": "lunch", "slot_index": 0, "name": "original-A", "menu_name": "original-A", "raw_zero": 0, "raw_null": None}]
    current = deepcopy(original)
    monkeypatch.setattr(order_service, "_build_position_menu_entries_safe", lambda *args, **kwargs: deepcopy(current))
    with session_scope() as session:
        session.get(Order, order_id).week_code = "2026-10@2026-10-04~2026-10-10"
    with session_scope() as session:
        receipt, captured = workflow._axis_source_receipt(session, session.get(Order, order_id), session.get(OrderWorkflowState, order_id))
    current[0]["name"] = current[0]["menu_name"] = "changed-B"
    monkeypatch.setattr(master.menu_service, "get_menu_for_facility", lambda *args: pytest.fail("workbook rereads B"))
    workbook = master.build_facility_template_workbook(facility_config=captured["configured_input"]["facility"], week_value=captured["order"]["week_code"], week_menu_entries=captured["physical_menu_entries"])
    values = [cell.value for sheet in workbook.worksheets for row in sheet for cell in row]
    assert "original-A" in values
    assert "changed-B" not in values
    current[:] = deepcopy(original)
    order_service._validate_geometry_source(receipt, receipt["selected_pdf_receipt"], captured["configured_input"], receipt["order_input_identity"])
    assert captured["physical_menu_entries"][0]["raw_zero"] == 0
    assert captured["physical_menu_entries"][0]["raw_null"] is None


@pytest.mark.parametrize("axis", ["header", "row"])
def test_axis_get_put_consume_get_keeps_original_proof(repair_order, monkeypatch, axis):
    from src.services import order_workflow_v2_service as workflow
    order_id, revision, _ = repair_order
    key = "corrected_xs" if axis == "header" else "corrected_ys"
    normalize = order_service._normalize_operator_header_axis_override if axis == "header" else order_service._normalize_operator_row_axis_override
    seen = []
    def image_provider(oid, *, source_receipt, start_input):
        # Only image generation is injected; original proof and saved override
        # consumption go through the actual shared validators.
        order_service._validate_geometry_source(source_receipt, source_receipt["selected_pdf_receipt"], start_input["configured_input"], source_receipt["order_input_identity"])
        saved = (start_input["workflow"]["secondary_actions_json"] or {}).get("workflow_v2", {}).get(axis + "_axis_override")
        consumed = normalize(saved, selected_pdf_receipt=source_receipt["selected_pdf_receipt"], configured_input=start_input["configured_input"], order_input_identity=source_receipt["order_input_identity"])
        seen.append(consumed)
        return {"status": "ready", "saved_override": consumed}
    monkeypatch.setattr(order_service, "build_hakodate_" + axis + "_axis_review", image_provider)
    get = workflow.get_header_axis_review if axis == "header" else workflow.get_row_axis_review
    save = workflow.save_header_axis_review_decision if axis == "header" else workflow.save_row_axis_review_decision
    first, error = get(order_id)
    assert error is None
    original = deepcopy(first["source_receipt"])
    _, error = save(order_id, **{key: [0, 50, 100]}, coordinate_space={"width": 100, "height": 100}, source_receipt=original, expected_revision=revision)
    assert error is None
    second, error = get(order_id)
    assert error is None
    assert seen[0] is None
    assert seen[1][key] == [0, 50, 100]
    assert second["saved_override"]["source_receipt"] == original
    assert second["source_receipt"]["geometry_source_digest"] == original["geometry_source_digest"]
    assert second["source_receipt"]["workflow_revision"] != original["workflow_revision"]


def test_grid_detector_failure_actual_api_is_not_equal_column_success(repair_order, monkeypatch):
    order_id, _, configured = repair_order
    configured["facility"]["fax_template"]["grid_expected_columns"] = 4
    _publish_repair_facility(configured, select_for_order=order_id)
    monkeypatch.setattr(order_service, "detect_table_grid", lambda *args, **kwargs: None)
    result, error, detail = order_service.detect_order_grid(order_id, grid_params={"grid_source": "pdf"})
    assert result is None
    assert error == "grid_not_found"
    assert detail["reason"] == "grid_not_found"
    with pytest.raises(HTTPException) as failure:
        orders.detect_grid(order_id, {"grid_source": "pdf"})
    assert failure.value.status_code == 404


@pytest.mark.parametrize("source", ["pdf", "layout_overlay", "ocr_overlay", "wrong_page", "changed_pdf"])
def test_actual_grid_uses_only_original_pdf_page(repair_order, monkeypatch, source):
    from pathlib import Path
    from types import SimpleNamespace
    order_id, _, configured = repair_order
    seen = []
    def detector(body, template):
        seen.append((body.selected_pdf_receipt, deepcopy(template)))
        return SimpleNamespace(table_box=[0, 0, 1, 1], column_edges=[0, 1], row_edges=[0, 1], confidence=1)
    monkeypatch.setattr(order_service, "detect_table_grid", detector)
    monkeypatch.setattr(order_service, "_load_order_ocr_cache", lambda *args: pytest.fail("stale overlay/cache source read"))
    if source == "wrong_page":
        configured["facility"]["fax_template"]["page"] = 2
        _publish_repair_facility(configured, select_for_order=order_id)
    elif source == "changed_pdf":
        with session_scope() as session:
            uri = session.get(Order, order_id).document_uri
        Path(uri).write_bytes(b"replaced same-name bytes")
    result, error, _ = order_service.detect_order_grid(order_id, grid_params={"grid_source": source if source in {"pdf", "layout_overlay", "ocr_overlay"} else "pdf"})
    if source == "pdf":
        assert error is None
        assert result["source"] == "pdf"
        assert result["fallback"] is False
        assert len(seen) == 1
        assert seen[0][1]["page"] == 1
    else:
        assert result is None
        assert error == {"layout_overlay": "grid_source_receipt_required", "ocr_overlay": "grid_source_receipt_required", "wrong_page": "selected_pdf_page_invalid", "changed_pdf": "document_load_failed"}[source]
        assert seen == []


def test_batch_template_uses_full_captured_config_not_current_master(repair_order, monkeypatch):
    from src.services import hakodate_cell_ocr_batch_service as batch
    _, _, configured = repair_order
    configured["facility"]["fax_template"]["unknown"] = None
    configured["facility"]["fax_template"]["zero"] = 0
    monkeypatch.setattr(config_service, "get_facility_config_for_template", lambda *args: pytest.fail("ambient current config"))
    result = batch._fax_template_for_manifest_item({"configured_input": deepcopy(configured), "fax_template_id": "own"}, "fixture")
    assert result == configured["facility"]["fax_template"]
    with pytest.raises(ValueError, match="hakodate_captured_config_required"):
        batch._fax_template_for_manifest_item({"fax_template_id": "own"}, "fixture")


def test_readonly_assignment_never_repairs_or_writes(repair_order, monkeypatch):
    order_id, _, _ = repair_order
    monkeypatch.setattr(order_service, "_augment_hakodate_ocr_payload_artifacts", lambda **_: pytest.fail("readonly provider repair"))
    monkeypatch.setattr(order_service, "_save_order_ocr_cache", lambda *args, **kwargs: pytest.fail("readonly cache write"))
    order_service.build_order_hakodate_assignment(order_id, allow_artifact_repair=False)
    with session_scope() as session:
        assert session.get(OrderOcrCache, order_id) is None


@pytest.mark.parametrize("superseded", [False, True])
def test_job_owned_cache_canonicalization_uses_original_receipt(repair_order, superseded):
    from src.models.ocr_job import OcrJob
    from src.services import workflow_attempt_guard_service as guard
    order_id, revision, configured = repair_order
    job_id, attempt = str(uuid4()), str(uuid4())
    with session_scope() as session:
        order = session.get(Order, order_id)
        captured = guard.capture_reparse_read_input(session, order)
        pdf = guard.capture_selected_pdf_receipt(session, order)
        session.add(OcrJob(id=job_id, order_id=order_id, status="running", input_reference=pdf["storage_uri"], metrics={"workflow_attempt_id": attempt, "reparse_read_input": captured, "selected_pdf_receipt": pdf}))
    if superseded:
        with session_scope() as session:
            job = session.get(OcrJob, job_id)
            job.metrics = {**job.metrics, "workflow_attempt_id": str(uuid4())}
    token_attempt = order_service._workflow_ocr_attempt_id.set(attempt)
    token_revision = order_service._workflow_ocr_revision.set(revision)
    try:
        with guard.workflow_execution_context(job_id, order_id, attempt, revision):
            if superseded:
                with pytest.raises(ValueError, match="configured_execution_input_changed"):
                    order_service._save_order_ocr_cache(order_id, {"own": 0}, job_id=job_id, augment_hakodate_artifacts=False, persist_evidence=False, refresh_workflow=False)
            else:
                assert guard.processing_start_read_input()["configured_input"] == configured
                result = order_service._save_order_ocr_cache(order_id, {"own": 0}, job_id=job_id, augment_hakodate_artifacts=False, persist_evidence=False, refresh_workflow=False)
                assert result is not None
    finally:
        order_service._workflow_ocr_attempt_id.reset(token_attempt)
        order_service._workflow_ocr_revision.reset(token_revision)
    with session_scope() as session:
        cache = session.get(OrderOcrCache, order_id)
        assert (cache is None) is superseded
        if cache is not None:
            assert cache.payload["own"] == 0


@pytest.mark.parametrize("failure", [None, "partial", "coordinated_shrink", "physical_identity", "wrong_page", "wrong_config", "missing_version", "missing_record_id", "conflicting_records", "changed_current_config"])
def test_real_template_targets_to_assignment_and_owned_cache_gate(repair_order, monkeypatch, failure):
    from src.services import master_order_form_template_service as master
    from src.services import hakodate_cell_ocr_batch_service as batch
    from src.services import hakodate_ocr_evidence_service as evidence_service
    order_id, _, configured = repair_order
    physical = [{"menu_date": "2026-10-07", "daypart": "昼", "daypart_key": "lunch", "slot_index": index, "name": name, "menu_name": name} for index, name in enumerate(["A", "B"])]
    monkeypatch.setattr(order_service, "_build_position_menu_entries_safe", lambda *args, **kwargs: deepcopy(physical))
    captured = order_service._capture_synchronous_cache_write_input(order_id)
    workbook = master.build_facility_template_workbook(facility_config=captured["configured_input"]["facility"], week_value=captured["input_versions"]["order"]["week_code"], week_menu_entries=captured["physical_menu_entries"])
    worksheet = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
    draft = batch.original_physical_draft(physical_menu_entries=captured["physical_menu_entries"], week_code=captured["input_versions"]["order"]["week_code"])
    menu_row_indexes = [row["worksheet_row"] - master.BODY_START_ROW for row in master.validated_body_row_plan(worksheet)]
    # Real geometry target producer and expectation freezer; recognizer values
    # are explicit provider-fixture outputs, never used to choose target count.
    regions, _ = batch._post_menu_target_regions(worksheet=worksheet, column_edges=[float(i * 10) for i in range(worksheet.max_column + 1)], row_edges=[float(i * 10) for i in range(59)], fax_template=captured["configured_input"]["facility"]["fax_template"], draft_sheet=draft, menu_body_row_indexes=menu_row_indexes)
    expected = batch.expected_sheet_cells_from_target_regions(regions)
    assert expected == ["E11", "E12"]
    source = order_service._hakodate_owned_source(captured["configured_input"], captured["selected_pdf_receipt"], captured)
    assert source["expected_sheet_cells"] == expected
    targets = evidence_service.target_cells_from_regions(regions)
    for target in targets:
        target.update(source="hakodate_best_method_pipeline", source_receipt=deepcopy(source))
    records = evidence_service.evidence_from_records(
        [{"text": "0", "bbox": region["bbox"], "confidence": 1, "engine_metadata": {"sheet_cell": region["sheet_cell"], "source_receipt": deepcopy(source), "recognizer_output": deepcopy(region)}} for region in regions],
        run_id="own-recognizer", engine="fixture-recognizer", source_scope="hakodate_cell_crop_batch",
    )
    payload = {"hakodate_repair_source_receipt": source, "hakodate_canonical_pipeline": {"status": "ready", "version": order_service.HAKODATE_CANONICAL_PIPELINE_VERSION, "expected_sheet_cells": expected}, "hakodate_preprocessing": {"target_cell_map": targets}, "hakodate_ocr_evidence": {"records": records}, "hakodate_ocr_evidence_records": deepcopy(records)}
    _attach_provider_artifact_packet(payload)
    if failure == "partial":
        targets.pop()
        records.pop()
    elif failure == "coordinated_shrink":
        source["expected_sheet_cells"] = expected[:1]
        payload["hakodate_canonical_pipeline"]["expected_sheet_cells"] = expected[:1]
        targets.pop()
        records.pop()
        for item in targets:
            item["source_receipt"] = deepcopy(source)
        for item in records:
            item["engine_metadata"]["source_receipt"] = deepcopy(source)
        payload["hakodate_ocr_evidence_records"] = deepcopy(records)
    elif failure == "physical_identity":
        source["physical_menu_digest"] = "other-physical-row"
        for item in targets:
            item["source_receipt"] = deepcopy(source)
        for item in records:
            item["engine_metadata"]["source_receipt"] = deepcopy(source)
        payload["hakodate_ocr_evidence_records"] = deepcopy(records)
    elif failure == "wrong_page":
        targets[0]["source_receipt"]["page"] = 2
    elif failure == "wrong_config":
        records[0]["engine_metadata"]["source_receipt"]["configured_input_digest"] = "other-config"
    elif failure == "missing_version":
        payload["hakodate_canonical_pipeline"].pop("version")
    elif failure == "missing_record_id":
        records[0].pop("evidence_id")
    elif failure == "conflicting_records":
        payload["hakodate_ocr_evidence"] = {"records": deepcopy(records)}
        payload["hakodate_ocr_evidence"]["records"][0]["raw_text"] = "foreign"
    elif failure == "changed_current_config":
        configured["policy"]["quantity_rules"]["zero_as_empty"] = True
        _write_repair_policy(configured)
    if failure and failure != "changed_current_config":
        with pytest.raises(ValueError):
            order_service._validate_hakodate_owned_artifacts(payload, order_service._hakodate_owned_source(captured["configured_input"], captured["selected_pdf_receipt"], captured))
    else:
        order_service._validate_hakodate_owned_artifacts(payload, order_service._hakodate_owned_source(captured["configured_input"], captured["selected_pdf_receipt"], captured))
        monkeypatch.setattr(config_service, "get_facility_config_for_template", lambda *args: pytest.fail("assignment recaptured current config"))
        assignment = order_service._build_hakodate_evidence_assignment_from_payload(order_id=order_id, facility_id="fixture", template_id="own", payload=payload, facility_config=deepcopy(captured["configured_input"]["facility"]))
        assert [item["assigned_value"] for item in assignment["assignments"]] == ["0", "0"]
    saved = order_service._save_order_ocr_cache(order_id, payload, synchronous_write_input=captured, augment_hakodate_artifacts=False, preserve_existing_artifacts=False, persist_evidence=False, refresh_workflow=False)
    assert (saved is None) is bool(failure)
    with session_scope() as session:
        assert (session.get(OrderOcrCache, order_id) is None) is bool(failure)


@pytest.mark.parametrize("regions", [[], [{"sheet_cell": None}], [{"sheet_cell": "E11"}, {"sheet_cell": "E11"}]])
def test_batch_expected_target_freezer_rejects_missing_and_duplicate(regions):
    from src.services import hakodate_cell_ocr_batch_service as batch
    with pytest.raises(ValueError, match="hakodate_expected_cells_missing"):
        batch.expected_sheet_cells_from_target_regions(regions)


@pytest.mark.parametrize("value", [0, None])
def test_original_workbook_expectation_owned_write_preserves_zero_and_null(repair_order, value):
    order_id, _, _ = repair_order
    start = order_service._capture_synchronous_cache_write_input(order_id)
    source = order_service._hakodate_owned_source(start["configured_input"], start["selected_pdf_receipt"], start)
    assert source["expected_sheet_cells"] == ["E11"]
    targets = [{"sheet_cell": cell, "worksheet_row": 11, "worksheet_col": 5, "bbox": [0, 0, 10, 10], "source_receipt": deepcopy(source), "value": value} for cell in source["expected_sheet_cells"]]
    records = [{"evidence_id": str(uuid4()), "source_scope": "hakodate_cell_crop_batch", "raw_text": value, "engine_metadata": {"sheet_cell": cell, "source_receipt": deepcopy(source)}} for cell in source["expected_sheet_cells"]]
    targets, records = _owned_artifact_fixture(source, value)
    payload = {"hakodate_repair_source_receipt": deepcopy(source), "hakodate_preprocessing": {"target_cell_map": targets}, "hakodate_ocr_evidence": {"records": records}, "hakodate_canonical_pipeline": {"status": "ready", "version": order_service.HAKODATE_CANONICAL_PIPELINE_VERSION, "expected_sheet_cells": deepcopy(source["expected_sheet_cells"])}}
    _attach_provider_artifact_packet(payload)
    order_service._validate_hakodate_owned_artifacts(payload, source)
    assert order_service._save_order_ocr_cache(order_id, payload, synchronous_write_input=start, augment_hakodate_artifacts=False, preserve_existing_artifacts=False, persist_evidence=False, refresh_workflow=False) is not None
    with session_scope() as session:
        saved = session.get(OrderOcrCache, order_id).payload
        assert saved["hakodate_preprocessing"]["target_cell_map"][0]["value"] is value
        assert saved["hakodate_ocr_evidence"]["records"][0]["raw_text"] is value


@pytest.mark.parametrize("canonical,alternative", [(None, [{"value": 0}]), ([None], [{"value": 0}]), ([{"value": 0}], [{"value": False}]), ([{"value": None}], None)])
def test_real_canonical_expression_reader_and_classifier_do_not_filter_or_choose(canonical, alternative):
    from src.services import ocr_evidence_service
    payload = {"hakodate_preprocessing": {"target_cell_map": canonical}, "hakodate_target_cell_map": alternative}
    with pytest.raises(ValueError):
        order_service._extract_hakodate_target_cells_from_payload(payload)
    result = ocr_evidence_service.classify_evidence_payload(payload)
    assert result["persistable"] is False


def test_missing_canonical_expression_and_record_identity_cannot_use_legacy_alternative():
    with pytest.raises(ValueError, match="canonical_representation_missing"):
        order_service._extract_hakodate_target_cells_from_payload({"hakodate_target_cell_map": [{"sheet_cell": "E11"}]})
    with pytest.raises(ValueError, match="own_record_identity_missing"):
        order_service._extract_hakodate_ocr_evidence_records_from_payload({"hakodate_ocr_evidence": {"records": [{"source_scope": "hakodate_cell_crop_batch", "raw_text": "0"}]}}, order_id="own")


def test_invalid_selected_evidence_preview_cannot_borrow_cache(repair_order, monkeypatch):
    order_id, _, _ = repair_order
    with session_scope() as session:
        session.add(OrderOcrCache(order_id=order_id, payload={"hakodate_assignment_preview": {"status": "ready", "assignments": [{"assigned_value": 9}]}}))
    monkeypatch.setattr(order_service, "_load_order_ocr_cache", lambda *args: pytest.fail("old cache preview adopted"))
    assert order_service.get_cached_hakodate_assignment_preview(order_id) is None


@pytest.mark.parametrize("value", [0, None])
def test_real_evidence_writer_keeps_independent_expectation_on_selected_replay(repair_order, value):
    from src.models.ocr_job import OcrJob
    from src.services import ocr_evidence_service, workflow_attempt_guard_service as guard
    order_id, revision, _ = repair_order
    job_id, attempt = str(uuid4()), str(uuid4())
    with session_scope() as session:
        order = session.get(Order, order_id)
        start = guard.capture_reparse_read_input(session, order)
        pdf = guard.capture_selected_pdf_receipt(session, order)
        session.add(OcrJob(id=job_id, order_id=order_id, status="running", input_reference=pdf["storage_uri"], metrics={"workflow_attempt_id": attempt, "reparse_read_input": start, "selected_pdf_receipt": pdf}))
    tokens = [(variable, variable.set(val)) for variable, val in ((order_service._workflow_ocr_job_id, job_id), (order_service._workflow_ocr_order_id, order_id), (order_service._workflow_ocr_attempt_id, attempt), (order_service._workflow_ocr_revision, revision))]
    try:
        with guard.workflow_execution_context(job_id, order_id, attempt, revision):
            source = order_service._hakodate_owned_source(start["configured_input"], pdf)
            assert source["expected_sheet_cells"] == ["E11"]
            targets = [{"target_cell_id": cell, "sheet_cell": cell, "worksheet_row": 11, "worksheet_col": 5, "semantic_field": "qty.regular_x", "bbox": [0, 0, 10, 10], "value": value, "source_receipt": deepcopy(source)} for cell in source["expected_sheet_cells"]]
            records = [{"evidence_id": str(uuid4()), "source_scope": "hakodate_cell_crop_batch", "raw_text": value, "normalized_value": "0" if value == 0 else "", "center": [5, 5], "engine_metadata": {"sheet_cell": cell, "source_receipt": deepcopy(source)}} for cell in source["expected_sheet_cells"]]
            targets, records = _owned_artifact_fixture(source, value)
            payload = {"hakodate_repair_source_receipt": deepcopy(source), "hakodate_preprocessing": {"target_cell_map": targets}, "hakodate_ocr_evidence": {"records": records}, "hakodate_canonical_pipeline": {"status": "ready", "version": order_service.HAKODATE_CANONICAL_PIPELINE_VERSION, "expected_sheet_cells": deepcopy(source["expected_sheet_cells"])}}
            _attach_provider_artifact_packet(payload)
            saved = ocr_evidence_service.persist_evidence_run(order_id, payload, status="ready" if value is not None else "diagnostic_only", producer_version="own-fixture/v1", source="own-recognizer", workflow_job_id=job_id, workflow_attempt_id=attempt, expected_workflow_revision=revision)
            assert saved is not None
            if value is not None:
                assert saved["artifact_manifest_json"]["original_hakodate_source"] == source
    finally:
        for variable, token in reversed(tokens):
            variable.reset(token)
    with session_scope() as session:
        row = session.get(OrderOcrEvidenceRun, saved["id"])
        assert ocr_evidence_service.evidence_run_is_selectable(row) is (value is not None)
        assert row.payload_json["hakodate_ocr_evidence"]["records"][0]["raw_text"] == value
        assert type(row.payload_json["hakodate_ocr_evidence"]["records"][0]["raw_text"]) is type(value)
        tampered = deepcopy(row.payload_json)
        tampered["hakodate_repair_source_receipt"]["expected_sheet_cells"] = []
        tampered["hakodate_canonical_pipeline"]["expected_sheet_cells"] = []
        tampered["hakodate_preprocessing"]["target_cell_map"] = []
        tampered["hakodate_ocr_evidence"]["records"] = []
        row.payload_json = tampered
        assert not ocr_evidence_service.evidence_run_is_selectable(row)
        if value is not None:
            assert row.artifact_manifest_json["original_hakodate_source"]["expected_sheet_cells"] == ["E11"]


def test_versionless_producer_hint_is_not_canonical():
    legacy = {"hakodate_canonical_pipeline": {"producer": "hakodate_cell_ocr_batch_service.build_hakodate_best_method_for_manifest_item", "source": "live_order_facility_source_workbook", "requested_order_id": "old", "status": "ready"}}
    assert order_service._payload_has_current_hakodate_canonical_pipeline(legacy) is False
    assert order_service._hakodate_version_matches_current(None) is False
    assert order_service._hakodate_version_matches_current(order_service.HAKODATE_CANONICAL_PIPELINE_VERSION) is True


@pytest.mark.parametrize("change", [None, "date", "bool_page", "bool_target", "read_race"])
def test_actual_canonical_producer_json_writer_and_selected_read(repair_order, monkeypatch, change):
    import json
    from datetime import date
    from pathlib import Path
    from types import SimpleNamespace
    from src.models.ocr_job import OcrJob
    from src.services import hakodate_cell_ocr_batch_service as batch, ocr_evidence_service, storage_service, workflow_attempt_guard_service as guard

    order_id, revision, configured = repair_order
    physical = [{"menu_date": date(2026, 10, 7), "daypart": "昼", "daypart_key": "lunch", "slot_index": 0, "name": "own", "menu_name": "own", "diet_type": "regular"}]
    monkeypatch.setattr(order_service.menu_service, "get_menu_entries_for_facility", lambda *args, **kwargs: deepcopy(physical))
    physical = order_service._build_position_menu_entries("2026-10@2026-10-05~2026-10-11", "fixture")
    assert physical[0]["menu_date"] == date(2026, 10, 7)
    monkeypatch.setattr(order_service, "_build_position_menu_entries_safe", lambda *args, **kwargs: deepcopy(physical))
    with session_scope() as session:
        order = session.get(Order, order_id)
        start = guard.capture_reparse_read_input(session, order)
        pdf = guard.capture_selected_pdf_receipt(session, order)
        job_id, attempt = str(uuid4()), str(uuid4())
        session.add(OcrJob(id=job_id, order_id=order_id, status="running", input_reference=pdf["storage_uri"], metrics={"workflow_attempt_id": attempt, "reparse_read_input": start, "selected_pdf_receipt": pdf}))
    verified = storage_service.load_selected_pdf_bytes(pdf["storage_uri"], pdf)
    draft = {"fields": ["date", "daypart", "menu"], "rows": [["2026-10-07", "昼", "own"]]}
    # External rendering/recognition are injected, not the production payload
    # constructor, record converter, classifier or actual fenced evidence writer.
    monkeypatch.setattr(order_service, "_hakodate_best_method_draft_sheet", lambda *args, **kwargs: deepcopy(draft))
    monkeypatch.setattr(order_service, "_build_live_hakodate_manifest_item", lambda **kwargs: {"configured_input": deepcopy(start["configured_input"]), "selected_pdf_receipt": deepcopy(pdf), "_verified_pdf_bytes": verified})
    def recognizer(*, item, output_dir, **kwargs):
        assert item["original_expected_sheet_cells"] == ["E11"]
        region = {**deepcopy(item["original_target_semantics"][0]), "bbox": [0, 0, 10, 10]}
        outputs = _provider_output_files(output_dir, regions=[region], records=[{**region, "ocr_text": "0", "recognizer_accepted_candidate": {"normalized_digits": "0"}, "own_null": None}], sheet_values={"E11": "0"})
        return SimpleNamespace(metrics={"expected_sheet_cells": ["E11"]}, expected_sheet_cells=("E11",), outputs=outputs, ocr_engine="fixture-recognizer"), None
    monkeypatch.setattr(batch, "build_hakodate_best_method_for_manifest_item", recognizer)
    variables = [(var, var.set(value)) for var, value in ((order_service._workflow_ocr_job_id, job_id), (order_service._workflow_ocr_order_id, order_id), (order_service._workflow_ocr_attempt_id, attempt), (order_service._workflow_ocr_revision, revision))]
    try:
        with guard.workflow_execution_context(job_id, order_id, attempt, revision):
            payload = order_service._hakodate_canonical_payload_from_manifest_item(order_id=order_id, facility_id="fixture", item={})
            assert order_service._classify_owned_evidence_payload(payload)["persistable"] is True
            assert payload["hakodate_ocr_evidence"]["records"] == payload["hakodate_ocr_evidence_records"]
            assert payload["hakodate_ocr_evidence"]["records"][0]["raw_text"] == "0"
            assert payload["hakodate_ocr_evidence"]["records"][0]["engine_metadata"]["recognizer_output"]["own_null"] is None
            saved = ocr_evidence_service.persist_evidence_run(order_id, payload, status="ready", source="own-recognizer", producer_version="own-fixture/v1", workflow_job_id=job_id, workflow_attempt_id=attempt, expected_workflow_revision=revision)
            assert saved is not None
    finally:
        for variable, token in reversed(variables):
            variable.reset(token)
    current = order_service._capture_synchronous_cache_write_input(order_id)
    assert guard.snapshots_equal(start["physical_menu_entries"], physical)
    evidence = json.loads(guard.canonical_snapshot_json(saved))
    if change == "read_race":
        with session_scope() as session:
            row = session.get(OrderOcrEvidenceRun, saved["id"])
            row.payload_json = {**row.payload_json, "hakodate_overlay": {"uri": "gs://fixture/own-overlay.png", "producer": "hakodate_best_method_pipeline", "version": order_service.HAKODATE_CANONICAL_PIPELINE_VERSION, "fingerprint": "own-provider-image"}}
            session.get(OrderWorkflowState, order_id).evidence_run_id = saved["id"]
        monkeypatch.setattr(order_service, "_build_hakodate_weekly_menu_base_sheet", lambda *args, **kwargs: (deepcopy(draft), None))
        signed_calls = []
        def signed(uri):
            signed_calls.append(uri)
            configured["policy"]["quantity_rules"]["zero_as_empty"] = True
            _write_repair_policy(configured)
            return "https://fixture/own-overlay.png"
        monkeypatch.setattr(order_service, "_signed_url_from_uri", signed)
        result = order_service.get_cached_hakodate_overlay_preview(order_id)
        assert signed_calls == ["gs://fixture/own-overlay.png"]
        assert result["status"] == "blocked" and result["assignment"] is None
        assert result["blockers"] == ["synchronous_input_conflict"]
        return
    if change == "date":
        current["physical_menu_entries"][0]["menu_date"] = "2026-10-08"
    elif change == "bool_page":
        evidence["payload_json"]["hakodate_repair_source_receipt"]["page"] = True
    elif change == "bool_target":
        evidence["payload_json"]["hakodate_preprocessing"]["target_cell_map"][0]["worksheet_row"] = True
    if change:
        with pytest.raises(ValueError):
            order_service._validate_selected_hakodate_artifacts(evidence, current)
    else:
        order_service._validate_selected_hakodate_artifacts(evidence, current)
        assert not guard.snapshots_equal({"page": True}, {"page": 1})


def test_shared_logical_plan_and_actual_geometry_use_identical_original_rows(repair_order):
    from src.services import master_order_form_template_service as master
    from src.services import hakodate_cell_ocr_batch_service as batch
    from src.services.hakodate_step_review_pipeline_service import canonical_logical_target_plan, _post_menu_target_regions
    order_id, _, _ = repair_order
    start = order_service._capture_synchronous_cache_write_input(order_id)
    workbook = master.build_facility_template_workbook(facility_config=start["configured_input"]["facility"], week_value=start["input_versions"]["order"]["week_code"], week_menu_entries=start["physical_menu_entries"])
    worksheet = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
    draft = batch.original_physical_draft(physical_menu_entries=start["physical_menu_entries"], week_code=start["input_versions"]["order"]["week_code"])
    menu_row_indexes = [row["worksheet_row"] - master.BODY_START_ROW for row in master.validated_body_row_plan(worksheet)]
    template = start["configured_input"]["facility"]["fax_template"]
    plan = canonical_logical_target_plan(worksheet=worksheet, fax_template=template, draft_sheet=draft)
    regions, _ = _post_menu_target_regions(worksheet=worksheet, fax_template=template, draft_sheet=draft, column_edges=[float(index) for index in range(worksheet.max_column + 1)], row_edges=[float(index) for index in range(50)], menu_body_row_indexes=menu_row_indexes)
    assert [item["sheet_cell"] for item in plan] == [item["sheet_cell"] for item in regions] == ["E11"]
    assert plan[0]["covered_sheet_cells"] == regions[0]["covered_sheet_cells"]
    missing_role = deepcopy(template)
    missing_role["columns"] = [col for col in missing_role["columns"] if col["role"] != "menu"]
    with pytest.raises(ValueError, match="menu_role_unresolved"):
        canonical_logical_target_plan(worksheet=worksheet, fax_template=missing_role, draft_sheet=draft)
    with pytest.raises(ValueError, match="row_map"):
        _post_menu_target_regions(worksheet=worksheet, fax_template=template, draft_sheet=draft, column_edges=[float(index) for index in range(worksheet.max_column + 1)], row_edges=[float(index) for index in range(50)])
    workbook.close()


@pytest.fixture
def connected_physical_pipeline(repair_order, monkeypatch, tmp_path, request):
    import cv2
    import numpy as np
    from src.services import hakodate_cell_ocr_batch_service as batch, hakodate_preprocessing_pipeline_service as collector, hakodate_step_review_pipeline_service as pipeline, master_order_form_template_service as master, order_workflow_v2_service as workflow
    order_id, _, configured = repair_order
    week = "2026-10@2026-10-01~2026-10-03"
    physical = [{"menu_date": f"2026-10-0{day}", "daypart": "昼", "daypart_key": "lunch", "slot_index": 0, "name": f"own-{day}", "diet_type": "regular"} for day in (1, 2, 3)]
    merged = getattr(request, "param", False)
    if merged:
        configured["facility"]["fax_template"]["columns"][-1]["body_merge"] = "daypart"
        week = "2026-10@2026-10-01~2026-10-01"
        physical = [{"menu_date": "2026-10-01", "daypart": "昼", "daypart_key": "lunch", "slot_index": slot, "name": f"own-{slot}", "diet_type": "regular"} for slot in range(3)]
        _publish_repair_facility(configured, select_for_order=order_id)
    raw_physical = deepcopy(physical)
    def entries_for_month(month_id, *args, **kwargs):
        return deepcopy([entry for entry in raw_physical if str(entry["menu_date"])[:7] == month_id])
    monkeypatch.setattr(order_service.menu_service, "get_menu_entries_for_facility", entries_for_month)
    monkeypatch.setattr(order_service, "_build_position_menu_entries_safe", order_service._build_position_menu_entries)
    monkeypatch.setattr(order_service.menu_service, "get_menu", lambda month_id, *args, **kwargs: {"entries": entries_for_month(month_id)})
    monkeypatch.setattr(order_service.menu_service, "get_menu_for_facility", lambda month_id, *args, **kwargs: {"entries": entries_for_month(month_id)})
    with session_scope() as session:
        order = session.get(Order, order_id)
        order.week_code = week
    with session_scope() as session:
        receipt, start = workflow._axis_source_receipt(session, session.get(Order, order_id), session.get(OrderWorkflowState, order_id))
    physical = deepcopy(start["physical_menu_entries"])
    configured = deepcopy(start["configured_input"])
    workbook = master.build_facility_template_workbook(facility_config=configured["facility"], week_value=week, week_menu_entries=physical)
    worksheet = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
    workbook_path = tmp_path / "own-structure.xlsx"
    workbook.save(workbook_path)
    plan = batch.capture_original_target_plan(configured_input=configured, physical_menu_entries=physical, week_code=week)
    image = np.full((400, 400, 3), 255, dtype=np.uint8)
    xs = [float(20 + index * 50) for index in range(worksheet.max_column + 1)]
    template_ys = [float(20 + index * 25) for index, _ in enumerate(master.validated_template_row_edges(worksheet))]
    fax_ys = [float(20 + index * 25) for index in range(12 if merged else 10)]
    for x in xs:
        cv2.line(image, (int(x), 20), (int(x), 290), (0, 0, 0), 1)
    for y in fax_ys:
        cv2.line(image, (20, int(y)), (int(xs[-1]), int(y)), (0, 0, 0), 1)
    image_path = tmp_path / "own-render.png"
    cv2.imwrite(str(image_path), image)
    draft = batch.original_physical_draft(physical_menu_entries=physical, week_code=week)
    item = {"order_id": order_id, "facility_code": "fixture", "configured_input": deepcopy(configured), "template_id": "own", "fax_template_id": "own", "fax_pdf": receipt["selected_pdf_receipt"]["storage_uri"], "template_pdf": receipt["selected_pdf_receipt"]["storage_uri"], "step2_png": str(image_path), "template_xlsx": str(workbook_path), "template_sheet_name": master.FACILITY_TEMPLATE_SHEET_NAME, "week_sheet_name": master.FACILITY_TEMPLATE_SHEET_NAME, "template_axes_x": xs, "template_axes_y": template_ys, "template_bbox": [xs[0], template_ys[0], xs[-1], template_ys[-1]], "original_physical_draft": draft, "original_expected_sheet_cells": batch.expected_sheet_cells_from_target_regions(plan), "original_target_semantics": pipeline.canonical_target_semantics(plan), "selected_pdf_receipt": deepcopy(receipt["selected_pdf_receipt"])}
    axis = {"row_intersection_y_match": {"structural_match": {"body_day_boundary_indexes": [0, 1, 2, 3, 4, 7, 8, 9] if merged else list(range(8)), "detected_body_band_count": 9 if merged else 7}, "corrected_ys": fax_ys}, "header_intersection_x_match": {"header_roi": [20, 20, 270, 80]}}
    # Inject only external raster/registration observations. Actual target plan,
    # row mapping, preprocessing, snap, runtime evaluation and adoption stay real.
    for module in (batch, collector):
        monkeypatch.setattr(module, "render_template_pdf_to_canvas", lambda *args, **kwargs: image.copy())
        monkeypatch.setattr(module, "render_pdf_page_to_bgr", lambda *args, **kwargs: image.copy())
        monkeypatch.setattr(module, "resolve_fixed_quad_px_for_manifest_item", lambda *args, **kwargs: ([[20, 20], [270, 20], [270, 290], [20, 290]], "explicit-fixture", {}))
        monkeypatch.setattr(module, "build_fixed_quad_template_registration", lambda **kwargs: (type("Registration", (), {"template_outer_grid_bbox_used": item["template_bbox"]})(), {}))
        monkeypatch.setattr(module, "_align_axes", lambda **kwargs: (xs, fax_ys, deepcopy(axis), None))
    def manifest(**kwargs):
        from src.services.storage_service import load_selected_pdf_bytes
        result = deepcopy(item)
        result["_verified_pdf_bytes"] = load_selected_pdf_bytes(item["fax_pdf"], item["selected_pdf_receipt"])
        return result
    monkeypatch.setattr(order_service, "_build_live_hakodate_manifest_item", manifest)
    yield order_id, receipt, start, item, tmp_path
    workbook.close()


def test_actual_header_row_and_normal_collector_share_partial_period_plan(connected_physical_pipeline):
    import json
    from src.services import hakodate_cell_ocr_batch_service as batch, hakodate_preprocessing_pipeline_service as collector
    order_id, receipt, start, item, output = connected_physical_pipeline
    assert batch._validate_menu_day_boundaries(item["original_physical_draft"], {"structural_match": {"body_day_boundary_indexes": list(range(8)), "detected_body_band_count": 7}}) == [4, 5, 6]
    for service in (order_service.build_hakodate_header_axis_review, order_service.build_hakodate_row_axis_review):
        result = service(order_id, source_receipt=receipt, start_input=start)
        assert result["status"] == "ready" and result["image_png_base64"]
    result, _ = collector.build_hakodate_preprocessing_for_manifest_item(item=item, page=1, output_dir=output)
    assert result.target_cell_count == len(item["original_expected_sheet_cells"])
    manifest = output / "own-manifest.json"
    manifest.write_text(json.dumps({"results": [item, deepcopy(item)]}))
    collected = collector.build_all_facility_hakodate_preprocessing_pdf(manifest_path=manifest, output_dir=output / "collected")
    assert [entry["page"] for entry in collected["results"]] == [1, 1]
    broken = deepcopy(item["original_physical_draft"])
    del broken["physical_menu_row_ids"][1]
    broken["physical_menu_row_count"] -= 1
    with pytest.raises(ValueError, match="physical_menu_days_disagree_with_order_period"):
        batch._validate_menu_day_boundaries(broken, {"structural_match": {"body_day_boundary_indexes": list(range(8)), "detected_body_band_count": 7}})


@pytest.mark.parametrize("value", [0, None])
@pytest.mark.parametrize("connected_physical_pipeline", [False, True], indirect=True)
@pytest.mark.parametrize("owner_kind", ["workflow", "nonworkflow"])
def test_connected_preprocess_runtime_producer_serialization_and_adoption(connected_physical_pipeline, monkeypatch, value, owner_kind):
    import json
    from PIL import Image
    from src.models.ocr_job import OcrJob
    from src.services import workflow_attempt_guard_service as guard, ocr_evidence_service
    from src.hakodate_best_method_runtime import render_best_method_overlay_all_facilities as runtime
    order_id, _receipt, _start, item, _output = connected_physical_pipeline
    calls = []
    def recognize(*, raw_rectified_bgr, regions, truth):
        calls.append(deepcopy(regions))
        assert len(regions) == len(item["original_expected_sheet_cells"])
        assert sum(len(region["logical_targets"]) for region in regions) == 3
        if item["configured_input"]["facility"]["fax_template"]["columns"][-1].get("body_merge"):
            assert any(len(region["logical_targets"]) > 1 for region in regions)
        assert all(region["logical_targets"] and region["covered_sheet_cells"] for region in regions)
        text = "0" if value == 0 else ""
        return [{**deepcopy(region), "truth": truth[region["sheet_cell"]], "expected_digits": truth[region["sheet_cell"]]["expected_digits"], "pred_digits": text, "raw_text": text, "score": 1.0, "recognizer_accepted_candidate": {"normalized_digits": text}, "recognizer_score": 1.0} for region in regions], {}, Image.new("RGB", (80, 80), "white")
    monkeypatch.setattr(runtime, "_run_text_recognizer_records", recognize)
    monkeypatch.setattr(order_service, "get_default_output_bucket", lambda: "fixture-output")
    monkeypatch.setattr(order_service, "save_artifact_bytes_to_gcs", lambda bucket, message, name, body, **kwargs: f"gs://{bucket}/{message}/{name}")
    with session_scope() as session:
        order = session.get(Order, order_id)
        start = guard.capture_reparse_read_input(session, order)
        pdf = guard.capture_selected_pdf_receipt(session, order)
        revision = session.get(OrderWorkflowState, order_id).revision
        job_id, attempt = str(uuid4()), str(uuid4())
        session.add(OcrJob(id=job_id, order_id=order_id, status="running", input_reference=pdf["storage_uri"], metrics={"workflow_attempt_id": attempt, "reparse_read_input": start, "selected_pdf_receipt": pdf}))
    if owner_kind == "nonworkflow":
        with session_scope() as session:
            session.get(OcrJob, job_id).metrics = {}
        boundary = order_service._capture_nonworkflow_order_execution(job_id, order_id, None)
        original = boundary["start_orders"][order_id]["processing_start_input"]
        assert guard.snapshots_equal({key: original[key] for key in start}, start)
        assert original["expected_sheet_cells"] == item["original_expected_sheet_cells"]
        assert guard.snapshots_equal(original["expected_target_semantics"], item["original_target_semantics"])
        start = original
        attempt, revision = None, None
        execution_context = guard.nonworkflow_execution_context(boundary)
    else:
        execution_context = guard.workflow_execution_context(job_id, order_id, attempt, revision)
    variables = [(var, var.set(data)) for var, data in ((order_service._workflow_ocr_job_id, job_id if owner_kind == "workflow" else None), (order_service._workflow_ocr_order_id, order_id if owner_kind == "workflow" else None), (order_service._workflow_ocr_attempt_id, attempt), (order_service._workflow_ocr_revision, revision))]
    try:
        with execution_context:
            with session_scope() as session:
                selected_uri, selected = order_service._captured_order_document_selection(session, session.get(Order, order_id), None)
                assert selected_uri == pdf["storage_uri"]
                assert selected["selected_document_id"] == pdf["document_id"]
                assert "error" not in selected
                rejected_uri, rejected = order_service._captured_order_document_selection(session, session.get(Order, order_id), "foreign-document")
                assert rejected_uri is None
                assert rejected["error"] == "selected_document_input_missing"
            payload = order_service._hakodate_canonical_payload_from_manifest_item(order_id=order_id, facility_id="fixture", item={})
            assert len(calls) == 1
            assert [cell["sheet_cell"] for cell in calls[0]] == item["original_expected_sheet_cells"]
            records = payload["hakodate_ocr_evidence"]["records"]
            assert len(records) == len(item["original_expected_sheet_cells"])
            assert all(record["raw_text"] == ("0" if value == 0 else "") for record in records)
            assessment = order_service._classify_owned_evidence_payload(payload)
            assert assessment["persistable"] is (value == 0)
            saved = ocr_evidence_service.persist_evidence_run(order_id, payload, status="ready" if value == 0 else "diagnostic_only", source="own-recognizer", producer_version="fixture-external-recognizer/v1", workflow_job_id=job_id, workflow_attempt_id=attempt, expected_workflow_revision=revision)
            assert saved is not None
            assert guard.snapshots_equal(saved["payload_json"]["_ocr_execution"]["processing_start_input"], start)
    finally:
        for variable, token in reversed(variables):
            variable.reset(token)
    assert saved["status"] == ("ready" if value == 0 else "diagnostic_only")
    from src.services import ocr_evidence_service
    with session_scope() as session:
        assert ocr_evidence_service.evidence_run_is_selectable(session.get(OrderOcrEvidenceRun, saved["id"])) is (value == 0)
    if value is None:
        return
    evidence = json.loads(guard.canonical_snapshot_json(saved))
    current = order_service._capture_synchronous_cache_write_input(order_id)
    order_service._validate_selected_hakodate_artifacts(evidence, current)
    for original_change in ("missing_original", "physical", "config", "expected"):
        changed = deepcopy(evidence)
        original = changed["payload_json"]["_ocr_execution"]["processing_start_input"]
        if original_change == "missing_original":
            changed["payload_json"]["_ocr_execution"].pop("processing_start_input")
        elif original_change == "physical":
            original["physical_menu_entries"][0]["slot_index"] = 99
        elif original_change == "expected":
            original["expected_sheet_cells"] = []
        else:
            original["configured_input"]["policy"]["quantity_rules"]["zero_as_empty"] = True
        with pytest.raises(ValueError):
            order_service._validate_selected_hakodate_artifacts(changed, current)
    changed_current = deepcopy(current)
    changed_current["configured_input"]["policy"]["quantity_rules"]["zero_as_empty"] = True
    with pytest.raises(ValueError):
        order_service._validate_selected_hakodate_artifacts(evidence, changed_current)
    saved_targets = evidence["payload_json"]["hakodate_preprocessing"]["target_cell_map"]
    assert sum(len(target["metadata"]["logical_targets"]) for target in saved_targets) == 3
    assert [target["metadata"]["covered_sheet_cells"] for target in saved_targets] == [region["covered_sheet_cells"] for region in calls[0]]
    for mutation in ("field", "role", "date", "daypart", "menu_name", "logical_targets", "covered_sheet_cells"):
        tampered = deepcopy(evidence)
        target = tampered["payload_json"]["hakodate_preprocessing"]["target_cell_map"][0]
        if mutation == "field":
            target["semantic_field"] = "qty.soft_x"
        elif mutation in {"date", "daypart", "menu_name"}:
            target["metadata"]["logical_targets"][0][mutation] = {"date": "2026-10-02", "daypart": "dinner", "menu_name": "foreign-menu"}[mutation]
        elif mutation == "role":
            target["metadata"][mutation] = "note"
        else:
            target["metadata"][mutation] = []
        with pytest.raises(ValueError, match="semantics_changed"):
            order_service._validate_selected_hakodate_artifacts(tampered, current)
    record_tamper = deepcopy(evidence)
    for expression in (record_tamper["payload_json"]["hakodate_ocr_evidence"]["records"], record_tamper["payload_json"]["hakodate_ocr_evidence_records"]):
        expression[0]["engine_metadata"]["recognizer_output"]["field"] = "qty.soft_x"
    with pytest.raises(ValueError, match="record_semantics_changed"):
        order_service._validate_selected_hakodate_artifacts(record_tamper, current)


def test_actual_master_roles_and_merged_logical_targets(repair_order):
    from src.services import master_order_form_template_service as master, hakodate_assignment_service as assignment, template_field_schema_service as schema, hakodate_cell_ocr_batch_service as batch
    from src.services.hakodate_step_review_pipeline_service import canonical_logical_target_plan, canonical_target_semantics, _post_menu_target_regions
    from src.services.workflow_attempt_guard_service import snapshots_equal
    order_id, _, configured = repair_order
    start = order_service._capture_synchronous_cache_write_input(order_id)
    facility = deepcopy(configured["facility"])
    facility["fax_template"]["columns"].extend([{"column_id": "soft", "index": 4, "source_index": 4, "role": "quantity", "diet_type": "soft", "area_id": "x"}, {"column_id": "remarks", "index": 5, "source_index": 5, "role": "remarks"}, {"column_id": "aux", "index": 6, "source_index": 6, "role": "aux"}])
    config_service._validate_authored_facility_config(facility, "own")
    assert schema.canonical_field_name_from_template_column({"role": "remarks"}) == schema.canonical_field_name_from_template_column({"role": "note"}) == "remarks"
    assert schema.canonical_field_name_from_template_column({"role": "menu"}) == schema.canonical_field_name_from_template_column({"role": "menu_name"}) == "menu"
    workbook = master.build_facility_template_workbook(facility_config=facility, week_value=start["input_versions"]["order"]["week_code"], week_menu_entries=start["physical_menu_entries"])
    worksheet = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
    slots = assignment.canonical_generated_column_slots(facility["fax_template"], col_count=worksheet.max_column)
    assert [(slot["role"], slot["slot_name"]) for slot in slots[1:]] == [("quantity", "qty.regular_x"), ("quantity", "qty.soft_x"), ("note", "remarks"), ("aux", "aux.col_6")]
    worksheet.merge_cells("E11:F11")
    draft = batch.original_physical_draft(physical_menu_entries=start["physical_menu_entries"], week_code=start["input_versions"]["order"]["week_code"])
    plan = canonical_logical_target_plan(worksheet=worksheet, fax_template=facility["fax_template"], draft_sheet=draft)
    regions, _ = _post_menu_target_regions(worksheet=worksheet, fax_template=facility["fax_template"], draft_sheet=draft, column_edges=[float(index * 30) for index in range(worksheet.max_column + 1)], row_edges=[float(index * 20) for index in range(50)], menu_body_row_indexes=[32])
    assert snapshots_equal(canonical_target_semantics(plan), canonical_target_semantics(regions))
    assert regions[0]["region_id"] == "E11:F11"
    assert regions[0]["covered_sheet_cells"] == ["E11", "F11"]
    assert [target["field"] for target in regions[0]["logical_targets"]] == ["qty.regular_x", "qty.soft_x"]
    unknown = deepcopy(facility["fax_template"])
    unknown["columns"][-1]["role"] = "unknown_authored_role"
    with pytest.raises(ValueError, match="role_unresolved|columns_changed"):
        assignment.canonical_generated_column_slots(unknown, col_count=worksheet.max_column)
    workbook.close()


def test_original_base_uses_captured_template_menu_and_json_timestamp(connected_physical_pipeline, monkeypatch):
    order_id, _, start, _, _ = connected_physical_pipeline
    monkeypatch.setattr(order_service, "_resolve_effective_sheet_template", lambda **kwargs: pytest.fail("late template selection"))
    monkeypatch.setattr(order_service, "_build_monthly_menu_diagnostics", lambda **kwargs: pytest.fail("late monthly source read"))
    config_now = deepcopy(start["configured_input"])
    config_now["facility"]["fax_template"]["columns"][-1]["diet_type"] = "soft"
    _publish_repair_facility(config_now)
    monkeypatch.setattr(order_service, "_build_position_menu_entries_safe", lambda *args, **kwargs: [])
    base, error = order_service._build_hakodate_weekly_menu_base_sheet(order_id, processing_read_input=start)
    assert error is None
    assert "qty.regular_x" in base["fields"] and "qty.soft_x" not in base["fields"]
    assert base["physical_menu_row_ids"] == ["2026-10-01__lunch__0", "2026-10-02__lunch__0", "2026-10-03__lunch__0"]
    menu_index = base["fields"].index("menu")
    assert [row[menu_index] for row in base["rows"]] == ["own-1", "own-2", "own-3"]


def _slot_draft():
    return {"fields": ["date", "daypart", "menu", "qty.regular_x"], "rows": [["2026-10-01", "lunch", "same-menu", 0], ["2026-10-01", "lunch", "same-menu", None]], "row_ids": ["operator-a", "operator-b"], "physical_menu_row_ids": ["2026-10-01__lunch__0", "2026-10-01__lunch__1"]}


def test_saved_draft_rebase_preserves_distinct_original_slots_and_edits():
    current = _slot_draft()
    fresh = deepcopy(current)
    fresh["rows"] = [fresh["rows"][1], fresh["rows"][0]]
    fresh["physical_menu_row_ids"].reverse()
    fresh["row_ids"] = ["generated-b", "generated-a"]
    for row in fresh["rows"]:
        row[-1] = 99
    merged = order_service._merge_saved_draft_into_semantic_sheet(current, fresh)
    assert merged["row_ids"] == ["operator-b", "operator-a"]
    assert merged["physical_menu_row_ids"] == fresh["physical_menu_row_ids"]
    assert [row[-1] for row in merged["rows"]] == [None, 0]
    assert current == _slot_draft()
    # Explicit manual rows are valid without inventing monthly physical slots.
    manual = deepcopy(current)
    del manual["physical_menu_row_ids"]
    assert order_service._merge_saved_draft_into_semantic_sheet(manual, deepcopy(manual))["rows"] == manual["rows"]


@pytest.mark.parametrize("fault", ["missing", "short", "duplicate_physical", "duplicate_row", "foreign", "null"])
def test_actual_rebase_identity_failure_returns_original_blocker_without_write(monkeypatch, fault):
    current = _slot_draft()
    fresh = deepcopy(current)
    fresh["fields"].append("remarks")
    for row in fresh["rows"]:
        row.append("")
    if fault == "missing":
        del current["row_ids"]
    elif fault == "short":
        fresh["physical_menu_row_ids"].pop()
    elif fault == "duplicate_physical":
        current["physical_menu_row_ids"][1] = current["physical_menu_row_ids"][0]
    elif fault == "duplicate_row":
        fresh["row_ids"][1] = fresh["row_ids"][0]
    elif fault == "foreign":
        fresh["physical_menu_row_ids"][1] = "2026-10-01__lunch__2"
    else:
        current["row_ids"][0] = None
    record = {"draft_sheet_json": current, "draft_state": "draft_ready", "blockers_json": []}
    original = deepcopy(record)
    monkeypatch.setattr(order_service, "_draft_record_requires_current_sheet_semantic_rebase", lambda *args, **kwargs: (True, fresh))
    monkeypatch.setattr(order_service, "_persist_sheet_draft_with_base_evidence", lambda **kwargs: pytest.fail("identity failure must not write"))
    result = order_service._rebase_draft_record_to_facility_schema("own-order", record, edited_by="current-sheet-refresh", synchronous_write_input={"draft": deepcopy(record), "evidence": {"id": "own-evidence"}})
    assert result["draft_state"] == "draft_blocked"
    assert result["blockers_json"][0].startswith("saved_draft_row_identity_")
    assert result["draft_sheet_json"] == original["draft_sheet_json"]
    assert record == original


def test_saved_draft_quantity_sibling_uses_same_original_identity():
    current = _slot_draft()
    fresh = deepcopy(current)
    fresh["rows"].reverse()
    fresh["row_ids"].reverse()
    fresh["physical_menu_row_ids"].reverse()
    for row in fresh["rows"]:
        row[-1] = 77
    current["rows"][1][-1] = 12
    merged = order_service._merge_current_draft_quantity_values_into_semantic_sheet(current, fresh, copy_shared_quantity_values=True)
    assert [row[-1] for row in merged["rows"]] == [12, 0]
    fresh["physical_menu_row_ids"][1] = fresh["physical_menu_row_ids"][0]
    with pytest.raises(ValueError, match="identity_duplicate"):
        order_service._merge_current_draft_quantity_values_into_semantic_sheet(current, fresh, copy_shared_quantity_values=True)


def test_current_context_refresh_stops_before_current_state_adoption(monkeypatch):
    sheet = _slot_draft()
    sheet["row_ids"][1] = sheet["row_ids"][0]
    record = {"draft_sheet_json": sheet, "draft_state": "draft_ready", "blockers_json": []}
    monkeypatch.setattr(order_service, "_capture_synchronous_cache_write_input", lambda *args: {"draft": deepcopy(record), "evidence": {"id": "own-evidence"}})
    monkeypatch.setattr(order_service, "get_order_by_id", lambda *args, **kwargs: {"id": "own-order"})
    monkeypatch.setattr(order_service, "_draft_record_is_authoritative_current_sheet", lambda draft: True)
    monkeypatch.setattr(order_service, "_draft_record_requires_current_sheet_semantic_rebase", lambda *args, **kwargs: (True, _slot_draft()))
    monkeypatch.setattr(order_service, "_persist_sheet_draft_with_base_evidence", lambda **kwargs: pytest.fail("must not write draft"))
    monkeypatch.setattr(order_service.order_current_state_service, "persist_current_state", lambda **kwargs: pytest.fail("must not adopt current state"))
    with pytest.raises(ValueError, match="identity_duplicate"):
        order_service.refresh_current_sheet_context("own-order")
    assert record["draft_sheet_json"] == sheet


@pytest.fixture
def authored_dirty_rebuild(connected_physical_pipeline, monkeypatch, request):
    from src.models.order_sheet_draft import OrderSheetDraft
    order_id, _, original, _, _ = connected_physical_pipeline
    monthly_entries = deepcopy(original["physical_menu_entries"])
    for entry in monthly_entries:
        entry["name"] = entry["menu_name"] = "same-menu"
    monkeypatch.setattr(order_service.menu_service, "get_menu_entries_for_facility", lambda *args, **kwargs: deepcopy(monthly_entries))
    original = order_service._capture_synchronous_cache_write_input(order_id)
    edit = getattr(request, "param", None)
    if edit is not None:
        authored = deepcopy(original["configured_input"])
        authored["facility"]["fax_template"]["columns"].append({"column_id": "remarks", "index": 4, "source_index": 4, "role": "remarks"})
        config_service._validate_authored_facility_config(authored["facility"], "own")
        monkeypatch.setattr(config_service, "capture_authored_processing_config", lambda *args: deepcopy(authored))
        original = order_service._capture_synchronous_cache_write_input(order_id)
    old_sheet, error = order_service._build_hakodate_weekly_menu_base_sheet(order_id, synchronous_write_input=original)
    assert error is None
    quantity_index = old_sheet["fields"].index("qty.regular_x")
    for row, value in zip(old_sheet["rows"], [0, None, ""]):
        row[quantity_index] = value
    if edit is not None:
        remarks_index = old_sheet["fields"].index("remarks")
        original_shell = {field: [row[old_sheet["fields"].index(field)] for row in old_sheet["rows"]] for field in ("date_mmdd", "daypart", "menu")}
        for row in old_sheet["rows"]:
            row[remarks_index] = edit["remarks"]
            for field in original_shell:
                row[old_sheet["fields"].index(field)] = "stale-operator-shell"
    draft_id = str(uuid4())
    with session_scope() as session:
        session.add(OrderSheetDraft(id=draft_id, order_id=order_id, base_evidence_run_id=original["evidence_run_id"], draft_sheet_json=old_sheet, draft_state="draft_ready", blockers_json=[], warnings_json=["operator_review_pending"], edited_by="operator"))
        session.flush()
        session.get(OrderWorkflowState, order_id).draft_id = draft_id
    authored = deepcopy(original["configured_input"])
    if edit is None:
        authored["facility"]["fax_template"]["columns"].append({"column_id": "remarks", "index": 4, "source_index": 4, "role": "remarks"})
    else:
        authored["facility"]["fax_template"]["columns"].append({"column_id": "aux", "index": 5, "source_index": 5, "role": "aux"})
    config_service._validate_authored_facility_config(authored["facility"], "own")
    monkeypatch.setattr(config_service, "capture_authored_processing_config", lambda *args: deepcopy(authored))
    captured = order_service._capture_synchronous_cache_write_input(order_id)
    assert captured["draft"]["id"] == draft_id
    assert not order_service.apply_gate_service.has_clean_saved_draft(captured["draft"])
    if edit is not None:
        actual_row_producer = order_service._build_rows_from_menu_entries

        def rows_with_external_remarks(**kwargs):
            rows, diagnostics = actual_row_producer(**kwargs)
            # External field-value adapter only: identities, shell, and rows
            # still come from the real authored monthly-menu producer.
            index = kwargs["fields"].index("remarks")
            for item in rows:
                item["values"][index] = "fresh-old-remarks"
            return rows, diagnostics

        monkeypatch.setattr(order_service, "_build_rows_from_menu_entries", rows_with_external_remarks)
    # The intake adapter calls the actual authored weekly row producer, rather
    # than returning a hand-built fresh sheet or replacing the rebase predicate.
    def own_sheet(request_order_id, **kwargs):
        assert request_order_id == order_id
        assert kwargs["evidence_run_override"] == captured["evidence"]
        assert kwargs["use_saved_draft"] is False
        return order_service._build_hakodate_weekly_menu_base_sheet(order_id, synchronous_write_input=captured)
    monkeypatch.setattr(order_service, "get_ocr_sheet", own_sheet)
    monkeypatch.setattr(order_service, "_capture_synchronous_cache_write_input", lambda *args: captured)
    writes = []
    def before_persist(**kwargs):
        sheet = kwargs["draft_sheet_json"]
        assert kwargs["synchronous_write_input"] is captured
        index = sheet["fields"].index("qty.regular_x")
        expected_quantities = ["", "", ""] if edit is not None and edit["blank_quantities"] else [0, None, ""]
        assert [row[index] for row in sheet["rows"]] == expected_quantities
        if edit is not None:
            remarks_index = sheet["fields"].index("remarks")
            for row in sheet["rows"]:
                assert type(row[remarks_index]) is type(edit["remarks"])
                assert row[remarks_index] == edit["remarks"]
            for field in ("date_mmdd", "daypart", "menu"):
                new_index = sheet["fields"].index(field)
                assert [row[new_index] for row in sheet["rows"]] == original_shell[field]
        assert sheet["physical_menu_row_ids"] == old_sheet["physical_menu_row_ids"]
        assert sheet["physical_menu_row_count"] == old_sheet["physical_menu_row_count"]
        assert sheet["row_ids"] == old_sheet["row_ids"]
        writes.append(deepcopy(kwargs))
        return {"draft_sheet_json": deepcopy(sheet), "draft_state": kwargs["draft_state"]}
    monkeypatch.setattr(order_service, "_persist_sheet_draft_with_base_evidence", before_persist)
    monkeypatch.setattr(order_service.workflow_state_service, "refresh_workflow_state", lambda *args: None)
    return order_id, captured, old_sheet, writes


@pytest.mark.parametrize("authored_dirty_rebuild", [{"remarks": value, "path": path, "blank_quantities": blank} for value in ("", None, 0, "operator-current-text") for path, blank in (("rebase", False), ("facility_rebuild", False), ("weekly_overwrite", False), ("facility_rebuild", True), ("weekly_overwrite", True))], indirect=True)
@pytest.mark.parametrize("connected_physical_pipeline", [True], indirect=True)
def test_actual_write_callers_preserve_explicit_remarks(authored_dirty_rebuild, request):
    edit = request.node.callspec.params["authored_dirty_rebuild"]
    path, blank_quantities = edit["path"], edit["blank_quantities"]
    order_id, captured, original, writes = authored_dirty_rebuild
    required, fresh = order_service._draft_record_requires_current_sheet_semantic_rebase(order_id, captured["draft"], evidence_run_override=captured["evidence"])
    assert required
    index = fresh["fields"].index("remarks")
    assert [row[index] for row in fresh["rows"]] == ["fresh-old-remarks"] * 3
    if path == "rebase":
        result = order_service._rebase_draft_record_to_facility_schema(order_id, captured["draft"], edited_by="current-sheet-refresh", synchronous_write_input=captured)
    elif path == "facility_rebuild":
        result = order_service._rebuild_current_sheet_for_facility_schema(order_id, edited_by="operator", synchronous_write_input=captured, blank_quantities=blank_quantities)
    else:
        result, error = order_service.force_overwrite_current_sheet_with_weekly_menu(order_id, blank_quantities=blank_quantities)
        assert error is None
    assert result and len(writes) == 1
    assert captured["draft"]["draft_sheet_json"] == original


@pytest.mark.parametrize("path", ["rebase", "facility_rebuild", "weekly_overwrite"])
@pytest.mark.parametrize("connected_physical_pipeline", [True], indirect=True)
def test_actual_authored_producer_dirty_edits_reach_preservation_boundary(authored_dirty_rebuild, path):
    order_id, captured, original, writes = authored_dirty_rebuild
    if path == "rebase":
        required, fresh = order_service._draft_record_requires_current_sheet_semantic_rebase(order_id, captured["draft"], evidence_run_override=captured["evidence"])
        assert required and "remarks" in fresh["fields"]
        result = order_service._rebase_draft_record_to_facility_schema(order_id, captured["draft"], edited_by="current-sheet-refresh", synchronous_write_input=captured)
    elif path == "facility_rebuild":
        result = order_service._rebuild_current_sheet_for_facility_schema(order_id, edited_by="operator", synchronous_write_input=captured)
    else:
        result, error = order_service.force_overwrite_current_sheet_with_weekly_menu(order_id)
        assert error is None
    assert result and len(writes) == 1
    assert captured["draft"]["draft_sheet_json"] == original
    assert len(set(original["physical_menu_row_ids"])) == 3
    for field in ("date_mmdd", "daypart", "menu"):
        index = original["fields"].index(field)
        assert len({row[index] for row in original["rows"]}) == 1


@pytest.mark.parametrize("fault", ["missing_ids", "duplicate", "missing_rows", "empty_rows", "foreign", "short_row", "unknown_edit"])
@pytest.mark.parametrize("path", ["facility_rebuild", "weekly_overwrite"])
@pytest.mark.parametrize("blank_quantities", [False, True])
@pytest.mark.parametrize("connected_physical_pipeline", [True], indirect=True)
def test_actual_facility_rebuild_rejects_invalid_identity_before_persist(authored_dirty_rebuild, fault, path, blank_quantities):
    order_id, captured, _, writes = authored_dirty_rebuild
    current = captured["draft"]["draft_sheet_json"]
    if fault == "missing_ids":
        del current["row_ids"]
    elif fault == "duplicate":
        current["physical_menu_row_ids"][1] = current["physical_menu_row_ids"][0]
    elif fault == "missing_rows":
        del current["rows"]
    elif fault == "empty_rows":
        current["rows"] = []
    elif fault == "foreign":
        current["physical_menu_row_ids"][1] = "2026-10-02__lunch__9"
    elif fault == "short_row":
        current["rows"][0].pop()
    else:
        current["fields"].append("operator.extra")
        for row in current["rows"]:
            row.append(None)
    with pytest.raises(ValueError, match="saved_draft_row_identity_"):
        if path == "facility_rebuild":
            order_service._rebuild_current_sheet_for_facility_schema(order_id, edited_by="operator", synchronous_write_input=captured, blank_quantities=blank_quantities)
        else:
            order_service.force_overwrite_current_sheet_with_weekly_menu(order_id, blank_quantities=blank_quantities)
    assert writes == []


@pytest.mark.parametrize("copy_shared", [False, True])
@pytest.mark.parametrize("fault", ["duplicate", "missing_ids", "missing_rows", "empty_rows", "foreign", "no_quantity", "unmapped_null", "short_row"])
def test_preservation_alternatives_validate_before_any_save(copy_shared, fault):
    current, fresh = _slot_draft(), _slot_draft()
    if fault == "duplicate":
        current["physical_menu_row_ids"][1] = current["physical_menu_row_ids"][0]
    elif fault == "missing_ids":
        del current["row_ids"]
    elif fault == "missing_rows":
        del current["rows"]
    elif fault == "empty_rows":
        fresh["rows"] = []
    elif fault == "foreign":
        fresh["physical_menu_row_ids"][0] = "2026-10-01__lunch__2"
    elif fault == "no_quantity":
        for sheet in (current, fresh):
            sheet["fields"].pop()
            for row in sheet["rows"]:
                row.pop()
        del fresh["row_ids"]
    elif fault == "unmapped_null":
        current["rows"][0][-1] = None
        fresh["fields"][-1] = "qty.soft_x"
    else:
        current["rows"][0].pop()
    with pytest.raises(ValueError, match="saved_draft_row_identity_"):
        order_service._merge_current_draft_quantity_values_into_semantic_sheet(current, fresh, copy_shared_quantity_values=copy_shared)


@pytest.mark.parametrize("copy_shared", [False, True])
def test_same_columns_preserve_raw_operator_zero_null_blank(copy_shared):
    current = _slot_draft()
    current["rows"].append(["2026-10-01", "lunch", "same-menu", ""])
    current["row_ids"].append("operator-c")
    current["physical_menu_row_ids"].append("2026-10-01__lunch__2")
    current["physical_menu_row_count"] = 3
    fresh = deepcopy(current)
    for row in fresh["rows"]:
        row[-1] = 99
    merged = order_service._merge_current_draft_quantity_values_into_semantic_sheet(current, fresh, copy_shared_quantity_values=copy_shared)
    assert [row[-1] for row in merged["rows"]] == [0, None, ""]
    assert merged["physical_menu_row_count"] == 3


@pytest.fixture
def accepted_output_writer(connected_physical_pipeline, monkeypatch, tmp_path):
    from src.models.order_sheet_draft import OrderSheetDraft
    from src.services import output_builder as output, order_workflow_v2_service as workflow, order_output_artifact_service as artifacts
    order_id, _, original, _, _ = connected_physical_pipeline
    draft_id = str(uuid4())
    configured = original["configured_input"]
    with session_scope() as session:
        state = session.get(OrderWorkflowState, order_id)
        template_id = session.get(Order, order_id).template_version_id
        lineage.verify_projection(lineage.read_projection(session, template_id))
        base_input = order_service._capture_synchronous_cache_write_input(order_id, write_session=session)
        base_sheet, error = order_service._build_hakodate_weekly_menu_base_sheet(order_id, synchronous_write_input=base_input)
        assert error is None
        quantity_index = base_sheet["fields"].index("qty.regular_x")
        for row, value in zip(base_sheet["rows"], (0, None, None)):
            row[quantity_index] = value
        session.add(OrderSheetDraft(id=draft_id, order_id=order_id, template_version_id=template_id, base_evidence_run_id=state.evidence_run_id, draft_sheet_json=base_sheet, draft_state="draft_ready"))
        session.flush()
        state.draft_id = draft_id
    # The real bagging producer constructs the accepted output receipt. Only
    # external menu-master values and file writers are adapters below.
    monkeypatch.setattr(order_service.menu_service, "resolve_menu_defaults", lambda *args, **kwargs: {})
    candidate = {"configured_input": deepcopy(configured), "lines": [{"date": entry["menu_date"], "daypart": entry["daypart_key"], "menu_name": entry["menu_name"], "menu_category": "主菜", "diet_type": "regular", "area_id": "x", "quantity_original": value, "quantity_corrected": value, "source_row_index": index, "menu_qty_per_serving": 1, "menu_unit_type": "g"} for index, (entry, value) in enumerate(zip(base_input["physical_menu_entries"], (0, None, None)))]}
    with session_scope() as session:
        start = order_service._capture_synchronous_cache_write_input(order_id, write_session=session)
        start.update(output_ocr_menu_entries=[], output_ocr_menu_defaults={})
        receipt = workflow._build_bagging_result_payload(order=session.get(Order, order_id), saved_sheet=session.get(OrderSheetDraft, draft_id), materialization_candidate=candidate, processing_input=start)
        receipt = order_service._json_snapshot(receipt)
        artifacts.save_bagging_result_artifact(session, payload=receipt)
        state = session.get(OrderWorkflowState, order_id)
        state.secondary_actions_json = {"workflow_v2": {"bagging_result_id": receipt["bagging_result_id"], "bagging_input_proof": {"bagging_result_id": receipt["bagging_result_id"], "materialization_digest": receipt["materialization_digest"], "owned_revision": state.revision, "start_input": order_service._json_snapshot({key: value for key, value in start.items() if key not in {"evidence", "draft"}})}}}
        assert workflow._finalize_workflow_mutation(session, state) is None
    monkeypatch.setattr(output, "OUTPUT_DIR", tmp_path / "outputs")
    generated = []
    def external_writer(path, *args, **kwargs):
        path.write_bytes(b"own-generated-artifact")
        generated.append(path)
    monkeypatch.setattr(output, "_write_label_csv", external_writer)
    monkeypatch.setattr(output, "_write_delivery_note", external_writer)
    monkeypatch.setattr(output, "_write_aggregate_csv", external_writer)
    return order_id, output, generated


@pytest.mark.parametrize("writer", ["outputs", "bags"])
@pytest.mark.parametrize("change", [None, "revision", "receipt", "sheet", "typed_snapshot", "typed_config"])
def test_actual_output_calculation_and_same_session_adoption(accepted_output_writer, writer, change):
    from src.models.output import Bag, DeliveryNote
    order_id, output, generated = accepted_output_writer
    with session_scope() as session:
        session.add(Bag(id="original-current-bag", order_id=order_id, quantity=0))
    with (pytest.raises(ValueError) if change is not None else nullcontext()):
        with session_scope() as session:
            original = output._capture_output_adoption(order_id, session)
            context = output._prepare_output_context(order_id, original_adoption=original)
            assert context["accepted_candidate"] is original["candidate"]
            assert [line["quantity_original"] for line in context["order_lines"]] == [0, None, None]
            if change == "typed_snapshot":
                original["write_input"]["workflow_revision"] = True
            elif change == "typed_config":
                original["write_input"]["configured_input"]["facility"]["fax_template"]["page"] = True
            elif change is not None:
                from src.models.order_output_artifact import OrderBaggingResult
                from src.models.order_sheet_draft import OrderSheetDraft
                with session_scope() as concurrent:
                    state = concurrent.get(OrderWorkflowState, order_id)
                    if change == "revision":
                        state.headline = "concurrent operator command"
                    elif change == "sheet":
                        concurrent.get(OrderSheetDraft, state.draft_id).draft_sheet_json = {"changed": None}
                    else:
                        artifact = concurrent.get(OrderBaggingResult, original["bagging_receipts"][0]["id"])
                        artifact.payload_json = {**artifact.payload_json, "status": "replaced"}
            if writer == "outputs":
                result = output.build_outputs(order_id, write_session=session, original_adoption=original)
                assert all(path.parent.name != order_id for path in generated)
                assert session.query(DeliveryNote).filter_by(order_id=order_id).one().file_uri == result["delivery_note"]
            else:
                result = output.rebuild_bags(order_id, write_session=session, original_adoption=original)
            assert result["order_id"] == order_id
    with session_scope() as session:
        assert (session.get(Bag, "original-current-bag") is not None) is (change is not None)


def test_output_publication_failure_and_preview_never_replace_current(accepted_output_writer, monkeypatch):
    from src.models.output import Bag, DeliveryNote
    order_id, output, generated = accepted_output_writer
    current = output.OUTPUT_DIR / f"{order_id}_delivery.xlsx"
    current.parent.mkdir(parents=True)
    current.write_bytes(b"previous-current")
    with session_scope() as session:
        session.add(Bag(id="original-current-bag", order_id=order_id, quantity=0))
    def failed_export(*args, **kwargs):
        raise OSError("external export failure")
    monkeypatch.setattr(output, "_write_delivery_note", failed_export)
    with pytest.raises(OSError), session_scope() as session:
        original = output._capture_output_adoption(order_id, session)
        output.build_outputs(order_id, write_session=session, original_adoption=original)
    assert current.read_bytes() == b"previous-current"
    with session_scope() as session:
        assert session.get(Bag, "original-current-bag") is not None
        assert session.query(DeliveryNote).filter_by(order_id=order_id).count() == 0
    preview = output.build_output_preview(order_id, "labels")
    assert preview["labels"] != str(output.OUTPUT_DIR / f"{order_id}_labels.csv")
    assert current.read_bytes() == b"previous-current"


def test_missing_output_ownership_or_acceptance_stops_before_writing(accepted_output_writer):
    order_id, output, generated = accepted_output_writer
    with pytest.raises(ValueError, match="original ownership input missing"):
        output.rebuild_bags(order_id)
    with session_scope() as session:
        session.get(OrderWorkflowState, order_id).secondary_actions_json = {}
    with pytest.raises(ValueError, match="acceptance receipt missing"):
        output.build_outputs(order_id)
    assert generated == []


@pytest.mark.parametrize("command", ["upsert", "bulk", "delete"])
def test_actual_override_api_blocks_before_any_committed_override(command, monkeypatch):
    from src.services import daily_output_override_service as overrides
    for method in ("upsert_override", "upsert_overrides", "delete_override"):
        monkeypatch.setattr(overrides, method, lambda *args, **kwargs: pytest.fail("override write before original multiorder ownership binding"))
    with pytest.raises(HTTPException) as blocked:
        if command == "delete":
            orders.delete_daily_output_override("original-override")
        elif command == "bulk":
            orders.upsert_daily_output_override_bulk(orders.DailyOutputOverrideBulkUpsertBody(date="2026-10-01", menu_name="own", unit_type="g", qty_per_serving=1))
        else:
            orders.upsert_daily_output_override(orders.DailyOutputOverrideUpsertBody(date="2026-10-01", facility_id="fixture", menu_name="own", unit_type="g", qty_per_serving=1))
    assert blocked.value.status_code == 409
    assert blocked.value.detail["code"] == "daily_output_original_ownership_required"


def test_actual_internal_override_scope_uses_caller_transaction(accepted_output_writer, monkeypatch):
    from datetime import date
    from src.models.output import Bag
    order_id, output, _ = accepted_output_writer
    # External order-index adapter, not an ownership/candidate validator.
    monkeypatch.setattr(order_service, "list_orders_by_line_date", lambda *args, **kwargs: [{"id": order_id}])
    with session_scope() as session:
        session.add(Bag(id="original-current-bag", order_id=order_id, quantity=0))
    with pytest.raises(RuntimeError, match="caller rollback"):
        with session_scope() as session:
            original = output._capture_output_adoption(order_id, session)
            affected = order_service._rebuild_bags_for_override_scope(date(2026, 10, 1), "fixture", write_session=session, original_adoptions=[original])
            assert affected == [order_id]
            assert session.get(Bag, "original-current-bag") is None
            raise RuntimeError("caller rollback")
    with session_scope() as session:
        assert session.get(Bag, "original-current-bag") is not None


def test_actual_output_export_race_cannot_publish_old_receipt(accepted_output_writer, monkeypatch):
    from src.models.output import Bag, DeliveryNote
    order_id, output, generated = accepted_output_writer
    with session_scope() as session:
        session.add(Bag(id="original-current-bag", order_id=order_id, quantity=0))
    external_writer = output._write_label_csv
    def changed_during_export(path, *args, **kwargs):
        external_writer(path, *args, **kwargs)
        with session_scope() as concurrent:
            concurrent.get(OrderWorkflowState, order_id).headline = "changed after calculation"
    monkeypatch.setattr(output, "_write_label_csv", changed_during_export)
    with pytest.raises(ValueError, match="ownership input conflict"):
        output.build_outputs(order_id)
    assert len(generated) == 3
    assert all(path.parent != output.OUTPUT_DIR for path in generated)
    with session_scope() as session:
        assert session.get(Bag, "original-current-bag") is not None
        assert session.query(DeliveryNote).filter_by(order_id=order_id).count() == 0


@pytest.fixture
def two_accepted_output_orders(accepted_output_writer, request, monkeypatch):
    from src.models.order_sheet_draft import OrderSheetDraft
    from src.models.output import Bag
    from src.services import order_workflow_v2_service as workflow, order_output_artifact_service as artifacts
    first_id, output, generated = accepted_output_writer
    options = getattr(request, "param", {})
    if options.get("setting"):
        with session_scope() as session:
            configured = config_service.capture_authored_processing_config(session.get(Order, first_id).facility_code)
        if options["setting"] == "mapping":
            configured["facility"]["fax_template"]["columns"].append({"column_id": "soft", "index": 4, "source_index": 4, "role": "quantity", "diet_type": "soft"})
        else:
            configured["facility"]["label_profile"] = {"label_fields": ["menu_name", "quantity"]}
        config_service._validate_authored_facility_config(configured["facility"], "own")
        monkeypatch.setattr(config_service, "capture_authored_processing_config", lambda _: deepcopy(configured))
    second_values = options.get("values", (0, None, None))
    second_id, evidence_id, draft_id = "zz-second-owned-output", str(uuid4()), str(uuid4())
    with session_scope() as session:
        first = session.get(Order, first_id)
        first_state = session.get(OrderWorkflowState, first_id)
        source = session.get(OrderOcrEvidenceRun, first_state.evidence_run_id)
        # Two explicit orders of the same authored fixture PDF/facility/period;
        # second receipt IDs are produced independently, never copied/relabelled.
        session.add(Order(id=second_id, facility_code=first.facility_code, week_code=first.week_code, received_at=first.received_at, current_document_id=first.current_document_id, document_uri=first.document_uri, message_id=first.message_id, template_version_id=first.template_version_id))
        session.flush()
        session.add(OrderOcrEvidenceRun(id=evidence_id, order_id=second_id, template_version_id=first_state.template_version_id, schema_version=source.schema_version, status=source.status, payload_json=deepcopy(source.payload_json), artifact_digest=source.artifact_digest))
        session.flush()
        state = OrderWorkflowState(order_id=second_id, evidence_run_id=evidence_id, template_version_id=first_state.template_version_id, state="output_review")
        session.add(state)
        session.flush()
        start = order_service._capture_synchronous_cache_write_input(second_id, write_session=session)
        sheet, error = order_service._build_hakodate_weekly_menu_base_sheet(second_id, synchronous_write_input=start)
        assert error is None
        qty_index = sheet["fields"].index("qty.regular_x")
        for row, value in zip(sheet["rows"], second_values):
            row[qty_index] = value
        draft = OrderSheetDraft(id=draft_id, order_id=second_id, template_version_id=state.template_version_id, base_evidence_run_id=evidence_id, draft_sheet_json=sheet, draft_state="draft_ready")
        session.add(draft)
        session.flush()
        state.draft_id = draft_id
        candidate = {"configured_input": deepcopy(start["configured_input"]), "lines": [{"date": entry["menu_date"], "daypart": entry["daypart_key"], "menu_name": entry["menu_name"], "menu_category": "主菜", "diet_type": "regular", "area_id": "x", "quantity_original": value, "quantity_corrected": value, "source_row_index": index, "menu_qty_per_serving": 1, "menu_unit_type": "g"} for index, (entry, value) in enumerate(zip(start["physical_menu_entries"], second_values))]}
        session.flush()
        start = order_service._capture_synchronous_cache_write_input(second_id, write_session=session)
        start.update(output_ocr_menu_entries=[], output_ocr_menu_defaults={})
        receipt = workflow._build_bagging_result_payload(order=session.get(Order, second_id), saved_sheet=draft, materialization_candidate=candidate, processing_input=start)
        artifacts.save_bagging_result_artifact(session, payload=order_service._json_snapshot(receipt))
        state.secondary_actions_json = {"workflow_v2": {"bagging_result_id": receipt["bagging_result_id"], "bagging_input_proof": {"bagging_result_id": receipt["bagging_result_id"], "materialization_digest": receipt["materialization_digest"], "owned_revision": state.revision, "start_input": order_service._json_snapshot({key: value for key, value in start.items() if key not in {"evidence", "draft"}})}}}
        assert workflow._finalize_workflow_mutation(session, state) is None
        for order_id in (first_id, second_id):
            session.add(Bag(id=f"prior-{order_id}", order_id=order_id, quantity=0))
    return sorted([first_id, second_id]), output


@pytest.mark.parametrize("later_conflict", [False, True])
def test_real_multiorder_scope_claims_each_original_once(two_accepted_output_orders, monkeypatch, later_conflict):
    from datetime import date
    from sqlalchemy import event
    from src.models.output import Bag
    order_ids, output = two_accepted_output_orders
    monkeypatch.setattr(order_service, "list_orders_by_line_date", lambda *args, **kwargs: [{"id": order_id} for order_id in reversed(order_ids)])
    claims = []
    def record_sql(connection, cursor, statement, parameters, context, executemany):
        normalized = " ".join(statement.lower().split())
        if normalized.startswith("update order_workflow_states set revision=order_workflow_states.revision"):
            claims.append(parameters)
    with (pytest.raises(ValueError, match="ownership input conflict") if later_conflict else nullcontext()):
        with session_scope() as session:
            originals = [output._capture_output_adoption(order_id, session) for order_id in order_ids]
            if later_conflict:
                with session_scope() as concurrent:
                    concurrent.get(OrderWorkflowState, order_ids[-1]).headline = "later order changed"
            event.listen(session.bind, "before_cursor_execute", record_sql)
            try:
                result = order_service._rebuild_bags_for_override_scopes(date(2026, 10, 1), ["fixture"], write_session=session, original_adoptions=originals)
                assert result == list(reversed(order_ids))
            finally:
                event.remove(session.bind, "before_cursor_execute", record_sql)
            assert len(claims) == len(order_ids)
    assert len(claims) == len(order_ids)
    with session_scope() as session:
        for order_id in order_ids:
            assert (session.get(Bag, f"prior-{order_id}") is not None) is later_conflict


@pytest.mark.parametrize("change", [None, "saved_sheet", "missing_proof", "typed_owned_revision"])
def test_output_rejects_preexisting_stale_bagging_lineage(accepted_output_writer, change):
    from src.models.order_sheet_draft import OrderSheetDraft
    order_id, output, generated = accepted_output_writer
    if change is not None:
        with session_scope() as session:
            state = session.get(OrderWorkflowState, order_id)
            if change == "saved_sheet":
                draft = session.get(OrderSheetDraft, state.draft_id)
                changed = deepcopy(draft.draft_sheet_json)
                changed["rows"][0][changed["fields"].index("qty.regular_x")] = 99
                draft.draft_sheet_json = changed
            else:
                meta = deepcopy(state.secondary_actions_json)
                if change == "missing_proof":
                    del meta["workflow_v2"]["bagging_input_proof"]
                else:
                    meta["workflow_v2"]["bagging_input_proof"]["owned_revision"] = True
                state.secondary_actions_json = meta
    with (pytest.raises(ValueError, match="original bagging") if change is not None else nullcontext()):
        output.build_outputs(order_id)
    if change is not None:
        assert generated == []


@pytest.mark.parametrize("change", [None, "snapshot", "saved_sheet"])
def test_actual_final_confirm_receipt_to_output_writer(accepted_output_writer, change):
    from src.models.order_confirmed_snapshot import OrderConfirmedSnapshot
    from src.models.order_sheet_draft import OrderSheetDraft
    from src.services import order_workflow_v2_service as workflow, order_output_artifact_service as artifacts
    order_id, output, generated = accepted_output_writer
    with session_scope() as session:
        state = session.get(OrderWorkflowState, order_id)
        bagging = workflow._current_bagging_result(session, state)
        bundle = workflow._build_output_bundle_payload(order=session.get(Order, order_id), bagging_result=bagging)
        artifacts.save_output_bundle_artifact(session, payload=bundle)
        meta = deepcopy(state.secondary_actions_json)
        meta["workflow_v2"]["output_bundle_id"] = bundle["output_bundle_id"]
        state.secondary_actions_json = meta
        state.state = "output_review"
        assert workflow._finalize_workflow_mutation(session, state) is None
    with session_scope() as session:
        revision = session.get(OrderWorkflowState, order_id).revision
    result, error = workflow.final_confirm(order_id, expected_revision=revision)
    assert error is None and result
    if change is not None:
        with session_scope() as session:
            state = session.get(OrderWorkflowState, order_id)
            if change == "snapshot":
                receipt = session.get(OrderConfirmedSnapshot, state.confirmed_snapshot_id)
                receipt.snapshot_json = {**receipt.snapshot_json, "saved_sheet": {"replaced": True}}
            else:
                draft = session.get(OrderSheetDraft, state.draft_id)
                draft.draft_sheet_json = {"replaced": None}
    with (pytest.raises(ValueError, match="original confirmed") if change is not None else nullcontext()):
        output.build_outputs(order_id)
    if change is not None:
        assert generated == []


@pytest.mark.parametrize("two_accepted_output_orders", [{"setting": "mapping"}, {"setting": "profile"}], indirect=True)
@pytest.mark.parametrize("bundle_type", ["delivery", "both", "labels", "labels_csv"])
def test_actual_daily_bundle_settings_mismatch_before_export(two_accepted_output_orders, monkeypatch, bundle_type):
    from src.api import outputs
    order_ids, output = two_accepted_output_orders
    monkeypatch.setattr(order_service, "list_orders_by_line_date", lambda *args, **kwargs: [{"id": identity} for identity in order_ids])
    monkeypatch.setattr(output.Workbook, "save", lambda *args: pytest.fail("incompatible settings reached export"))
    with pytest.raises(HTTPException) as rejected:
        outputs.download_daily_bundle("2026-10-01", bundle_type=bundle_type)
    assert rejected.value.status_code == 400
    assert "daily_bundle_accepted_settings_mismatch" in rejected.value.detail
    assert not output.OUTPUT_DIR.exists()


@pytest.mark.parametrize("two_accepted_output_orders", [{"values": (7, 0, None)}], indirect=True)
@pytest.mark.parametrize("bundle_type", ["delivery", "labels_csv"])
@pytest.mark.parametrize("overlap", [False, True])
def test_actual_daily_bundle_api_owned_files_and_all_orders(two_accepted_output_orders, monkeypatch, bundle_type, overlap):
    from pathlib import Path
    from src.api import outputs
    from openpyxl import Workbook
    from datetime import datetime as real_datetime
    from types import SimpleNamespace
    order_ids, output = two_accepted_output_orders
    monkeypatch.setattr(output, "datetime", SimpleNamespace(utcnow=lambda: real_datetime(2026, 10, 7, 12, 0, 0)))
    monkeypatch.setattr(order_service, "list_orders_by_line_date", lambda *args, status=None, **kwargs: [{"id": identity} for identity in (order_ids if status == "all" else order_ids[:1])])
    exported = []
    nested = []
    def second_request():
        response = outputs.download_daily_bundle("2026-10-01", bundle_type=bundle_type, status="first")
        nested.append((response, Path(response.path).read_bytes()))
    def delivery_export(path, rows, columns, *args):
        if overlap and not nested and not exported:
            exported.append((path, deepcopy(rows), deepcopy(columns)))
            second_request()
        else:
            exported.append((path, deepcopy(rows), deepcopy(columns)))
        workbook = Workbook()
        workbook.active.append([column["name"] for column in columns])
        for row in rows:
            workbook.active.append([str(row.get(column["name"])) for column in columns])
        workbook.save(path)
    monkeypatch.setattr(output, "_write_delivery_note", delivery_export)
    if overlap and bundle_type == "labels_csv":
        original_close = output.zipfile.ZipFile.close
        active = False
        def close_archive(archive):
            nonlocal active
            original_close(archive)
            if not active and not nested and archive.mode == "w":
                active = True
                second_request()
                active = False
        monkeypatch.setattr(output.zipfile.ZipFile, "close", close_archive)
    first = outputs.download_daily_bundle("2026-10-01", bundle_type=bundle_type, status="all")
    if not overlap:
        second_request()
    second, second_bytes = nested[0]
    assert Path(first.path).parent != Path(second.path).parent
    assert Path(first.path).name == Path(second.path).name
    assert first.headers["content-disposition"] == second.headers["content-disposition"]
    assert Path(second.path).read_bytes() == second_bytes
    assert Path(first.path).read_bytes() != second_bytes
    assert first.headers["x-daily-bundle-error-orders"] == "0"
    if bundle_type == "delivery":
        own = [entry for entry in exported if entry[0].parent == Path(first.path).parent]
        assert len(own) == 1
        quantity_names = [column["name"] for column in own[0][2] if column.get("source") == "quantity"]
        assert sum(row.get(name) or 0 for row in own[0][1] for name in quantity_names) == 7
    for identity in order_ids:
        ctx = output._prepare_output_context_for_bundle(identity, include_bags=True, include_ocr_menu_meta=False, include_expanded_copy=False)
        expected = [0, None, None] if identity == order_ids[0] else [7, 0, None]
        assert [line["quantity_original"] for line in ctx["order_lines"]] == expected
    from datetime import date
    path, summary = output.build_daily_output_bundle(date(2026, 10, 1), bundle_type=bundle_type, status="all")
    assert path not in {Path(first.path), Path(second.path)}
    assert summary["status_filter"] == "all" and summary["bundle_type"] == bundle_type
    assert summary["error_orders"] == 0
    assert summary["items"][0]["order_ids"] == order_ids


def test_bundle_context_typeerror_does_not_retry_default_options(monkeypatch):
    from src.services import output_builder as output
    calls = []
    def obsolete_contract(order_id, **kwargs):
        calls.append((order_id, kwargs))
        raise TypeError("include_ocr_menu_meta rejected")
    monkeypatch.setattr(output, "_prepare_output_context", obsolete_contract)
    with pytest.raises(TypeError, match="include_ocr_menu_meta rejected"):
        output._prepare_output_context_for_bundle("own", include_bags=False, include_ocr_menu_meta=False, include_expanded_copy=False, require_invoice=False)
    assert calls == [("own", {"include_bags": False, "include_ocr_menu_meta": False, "include_expanded_copy": False, "timings": None, "require_invoice": False})]
