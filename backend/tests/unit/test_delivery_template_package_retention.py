"""Actual writer/package preservation with disposable authored workbooks."""
from datetime import date
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import quote_sheetname
from openpyxl.drawing.image import Image as WorkbookImage
from openpyxl.worksheet.table import Table
from PIL import Image
import pytest

from src.services import output_builder as output


@pytest.mark.parametrize("slot", [False, True])
@pytest.mark.parametrize("days", [1, 2])
def test_writer_keeps_support_sheets_images_and_print_settings(tmp_path, slot, days):
    template = tmp_path / "authored.xlsx"
    picture = tmp_path / "authored.png"
    Image.new("RGB", (12, 12), (20, 120, 60)).save(picture)
    workbook = Workbook()
    helper = workbook.active
    helper.title = "Rules"
    helper["A1"] = "Authored supporting content"
    helper["B1"] = "=1+2"
    invoice = workbook.create_sheet("Invoice")
    invoice.append(["Date", "Meal", "Menu", "Qty"])
    invoice.append([None, "朝", "主" if slot else None, None])
    invoice.append([None, "昼", "主" if slot else None, None])
    invoice["G1"] = "='Invoice'!D2"
    invoice["H1"] = '="Invoice!D2"'
    invoice.print_area = "A1:D3"
    invoice.print_title_rows = "1:1"
    invoice.add_image(WorkbookImage(picture), "F1")
    workbook.active = 1
    helper.sheet_state = "hidden"
    workbook.save(template)
    workbook.close()
    original = template.read_bytes()
    original_digest = sha256(original).hexdigest()
    with ZipFile(BytesIO(original)) as archive:
        original_parts = {name: archive.read(name) for name in archive.namelist()}
    helper_path = output._worksheet_paths_by_name(original_parts)["Rules"]
    columns = [
        {"name": "Date", "source": "date"},
        {"name": "Meal", "source": "daypart"},
        {"name": "Menu", "source": "menu_name"},
        {"name": "Qty", "source": "quantity", "diet_type": "regular", "area_id": "X"},
    ]
    rows = [{"date": date(2026, 10, day), "daypart": "朝", "menu_name": "Authored meal",
             "menu_category": "主菜", "Qty": 7 + day} for day in range(1, days + 1)]
    target = tmp_path / "generated.xlsx"
    output._write_delivery_note(target, rows, columns, str(template), True, "Invoice", None)
    assert sha256(template.read_bytes()).hexdigest() == original_digest
    with ZipFile(target) as archive:
        assert archive.testzip() is None
        parts = {name: archive.read(name) for name in archive.namelist()}
    assert set(original_parts) - {"xl/calcChain.xml"} <= set(parts)
    assert parts[helper_path] == original_parts[helper_path]
    for name, content in original_parts.items():
        if name.startswith("xl/media/") or name.startswith("xl/drawings/"):
            assert parts[name] == content
    for name, content in parts.items():
        if name.endswith((".xml", ".rels")):
            ET.fromstring(content)
    result = load_workbook(target)
    try:
        assert result["Rules"].sheet_state == "hidden"
        assert result["Rules"]["B1"].value == "=1+2"
        invoices = [sheet for sheet in result.worksheets if sheet.title != "Rules"]
        assert len(invoices) == (days if slot else 1)
        for sheet in invoices:
            assert str(sheet.print_area) == quote_sheetname(sheet.title) + "!$A$1:$D$3"
            assert sheet.print_title_rows.replace("$", "") == "1:1"
            assert len(sheet._images) == 1
            assert sheet["G1"].value == "=" + quote_sheetname(sheet.title) + "!D2"
            assert sheet["H1"].value == '="Invoice!D2"'
        quantities = [sheet["D2"].value for sheet in invoices] if slot else [result["Invoice"].cell(2 + i, 4).value for i in range(days)]
        assert quantities == [7 + day for day in range(1, days + 1)]
    finally:
        result.close()


def test_table_clone_is_rejected_before_publishing_corrupt_package(tmp_path):
    template = tmp_path / "table.xlsx"
    workbook = Workbook()
    source = workbook.active
    source.title = "Invoice"
    source.append(["Menu", "Qty"])
    source.append(["Authored meal", 8])
    source.add_table(Table(displayName="AuthoredInvoice", ref="A1:B2"))
    workbook.save(template)
    original = template.read_bytes()
    target = tmp_path / "output.xlsx"
    target.write_bytes(b"existing-output-must-survive")
    workbook.copy_worksheet(source).title = "Second date"
    with pytest.raises(ValueError, match="^template_table_worksheet_clone_unsupported$"):
        output._patch_template_package_with_workbook_values(
            workbook, template_bytes=original, output_path=target,
            worksheet_sources={"Invoice": "Invoice", "Second date": "Invoice"},
        )
    assert template.read_bytes() == original
    assert target.read_bytes() == b"existing-output-must-survive"
    workbook.remove(workbook["Second date"])
    output._patch_template_package_with_workbook_values(
        workbook, template_bytes=original, output_path=target,
        worksheet_sources={"Invoice": "Invoice"},
    )
    result = load_workbook(target)
    try:
        assert list(result["Invoice"].tables) == ["AuthoredInvoice"]
        assert result["Invoice"]["B2"].value == 8
    finally:
        result.close()
        workbook.close()
