"""Synthetic period contracts with isolated outputs and no provider/user writes."""
from copy import deepcopy
from datetime import date, datetime, timedelta
from io import BytesIO
from unittest.mock import Mock

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy.engine import Engine

from src.api import outputs
from src.services import output_builder as builder


@pytest.fixture(autouse=True)
def no_database(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("test must not connect to any database")
    monkeypatch.setattr(Engine, "connect", forbidden)


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(outputs.router, prefix="/outputs")
    route = next(route for route in app.routes if route.path == "/outputs/weekly-weight")
    for dependency in route.dependant.dependencies:
        app.dependency_overrides[dependency.call] = lambda: {"role": "operator"}
    with TestClient(app) as client:
        yield client


@pytest.mark.parametrize("start,end", [
    ("2026-09-13", "2026-09-19"), ("2026-09-27", "2026-09-30"),
    ("2026-10-01", "2026-10-03"), ("2020-12-27", "2021-01-02"),
    ("2024-02-29", "2024-03-02"), ("2026-09-13", "2026-09-13"),
])
def test_exact_period_keeps_bounds(start, end):
    bounds = (date.fromisoformat(start), date.fromisoformat(end))
    assert builder.weekly_weight_period(start=bounds[0], end=bounds[1]) == bounds


@pytest.mark.parametrize("params", [
    {}, {"start": "2026-09-13"}, {"end": "2026-09-19"},
    {"start": "2026-09-19", "end": "2026-09-13"},
    {"start": "2026-09-13", "end": "2026-09-20"},
    {"start": "2026-09-19", "end": "2026-09-20"},
    {"start": "2026-02-29", "end": "2026-03-01"},
    {"start": "20260913", "end": "2026-09-19"},
    {"start": "", "end": "2026-09-19"},
    {"start": "2026-09-13T00:00:00Z", "end": "2026-09-19"},
    {"start": "2026-09-13", "end": "2026-09-19", "date": "2026-09-12"},
    {"start": "2026-10-01", "end": "2026-10-03", "date": "2026-09-27"},
    {"date": "2026-W38-1"}, {"date": "2026-09-31"}, {"date": "0000-01-01"},
])
def test_api_invalid_or_conflicting_period_stops_before_build(client, monkeypatch, params):
    build = Mock(side_effect=AssertionError("must reject before output"))
    monkeypatch.setattr(builder, "build_weekly_weight_summary_workbook", build)
    response = client.get("/outputs/weekly-weight", params=params)
    assert response.status_code == 400
    build.assert_not_called()


def test_api_range_and_consistent_date_are_not_silently_discarded(client, monkeypatch, tmp_path):
    file = tmp_path / "synthetic.xlsx"
    file.write_bytes(b"synthetic")
    build = Mock(return_value=file)
    monkeypatch.setattr(builder, "build_weekly_weight_summary_workbook", build)
    response = client.get("/outputs/weekly-weight", params={
        "start": "2026-10-01", "end": "2026-10-03", "date": "2026-10-02", "status": "確定",
    })
    assert response.status_code == 200
    build.assert_called_once_with(date(2026, 10, 2), status="確定", start=date(2026, 10, 1), end=date(2026, 10, 3))


def test_date_only_call_signature_and_sunday_week_stay_unchanged(client, monkeypatch, tmp_path):
    file = tmp_path / "synthetic.xlsx"
    file.write_bytes(b"synthetic")
    build = Mock(return_value=file)
    monkeypatch.setattr(builder, "build_weekly_weight_summary_workbook", build)
    assert client.get("/outputs/weekly-weight", params={"date": "2026-09-16"}).status_code == 200
    build.assert_called_once_with(date(2026, 9, 16), status=None)
    assert builder.weekly_weight_period(date(2026, 9, 16)) == (date(2026, 9, 13), date(2026, 9, 19))


def synthetic_rows(monkeypatch):
    dates = [date(2026, 9, 30), date(2026, 10, 1), date(2026, 10, 2), date(2026, 10, 3), date(2026, 10, 4)]
    lines = [{"date": day, "daypart": "昼", "menu_category": "主Ａ", "menu_name": "synthetic",
              "diet_type": "regular", "quantity_corrected": quantity,
              "menu_unit_type": "g", "menu_qty_per_serving": 100}
             for day, quantity in zip(dates, [99, 6, 0, 9, 99])]
    original = deepcopy(lines)
    calls = []
    def orders(day, status=None):
        calls.append((day, status))
        return [{"id": "synthetic-only"}]
    monkeypatch.setattr(builder.order_service, "list_orders_by_line_date", orders)
    monkeypatch.setattr(builder, "_prepare_output_context_for_bundle", lambda *args, **kwargs: {"order_lines": lines})
    return lines, original, calls


def test_real_collector_limits_dates_without_quantity_mutation(monkeypatch):
    lines, original, calls = synthetic_rows(monkeypatch)
    rows = builder._weekly_weight_collect_rows(None, start=date(2026, 10, 1), end=date(2026, 10, 3), status="確定")
    assert calls == [(date(2026, 10, day), "確定") for day in (1, 2, 3)]
    assert rows[(date(2026, 10, 1), "昼", "主Ａ")]["regular_quantity"] == 6
    assert rows[(date(2026, 10, 3), "昼", "主Ａ")]["regular_quantity"] == 9
    assert {key[0] for key in rows} == {date(2026, 10, 1), date(2026, 10, 3)}
    assert rows[(date(2026, 10, 1), "昼", "主Ａ")]["regular_amounts"] == {"g": 600.0}
    assert lines == original


@pytest.mark.parametrize("reference", [False, True])
def test_excel_exact_period_title_dates_and_quantities(client, monkeypatch, tmp_path, reference):
    lines, original, calls = synthetic_rows(monkeypatch)
    monkeypatch.setattr(builder, "OUTPUT_DIR", tmp_path)
    if reference:
        # A disposable real-layout copy proves old formula/date slots are cleared.
        shell = builder._build_weekly_weight_workbook_shell(date(2026, 9, 27))
        shell.worksheets[0]["A59"] = "=A51+1"
        template = tmp_path / "reference.xlsx"
        shell.save(template)
        monkeypatch.setattr(builder, "_weekly_weight_reference_layout_path", lambda: template)
    else:
        monkeypatch.setattr(builder, "_weekly_weight_reference_layout_path", lambda: None)
    response = client.get("/outputs/weekly-weight", params={"start": "2026-10-01", "end": "2026-10-03", "status": "確定"})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    sheet = workbook.worksheets[0]
    assert sheet.title == "10月1日～10月3日"
    assert [sheet.cell(row, 1).value for row in (11, 19, 27)] == [datetime(2026, 10, day) for day in (1, 2, 3)]
    assert all(sheet.cell(row, 1).value is None for row in (35, 43, 51, 59))
    assert sheet.cell(13, 5).value == 6
    assert sheet.cell(29, 5).value == 9
    assert sum(value for row in sheet.iter_rows(min_row=11, max_row=66, min_col=5, max_col=5)
               for cell in row if isinstance(value := cell.value, (int, float))) == 15
    assert calls == [(date(2026, 10, day), "確定") for day in (1, 2, 3)]
    assert lines == original


def test_builder_pair_guard_runs_before_any_collection(monkeypatch):
    collect = Mock(side_effect=AssertionError("no collection on bad period"))
    monkeypatch.setattr(builder, "_weekly_weight_collect_rows", collect)
    with pytest.raises(ValueError, match="period_pair_required"):
        builder.build_weekly_weight_summary_workbook(start=date(2026, 9, 13))
    collect.assert_not_called()
