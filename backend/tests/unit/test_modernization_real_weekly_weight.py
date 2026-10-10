"""Real confirmed synthetic SQL input to mounted weekly-weight HTTP.

Manual-seeded OCR, AUTH_DISABLED and no startup lifespan: not native/auth/PG/
UI/live proof. No business reader, collector, renderer or validator is doubled.
The existing fixture's breakfast main category is not changed to fit the grid.
"""

from copy import deepcopy
from hashlib import sha256
from io import BytesIO

from openpyxl import load_workbook

from test_modernization_real_bagging_confirmation import (
    actual_bagging, owned_bagging as owned_bagging, revision,
)
from test_modernization_workflow_surface_parity import (
    database_snapshot, surface as surface,
)
from src.services import master_order_form_template_service as master
from src.services import order_workflow_v2_service as workflow
from src.services import output_builder as output


def test_real_confirmed_weekly_weight_receives_seven_and_point_seven_kg(
    surface, owned_bagging, monkeypatch, capsys,
):
    fixture, client = surface
    assert fixture is owned_bagging
    bagged = actual_bagging(fixture)
    candidate = deepcopy(bagged["bagging_result"]["materialization_candidate"])
    assert [(line["menu_name"], line["quantity_original"])
            for line in candidate["lines"]] == [
        ("選択献立0", 7), ("選択献立3", 0),
    ]
    days = {output._ensure_date(line["date"]) for line in candidate["lines"]}
    assert len(days) == 1 and None not in days
    target = days.pop()
    confirmed, error = workflow.final_confirm(
        fixture.order_id, confirmed_by="synthetic-weekly-operator",
        expected_revision=revision(fixture),
    )
    assert error is None, error
    assert confirmed is not None
    before = database_snapshot()
    saved = deepcopy(fixture.saved())
    assert [row[4] for row in saved["rows"]] == ["7", "", "", "0"]
    hashes = {path: sha256(path.read_bytes()).hexdigest() for path in (
        fixture.pdf, fixture.source, fixture.master_path, master.MASTER_TEMPLATE_PATH,
    )}
    persisted = [row for row in before["order_lines"] if row["order_id"] == fixture.order_id]
    assert sorted((row["menu_name"], row["quantity_original"]) for row in persisted) == [
        ("選択献立0", 7), ("選択献立3", 0),
    ]
    monkeypatch.setattr(output, "OUTPUT_DIR", fixture.output)
    context = output._prepare_output_context_for_bundle(
        fixture.order_id, require_invoice=False, include_bags=False,
        include_ocr_menu_meta=False, include_expanded_copy=True,
    )
    assert [(line["menu_name"], line["quantity_original"], line["menu_qty_per_serving"])
            for line in context["order_lines"]] == [
        ("選択献立0", 7, 100), ("選択献立3", 0, 100),
    ]
    collected = output._weekly_weight_collect_rows(target)
    assert len(collected) == 1
    (day, daypart, slot), values = next(iter(collected.items()))
    assert (day, daypart) == (target, "朝")
    assert values["regular_menu"] == "選択献立0"
    assert values["regular_quantity"] == 7
    assert values["regular_amounts"] == {"g": 700}
    with capsys.disabled():
        print(f"weekly-source categories={[line['menu_category'] for line in context['order_lines']]} slot={slot}")
    response = client.get("/outputs/weekly-weight", params={"date": target.isoformat()})
    assert response.status_code == 200, response.text
    body = response.content
    assert len(body) > 0 and int(response.headers["content-length"]) == len(body)
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["content-disposition"].startswith("attachment;")
    assert ".xlsx" in response.headers["content-disposition"]
    received = fixture.root / "http-received"
    received.mkdir()
    path = received / "weekly-weight.xlsx"
    path.write_bytes(body)
    assert path.read_bytes() == body
    with capsys.disabled():
        print("definition-source-base=9a5310bf4a5eea2420c267e4847a1815a345e3c1")
        print(f"http-received artifact={path} sha256={sha256(body).hexdigest()}")
    assert database_snapshot() == before
    assert fixture.saved() == saved
    assert all(sha256(path.read_bytes()).hexdigest() == digest for path, digest in hashes.items())
    workbook = load_workbook(BytesIO(body), read_only=True, data_only=False)
    try:
        assert len(workbook.worksheets) == 1
        rows = list(workbook.worksheets[0].iter_rows(min_row=11, max_row=66, values_only=True))
        # Source formats grams as kg with one decimal. Do not substitute a
        # breakfast side category merely to make this main-course row appear.
        assert [(row[2], row[3], row[4], row[5]) for row in rows if row[3] == "選択献立0"] == [
            (slot, "選択献立0", 7, 0.7),
        ]
        assert all(row[3] != "選択献立3" for row in rows)
    finally:
        workbook.close()
    assert database_snapshot() == before
    assert fixture.saved() == saved
    assert bagged["bagging_result"]["materialization_candidate"] == candidate
    assert all(sha256(path.read_bytes()).hexdigest() == digest for path, digest in hashes.items())
