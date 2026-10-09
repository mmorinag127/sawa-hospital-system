"""Local service/producer boundaries, not operator FAX, OCR execution, or live proof.

Authored logical targets come from canonical_logical_target_plan and a generated
unit workbook. Saved selection is manually seeded by the unchanged
SavedSheetFixture; PDF receipts and JSON projections are produced by real code.
Only database/session factories and local source/output paths are redirected.
No correctness guard, quantity writer, ownership, or config function is mocked.
"""

from copy import deepcopy
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
from shutil import copyfile
from tempfile import TemporaryDirectory

import fitz
from openpyxl import load_workbook
import pytest

from modernization_saved_sheet_export_fixtures import SavedSheetFixture
from src import db
from src.models.document import OrderDocument
from src.models.menu import MenuMaster, MonthlyMenu, MonthlyMenuEntry, MonthlyMenuItem
from src.models.order import Order
from src.models.order_ocr_evidence_run import OrderOcrEvidenceRun
from src.models.order_sheet_draft import OrderSheetDraft
from src.services import config_service, menu_service, order_service as orders
from src.services import hakodate_cell_ocr_batch_service as batch
from src.services import hakodate_ocr_evidence_service as evidence
from src.services import hakodate_step_review_pipeline_service as pipeline
from src.services import master_order_form_template_service as master
from src.services.storage_service import load_selected_pdf_bytes
from src.services.workflow_attempt_guard_service import capture_selected_pdf_receipt
from src.models.uploaded_pdf import UploadedPdf


FIELDS = ["date", "daypart", "category", "menu_name", "qty.regular_x",
          "qty.no_meat_x", "qty.no_fish_x", "qty.staff_x"]


def authored_template(*, notes=False):
    columns = [{"index": index, "source_index": index, "role": role}
               for index, role in enumerate(("date", "daypart", "category", "menu_name"))]
    columns.extend({"index": index, "source_index": index, "role": "quantity",
                    "diet_type": diet, "area_id": "X"}
                   for index, diet in enumerate(("regular", "no_meat", "no_fish", "staff"), 4))
    if notes:
        columns.append({"index": 8, "source_index": 8, "role": "note"})
    return {"template_id": "fax_layout_regular_forbidden_v1", "units": "normalized",
            "table_box": [0, 0, 1, 1], "header_rows": 0, "columns": columns}


def actual_plan(count=4, *, merged=True, notes=False, note_text="備考: 変更なし 007"):
    """Real authored Excel spans, with observation-only unit pixel geometry."""
    fields = FIELDS + (["remarks"] if notes else [])
    rows = [["3/1", "朝", "主", f"選択献立{i}", "", "", "", ""]
            + ([note_text] if notes else []) for i in range(count)]
    base = {"fields": fields, "rows": deepcopy(rows)}
    template = authored_template(notes=notes)
    if merged:
        template["body_merge_policy"] = {"mode": "daypart", "columns": ["regular"], "required": True}
    entries = [{"menu_date": "2026-03-01", "daypart": row[1], "category": row[2],
                "name": row[3], "sort_order": index + 1} for index, row in enumerate(rows)]
    canonical_base = deepcopy(base)
    for row in canonical_base["rows"]:
        row[0] = "2026-03-01"
    original_master = master.MASTER_TEMPLATE_PATH.read_bytes()
    with TemporaryDirectory(prefix="explicit-quantity-plan-") as directory:
        copied_master = Path(directory) / "copied-master.xlsx"
        copyfile(master.MASTER_TEMPLATE_PATH, copied_master)
        output = Path(directory) / "generated.xlsx"
        master.build_facility_template_xlsx(
            facility_config={"facility_id": "F", "facility_name": "合成数量契約", "fax_template": template},
            master_template_path=copied_master, output_path=output,
            week_value="2026-03@2026-03-01~2026-03-07", week_menu_entries=entries,
        )
        workbook = load_workbook(output)
        try:
            plan = pipeline.canonical_logical_target_plan(
                worksheet=workbook[master.FACILITY_TEMPLATE_SHEET_NAME],
                fax_template=template, draft_sheet=canonical_base,
            )
        finally:
            workbook.close()
        assert copied_master.read_bytes() == original_master
    assert master.MASTER_TEMPLATE_PATH.read_bytes() == original_master
    for region in plan:
        # Full dates were validated above; preserve this producer unit's display-date contract.
        region["date"] = rows[region["worksheet_row"] - 11][0]
        for target in region["logical_targets"]:
            target["date"] = rows[target["worksheet_row"] - 11][0]
        row, col = region["worksheet_row"], region["worksheet_col"]
        # Geometry belongs only to this local producer boundary, not an OCR receipt.
        region["bbox"] = [col * 10, row * 10, col * 10 + 8, row * 10 + 8]
        region["center"] = [col * 10 + 4, row * 10 + 4]
        region["ocr_text"] = note_text if region["role"] == "note" else ""
        region["ocr_normalized"] = region["ocr_text"]
    return base, plan


def regular_region(plan):
    matches = [region for region in plan if region["field"] == "qty.regular_x"]
    assert len(matches) == 1
    return matches[0]


def records_for(plan, value="7", confidence=.95):
    records = []
    for region in plan:
        raw = value if region["field"] == "qty.regular_x" else region["ocr_text"]
        records.append({"raw_text": raw, "normalized_value": raw,
                        "confidence": confidence, "bbox": deepcopy(region["bbox"]),
                        "center": deepcopy(region["center"]),
                        "evidence_id": "unit-" + region["region_id"],
                        "role": region["role"], "field": region["field"]})
    return records


def project(base, plan, records):
    """Actual manifest-style double conversion, assignment, sheet output, projection."""
    targets = evidence.target_cells_from_regions(plan)
    original_targets = deepcopy(targets)
    normalized = evidence.evidence_from_records(
        records, run_id="unit-selected", engine="unit-provider",
        source_scope="unit-local-cell-boundary", retain_empty_records=True,
    )
    assigned = evidence.assign_evidence_to_target_cells(evidence_records=normalized, target_cells=targets)
    sheet = evidence.sheet_output_from_assigned_results(
        assignments=assigned["assignments"], blockers=assigned["blockers"],
        unassigned_evidence=assigned["unassigned_evidence"],
    )
    status = "auto_assignable" if not sheet["blockers"] else "blocked"
    result = orders._apply_hakodate_sheet_output_to_sheet_payload(
        base_sheet=base, assignment={**assigned, "target_cells": targets,
                                    "sheet_output": sheet, "status": status},
        facility_config={"facility_id": "F", "fax_template": authored_template(notes="remarks" in base["fields"])},
    )
    assert targets == original_targets
    assert evidence.target_cells_from_regions(targets) == targets
    return result, normalized, assigned, sheet


@pytest.mark.parametrize("count", [4, 5, 6, 3])
@pytest.mark.parametrize("value", ["7", "0", "", "20"])
def test_actual_authored_spans_double_conversion_and_batch_parity(count, value):
    base, plan = actual_plan(count)
    region = regular_region(plan)
    assert len(region["logical_targets"]) == count
    originals = deepcopy((base, plan))
    result, _, assigned, sheet = project(base, plan, records_for(plan, value))
    batch_regions = deepcopy(plan)
    regular_region(batch_regions).update(ocr_text=value, ocr_normalized=value)
    assignments = batch.sheet_assignments_from_ocr_regions(batch_regions)
    grid = batch.sheet_value_grid_from_assignments(assignments)
    assert [row[4] for row in result["rows"]] == [value] * count
    assert result["blockers"] == [] and result["review_state"] == "auto_assignable"
    assert result["hakodate_evidence_projection"]["metrics"]["expanded_cell_same_daypart_filled_count"] == 0
    for target in region["logical_targets"]:
        cell = target["sheet_cell"]
        assert grid["cells"][cell]["value_normalized"] == sheet["cells"][cell]["value_normalized"] == value
        assert sheet["cells"][cell]["date"] == target["date"]
        assert sheet["cells"][cell]["menu_name"] == target["menu_name"]
    assert assigned["blockers"] == [] and (base, plan) == originals


@pytest.mark.parametrize("note", ["備考: 減塩 007", "0", "", "７? -7"])
def test_real_mixed_quantity_note_targets_preserve_note_text_in_all_paths(note):
    base, plan = actual_plan(notes=True, note_text=note)
    result, _, _, sheet = project(base, plan, records_for(plan))
    batch_regions = deepcopy(plan)
    regular_region(batch_regions).update(ocr_text="7", ocr_normalized="7")
    grid = batch.sheet_value_grid_from_assignments(batch.sheet_assignments_from_ocr_regions(batch_regions))
    assert [row[8] for row in result["rows"]] == [note] * 4
    assert [row[4] for row in result["rows"]] == ["7"] * 4
    assert not result["blockers"]
    for region in plan:
        if region["role"] == "note":
            key = region["sheet_cell"]
            assert grid["cells"][key]["value_normalized"] == sheet["cells"][key]["value_text"] == note
    assert result["ocr_numeric_cell_summary"]["raw_ocr_numeric_count"] == 4


@pytest.mark.parametrize("raw", ["?", "7?", "-7", "1.5"])
def test_uncertain_quantity_not_rescued_by_valid_sibling_or_notes(raw):
    base, plan = actual_plan(notes=True)
    records = records_for(plan, raw)
    valid = deepcopy(records[0])
    assert records[0]["field"] == "qty.regular_x"
    valid.update(raw_text="7", normalized_value="7", evidence_id="unit-valid-sibling")
    result, normalized, _, _ = project(base, plan, records + [valid])
    assert normalized[0]["raw_text"] == raw
    assert result["review_state"] == "blocked" and result["blockers"]
    assert [row[4] for row in result["rows"]] != ["7"] * 4
    assert [row[8] for row in result["rows"]] == ["備考: 変更なし 007"] * 4
    bad_batch = deepcopy(plan)
    regular_region(bad_batch).update(ocr_text=raw, ocr_normalized="7")
    with pytest.raises(ValueError, match="hakodate_ocr_quantity"):
        batch.sheet_assignments_from_ocr_regions(bad_batch)


@pytest.mark.parametrize("raw, normalized", [("7?", "7"), ("-7", "7"), ("", "7"), ("0", "7")])
def test_raw_normalized_alias_conflicts_stay_visible(raw, normalized):
    base, plan = actual_plan()
    records = records_for(plan, raw)
    records[0]["normalized_value"] = normalized
    result, saved, _, _ = project(base, plan, records)
    assert saved[0]["raw_text"] == raw
    assert saved[0]["engine_metadata"]["source_values"]["normalized_value"] == normalized
    assert result["blockers"] and result["review_state"] == "blocked"


@pytest.mark.parametrize("raw, expected", [("１，００７", "1007"), ("1,007", "1007"), ("7.0", "7"), ("０", "0")])
def test_exact_canonical_count_formats_not_decimal_guess(raw, expected):
    base, plan = actual_plan()
    result, saved, _, _ = project(base, plan, records_for(plan, raw))
    assert saved[0]["raw_text"] == raw
    assert [row[4] for row in result["rows"]] == [expected] * 4
    assert not result["blockers"]


@pytest.mark.parametrize("value", ["0", "7"])
@pytest.mark.parametrize("confidence", [.89, .15, .05, .0, None, True, 2.0, -1.0, float("nan"), float("inf"), "invalid"])
def test_nonblank_unaccepted_confidence_deferred_with_real_diagnostics(value, confidence):
    base, plan = actual_plan()
    result, saved, assigned, _ = project(base, plan, records_for(plan, value, confidence))
    assert result["review_state"] == "blocked" and result["blockers"]
    deferred = [cell for cell in result["hakodate_evidence_projection"]["deferred"] if cell["field"] == "qty.regular_x"]
    assert len(deferred) == 4 and {cell["value"] for cell in deferred} == {value}
    assert all(cell["classification"] != "accepted" for cell in deferred)
    assert not any(cell["field"] == "qty.regular_x" for cell in result["hakodate_evidence_projection"]["applied"])
    original = saved[0]["engine_metadata"]["raw_confidence"]
    if isinstance(confidence, float) and confidence != confidence:
        assert original != original
    else:
        assert original == confidence and type(original) is type(confidence)
    if confidence is None or isinstance(confidence, bool) or confidence in (2.0, -1.0, float("inf"), "invalid") or (isinstance(confidence, float) and confidence != confidence):
        assert saved[0]["confidence"] is None
        assert assigned["assignments"][0]["assignment_confidence"] is None


@pytest.mark.parametrize("confidence", [.95, .90, .89, .15, .05, .0, None, True, 2.0, -1.0, float("nan"), float("inf"), "invalid"])
def test_genuine_blank_no_confidence_only_block_or_fake_high(confidence):
    base, plan = actual_plan()
    result, _, _, _ = project(base, plan, records_for(plan, "", confidence))
    assert [row[4] for row in result["rows"]] == [""] * 4
    assert result["blockers"] == []
    if evidence.normalize_ocr_confidence(confidence) is None:
        assert all(tier != "high" for row in result["cell_confidence_rows"] for tier in row)


@pytest.mark.parametrize("value", ["0", "7"])
@pytest.mark.parametrize("confidence", [.90, .95, 1.0])
def test_actual_accepted_threshold_applies_without_confidence_fabrication(value, confidence):
    base, plan = actual_plan()
    result, _, _, _ = project(base, plan, records_for(plan, value, confidence))
    assert [row[4] for row in result["rows"]] == [value] * 4
    assert not result["blockers"]
    assert [row[4] for row in result["cell_score_rows"]] == [confidence] * 4


@pytest.mark.parametrize("invalid", [True, None, 2.0, -1.0, float("nan"), float("inf")])
def test_actual_evidence_mean_cannot_launder_invalid_component(invalid):
    base, plan = actual_plan()
    records = records_for(plan, "7", .95)
    extra = deepcopy(records[0])
    extra.update(evidence_id="unit-second-component", confidence=invalid)
    result, saved, assigned, _ = project(base, plan, records + [extra])
    assert assigned["assignments"][0]["assignment_confidence"] is None
    assert saved[-1]["confidence"] is None
    assert result["review_state"] == "blocked"


@pytest.mark.parametrize("scores, expected", [([.9, 1.0], .95), ([0., 1.], .5)])
def test_actual_valid_component_mean_unchanged(scores, expected):
    base, plan = actual_plan()
    records = records_for(plan, "7", scores[0])
    extra = deepcopy(records[0])
    extra.update(evidence_id="unit-second-component", confidence=scores[1])
    result, _, assigned, _ = project(base, plan, records + [extra])
    assert assigned["assignments"][0]["assignment_confidence"] == expected
    assert (result["review_state"] == "blocked") == (expected < .90)


@pytest.mark.parametrize("scores", [[2.0, 0.0], [None, .99], [True, .99]])
def test_invalid_plus_valid_exact_reported_confidence_pairs(scores):
    base, plan = actual_plan()
    records = records_for(plan, "0", scores[0])
    extra = deepcopy(records[0])
    extra.update(evidence_id="unit-reported-pair", confidence=scores[1])
    result, _, assigned, _ = project(base, plan, records + [extra])
    assert assigned["assignments"][0]["assignment_confidence"] is None
    assert result["review_state"] == "blocked" and result["blockers"]


def test_absent_confidence_key_not_averaged_away():
    base, plan = actual_plan()
    records = records_for(plan, "7", .99)
    extra = deepcopy(records[0])
    extra["evidence_id"] = "unit-missing-key"
    del extra["confidence"]
    result, normalized, assigned, _ = project(base, plan, records + [extra])
    assert normalized[-1]["engine_metadata"]["raw_confidence"] is None
    assert assigned["assignments"][0]["assignment_confidence"] is None
    assert result["review_state"] == "blocked"


def test_actual_contact_word_producer_note_normalization_and_quantity_uncertainty():
    base, plan = actual_plan(notes=True, note_text="備考 ００７")
    words = []
    for region in plan:
        region["ocr_contact_slot"] = deepcopy(region["bbox"])
        raw = "備考 ００７" if region["role"] == "note" else "0"
        words.append({"text": raw, "x": region["center"][0] / 1000,
                      "y": region["center"][1] / 1000})
    original = deepcopy(words)
    produced = batch.assign_yomitoku_words_to_contact_regions(words=words, regions=plan, sheet_size=(1000, 1000))
    assignments = batch.sheet_assignments_from_ocr_regions(produced)
    grid = batch.sheet_value_grid_from_assignments(assignments)
    for region in produced:
        expected = "備考 ００７" if region["role"] == "note" else "0"
        assert region["ocr_text"] == region["ocr_normalized"] == expected
        for target in region["logical_targets"]:
            assert grid["cells"][target["sheet_cell"]]["value_normalized"] == expected
    records = [{**region, "raw_text": region["ocr_text"], "confidence": .95} for region in produced]
    result, _, _, _ = project(base, produced, records)
    assert [row[8] for row in result["rows"]] == ["備考 ００７"] * 4
    assert not result["blockers"] and words == original
    regular_region(produced).update(ocr_text="?", ocr_normalized="?")
    with pytest.raises(ValueError, match="hakodate_ocr_quantity_unresolved"):
        batch.sheet_assignments_from_ocr_regions(produced)


def test_authored_singleton_keeps_identity_and_note_grid_metadata():
    base, plan = actual_plan(1, merged=False, notes=True)
    for region in plan:
        region.pop("logical_targets")
        region.pop("covered_sheet_cells")
        region["merged_cell"] = None
        region.update(grid_row_index=10, grid_col_index=region["worksheet_col"] - 1)
    result, _, _, sheet = project(base, plan, records_for(plan, "0"))
    assert result["rows"][0][4] == "0" and not result["blockers"]
    normalized = evidence.target_cells_from_regions(plan)
    for target in normalized:
        logical = evidence.logical_sheet_targets_from_region(target)[0]
        assert logical["date"] == "3/1" and logical["daypart"] == "朝"
        assert logical["menu_name"] == "選択献立0"
        assert logical["grid_row_index"] == 10
        assert logical["grid_col_index"] == target["worksheet_col"] - 1
        assert logical["field_label"] == target["metadata"]["field_label"]
        assert sheet["cells"][target["sheet_cell"]]["menu_name"] == "選択献立0"


@pytest.mark.parametrize("fault", ["bbox", "missing_bbox", "center", "identity", "metadata", "missing_merged", "empty_targets", "bad_targets", "bad_target", "duplicate", "covered", "missing_role", "conflict_role", "mixed_roles"])
def test_actual_producer_and_batch_reject_malformed_required_targets(fault):
    _, plan = actual_plan()
    region = regular_region(plan)
    if fault == "bbox":
        region["bbox"] = [0, 0, float("nan"), 1]
    elif fault == "missing_bbox":
        region.pop("bbox")
    elif fault == "center":
        region["center"] = {"x": 1}
    elif fault == "identity":
        for key in ("target_cell_id", "region_id", "sheet_cell"):
            region.pop(key, None)
    elif fault == "metadata":
        region["metadata"] = {"metadata": {"logical_targets": region["logical_targets"]}}
    elif fault == "missing_merged":
        region.pop("logical_targets")
    elif fault == "empty_targets":
        region["logical_targets"] = []
    elif fault == "bad_targets":
        region["logical_targets"] = {}
    elif fault == "bad_target":
        region["logical_targets"][0]["sheet_cell"] = []
    elif fault == "duplicate":
        region["logical_targets"].append(deepcopy(region["logical_targets"][0]))
    elif fault == "covered":
        region["covered_sheet_cells"] = ["E99"]
    elif fault == "missing_role":
        region["logical_targets"][0].pop("role")
    elif fault == "conflict_role":
        region["logical_targets"][0]["role"] = "note"
    else:
        region["logical_targets"][0].update(role="note", field="remarks")
    original = deepcopy(plan)
    for call in (lambda: evidence.target_cells_from_regions(plan),
                 lambda: evidence.assign_evidence_to_target_cells(evidence_records=[], target_cells=plan),
                 lambda: batch.sheet_assignments_from_ocr_regions(plan)):
        with pytest.raises(ValueError, match="hakodate_(logical_targets|target_)"):
            call()
    # NaN is deliberately non-reflexive; compare all other raw contracts directly.
    if fault != "bbox":
        assert plan == original


@pytest.mark.parametrize("fault", ["duplicate", "missing", "covered", "role", "cell_type"])
def test_sheet_output_revalidates_actual_assigned_required_targets(fault):
    base, plan = actual_plan()
    _, _, assigned, _ = project(base, plan, records_for(plan))
    assignments = deepcopy(assigned["assignments"])
    metadata = assignments[0]["target_metadata"]
    required = [target["sheet_cell"] for target in metadata["logical_targets"]]
    if fault == "duplicate":
        metadata["logical_targets"].append(deepcopy(metadata["logical_targets"][0]))
    elif fault == "missing":
        del metadata["logical_targets"]
    elif fault == "covered":
        metadata["covered_sheet_cells"] = ["E99"]
    elif fault == "role":
        metadata["logical_targets"][0].update(role="note", field="remarks")
    else:
        metadata["logical_targets"][0]["sheet_cell"] = {}
    output = evidence.sheet_output_from_assigned_results(assignments=assignments)
    assert output["blockers"]
    assert all(cell not in output["cells"] for cell in required)


def test_recognizer_note_claim_cannot_override_authored_quantity_target():
    base, plan = actual_plan(notes=True)
    records = records_for(plan, "7?")
    records[0]["engine_metadata"] = {"recognizer_output": {
        "role": "note", "field": "remarks", "raw_text": "7?",
    }}
    result, saved, _, _ = project(base, plan, records)
    assert saved[0]["raw_text"] == "7?"
    assert result["review_state"] == "blocked"
    assert any("hakodate_target_role_conflict" in blocker for blocker in result["blockers"])
    assert [row[4] for row in result["rows"]] != ["7"] * 4


class PublicSavedCandidate(SavedSheetFixture):
    """Real owned SQLite/PDF/dated JSON capture, manually seeded OCR selection."""

    def __init__(self, tmp_path, monkeypatch):
        super().__init__(tmp_path, monkeypatch)
        monkeypatch.setattr(db, "SessionLocal", self.sessions)
        monkeypatch.setattr(menu_service, "engine", self.engine)
        self.authored["fax_template"] = authored_template()
        self.republish()
        path = tmp_path / "synthetic-local-unit.pdf"
        document = fitz.open()
        document.new_page().insert_text((50, 50), "Synthetic unit F / 2026-03-01..07; not operator FAX")
        document.save(path)
        document.close()
        self.pdf = path
        names = [f"選択献立{i}" for i in range(4)]
        with self.session_scope() as session:
            session.add(MonthlyMenu(id="2026-03", month_start=date(2026, 3, 1)))
            for index, name in enumerate(names):
                master_id = f"UNIT-MASTER-{index}"
                session.add(MenuMaster(id=master_id, name=name,
                                       normalized_name=menu_service._normalize_menu_name(name),
                                       unit_type="g", qty_per_serving=100, category="主", daypart="朝"))
                session.flush()
                session.add(MonthlyMenuItem(id=f"UNIT-ITEM-{index}", monthly_menu_id="2026-03",
                                           menu_master_id=master_id, name=name, diet_type="regular",
                                           category="主", daypart="朝", unit_type="g", qty_per_serving=100))
                session.add(MonthlyMenuEntry(id=f"UNIT-ENTRY-{index}", monthly_menu_id="2026-03",
                                            menu_date=date(2026, 3, 1), daypart="朝", category="主",
                                            diet_type="regular", name=name, slot_index=index))
            session.add(OrderDocument(id="UNIT-DOC", facility_code="F", week_code="2026-03@2026-03-01~2026-03-07",
                                      storage_uri=str(path), source_email_id=self.order_id,
                                      received_at=datetime(2026, 3, 1, 9), status="done"))
            session.flush()
            order = session.get(Order, self.order_id)
            order.current_document_id = "UNIT-DOC"
            order.document_uri = str(path)
            session.add(UploadedPdf(id="UNIT-PDF", message_id=self.order_id, storage_uri=str(path),
                                    original_filename=path.name, content_sha256=sha256(path.read_bytes()).hexdigest(),
                                    received_at=datetime(2026, 3, 1, 9), page_count=1,
                                    current_document_id="UNIT-DOC", current_order_id=self.order_id))
            session.flush()
            self.receipt = capture_selected_pdf_receipt(session, order)
        assert load_selected_pdf_bytes(str(path), self.receipt) == path.read_bytes()

    def save(self, quantities, *, enabled):
        self.authored["expanded_cell_same_daypart_copy_enabled"] = enabled
        self.authored["fax_template"]["body_merge_policy"] = {"mode": "daypart", "columns": ["qty.regular_x"]}
        workbook = load_workbook(self.source)
        worksheet = workbook.active
        for index in range(4):
            for column, value in enumerate(["3/1", "朝", "主", f"選択献立{index}"], 1):
                worksheet.cell(11 + index, column, value)
        worksheet.merge_cells("E11:E14")
        workbook.save(self.source)
        workbook.close()
        self.republish()
        payload = deepcopy(self.payload)
        payload["rows"] = [["3/1", "朝", "主", f"選択献立{i}", value, "", "", ""]
                           for i, value in enumerate(quantities)]
        self.payload = payload
        with self.session_scope() as session:
            session.get(OrderSheetDraft, self.draft_id).draft_sheet_json = deepcopy(payload)
            selected = session.get(OrderOcrEvidenceRun, self.evidence_id)
            selected.payload_json = {"tables": [{"rows": deepcopy(payload["rows"])}]}
            selected.artifact_digest = sha256(json.dumps(selected.payload_json).encode()).hexdigest()
        return deepcopy(payload)

    def saved(self):
        with self.session_scope() as session:
            return deepcopy(session.get(OrderSheetDraft, self.draft_id).draft_sheet_json)


@pytest.fixture
def public_saved(tmp_path, monkeypatch):
    fixture = PublicSavedCandidate(tmp_path, monkeypatch)
    yield fixture
    fixture.engine.dispose()


@pytest.mark.parametrize("enabled", [True, False])
@pytest.mark.parametrize("quantities", [["7", "?", "", "0"], ["7", "", "", ""]])
def test_public_candidate_real_capture_original_saved_quantities(public_saved, enabled, quantities):
    original = public_saved.save(quantities, enabled=enabled)
    captured = orders._capture_synchronous_cache_write_input(public_saved.order_id)
    assert captured["selected_pdf_receipt"] == public_saved.receipt
    assert captured["draft_id"] == public_saved.draft_id
    assert len(captured["physical_menu_entries"]) == 4
    assert captured["configured_input"]["selected_authored_version"]["version_id"] == "authored-March"
    result = orders.build_confirm_materialization_candidate(public_saved.order_id)
    assert public_saved.saved() == original
    if "?" in quantities:
        assert result["error"] == "sheet_ocr_review_required" and result["lines"] == []
        assert [(issue["row_index"], issue["column_index"], issue["raw"]) for issue in result["cell_issues"]] == [(1, 4, "?")]
    else:
        assert result["error"] is None
        assert result["line_count"] == 1
        assert [(line["source_row_index"], line["source_col_index"], line["source_field"], line["quantity_original"]) for line in result["lines"]] == [(0, 4, "qty.regular_x", 7)]
        assert result["synchronous_write_input"]["draft_id"] == public_saved.draft_id
        assert orders._build_materialization_quantity_guard_detail(fields=FIELDS, rows_payload=original["rows"], candidate_lines=result["lines"]) is None


@pytest.mark.parametrize("enabled", [True, False])
def test_public_candidate_zero_not_backfilled_and_raw_guard_original(public_saved, enabled):
    original = public_saved.save(["7", "", "", "0"], enabled=enabled)
    result = orders.build_confirm_materialization_candidate(public_saved.order_id)
    assert result["error"] is None
    assert [(line["source_row_index"], line["quantity_original"]) for line in result["lines"]] == [(0, 7), (3, 0)]
    assert public_saved.saved() == original
    assert orders._build_materialization_quantity_guard_detail(fields=FIELDS, rows_payload=original["rows"], candidate_lines=result["lines"]) is None
    invented = deepcopy(result["lines"])
    invented.append({**invented[0], "source_row_index": 1})
    guard = orders._build_materialization_quantity_guard_detail(fields=FIELDS, rows_payload=original["rows"], candidate_lines=invented)
    assert guard["mismatch_count"] == 1 and guard["mismatches"][0]["source_row_index"] == 1


def test_public_candidate_missing_pdf_real_guard_stops_without_raw_change(public_saved):
    original = public_saved.save(["7", "", "", "0"], enabled=True)
    with public_saved.session_scope() as session:
        session.get(Order, public_saved.order_id).current_document_id = None
    with pytest.raises(ValueError, match="selected_pdf_document_required"):
        orders.build_confirm_materialization_candidate(public_saved.order_id)
    assert public_saved.saved() == original


def test_public_candidate_foreign_selected_evidence_stops_without_raw_change(public_saved):
    original = public_saved.save(["7", "", "", "0"], enabled=True)
    with public_saved.session_scope() as session:
        session.get(OrderOcrEvidenceRun, public_saved.evidence_id).order_id = "ORD-FOREIGN"
    result = orders.build_confirm_materialization_candidate(public_saved.order_id)
    assert result == {"error": "synchronous_input_missing", "lines": []}
    assert public_saved.saved() == original


def test_public_candidate_foreign_pdf_message_stops_real_receipt_capture(public_saved):
    original = public_saved.save(["7", "", "", "0"], enabled=True)
    with public_saved.session_scope() as session:
        session.get(OrderDocument, "UNIT-DOC").source_email_id = "ORD-FOREIGN"
    with pytest.raises(ValueError, match="selected_pdf_document_required"):
        orders.build_confirm_materialization_candidate(public_saved.order_id)
    assert public_saved.saved() == original


def test_real_captured_config_retirement_stops_fence_without_raw_change(public_saved):
    original = public_saved.save(["7", "", "", "0"], enabled=False)
    captured = orders._capture_synchronous_cache_write_input(public_saved.order_id)
    assert config_service.captured_processing_config_matches("F", captured["configured_input"])
    public_saved.authored["facility_name"] = "明示的な後続変更"
    public_saved.write_master()
    assert not config_service.captured_processing_config_matches("F", captured["configured_input"])
    with public_saved.session_scope() as session:
        assert not orders._synchronous_cache_write_fence(session, public_saved.order_id, captured)
    assert public_saved.saved() == original
