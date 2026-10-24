"""Synthetic-native numeric FAX definitions, not customer/operator UI/stg proof.

The unchanged OwnedProducer supplies the authored facility/week, copied master,
SQLite and real ownership services. This separate nonblank scenario does not
replace or waive the original blank-containing native or artifact-native tests.
No recognition, geometry, assignment, materialization or receipt doubles.
"""

from copy import deepcopy
from datetime import date, timedelta
from hashlib import sha256
from pathlib import Path
import shutil

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
import pytest

from test_modernization_owned_ocr_workflow import (
    END, FACILITY, JOB, ORDER, START, TEMPLATE, WEEK, OwnedProducer,
)
from src import db
from src.models.ocr_job import OcrJob
from src.models.order_ocr_evidence_run import OrderOcrEvidenceRun
from src.models.order_sheet_draft import OrderSheetDraft
from src.models.order_workflow_state import OrderWorkflowState
from src.services import canonical_ocr_artifact_service as artifacts
from src.services import evidence_manifest_service as manifests
from src.services import hakodate_cell_ocr_batch_service as batch
from src.services import master_order_form_template_service as master
from src.services import ocr_evidence_service, order_service, order_workflow_v2_service as workflow
from src.services.storage_service import load_selected_pdf_bytes
from src.services.workbook_pdf_renderer import render_workbook_path_to_pdf


@pytest.fixture
def numeric_fax(tmp_path, monkeypatch, isolated_formal_sqlite_template):
    f = OwnedProducer(tmp_path, monkeypatch, formal_template=isolated_formal_sqlite_template)
    try:
        monkeypatch.setenv("OCR_ARTIFACT_DIR", str(tmp_path / "numeric-published"))
        original_pdf = f.pdf
        original_receipt = deepcopy(f.receipt)
        originals = {path: path.read_bytes() for path in (
            original_pdf, tmp_path / "synthetic-fax.xlsx", f.master_copy, f.master_json,
        )}
        entries = order_service._build_position_menu_entries_safe(WEEK, FACILITY)
        identity = batch.original_physical_draft(physical_menu_entries=entries, week_code=WEEK)
        assert len(entries) == identity["physical_menu_row_count"] == len(f.quantity_plan) == 56
        assert len(f.notes_plan) == 56
        assert identity["week_id"] == WEEK
        assert {row[0] for row in identity["rows"]} == {
            (START + timedelta(days=offset)).isoformat() for offset in range(7)
        }
        cells = [target["sheet_cell"] for target in f.quantity_plan]
        assert len(set(cells)) == 56
        pattern = ("7", "0", "6", "9")
        expected = dict(zip(cells, (pattern[index % 4] for index in range(56)), strict=True))
        xlsx = tmp_path / "synthetic-numeric-fax.xlsx"
        shutil.copyfile(tmp_path / "synthetic-fax.xlsx", xlsx)
        workbook = load_workbook(xlsx)
        try:
            ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
            before = {cell.coordinate: deepcopy(cell.value) for row in ws for cell in row}
            assert [ws[cell].row for cell in cells] == sorted(ws[cell].row for cell in cells)
            for cell, value in expected.items():
                ws[cell] = value
                ws[cell].font = Font(name="DejaVu Sans", size=18, color="000000")
                ws[cell].alignment = Alignment(horizontal="center", vertical="center")
            workbook.save(xlsx)
        finally:
            workbook.close()
        workbook = load_workbook(xlsx)
        try:
            ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
            assert {cell: ws[cell].value for cell in cells} == expected
            assert all(ws[cell].value == value for cell, value in before.items() if cell not in expected)
            assert all(ws[cell].font.name == "DejaVu Sans" and ws[cell].font.sz == 18 for cell in cells)
        finally:
            workbook.close()
        f.pdf = render_workbook_path_to_pdf(
            xlsx, output_path=tmp_path / "synthetic-numeric-fax.pdf",
            sheet_name=master.FACILITY_TEMPLATE_SHEET_NAME,
        )
        pdf_bytes = f.pdf.read_bytes()
        assert pdf_bytes.startswith(b"%PDF-")
        assert f.pdf != original_pdf
        assert sha256(pdf_bytes).hexdigest() != original_receipt["sha256"]
        with db.session_scope() as session:
            f.add_document(session, "DOC-NATIVE-NUMERIC", "PDF-NATIVE-NUMERIC", f.pdf,
                           message_id="UNIT-NATIVE-NUMERIC-MESSAGE")
        assert f.receipt["document_id"] == "DOC-NATIVE-NUMERIC"
        assert f.receipt["uploaded_pdf_id"] == "PDF-NATIVE-NUMERIC"
        assert f.receipt["sha256"] == sha256(pdf_bytes).hexdigest()
        assert load_selected_pdf_bytes(str(original_pdf), original_receipt) == originals[original_pdf]
        initial, error = workflow.get_workflow_readonly(ORDER)
        assert error is None, error
        f.context, error = workflow.confirm_context(
            order_id=ORDER, facility_id=FACILITY, week_start=START.isoformat(),
            week_end=END.isoformat(), template_id=TEMPLATE, expected_revision=initial["revision"],
        )
        assert error is None, (error, f.context)
        f.expected = {target["sheet_cell"]: expected.get(target["sheet_cell"], "") for target in f.plan}
        yield f, identity
        assert all(path.read_bytes() == data for path, data in originals.items())
        assert f.pdf.read_bytes() == pdf_bytes
    finally:
        f.engine.dispose()
        assert sha256(master.MASTER_TEMPLATE_PATH.read_bytes()).hexdigest() == f.locked_hash


def test_synthetic_native_numeric_fax_adoption_save_and_four_packet_stops(numeric_fax):
    f, identity = numeric_fax
    evidence = f.produce()
    payload = evidence["payload_json"]
    source = payload["hakodate_repair_source_receipt"]
    assert source["selected_pdf_receipt"] == f.receipt
    assert evidence["artifact_manifest_json"]["original_hakodate_source"] == source
    execution = payload["_ocr_execution"]
    assert execution["job_id"] == JOB
    assert execution["evidence_run_id"] == evidence["id"]
    assert execution["nonworkflow_execution_id"] == source["execution_id"]
    records = payload["hakodate_ocr_evidence_records"]
    assert len(records) == len(f.quantity_plan) + len(f.notes_plan) == 112
    by_cell = {record["engine_metadata"]["sheet_cell"]: record for record in records}
    assert set(by_cell) == set(f.expected)
    for cell, expected in f.expected.items():
        assert by_cell[cell]["raw_text"] == expected, (cell, by_cell[cell])
        assert by_cell[cell]["normalized_value"] == expected
    pipeline = payload["hakodate_canonical_pipeline"]
    assert pipeline["expected_sheet_cells"] == list(f.expected)
    packet = pipeline["artifact_packet"]
    assert set(packet["outputs"]) == set(artifacts.REQUIRED_OUTPUTS)
    assert len(packet["outputs"]) == 8
    assert packet["source_receipt"] == source
    artifacts.validate_payload(payload)
    for key, receipt in packet["outputs"].items():
        path = Path(receipt["uri"])
        assert path.is_relative_to(f.root / "numeric-published"), (key, path)
        data = path.read_bytes()
        assert len(data) == receipt["size_bytes"] > 0
        assert sha256(data).hexdigest() == receipt["sha256"]
        assert pipeline["outputs"][key] == str(path)
    assert all(record["raw_payload_ref"] == pipeline["outputs"]["records"] for record in records)
    manifest = manifests.ensure_evidence_manifest(payload)["evidence_manifest"]
    assert manifest["contract"] == artifacts.CONTRACT
    assert manifest["required_artifacts"] == list(artifacts.REQUIRED_OUTPUTS)
    assert manifest["artifacts_complete"] is True
    assert manifest["missing_artifacts"] == manifest["validation_errors"] == []
    assert ocr_evidence_service.classify_evidence_payload(payload)["persistable"] is True
    with db.session_scope() as session:
        row = session.get(OrderOcrEvidenceRun, evidence["id"])
        assert ocr_evidence_service.evidence_run_is_selectable(row) is True
        assert row.payload_json == payload
        assert session.get(OcrJob, JOB).status == "done"
        assert session.get(OrderWorkflowState, ORDER).evidence_run_id is None
    f.select()
    captured = order_service._capture_synchronous_cache_write_input(ORDER)
    assert captured is not None
    order_service._validate_selected_hakodate_artifacts(evidence, captured)
    projected, error = workflow.build_sheet_from_selected_ocr(ORDER)
    assert error is None, (error, projected)
    assert projected["selected_ocr_result_id"] == evidence["id"]
    assert projected["blockers"] == []
    sheet = projected["sheet"]
    assert len(sheet["rows"]) == 56
    assert sheet["order_id"] == ORDER
    assert sheet["facility_id"] == FACILITY
    assert sheet["week_id"] == WEEK
    assert sheet["physical_menu_row_ids"] == identity["physical_menu_row_ids"]
    assert sheet["physical_menu_row_count"] == 56
    assert [[row[sheet["fields"].index(field)] for field in ("date_mmdd", "daypart", "menu")]
            for row in sheet["rows"]] == [
                [date.fromisoformat(day).strftime("%m/%d"), meal, menu]
                for day, meal, menu in identity["rows"]
            ]
    quantity = sheet["fields"].index("qty.regular_x")
    assert [row[quantity] for row in sheet["rows"]] == [f.expected[target["sheet_cell"]] for target in f.quantity_plan]
    assert all(row[sheet["fields"].index("remarks")] == "" for row in sheet["rows"])

    uri = Path(pipeline["outputs"]["records"])
    original_bytes = uri.read_bytes()
    for change, blocker in (
        ("missing_packet", "canonical_artifact_packet_required"),
        ("tampered_bytes", "canonical_artifact_bytes_changed"),
        ("deleted_bytes", "canonical_artifact_bytes_unavailable"),
        ("changed_source", "canonical_artifact_source_changed"),
    ):
        changed = deepcopy(evidence)
        changed_pipeline = changed["payload_json"]["hakodate_canonical_pipeline"]
        with db.session_scope() as session:
            state = session.get(OrderWorkflowState, ORDER)
            state_before = (state.revision, state.evidence_run_id, state.draft_id)
        try:
            if change == "missing_packet":
                changed_pipeline.pop("artifact_packet")
            elif change == "tampered_bytes":
                uri.write_bytes(b"[]")
            elif change == "deleted_bytes":
                uri.unlink()
            else:
                digest = changed_pipeline["artifact_packet"]["source_receipt"]["configured_input_digest"]
                changed_pipeline["artifact_packet"]["source_receipt"]["configured_input_digest"] = (
                    "1" if digest[0] != "1" else "0"
                ) + digest[1:]
            assessment = ocr_evidence_service.classify_evidence_payload(changed["payload_json"])
            assert assessment["persistable"] is False, change
            assert blocker in assessment["message"], (change, assessment)
            with db.session_scope() as session:
                session.get(OrderOcrEvidenceRun, evidence["id"]).payload_json = deepcopy(changed["payload_json"])
            with db.session_scope() as session:
                assert ocr_evidence_service.evidence_run_is_selectable(
                    session.get(OrderOcrEvidenceRun, evidence["id"])
                ) is False, change
            with pytest.raises(ValueError, match=f"^{blocker}$"):
                order_service._validate_selected_hakodate_artifacts(changed, captured)
            rejected, error = workflow.build_sheet_from_selected_ocr(ORDER)
            assert rejected is None, (change, rejected)
            assert error == blocker, (change, error)
            with db.session_scope() as session:
                state = session.get(OrderWorkflowState, ORDER)
                assert (state.revision, state.evidence_run_id, state.draft_id) == state_before
        finally:
            uri.write_bytes(original_bytes)
            with db.session_scope() as session:
                session.get(OrderOcrEvidenceRun, evidence["id"]).payload_json = deepcopy(payload)
        artifacts.validate_payload(payload)
        with db.session_scope() as session:
            assert ocr_evidence_service.evidence_run_is_selectable(
                session.get(OrderOcrEvidenceRun, evidence["id"])
            ) is True
        order_service._validate_selected_hakodate_artifacts(evidence, captured)

    with db.session_scope() as session:
        revision = session.get(OrderWorkflowState, ORDER).revision
    original_sheet = deepcopy(sheet)
    saved, error = workflow.save_sheet(
        order_id=ORDER, sheet=sheet, edited_by="synthetic-native-numeric-unit",
        expected_revision=revision,
    )
    assert error is None, (error, saved)
    assert sheet == original_sheet
    with db.session_scope() as session:
        state = session.get(OrderWorkflowState, ORDER)
        draft = session.get(OrderSheetDraft, state.draft_id)
        assert state.revision == revision + 1
        assert state.evidence_run_id == draft.base_evidence_run_id == evidence["id"]
        assert draft.template_version_id == f.projection_id
        assert draft.draft_sheet_json == original_sheet
        assert session.get(OrderOcrEvidenceRun, evidence["id"]).payload_json == payload
