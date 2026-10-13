"""Actual synthetic workbook/PNG/OpenCV component proof, not native/live proof."""

from collections import Counter
from copy import deepcopy

import cv2
import numpy as np
from openpyxl import load_workbook
import pytest

from test_modernization_dynamic_form_layout import copied_master, entries_for, WEEK
from test_modernization_printed_header_boundaries import (
    _assert_body, _observe, _sheet, form_factory,
)
from src.services import hakodate_cell_ocr_batch_service as batch
from src.services import hakodate_step_review_pipeline_service as review
from src.services import master_order_form_template_service as master
from src.services import workbook_pdf_renderer as renderer


PHYSICAL_8_8 = [0, 7, 8, 15, 16, 23, 24, 31, 32, 39, 40, 47, 48, 55, 56]
SEMANTIC_8_8 = [0, 8, 16, 24, 32, 40, 48, 56]


def _generated_fax(tmp_path, copied_master, canonical, entries, label):
    """Generate a different, internally valid physical source; never fake axes."""
    path = tmp_path / f"{label}.xlsx"
    master.build_facility_template_xlsx(
        facility_config=canonical.config, master_template_path=copied_master,
        output_path=path, week_value=WEEK, week_menu_entries=entries,
    )
    workbook = load_workbook(path)
    try:
        ws = workbook[master.FACILITY_TEMPLATE_SHEET_NAME]
        for offset in range(len(entries)):
            ws.cell(11 + offset, 5).value = [7, 0, None, "?"][offset % 4]
        workbook.save(path)
    finally:
        workbook.close()
    reloaded = load_workbook(path)
    try:
        ws = reloaded[master.FACILITY_TEMPLATE_SHEET_NAME]
        plan = master.validated_body_row_plan(ws)
        assert [row["date"] for row in plan] == [entry["menu_date"] for entry in entries]
        assert [row["name"] for row in plan] == [entry["name"] for entry in entries]
        assert [row["source_index"] for row in plan] == list(range(len(entries)))
        geometry = renderer.worksheet_render_geometry(ws, dpi=144)
        image = renderer.render_worksheet_to_image(ws, dpi=144)
        try:
            assert image.size == (geometry["image_width"], geometry["image_height"])
            image.save(tmp_path / f"{label}.png")
            bgr = cv2.cvtColor(np.asarray(image.convert("RGB")), cv2.COLOR_RGB2BGR)
        finally:
            image.close()
        xs = [float(geometry["x_positions"][col]) for col in range(1, ws.max_column + 2)]
        ys = [geometry["y_positions"][row] for row in (7, 9, *range(11, len(entries) + 12))]
        corrected, evidence = review._row_intersection_correct_ys(
            rectified_fax=bgr, corrected_xs=xs, template_ys=ys, worksheet=ws,
        )
        assert corrected is not None
        assert evidence["used"] is True
        assert evidence["structural_match"]["detected_body_band_count"] == len(entries)
        assert corrected[2:] == pytest.approx(ys[2:], abs=3)
        assert [ws.cell(11 + offset, 5).value for offset in range(len(entries))] == [
            [7, 0, None, "?"][offset % 4] for offset in range(len(entries))
        ]
        return bgr, evidence
    finally:
        reloaded.close()


@pytest.mark.parametrize("first,last,week", [
    (0, 56, WEEK),
    (0, 32, "2026-09@2026-09-27~2026-09-30"),
    (32, 56, "2026-10@2026-10-01~2026-10-03"),
], ids=["full-week", "September-target", "October-target"])
def test_weekday_separators_remain_physical_not_semantic_and_month_mapping_keeps_all_rows(
    form_factory, first, last, week,
):
    form = form_factory()
    quantities = deepcopy(form.quantities)
    corrected, evidence = _observe(form)
    _assert_body(form, corrected, evidence)
    structural = evidence["structural_match"]
    assert structural["body_day_boundary_indexes"] == SEMANTIC_8_8
    provenance = structural["canonical_day_source"]
    assert provenance["source"] == "validated_generated_template_schema.week_menu_physical_row_plan"
    assert provenance["date_column_boundary_source"] == "actual_worksheet_renderer.merge_aware_borders.column_A"
    assert provenance["canonical_date_column_boundary_indexes"] == PHYSICAL_8_8
    assert provenance["observed_date_column_boundary_indexes"] == PHYSICAL_8_8
    assert provenance["canonical_physical_row_count"] == provenance["observed_physical_row_count"] == 56
    subdivisions = structural["date_column_subdivision_boundaries"]
    assert [item["body_boundary_index"] for item in subdivisions] == [7, 15, 23, 31, 39, 47, 55]
    assert [item["worksheet_row"] for item in subdivisions] == [18, 26, 34, 42, 50, 58, 66]
    assert [item["date"] for item in subdivisions] == [
        "2026-09-27", "2026-09-28", "2026-09-29", "2026-09-30",
        "2026-10-01", "2026-10-02", "2026-10-03",
    ]
    for item in subdivisions:
        assert item["source"] == "canonical_rendered_weekday_or_date_cell_subdivision"
        assert item["observed_y"] == corrected[2 + item["body_boundary_index"]]
        assert item["body_boundary_index"] not in structural["body_day_boundary_indexes"]
    assert [day["first_body_row_index"] for day in provenance["days"]] == SEMANTIC_8_8[:-1]
    assert [day["stop_body_row_index"] for day in provenance["days"]] == SEMANTIC_8_8[1:]
    sheet = _sheet(form.entries[first:last], week)
    original = deepcopy(sheet)
    assert batch._validate_menu_day_boundaries(sheet, evidence) == list(range(first, last))
    assert sheet == original
    assert form.quantities == quantities
    assert len(corrected[2:]) == 57


def test_same_total_valid_generated_nine_seven_fax_cannot_use_eight_eight_canonical_source(
    form_factory, tmp_path, copied_master,
):
    canonical = form_factory()
    entries = []
    for day, sizes in enumerate([(3, 3, 3), (1, 3, 3), *([(2, 3, 3)] * 5)]):
        entries.extend(entries_for([day], sizes))
    assert list(Counter(entry["menu_date"] for entry in entries).values()) == [9, 7, 8, 8, 8, 8, 8]
    ids = _sheet(entries, WEEK)["physical_menu_row_ids"]
    assert len(ids) == len(set(ids)) == len(canonical.entries) == 56
    fax, own_evidence = _generated_fax(tmp_path, copied_master, canonical, entries, "valid-nine-seven")
    assert own_evidence["structural_match"]["body_day_boundary_indexes"] == [0, 9, 16, 24, 32, 40, 48, 56]
    assert batch._validate_menu_day_boundaries(_sheet(entries, WEEK), own_evidence) == list(range(56))
    original_quantities = deepcopy(canonical.quantities)
    corrected, evidence = _observe(canonical, image=fax)
    assert corrected is None
    assert evidence["used"] is False
    assert evidence["reason"] == "fax_date_column_physical_boundaries_changed"
    assert evidence["missing_date_column_boundary_indexes"] == []
    assert evidence["extra_date_column_boundary_indexes"] == [9]
    source = evidence["canonical_day_source"]
    assert source["canonical_physical_row_count"] == source["observed_physical_row_count"] == 56
    assert source["canonical_date_column_boundary_indexes"] == PHYSICAL_8_8
    assert source["observed_date_column_boundary_indexes"] == [0, 8, 9, 15, 16, 23, 24, 31, 32, 39, 40, 47, 48, 55, 56]
    assert [[canonical.ws.cell(row, 5).value] for row in range(11, 67)] == original_quantities


@pytest.mark.parametrize("missing", [[7], [7, 23, 55], [7, 15, 23, 31, 39, 47, 55]],
                         ids=["one-weekday", "some-weekdays", "all-weekdays"])
def test_approved_missing_weekday_rulings_accepts_real_days_without_fabricated_observations(
    form_factory, tmp_path, missing,
):
    # Approved change: missing optional weekday rulings no longer reject a valid FAX.
    form = form_factory()
    image = form.image.copy()
    left, right = map(round, form.xs[:2])
    for index in missing:
        y = form.ys[2 + index]
        image[y - 4:y + 5, left - 4:right - 4] = 255
        cv2.line(image, (left, y - 4), (left, y + 4), (0, 0, 0), 3)
    assert cv2.imwrite(str(tmp_path / "synthetic-optional-weekday-lines.png"), image)
    corrected, evidence = _observe(form, image=image)
    _assert_body(form, corrected, evidence)
    structural = evidence["structural_match"]
    assert structural["body_day_boundary_indexes"] == SEMANTIC_8_8
    source = structural["canonical_day_source"]
    assert source["canonical_physical_row_count"] == source["observed_physical_row_count"] == 56
    assert source["canonical_date_column_boundary_indexes"] == PHYSICAL_8_8
    assert source["observed_date_column_boundary_indexes"] == sorted(set(PHYSICAL_8_8) - set(missing))
    subdivisions = structural["date_column_subdivision_boundaries"]
    assert [item["body_boundary_index"] for item in subdivisions] == sorted(set(PHYSICAL_8_8) - set(SEMANTIC_8_8) - set(missing))
    for item in subdivisions:
        assert item["observed_y"] == corrected[2 + item["body_boundary_index"]]
    for first, last, week in ((0, 56, WEEK), (0, 32, "2026-09@2026-09-27~2026-09-30"),
                              (32, 56, "2026-10@2026-10-01~2026-10-03")):
        sheet = _sheet(form.entries[first:last], week)
        before = deepcopy(sheet)
        assert batch._validate_menu_day_boundaries(sheet, evidence) == list(range(first, last))
        assert sheet == before
    assert [[form.ws.cell(row, 5).value] for row in range(11, 67)] == form.quantities


@pytest.mark.parametrize("damage,index,missing,extra,optional_missing", [
    ("missing-day-end", 8, [8], [], []),
    ("extra-subdivision", 4, [], [4], []),
    ("missing-day-and-all-weekdays", 8, [8], [], [7, 15, 23, 31, 39, 47, 55]),
])
def test_date_column_line_loss_or_extra_line_stops_without_changing_body_count(
    form_factory, tmp_path, damage, index, missing, extra, optional_missing,
):
    form = form_factory()
    image = form.image.copy()
    y = form.ys[2 + index]
    left, right = map(round, form.xs[:2])
    if extra:
        cv2.line(image, (left, y), (right, y), (0, 0, 0), 3)
    else:
        for removed in [index, *optional_missing]:
            y = form.ys[2 + removed]
            image[y - 4:y + 5, left - 4:right - 4] = 255
            # Other columns still carry this exact observed body edge.
            cv2.line(image, (left, y - 4), (left, y + 4), (0, 0, 0), 3)
    assert cv2.imwrite(str(tmp_path / f"synthetic-{damage}.png"), image)
    corrected, evidence = _observe(form, image=image)
    assert corrected is None
    assert evidence["used"] is False
    assert evidence["reason"] == "fax_date_column_physical_boundaries_changed"
    assert evidence["missing_date_column_boundary_indexes"] == missing
    assert evidence["extra_date_column_boundary_indexes"] == extra
    source = evidence["canonical_day_source"]
    assert source["canonical_physical_row_count"] == source["observed_physical_row_count"] == 56
    assert source["canonical_date_column_boundary_indexes"] == PHYSICAL_8_8
    assert source["observed_date_column_boundary_indexes"] == sorted((set(PHYSICAL_8_8) - set(missing) - set(optional_missing)) | set(extra))
    assert [[form.ws.cell(row, 5).value] for row in range(11, 67)] == form.quantities


def test_valid_generated_shorter_physical_footprint_stops_instead_of_padding_or_truncating(
    form_factory, tmp_path, copied_master,
):
    canonical = form_factory()
    entries = canonical.entries[:-1]
    assert len(entries) == 55
    fax, own_evidence = _generated_fax(tmp_path, copied_master, canonical, entries, "valid-fifty-five")
    assert own_evidence["structural_match"]["body_day_boundary_indexes"] == [0, 8, 16, 24, 32, 40, 48, 55]
    assert batch._validate_menu_day_boundaries(_sheet(entries, WEEK), own_evidence) == list(range(55))
    corrected, evidence = _observe(canonical, image=fax)
    assert corrected is None
    assert evidence["used"] is False
    assert evidence["reason"] == "fax_physical_body_count_changed"
    assert evidence["canonical_day_source"]["canonical_physical_row_count"] == 56
    assert evidence["canonical_day_source"]["observed_physical_row_count"] == 55
    assert [[canonical.ws.cell(row, 5).value] for row in range(11, 67)] == canonical.quantities
