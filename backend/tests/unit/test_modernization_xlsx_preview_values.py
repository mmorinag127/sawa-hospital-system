"""Generated workbook/HTML unit boundary only; not operator, OCR, or staging proof."""

from datetime import date, datetime
from decimal import Decimal
from html.parser import HTMLParser

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from src.api.outputs import _render_xlsx_preview


class PreviewTable(HTMLParser):
    """Observe actual HTML cells and spans without replacing rendering logic."""

    def __init__(self, html):
        super().__init__(convert_charrefs=True)
        self.rows = []
        self.row = None
        self.cell = None
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        elif tag == "td":
            self.cell = {"text": "", "attrs": dict(attrs)}

    def handle_data(self, data):
        if self.cell is not None:
            self.cell["text"] += data

    def handle_endtag(self, tag):
        if tag == "td":
            self.row.append(self.cell)
            self.cell = None
        elif tag == "tr":
            self.rows.append(self.row)
            self.row = None

    def texts(self):
        return [[cell["text"] for cell in row] for row in self.rows]


def save_and_render(workbook, tmp_path, *, title="Unit workbook"):
    path = tmp_path / "unit-preview.xlsx"
    workbook.save(path)
    workbook.close()
    return path, _render_xlsx_preview(path, title)


def test_integer_and_decimal_zero_and_signed_numbers_are_visible(tmp_path):
    workbook = Workbook()
    workbook.active.append([0, 0.0, Decimal("0.0000"), 12, -3, 4.25, -0.125])
    path, html = save_and_render(workbook, tmp_path)
    persisted = load_workbook(path, data_only=True)
    try:
        # XLSX stores numeric values, not the original Python Decimal scale.
        assert [cell.value for cell in persisted.active[1]] == [0, 0, 0, 12, -3, 4.25, -0.125]
    finally:
        persisted.close()
    assert PreviewTable(html).texts() == [["0", "0", "0", "12", "-3", "4.25", "-0.125"]]


def test_none_is_blank_but_native_false_is_not_missing(tmp_path):
    workbook = Workbook()
    workbook.active.append([None, False, True, None, 0, "end"])
    path, html = save_and_render(workbook, tmp_path)
    persisted = load_workbook(path, data_only=True)
    try:
        assert persisted.active["A1"].value is None
        assert persisted.active["B1"].value is False
        assert persisted.active["C1"].value is True
        assert persisted.active["D1"].value is None
    finally:
        persisted.close()
    assert PreviewTable(html).texts() == [["", "False", "True", "", "0", "end"]]


def test_dates_and_text_keep_date_format_and_html_escaping(tmp_path):
    workbook = Workbook()
    workbook.active.title = "Dates & <text>"
    raw = '<script>alert("x")</script> & \'quoted\''
    workbook.active.append([date(2026, 3, 1), datetime(2026, 3, 7, 15, 30), raw, "0"])
    _, html = save_and_render(workbook, tmp_path, title="Preview <&>")
    assert PreviewTable(html).texts() == [["2026-03-01", "2026-03-07", raw, "0"]]
    assert "<h2>Dates &amp; &lt;text&gt;</h2>" in html
    assert "<h1>Preview &lt;&amp;&gt; プレビュー</h1>" in html
    assert "&lt;script&gt;alert(&quot;x&quot;)&lt;/script&gt; &amp; &#x27;quoted&#x27;" in html
    assert "<script>" not in html


def test_merged_spans_styles_and_hidden_sheet_layout_remain_intact(tmp_path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Visible & merged"
    worksheet.merge_cells("A1:B2")
    worksheet["A1"] = "merged anchor"
    worksheet["A1"].font = Font(bold=True)
    worksheet["A1"].alignment = Alignment(horizontal="center", wrap_text=True)
    worksheet["C1"] = 0
    worksheet["C2"] = -1
    worksheet.append(["last row", None, False])
    hidden = workbook.create_sheet("Hidden sentinel")
    hidden["A1"] = "hidden-content-must-not-render"
    hidden.sheet_state = "hidden"
    very_hidden = workbook.create_sheet("Very hidden sentinel")
    very_hidden["A1"] = "very-hidden-content-must-not-render"
    very_hidden.sheet_state = "veryHidden"
    second = workbook.create_sheet("Second visible")
    second["A1"] = 0
    _, html = save_and_render(workbook, tmp_path)
    table = PreviewTable(html)
    assert table.texts() == [["merged anchor", "0"], ["-1"], ["last row", "", "False"], ["0"]]
    anchor = table.rows[0][0]["attrs"]
    assert anchor["rowspan"] == "2" and anchor["colspan"] == "2"
    assert "font-weight:700" in anchor["style"]
    assert "text-align:center" in anchor["style"]
    assert "white-space:pre-wrap" in anchor["style"]
    assert all("rowspan" not in cell["attrs"] and "colspan" not in cell["attrs"]
               for row in table.rows for cell in row if cell["text"] != "merged anchor")
    assert html.count("<section>") == 2
    assert "<h2>Visible &amp; merged</h2>" in html
    assert "<h2>Second visible</h2>" in html
    assert "Hidden sentinel" not in html and "hidden-content-must-not-render" not in html
    assert "Very hidden sentinel" not in html and "very-hidden-content-must-not-render" not in html
