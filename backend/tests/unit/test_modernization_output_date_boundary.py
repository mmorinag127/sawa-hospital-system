"""Mounted output HTTP boundaries with owned SQLite and manually seeded OCR.

The actual app uses the existing TestClient without startup lifespan and with
AUTH_DISABLED. Real accepted CSV/XLSX production is separate from raising-only
producer fault adapters; neither proves UI/auth/native/formal PG/live/stg.
Omitted-date CSV faults test its existing unwrapped preview mode, not success.
The old 43 modernization nonpasses are not waived by these definitions.
"""

from copy import deepcopy
import codecs
import csv
from datetime import date
from hashlib import sha256
from io import BytesIO, StringIO

from openpyxl import load_workbook
import pytest
from sqlalchemy import event

from test_modernization_real_bagging_confirmation import (
    actual_bagging, owned_bagging as owned_bagging,
)
from test_modernization_workflow_surface_parity import (
    database_snapshot, surface as surface,
)
from src.api import outputs
from src.services import master_order_form_template_service as master
from src.services import output_builder


def _inputs(fixture):
    return (
        database_snapshot(), deepcopy(fixture.saved()),
        {path: sha256(path.read_bytes()).hexdigest() for path in (
            fixture.pdf, fixture.source, fixture.master_path, master.MASTER_TEMPLATE_PATH,
        )},
    )


def _assert_inputs(fixture, before):
    database, sheet, hashes = before
    assert database_snapshot() == database
    assert fixture.saved() == sheet
    assert all(sha256(path.read_bytes()).hexdigest() == digest for path, digest in hashes.items())


def _output_files(fixture):
    return {path: path.read_bytes() for path in fixture.output.rglob("*") if path.is_file()}


def _received_csv(response):
    assert response.status_code == 200, response.text
    body = response.content
    assert len(body) > 0
    assert int(response.headers["content-length"]) == len(body)
    assert response.headers["content-type"] == "text/csv; charset=ms932"
    assert codecs.lookup(response.encoding) == codecs.lookup("cp932")
    text = body.decode("cp932")
    assert response.text == text
    assert "\ufffd" not in text
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment;") and ".csv" in disposition
    return text


@pytest.mark.parametrize("route", [
    "/outputs/labels", "/outputs/labels.xlsx", "/outputs/delivery-notes/html",
    "/outputs/daily-delivery-notes/html", "/outputs/daily-bundle", "/outputs/weekly-weight",
])
@pytest.mark.parametrize("invalid", ["bad", "2026-02-30"])
def test_invalid_date_stops_before_sql_or_output(surface, owned_bagging, route, invalid):
    fixture, client = surface
    assert owned_bagging is fixture
    before = _inputs(fixture)
    files = _output_files(fixture)
    statements = []

    def observe(_connection, _cursor, statement, _parameters, _context, _many):
        statements.append(statement)

    params = {"date": invalid}
    if route in {"/outputs/labels", "/outputs/labels.xlsx", "/outputs/delivery-notes/html"}:
        params["order_id"] = fixture.order_id
    # Snapshot SELECTs are deliberately outside this request-only observation.
    event.listen(fixture.engine, "before_cursor_execute", observe)
    try:
        response = client.get(route, params=params)
    finally:
        event.remove(fixture.engine, "before_cursor_execute", observe)
    assert response.status_code == 400, response.text
    detail = "weekly_weight_date_invalid" if route == "/outputs/weekly-weight" and invalid == "bad" else "date must be YYYY-MM-DD"
    assert response.json() == {"detail": detail}
    assert statements == []
    assert _output_files(fixture) == files
    _assert_inputs(fixture, before)


def test_real_http_accepted_csv_xlsx_receive_full_body_and_preserve_inputs(
    surface, owned_bagging, monkeypatch, capsys,
):
    fixture, client = surface
    assert owned_bagging is fixture
    produced = actual_bagging(fixture)
    produced_before = deepcopy(produced)
    lines = produced["bagging_result"]["materialization_candidate"]["lines"]
    candidate = produced["bagging_result"]["materialization_candidate"]
    assert candidate["configured_input"]["policy"]["quantity_rules"]["zero_as_empty"] is True
    assert [(line["menu_name"], line["quantity_original"]) for line in lines] == [
        ("選択献立0", 7), ("選択献立3", 0),
    ]
    days = {output_builder._ensure_date(line["date"]) for line in lines}
    assert len(days) == 1 and None not in days
    target = days.pop()
    before = _inputs(fixture)
    assert [row[4] for row in before[1]["rows"]] == ["7", "", "", "0"]
    monkeypatch.setattr(output_builder, "OUTPUT_DIR", fixture.output)
    bodies = {}
    for suffix, media_type, extension in (
        ("labels", "text/csv", ".csv"),
        ("labels.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx"),
    ):
        response = client.get("/outputs/" + suffix, params={
            "order_id": fixture.order_id, "date": target.isoformat(),
        })
        assert response.status_code == 200, response.text
        body = response.content
        assert len(body) > 0
        assert int(response.headers["content-length"]) == len(body)
        assert response.headers["content-type"].split(";")[0] == media_type
        disposition = response.headers["content-disposition"]
        assert disposition.startswith("attachment;")
        assert extension in disposition
        assert target.isoformat() in disposition
        bodies[suffix] = body
        if suffix == "labels":
            assert "①" in _received_csv(response)
    reader = csv.DictReader(StringIO(bodies["labels"].decode("cp932"), newline=""))
    csv_rows = list(reader)
    assert len(csv_rows) == 2
    assert [row["メニュー"] for row in csv_rows] == ["副菜①", "副菜①"]
    assert {row["商品名１"]: (row[""], row["内容量"], row["内容詳細"])
            for row in csv_rows} == {
        "選択献立0": ("7人前", "700g", "100g"),
        "選択献立3": ("0人前", "0g", "100g"),
    }
    workbook = load_workbook(BytesIO(bodies["labels.xlsx"]), read_only=True, data_only=False)
    try:
        assert len(workbook.worksheets) == 1
        sheet = workbook.worksheets[0]
        assert (sheet.max_row, sheet.max_column) == (3, 11)
        assert [sheet.cell(2, col).value for col in (7, 9, 10, 11)] == [
            "選択献立0", "700g", "100g", "7人前",
        ]
        assert [sheet.cell(3, col).value for col in (7, 9, 10, 11)] == [
            "選択献立3", "0g", "100g", "0人前",
        ]
        # Empty XLSX cells/headers roundtrip as None, CSV as empty strings.
        headers = ["" if cell.value is None else str(cell.value) for cell in sheet[1]]
        assert headers == reader.fieldnames
        rows = [dict(zip(headers, ("" if value is None else str(value) for value in row), strict=True))
                for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert rows == csv_rows
    finally:
        workbook.close()
    # The existing no-date preview modes do not append the explicit zero label;
    # aggregate also omits zero under the already captured zero_as_empty policy.
    for route, artifact_name in (
        ("labels", "labels-omitted-date.csv"),
        ("manufacturing-aggregate", "aggregate-omitted-date.csv"),
    ):
        response = client.get("/outputs/" + route, params={"order_id": fixture.order_id})
        text = _received_csv(response)
        rows = list(csv.DictReader(StringIO(text, newline="")))
        assert len(rows) == 1
        assert [(row["商品名１"], row[""], row["内容量"], row["内容詳細"]) for row in rows] == [
            ("選択献立0", "7人前", "700g", "100g"),
        ]
        assert "選択献立3" not in text
        if route == "labels":
            assert rows[0]["メニュー"] == "副菜①"
            assert "①" in text
        bodies[artifact_name] = response.content
        _assert_inputs(fixture, before)
    # These are HTTP-received bytes, distinct from server-generated files.
    received = fixture.root / "http-received"
    received.mkdir()
    for name, body in bodies.items():
        filename = {"labels": "labels-dated.csv", "labels.xlsx": "labels-dated.xlsx"}.get(name, name)
        path = received / filename
        path.write_bytes(body)
        assert path.read_bytes() == body
        with capsys.disabled():
            print(f"http-received artifact={path} sha256={sha256(body).hexdigest()}")
    assert produced == produced_before
    _assert_inputs(fixture, before)


@pytest.mark.parametrize("route,producer,prefix,dated", [
    ("labels", "build_label_csv_for_order", "label csv build failed", True),
    ("labels.xlsx", "build_label_workbook_for_order", "label workbook build failed", True),
    ("labels.xlsx", "build_label_workbook_for_order", "label workbook build failed", False),
    ("delivery-notes/html", "build_delivery_preview", "delivery note html build failed", True),
    ("delivery-notes/html", "build_delivery_preview", "delivery note html build failed", False),
])
@pytest.mark.parametrize("exception_type,status", [(ValueError, 400), (RuntimeError, 500)])
def test_raising_producer_boundary_keeps_status_and_exact_arguments(
    surface, monkeypatch, route, producer, prefix, dated, exception_type, status,
):
    fixture, client = surface
    before = _inputs(fixture)
    files = _output_files(fixture)
    calls = []
    marker = exception_type("synthetic-output-boundary-fault")

    def fail(*args, **kwargs):
        calls.append((args, kwargs))
        raise marker

    monkeypatch.setattr(outputs, producer, fail)
    params = {"order_id": fixture.order_id}
    if dated:
        params["date"] = "2026-03-01"
    response = client.get("/outputs/" + route, params=params)
    assert response.status_code == status, response.text
    assert response.json() == {"detail": str(marker) if status == 400 else f"{prefix}: {marker}"}
    kwargs = {"target_date": date(2026, 3, 1) if dated else None}
    if producer == "build_delivery_preview":
        kwargs["include_diagnostics"] = False
    assert calls == [((fixture.order_id,), kwargs)]
    assert _output_files(fixture) == files
    _assert_inputs(fixture, before)


@pytest.mark.parametrize("exception_type", [ValueError, RuntimeError])
def test_omitted_date_csv_uses_unwrapped_preview_fault_contract(surface, monkeypatch, exception_type):
    fixture, client = surface
    before = _inputs(fixture)
    files = _output_files(fixture)
    calls = []
    marker = exception_type("synthetic-preview-fault")

    def fail(*args, **kwargs):
        calls.append((args, kwargs))
        raise marker

    monkeypatch.setattr(outputs, "build_output_preview", fail)
    with pytest.raises(exception_type) as caught:
        client.get("/outputs/labels", params={"order_id": fixture.order_id})
    assert caught.value is marker
    assert calls == [((fixture.order_id, "labels"), {})]
    assert _output_files(fixture) == files
    _assert_inputs(fixture, before)
