"""Owned JSON/SQLite and real in-process monthly API definitions.

Local AUTH_DISABLED mode and no lifespan: not production auth, formal schema,
native OCR, operator-flow, or staging proof. Race stimuli change actual owned
files after the real loader/writer runs; no selector or receipt guard is mocked.
"""

from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from modernization_saved_sheet_export_fixtures import FAMILY, SavedSheetFixture, WEEK
from test_modernization_dynamic_source_forms import database_snapshot
from src.models.menu import MonthlyMenu, MonthlyMenuEntry, MonthlyMenuItem
from src.services import menu_service
from src.services import order_form_service as forms


@pytest.fixture
def monthly(tmp_path, monkeypatch):
    saved = SavedSheetFixture(tmp_path, monkeypatch)
    monkeypatch.setenv("AUTH_DISABLED", "true")
    monkeypatch.setattr(menu_service, "engine", saved.engine)
    monkeypatch.setattr(menu_service, "session_scope", saved.session_scope)
    monkeypatch.setattr(menu_service, "_MENU_SCHEMA_INITIALIZED", False)
    with saved.session_scope() as session:
        session.add(MonthlyMenu(id="2026-03", month_start=date(2026, 3, 1)))
        session.flush()
        for day in range(1, 32):
            name = f"dated-month-menu-{day}"
            session.add(MonthlyMenuEntry(
                id=f"dated-entry-{day}", monthly_menu_id="2026-03",
                menu_date=date(2026, 3, day), daypart="朝", category="主",
                name=name, diet_type="regular", facility_override="F", slot_index=1))
            session.add(MonthlyMenuItem(
                id=f"dated-item-{day}", monthly_menu_id="2026-03", name=name,
                daypart="朝", category="主", diet_type="regular", facility_override="F",
                unit_type="g", qty_per_serving=100, master_resolution_mode="month_only"))
    try:
        yield saved
    finally:
        saved.close()


def generate(params=None):
    from src.main import app

    client = TestClient(app)  # Deliberately no lifespan/startup-schema claim.
    try:
        return client.post("/order-forms/generate", params={
            "facility_id": "F", "month_id": "2026-03", **(params or {}),
        })
    finally:
        client.close()


def assert_no_publication(saved, before):
    assert database_snapshot(saved) == before
    assert list(saved.output.iterdir()) == []


def test_conflicting_full_root_never_overrides_sole_dated_month(monthly):
    saved = monthly
    root = saved.master["facilities"][0]
    root.update(deepcopy(saved.authored))
    root.update(facility_name="DO NOT USE ROOT", order_form_source_workbook="missing-root-source.xlsx",
                fax_template_id="wrong-root-template")
    root["fax_template"]["template_id"] = "wrong-root-template"
    saved.write_master()  # Versions and their genuine existing projections stay unchanged.
    before = database_snapshot(saved)
    source_before = saved.source.read_bytes()
    response = generate()
    assert response.status_code == 200, response.text if response.status_code != 200 else ""
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f"order_form_F_2026-03_{FAMILY}_" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content))
    try:
        meta = {row[0]: row[1] for row in workbook["設定"].iter_rows(values_only=True) if row[0] is not None}
        assert meta["facility_name"] == "保存済み施設"
        assert meta["source_workbook"] == saved.source.name
        assert meta["fax_template_id"] == FAMILY
        assert meta["entry_count"] == 31
        assert workbook[WEEK]["D11"].value == "dated-month-menu-1"
        assert workbook[WEEK]["J5"].value == "authored-target"
    finally:
        workbook.close()
    assert saved.source.read_bytes() == source_before
    assert database_snapshot(saved) == before


@pytest.mark.parametrize("case,status,detail", [
    ("missing-day", 409, {"error": "facility_version_date_unrecorded", "facility_id": "F", "effective_date": "2026-03-31"}),
    ("gap", 409, {"error": "facility_version_date_unrecorded", "facility_id": "F", "effective_date": "2026-03-16"}),
    ("overlap", 409, {"error": "facility_version_period_overlap"}),
    ("invalid-snapshot", 422, {"error": "facility_version_snapshot_incomplete", "version_id": "authored-March"}),
    ("mixed-month", 409, {"error": "order_form_month_multiple_versions_unresolved", "facility_id": "F", "month_id": "2026-03", "version_ids": ["authored-March", "authored-March-second"]}),
])
def test_actual_api_canonical_month_rejections_never_fall_back(monthly, case, status, detail):
    saved = monthly
    versions = saved.master["facilities"][0]["versions"]
    if case == "missing-day":
        versions[0]["valid_to"] = "2026-03-30"
    elif case == "invalid-snapshot":
        del versions[0]["facility"]["areas"]
    else:
        second = deepcopy(versions[0])
        second["version_id"] = "authored-March-second"
        versions[0]["valid_to"] = "2026-03-15"
        second["valid_from"] = {"gap": "2026-03-17", "overlap": "2026-03-15", "mixed-month": "2026-03-16"}[case]
        versions.append(second)
    # Raw authored tamper only: never materialize a valid receipt for invalid JSON.
    saved.write_master()
    before = database_snapshot(saved)
    response = generate()
    assert response.status_code == status, response.text
    assert response.json() == {"detail": detail}
    assert_no_publication(saved, before)
    # Mixed-month split policy remains unanswered; this proves only the explicit stop.


@pytest.mark.parametrize("case,detail", [
    ("unregistered", "order form pattern not found"),
    ("duplicate", "order form pattern ambiguous"),
    ("facility-unregistered", "order form pattern not found"),
])
def test_explicit_patterns_do_not_select_first_or_default(monthly, case, detail):
    saved = monthly
    saved.master["order_form_patterns"] = [{"pattern_id": "first-unused", "label": "must not select"}]
    params = {"pattern_id": "explicit"}
    if case == "duplicate":
        saved.master["order_form_patterns"].extend([
            {"pattern_id": "explicit", "label": "A"}, {"pattern_id": "explicit", "label": "B"}])
    elif case == "facility-unregistered":
        saved.authored["order_form_pattern_id"] = "explicit"
        params = {}
    saved.write_master()
    before = database_snapshot(saved)
    response = generate(params)
    assert response.status_code == 400, response.text
    assert response.json() == {"detail": detail}
    assert_no_publication(saved, before)


@pytest.mark.parametrize("change,boundary", [
    ("json", "load"), ("source", "load"), ("json", "write"), ("source", "write"),
])
def test_actual_input_changes_before_or_after_owned_output_write_block_publication(monthly, monkeypatch, change, boundary):
    saved = monthly
    before = database_snapshot(saved)
    json_before, source_before = saved.master_path.read_bytes(), saved.source.read_bytes()
    actual_load, actual_open = forms.load_workbook, Path.open
    observed = []

    def change_actual_input():
        if change == "json":
            saved.authored["facility_name"] = "actual concurrent authored change"
            saved.write_master()
        else:
            workbook = actual_load(saved.source)
            try:
                workbook[WEEK]["J5"] = "actual concurrent source change"
                workbook.save(saved.source)
            finally:
                workbook.close()

    if boundary == "load":
        def load_then_change(*args, **kwargs):
            workbook = actual_load(*args, **kwargs)
            if not observed:
                observed.append("real source loaded")
                change_actual_input()
            return workbook

        monkeypatch.setattr(forms, "load_workbook", load_then_change)
    else:
        class ChangeAfterRealWrite:
            def __init__(self, path, stream):
                self.path, self.stream = path, stream

            def __enter__(self):
                return self

            def write(self, data):
                written = self.stream.write(data)
                self.stream.flush()
                assert self.path.is_file() and self.path.stat().st_size > 0
                observed.append("real owned output written")
                change_actual_input()
                return written

            def __exit__(self, *exc):
                return self.stream.__exit__(*exc)

        def open_then_observe(path, mode="r", *args, **kwargs):
            stream = actual_open(path, mode, *args, **kwargs)
            if mode == "xb" and path.is_relative_to(saved.output):
                assert path.parent.parent == saved.output
                return ChangeAfterRealWrite(path, stream)
            return stream

        monkeypatch.setattr(Path, "open", open_then_observe)
    response = generate()
    if change == "json":
        assert response.status_code == 409, response.text
        assert response.json() == {"detail": {"error": "facility_master_revision_conflict"}}
        assert saved.master_path.read_bytes() != json_before
        assert saved.source.read_bytes() == source_before
    else:
        assert response.status_code == 400, response.text
        assert response.json() == {"detail": "monthly order form source workbook changed"}
        assert saved.source.read_bytes() != source_before
        assert saved.master_path.read_bytes() == json_before
    assert observed == ["real source loaded" if boundary == "load" else "real owned output written"]
    assert_no_publication(saved, before)
