"""Owned synthetic 2099-06-07 inputs; real lineage, PDF receipt and sheet capture.

Pattern: tests/modernization_history_input_fixture.py. No operator data, OCR
recognizer substitution, or precomputed successful materialization receipt.
"""
from copy import deepcopy
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
import os

import fitz
from openpyxl import Workbook

from src.models.document import OrderDocument
from src.models.facility import Facility
from src.models.order_ocr_evidence_run import OrderOcrEvidenceRun
from src.models.order_sheet_draft import OrderSheetDraft
from src.models.order_workflow_state import OrderWorkflowState
from src.models.uploaded_pdf import UploadedPdf
from src.services import canonical_facility_lineage as lineage
from src.services import order_form_service, order_service, order_workflow_v2_service as workflow

FACILITY = "readonly-boundary-facility"
WEEK = "2099-06@2099-06-07~2099-06-13"
DAY = date(2099, 6, 7)
FAMILY = "fax_layout_regular_forbidden_v1"


def configure(root, monkeypatch):
    root = Path(root)
    source = root / "readonly-source.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "6月7日～6月13日"
    sheet["A2"] = "【発　注　連　絡　表】"
    sheet["D7"] = "献立"
    for cell, label in (("E7", "常食"), ("F7", "肉禁"), ("G7", "魚禁"), ("H7", "職員")):
        sheet[cell] = label
    for index, value in enumerate(["2099-06-07", "朝食", "主菜", "Dish", 17, "", "", ""], 1):
        sheet.cell(11, index, value)
    book.save(source)
    book.close()
    columns = [{"index": i, "source_index": i, "role": role}
               for i, role in enumerate(("date", "daypart", "category", "menu_name"))]
    columns += [{"index": i, "source_index": i, "role": "quantity", "diet_type": diet, "area_id": "X"}
                for i, diet in enumerate(("regular", "no_meat", "no_fish", "staff"), 4)]
    facility = {"facility_id": FACILITY, "facility_name": "Synthetic readonly boundary facility",
        "areas": None, "fax_template_id": FAMILY, "order_form_source_workbook": source.name,
        "fax_template": {"template_id": FAMILY, "units": "normalized", "table_box": [0, 0, 1, 1],
                         "header_rows": 0, "columns": columns},
        "bag_types": [{"bag_type_id": "READONLY-1000G", "label": "1000g", "max_qty": 1000, "unit": "g"}],
        "packaging_policy": {"split_key": ["facility", "date", "daypart", "menu_name", "diet_type", "area_id", "bag_type"]}}
    master = root / "readonly-facilities.json"
    master.write_text(json.dumps({"facilities": [{"facility_id": FACILITY,
        "facility_name": facility["facility_name"], "versions": [{"version_id": "readonly-2099-06",
        "valid_from": "2099-06-01", "valid_to": "2099-06-30", "facility": facility}]}]}, ensure_ascii=False))
    output = root / "readonly-output"
    output.mkdir()
    monkeypatch.setenv("FACILITY_MASTER_PATH", str(master))
    monkeypatch.setenv("ORDERS_OUTPUT_DIR", str(output))
    monkeypatch.setattr(order_form_service, "_FAX_SOURCE_TEMPLATE_DIR", root)
    monkeypatch.setattr(order_form_service, "_OUTPUT_DIR", output)


def bind_order(session, order):
    root = Path(os.environ["FACILITY_MASTER_PATH"]).parent
    if session.get(Facility, FACILITY) is None:
        session.add(Facility(id=FACILITY, name="Synthetic readonly boundary facility"))
        session.flush()
    selection, _ = lineage.resolve(session, FACILITY, DAY)
    _, version = lineage.materialize(session, selection)
    order.facility_code = FACILITY
    order.template_version_id = version.id
    order.received_at = datetime(2099, 6, 7, 9)
    pdf = root / "operator-order.pdf"
    document = fitz.open()
    document.new_page().insert_text((50, 50),
        "Synthetic readonly boundary facility / 2099-06-07..13 / operator-order; original 17, corrected 19")
    document.save(pdf)
    document.close()
    order.document_uri = str(pdf)
    session.add(OrderDocument(id="operator-document", facility_code=FACILITY, week_code=WEEK,
        storage_uri=str(pdf), source_email_id=order.message_id, received_at=order.received_at, status="done"))
    session.flush()
    order.current_document_id = "operator-document"
    session.add(UploadedPdf(id="operator-pdf", message_id=order.message_id, storage_uri=str(pdf),
        original_filename=pdf.name, content_sha256=sha256(pdf.read_bytes()).hexdigest(),
        received_at=order.received_at, page_count=1, current_document_id=order.current_document_id, current_order_id=order.id))
    evidence = {"tables": [{"rows": [["2099-06-07", "朝食", "主菜", "Dish", 17, "", "", ""]]}]}
    session.add(OrderOcrEvidenceRun(id="operator-evidence", order_id=order.id, template_version_id=version.id,
        schema_version="v2", producer_version="synthetic-readonly/v1", source="unit-fixture", status="ready",
        payload_json=evidence, artifact_digest=sha256(json.dumps(evidence).encode()).hexdigest()))
    return selection


def state(order, *, draft_id="draft"):
    return OrderWorkflowState(order_id=order.id, state="sheet_saved", draft_id=draft_id,
        evidence_run_id="operator-evidence", template_version_id=order.template_version_id, blockers_json=[],
        secondary_actions_json={workflow.WORKFLOW_V2_META_KEY: {"facility_id": FACILITY,
            "week_start": "2099-06-07", "week_end": "2099-06-13", "template_version_id": order.template_version_id}})


def save_source(session, order, name="Dish", quantity=19):
    order.week_code = WEEK
    rows = [["2099-06-07", "朝食", "主菜", name, quantity, "", "", ""]]
    draft = OrderSheetDraft(id="draft", order_id=order.id, template_version_id=order.template_version_id,
        base_evidence_run_id="operator-evidence", base_template_resolution_id=FAMILY, draft_state="saved",
        draft_sheet_json={"base_evidence_run_id": "operator-evidence",
            "fields": ["date", "daypart", "category", "menu_name", "qty.regular_x", "qty.no_meat_x", "qty.no_fish_x", "qty.staff_x"],
            "rows": rows})
    session.add(draft)
    session.flush()
    session.add(state(order))
    session.flush()
    order_service._record_current_order_version(session, order, session.get(OrderDocument, order.current_document_id))
    return draft


def capture(session, order, draft):
    result = workflow._build_materialization_candidate_for_saved_sheet(session=session, order=order, saved_sheet=draft)
    assert not result.get("error"), result
    assert result["configured_input"]["facility"]["facility_id"] == FACILITY
    assert result["lines"]
    return deepcopy(result)
