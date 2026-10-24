"""Physical tabular invoice specimens; never a production template fallback."""
from copy import deepcopy
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Border, Side


def declare_invoice(root, template):
    declared = deepcopy(template)
    assert declared["columns"], "explicit component columns required"
    assert "template_uri" not in declared, "never replace a declared original template"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "component-invoice"
    edge = Side(style="thin", color="000000")
    for index, column in enumerate(declared["columns"], 1):
        assert column["name"] and column["source"]
        sheet.cell(1, index, column.get("header") or column["name"])
        sheet.cell(2, index).border = Border(left=edge, right=edge, top=edge, bottom=edge)
    path = root / ("authored-invoice-" + uuid4().hex + ".xlsx")
    workbook.save(path)
    workbook.close()
    declared.update(template_uri=str(path), sheet_name=sheet.title)
    return declared
